"""Build docs/CloseDesk-how-it-works.pptx. Run from the repo root."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x12, 0x24, 0x3F)
NAVY_2 = RGBColor(0x1B, 0x36, 0x5D)
INK = RGBColor(0x1B, 0x24, 0x16)
MUTED = RGBColor(0x5E, 0x6A, 0x59)
PAPER = RGBColor(0xF6, 0xF4, 0xEF)
WHITE = RGBColor(0xFF, 0xFC, 0xF7)
GOLD = RGBColor(0xC4, 0xA1, 0x5A)
DANGER = RGBColor(0x9B, 0x23, 0x35)

W = Inches(13.333)
H = Inches(7.5)


def _fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _box(slide, x, y, w, h, color: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    _fill(shape, color)
    return shape


def _text(slide, text, x, y, w, h, *, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri"):
    shape = slide.shapes.add_textbox(x, y, w, h)
    frame = shape.text_frame
    frame.word_wrap = True
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return shape


def _bullets(slide, lines, x, y, w, h, *, size=18, color=INK):
    shape = slide.shapes.add_textbox(x, y, w, h)
    frame = shape.text_frame
    frame.word_wrap = True
    for index, line in enumerate(lines):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = "Calibri"
    return shape


def _chrome(slide, kicker: str, title: str) -> None:
    _box(slide, 0, 0, W, H, PAPER)
    _box(slide, 0, 0, Inches(0.16), H, GOLD)
    _text(slide, kicker.upper(), Inches(0.55), Inches(0.32), Inches(12), Inches(0.32), size=12, bold=True, color=GOLD)
    _text(slide, title, Inches(0.55), Inches(0.58), Inches(12.2), Inches(0.7), size=32, bold=True, color=NAVY, font="Georgia")


def _footer(slide, page: int, pages: int) -> None:
    _text(slide, "CloseDesk  ·  scripts extract, Bionic decides", Inches(0.55), Inches(7.05), Inches(9), Inches(0.3), size=11, color=MUTED)
    _text(slide, f"{page}  /  {pages}", Inches(11.4), Inches(7.05), Inches(1.4), Inches(0.3), size=11, color=MUTED, align=PP_ALIGN.RIGHT)


def build(path: Path) -> None:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "CloseDesk — how it works"
    prs.core_properties.subject = "Overnight inbox reading with a local Bionic model"
    blank = prs.slide_layouts[6]
    pages = 13

    # 1 title
    slide = prs.slides.add_slide(blank)
    _box(slide, 0, 0, W, H, NAVY)
    _box(slide, 0, 0, Inches(0.16), H, GOLD)
    _text(slide, "CLOSEDESK", Inches(0.7), Inches(1.7), Inches(11), Inches(0.4), size=14, bold=True, color=GOLD)
    _text(slide, "The inbox reads itself overnight.", Inches(0.7), Inches(2.15), Inches(11), Inches(1.4), size=44, bold=True, color=WHITE, font="Georgia")
    _text(
        slide,
        "Scripts pull the text, the amounts, and the dates.\nA local model in Bionic decides what each message is and where it goes.",
        Inches(0.7),
        Inches(4.0),
        Inches(10),
        Inches(1.2),
        size=20,
        color=RGBColor(0xD5, 0xDC, 0xCF),
    )
    _text(slide, "Daily digest   ·   Important   ·   Action items   ·   Informational   ·   Reference", Inches(0.7), Inches(6.3), Inches(11), Inches(0.4), size=16, color=GOLD)

    # 2 problem
    slide = prs.slides.add_slide(blank)
    _chrome(slide, "The problem", "A heavy inbox all looks the same.")
    _bullets(
        slide,
        [
            "Invoices, wires, bank statements, payroll, audit requests, and newsletters arrive in one list.",
            "Opening each one to find out what it is takes longer than the work itself.",
            "The dangerous note — “our bank details changed, please wire today” — looks like the rest.",
            "You want a morning pile, not another inbox to scroll.",
        ],
        Inches(0.55),
        Inches(1.7),
        Inches(12),
        Inches(3.2),
        size=22,
    )
    _box(slide, Inches(0.55), Inches(5.2), Inches(12.1), Inches(1.45), NAVY)
    _text(
        slide,
        "CloseDesk’s morning is five places: the digest, Important, the action list, Informational, and Reference.",
        Inches(0.8),
        Inches(5.5),
        Inches(11.6),
        Inches(0.9),
        size=20,
        color=WHITE,
    )
    _footer(slide, 2, pages)

    # 3 split
    slide = prs.slides.add_slide(blank)
    _chrome(slide, "The split", "Scripts are the tools. Bionic is the reader.")
    _box(slide, Inches(0.55), Inches(1.7), Inches(5.9), Inches(4.9), WHITE)
    _text(slide, "Scripts", Inches(0.8), Inches(1.9), Inches(5.4), Inches(0.45), size=22, bold=True, color=NAVY, font="Georgia")
    _bullets(
        slide,
        [
            "Read PDF, Excel, Word, PowerPoint, CSV, ZIP, .msg, and .eml.",
            "Pull invoice numbers, amounts, due dates, and last-4 of account numbers.",
            "Flag a payment-instruction change.",
            "Draft a folder so the board is never empty.",
        ],
        Inches(0.8),
        Inches(2.55),
        Inches(5.3),
        Inches(3.6),
        size=16,
    )
    _box(slide, Inches(6.7), Inches(1.7), Inches(6.0), Inches(4.9), NAVY)
    _text(slide, "Bionic", Inches(6.95), Inches(1.9), Inches(5.5), Inches(0.45), size=22, bold=True, color=GOLD, font="Georgia")
    _bullets(
        slide,
        [
            "Reads the packet. It does not re-open the PDF.",
            "Chooses the category, the folder, and the importance.",
            "Writes the one-line summary you scan at breakfast.",
            "Writes the action items, or keeps the script’s list.",
        ],
        Inches(6.95),
        Inches(2.55),
        Inches(5.4),
        Inches(3.6),
        size=16,
        color=WHITE,
    )
    _footer(slide, 3, pages)

    # 4 drop folder
    slide = prs.slides.add_slide(blank)
    _chrome(slide, "On the laptop", "A folder is enough. Outlook login is optional.")
    cards = [
        ("inbox/incoming", "Drop .msg and .eml files here. A loose PDF or workbook is read on its own."),
        ("inbox/attachments", "If the files were saved separately, put them in a folder named like the message."),
        ("inbox/processed", "Originals move here after a successful read, under the day’s date."),
        ("inbox/extracted", "Unpacked copies, so you can open the file without hunting through mail."),
    ]
    for index, (title, body) in enumerate(cards):
        col = index % 2
        row = index // 2
        x = Inches(0.55 + col * 6.3)
        y = Inches(1.75 + row * 2.4)
        _box(slide, x, y, Inches(6.0), Inches(2.15), WHITE)
        _text(slide, title, x + Inches(0.25), y + Inches(0.25), Inches(5.5), Inches(0.4), size=20, bold=True, color=NAVY, font="Georgia")
        _text(slide, body, x + Inches(0.25), y + Inches(0.85), Inches(5.5), Inches(1.0), size=16, color=INK)
    _footer(slide, 4, pages)

    # 5 what bionic returns
    slide = prs.slides.add_slide(blank)
    _chrome(slide, "One reading", "What Bionic sends back for each message.")
    fields = [
        ("Category", "Invoice, remittance, statement, tax notice, newsletter, and the rest."),
        ("Folder", "Important, informational, or reference."),
        ("Importance", "Critical, high, medium, or low."),
        ("Summary", "One sentence. This is what the list shows."),
        ("Actions", "Real tasks, with a due date when the packet has one."),
        ("Why", "A sentence you can check when the filing looks odd."),
    ]
    for index, (title, body) in enumerate(fields):
        col = index % 3
        row = index // 3
        x = Inches(0.55 + col * 4.2)
        y = Inches(1.8 + row * 2.35)
        _box(slide, x, y, Inches(4.0), Inches(2.1), WHITE)
        _text(slide, title, x + Inches(0.22), y + Inches(0.28), Inches(3.5), Inches(0.4), size=20, bold=True, color=NAVY, font="Georgia")
        _text(slide, body, x + Inches(0.22), y + Inches(0.9), Inches(3.5), Inches(0.9), size=15, color=INK)
    _footer(slide, 5, pages)

    # 6 morning places
    slide = prs.slides.add_slide(blank)
    _chrome(slide, "In the morning", "Five places. You do not start at the top of the inbox.")
    places = [
        ("Daily digest", "Fraud first, then overdue, due today, invoices, and cash."),
        ("Important", "Needs a decision, a payment check, or a close task."),
        ("Action items", "The task list. Mark them done, or export the CSV."),
        ("Informational", "Newsletters and FYI. Nothing is waiting on you."),
        ("Reference", "Statements, POs, contracts. Keep them. Not tonight’s work."),
    ]
    for index, (title, body) in enumerate(places):
        y = Inches(1.6 + index * 1.0)
        _box(slide, Inches(0.55), y, Inches(12.2), Inches(0.88), WHITE if index % 2 == 0 else RGBColor(0xFF, 0xF8, 0xEE))
        _text(slide, title, Inches(0.75), y + Inches(0.18), Inches(3.2), Inches(0.5), size=18, bold=True, color=NAVY, font="Georgia")
        _text(slide, body, Inches(4.1), y + Inches(0.2), Inches(8.3), Inches(0.5), size=16, color=INK)
    _footer(slide, 6, pages)

    # 7 overnight
    slide = prs.slides.add_slide(blank)
    _chrome(slide, "While you sleep", "One command runs the night.")
    steps = [
        ("1", "Read", "New files in the drop folder, and Outlook if you connected it."),
        ("2", "Packet", "Text, amounts, dates, filenames. The raw file stays closed."),
        ("3", "File", "The local model files up to 40 drafts."),
        ("4", "Leave", "Folders, the digest, and data/overnight/YYYY-MM-DD.md."),
    ]
    for index, (num, title, body) in enumerate(steps):
        x = Inches(0.55 + index * 3.2)
        _box(slide, x, Inches(2.1), Inches(3.0), Inches(3.6), NAVY if index % 2 == 0 else NAVY_2)
        _text(slide, num, x + Inches(0.2), Inches(2.3), Inches(2.6), Inches(0.6), size=28, bold=True, color=GOLD, font="Georgia")
        _text(slide, title, x + Inches(0.2), Inches(3.05), Inches(2.6), Inches(0.45), size=20, bold=True, color=WHITE, font="Georgia")
        _text(slide, body, x + Inches(0.2), Inches(3.6), Inches(2.6), Inches(1.7), size=15, color=RGBColor(0xE8, 0xEA, 0xDF))
    _text(slide, "python -m controller_inbox overnight", Inches(0.55), Inches(6.0), Inches(12), Inches(0.4), size=18, bold=True, color=NAVY, font="Consolas")
    _footer(slide, 7, pages)

    # 8 two ways
    slide = prs.slides.add_slide(blank)
    _chrome(slide, "Two ways to run it", "Same model. Same tools. Pick who drives.")
    _box(slide, Inches(0.55), Inches(1.75), Inches(6.0), Inches(4.8), WHITE)
    _text(slide, "Unattended", Inches(0.8), Inches(2.0), Inches(5.5), Inches(0.45), size=22, bold=True, color=NAVY, font="Georgia")
    _bullets(
        slide,
        [
            "Leave LM Studio’s local server on.",
            "Run the overnight command before you stop.",
            "In the morning, open the dashboard.",
            "Rows still waiting say so. Run it again if the server was off.",
        ],
        Inches(0.8),
        Inches(2.7),
        Inches(5.4),
        Inches(3.4),
        size=16,
    )
    _box(slide, Inches(6.8), Inches(1.75), Inches(6.0), Inches(4.8), WHITE)
    _text(slide, "In Bionic Studio", Inches(7.05), Inches(2.0), Inches(5.5), Inches(0.45), size=22, bold=True, color=NAVY, font="Georgia")
    _bullets(
        slide,
        [
            "Install the closedesk-inbox skill.",
            "Ask Bionic to process the waiting queue.",
            "It calls prepare_queue, then save_reading, then the digest.",
            "Use this when you want to watch a batch, or correct course.",
        ],
        Inches(7.05),
        Inches(2.7),
        Inches(5.4),
        Inches(3.4),
        size=16,
    )
    _footer(slide, 8, pages)

    # 9 install
    slide = prs.slides.add_slide(blank)
    _chrome(slide, "Ship it into Bionic", "The skill is a folder. Bionic already knows this format.")
    _bullets(
        slide,
        [
            "The skill lives at bionic/closedesk-inbox/SKILL.md.",
            "In Bionic: Settings → Skills, and add that file. Or type @Install Skill.",
            "Open this project as Bionic’s working folder. The tools are python -m controller_inbox tool …",
            "Set CONTROLLER_INBOX_LLM=true and point LLM_BASE_URL at http://127.0.0.1:1234/v1.",
            "local-model means “whatever LM Studio has loaded.” Pin a model id only if you want one.",
            "The written steps are in docs/BIONIC_GUIDE.md.",
        ],
        Inches(0.55),
        Inches(1.7),
        Inches(12.2),
        Inches(4.8),
        size=20,
    )
    _footer(slide, 9, pages)

    # 10 corrections
    slide = prs.slides.add_slide(blank)
    _chrome(slide, "When it is wrong", "Say why. That reason is the next lesson.")
    _bullets(
        slide,
        [
            "Open the message and use Wrong category? A sentence is required.",
            "The next message from that sender follows the correction.",
            "The overnight run will not overwrite a message you already fixed.",
            "The reason is handed to Bionic inside the packet for later mail from that sender.",
            "The same line is saved in data/training/corrections.jsonl.",
            "A saved correction cannot clear a new payment-instruction warning.",
        ],
        Inches(0.55),
        Inches(1.7),
        Inches(12.2),
        Inches(4.8),
        size=20,
    )
    _footer(slide, 10, pages)

    # 11 fraud
    slide = prs.slides.add_slide(blank)
    _chrome(slide, "The one rule the model cannot bend", "A changed bank account is not informational.")
    _box(slide, Inches(0.55), Inches(1.75), Inches(12.2), Inches(2.3), DANGER)
    _text(slide, "Do not process. Verify by phone.", Inches(0.85), Inches(2.0), Inches(11.5), Inches(0.55), size=28, bold=True, color=WHITE, font="Georgia")
    _text(
        slide,
        "Call a number you already have. Do not pay, and do not update the vendor from the email.",
        Inches(0.85),
        Inches(2.7),
        Inches(11.5),
        Inches(0.9),
        size=18,
        color=WHITE,
    )
    _bullets(
        slide,
        [
            "If Bionic files that packet as a newsletter or as informational, CloseDesk puts it back in Important.",
            "Any action that says to pay the new account is dropped. The phone-verify task stays.",
            "A real wire that says it is not a change of account stays a wire request, not a fraud flag.",
        ],
        Inches(0.55),
        Inches(4.3),
        Inches(12.2),
        Inches(2.3),
        size=18,
    )
    _footer(slide, 11, pages)

    # 12 setup
    slide = prs.slides.add_slide(blank)
    _chrome(slide, "Easy setup", "Five steps. The written page is docs/SETUP.md.")
    steps = [
        "1    Install.  python -m venv .venv, then pip install -e \".[dev]\"",
        "2    Prove it.  python -m controller_inbox demo --serve",
        "3    Start LM Studio’s local server, and set CONTROLLER_INBOX_LLM=true",
        "4    Drop .msg files into inbox/incoming",
        "5    python -m controller_inbox overnight, then serve in the morning",
    ]
    _bullets(slide, steps, Inches(0.55), Inches(1.7), Inches(12.2), Inches(4.2), size=22)
    _text(
        slide,
        "Outlook login is optional. The drop folder is enough for the first night.",
        Inches(0.55),
        Inches(6.15),
        Inches(12),
        Inches(0.45),
        size=18,
        color=MUTED,
    )
    _footer(slide, 12, pages)

    # 13 ready
    slide = prs.slides.add_slide(blank)
    _chrome(slide, "Ready to run", "Three checks, then leave it on overnight.")
    checks = [
        ("1", "The sample opens", "demo --serve shows Morning, Important, Informational, and Reference."),
        ("2", "Your file lands", "A .msg in inbox/incoming shows up after ingest or overnight."),
        ("3", "Bionic reads", "With the LM Studio server on, rows can say Read by Bionic."),
    ]
    for index, (num, title, body) in enumerate(checks):
        y = Inches(1.7 + index * 1.45)
        _box(slide, Inches(0.55), y, Inches(12.2), Inches(1.3), WHITE)
        _text(slide, num, Inches(0.75), y + Inches(0.28), Inches(0.7), Inches(0.7), size=28, bold=True, color=GOLD, font="Georgia")
        _text(slide, title, Inches(1.6), y + Inches(0.18), Inches(10.5), Inches(0.45), size=22, bold=True, color=NAVY, font="Georgia")
        _text(slide, body, Inches(1.6), y + Inches(0.68), Inches(10.5), Inches(0.4), size=16, color=INK)
    _footer(slide, 13, pages)

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)


if __name__ == "__main__":
    out = Path(__file__).resolve().parents[1] / "docs" / "CloseDesk-how-it-works.pptx"
    build(out)
    print(out)
