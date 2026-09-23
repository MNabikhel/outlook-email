"""Build docs/CloseDesk-quick-setup.pptx from docs/images/. Run from scripts/: ../.venv/bin/python build_setup_deck.py"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from build_closedesk_deck import DANGER, GOLD, H, INK, MUTED, NAVY, NAVY_2, PAPER, W, WHITE, _box, _text

ROOT = Path(__file__).resolve().parents[1]
IMAGES = ROOT / "docs" / "images"
PAGES = 12
LIGHT = WHITE
SOFT = PAPER
RGB_SOFT = RGBColor(0xD5, 0xDC, 0xCF)


def _rich(slide, lines, x, y, w, h, *, size=17, color=INK, gap=9):
    """Bullets where **text** is bold."""
    shape = slide.shapes.add_textbox(x, y, w, h)
    frame = shape.text_frame
    frame.word_wrap = True
    for index, line in enumerate(lines):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(gap)
        for n, part in enumerate(line.split("**")):
            if not part:
                continue
            run = p.add_run()
            run.text = part
            run.font.size = Pt(size)
            run.font.bold = n % 2 == 1
            run.font.color.rgb = color
            run.font.name = "Calibri"
    return shape


def _page(prs, kicker: str, title: str, page: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _box(slide, 0, 0, W, H, SOFT)
    _box(slide, 0, 0, Inches(0.16), H, GOLD)
    _text(slide, kicker.upper(), Inches(0.55), Inches(0.32), Inches(12), Inches(0.32), size=12, bold=True, color=GOLD)
    _text(slide, title, Inches(0.55), Inches(0.58), Inches(12.2), Inches(0.7), size=30, bold=True, color=NAVY, font="Georgia")
    _text(slide, "CloseDesk  ·  quick setup", Inches(0.55), Inches(7.05), Inches(9), Inches(0.3), size=11, color=MUTED)
    _text(slide, f"{page}  /  {PAGES}", Inches(11.4), Inches(7.05), Inches(1.4), Inches(0.3), size=11, color=MUTED, align=PP_ALIGN.RIGHT)
    return slide


def _shot(slide, name: str, caption: str = "") -> None:
    """A screenshot on the right two-thirds of the slide, with a thin frame."""
    x, y, w = Inches(4.75), Inches(1.45), Inches(8.2)
    h = int(w / 1.6)
    _box(slide, x - Inches(0.05), y - Inches(0.05), w + Inches(0.1), h + Inches(0.1), NAVY)
    slide.shapes.add_picture(str(IMAGES / name), x, y, width=w, height=h)
    if caption:
        _text(slide, caption, x, y + h + Inches(0.1), w, Inches(0.35), size=12, color=MUTED)


def _step(slide, number: str) -> None:
    _box(slide, Inches(11.95), Inches(0.35), Inches(0.9), Inches(0.9), NAVY)
    _text(slide, number, Inches(11.95), Inches(0.43), Inches(0.9), Inches(0.7), size=30, bold=True, color=GOLD, align=PP_ALIGN.CENTER, font="Georgia")


def build(path: Path) -> None:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "CloseDesk — quick setup"
    prs.core_properties.subject = "Set up CloseDesk on a laptop, with or without a local AI model"

    # 1 title
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _box(slide, 0, 0, W, H, NAVY)
    _box(slide, 0, 0, Inches(0.16), H, GOLD)
    _text(slide, "CLOSEDESK  ·  QUICK SETUP", Inches(0.7), Inches(1.6), Inches(11), Inches(0.4), size=14, bold=True, color=GOLD)
    _text(slide, "Your inbox, sorted into one short page each morning.", Inches(0.7), Inches(2.1), Inches(11.5), Inches(1.6), size=40, bold=True, color=WHITE, font="Georgia")
    _text(
        slide,
        "About 10 minutes to set up. Runs on your laptop.\nNo Outlook add-in, no IT ticket, no cloud. A local AI model is optional.",
        Inches(0.7),
        Inches(4.0),
        Inches(11),
        Inches(1.2),
        size=20,
        color=RGB_SOFT,
    )
    _text(slide, "Install Python   ·   Download   ·   Double-click   ·   Drag mail in   ·   (Optional) LM Studio", Inches(0.7), Inches(6.3), Inches(12), Inches(0.4), size=16, color=GOLD)

    # 2 what you need
    slide = _page(prs, "Before you start", "What you need, and what you don't.", 2)
    _box(slide, Inches(0.55), Inches(1.55), Inches(6.0), Inches(5.1), LIGHT)
    _text(slide, "You need", Inches(0.8), Inches(1.75), Inches(5.5), Inches(0.45), size=22, bold=True, color=NAVY, font="Georgia")
    _rich(
        slide,
        [
            "A **Windows or Mac** laptop.",
            "**Python 3.11 or newer** (free, from python.org).",
            "The **CloseDesk folder** from GitHub.",
            "Outlook: classic, new, web, or Mac. Any of them works.",
            "**Optional:** LM Studio and a small model, for written summaries and chat answers.",
        ],
        Inches(0.8),
        Inches(2.4),
        Inches(5.5),
        Inches(4.0),
    )
    _box(slide, Inches(6.8), Inches(1.55), Inches(6.0), Inches(5.1), NAVY)
    _text(slide, "You don't need", Inches(7.05), Inches(1.75), Inches(5.5), Inches(0.45), size=22, bold=True, color=GOLD, font="Georgia")
    _rich(
        slide,
        [
            "An Outlook add-in or admin rights to Outlook.",
            "IT approval, Azure, or a Microsoft 365 app registration.",
            "A powerful laptop. **No model? Everything still works:** sorting, summaries, tasks, fraud warnings, search, chat lookups, and the daily digest.",
            "To send your mail anywhere. It stays on this computer.",
        ],
        Inches(7.05),
        Inches(2.4),
        Inches(5.5),
        Inches(4.0),
        color=WHITE,
    )

    # 3 python
    slide = _page(prs, "Step 1", "Install Python.", 3)
    _step(slide, "1")
    _rich(
        slide,
        [
            "Go to **python.org/downloads** and click the big yellow **Download Python** button.",
            "Run the installer.",
            "**Windows: tick “Add python.exe to PATH”** at the bottom of the first screen, then click **Install Now**. This one box is the most common setup problem.",
            "Mac: open the downloaded .pkg and click through.",
            "You only do this once.",
        ],
        Inches(0.55),
        Inches(1.6),
        Inches(7.6),
        Inches(4.6),
        size=20,
    )
    _box(slide, Inches(8.6), Inches(1.6), Inches(4.2), Inches(3.2), NAVY)
    _text(slide, "Windows installer", Inches(8.85), Inches(1.8), Inches(3.8), Inches(0.4), size=14, bold=True, color=GOLD)
    _text(slide, "☐  Use admin privileges when installing py.exe", Inches(8.85), Inches(2.35), Inches(3.8), Inches(0.7), size=16, color=RGB_SOFT)
    _box(slide, Inches(8.75), Inches(3.1), Inches(3.9), Inches(0.55), GOLD)
    _text(slide, "☑  Add python.exe to PATH", Inches(8.85), Inches(3.17), Inches(3.8), Inches(0.45), size=16, bold=True, color=NAVY)
    _text(slide, "Tick this one.", Inches(8.85), Inches(4.05), Inches(3.8), Inches(0.5), size=17, bold=True, color=GOLD)

    # 4 download
    slide = _page(prs, "Step 2", "Get CloseDesk from GitHub.", 4)
    _step(slide, "2")
    _rich(
        slide,
        [
            "Open the CloseDesk page on GitHub.",
            "Click the green **Code** button, then **Download ZIP**.",
            "**Unzip it** somewhere permanent, for example **Documents**. (Windows: right-click the ZIP → **Extract All**.)",
            "The folder is called **outlook-email-main**. Rename it **CloseDesk** if you like.",
            "Don't run it from inside the ZIP; unzip first.",
            "Your mail and settings live in this folder, in **inbox** and **data**. Keep the folder when you update.",
        ],
        Inches(0.55),
        Inches(1.6),
        Inches(7.6),
        Inches(4.8),
        size=19,
    )
    _box(slide, Inches(8.6), Inches(1.6), Inches(4.2), Inches(4.3), NAVY)
    _text(slide, "Inside the folder", Inches(8.85), Inches(1.8), Inches(3.8), Inches(0.4), size=14, bold=True, color=GOLD)
    tree = [
        ("CloseDesk.bat", "double-click (Windows)"),
        ("CloseDesk.command", "double-click (Mac)"),
        ("inbox\\incoming", "drop mail here"),
        ("data", "appears after the first run"),
        ("README.md", "the short guide"),
    ]
    for index, (name, note) in enumerate(tree):
        y = Inches(2.35 + index * 0.68)
        _text(slide, name, Inches(8.85), y, Inches(3.8), Inches(0.35), size=16, bold=True, color=WHITE)
        _text(slide, note, Inches(8.85), y + Inches(0.3), Inches(3.8), Inches(0.3), size=12, color=RGB_SOFT)

    # 5 first launch
    slide = _page(prs, "Step 3", "Double-click CloseDesk.", 5)
    _step(slide, "3")
    _rich(
        slide,
        [
            "Windows: **CloseDesk.bat**.\nMac: **CloseDesk.command**.",
            "The first run sets itself up (a minute or two), then opens the dashboard in your browser.",
            "**Leave the black window open** while you use it.",
            "Windows “protected your PC”? Click **More info → Run anyway**.",
            "Mac “can't be opened”? **Right-click → Open → Open** once.",
            "Click **Load sample mailbox** to look around first.",
        ],
        Inches(0.55),
        Inches(1.5),
        Inches(4.0),
        Inches(5.4),
        size=15,
        gap=7,
    )
    _shot(slide, "get_started.jpg", "The first page. The sample is removed as soon as your own mail comes in.")

    # 6 add mail
    slide = _page(prs, "Step 4", "Drag your mail in, then Process new mail.", 6)
    _step(slide, "4")
    _rich(
        slide,
        [
            "**Classic Outlook (Windows):** select yesterday's mail (click the first, Shift+click the last) and drag it onto the **inbox\\incoming** folder.",
            "**New Outlook / web:** open a message → **… → Save as** → save the .eml into **inbox/incoming**.",
            "**Mac:** drag messages into **inbox/incoming** in Finder.",
            "Click **Process new mail**.",
            "Tip: pin **inbox\\incoming** to Quick Access.",
        ],
        Inches(0.55),
        Inches(1.5),
        Inches(4.0),
        Inches(5.4),
        size=15,
        gap=7,
    )
    _shot(slide, "today.jpg", "Today: payment-change warnings first, then your ranked focus list.")

    # 7 local AI
    slide = _page(prs, "Step 5 (optional)", "Add a local AI model, if the laptop can run one.", 7)
    _step(slide, "5")
    _rich(
        slide,
        [
            "Install **LM Studio** (lmstudio.ai).",
            "Download a small **instruct** model: 3B–8B is plenty.",
            "Load it, open the **Developer** tab, click **Start server**.",
            "Click **Process new mail**. CloseDesk finds the model on its own: **Local AI: Off** disappears and the bottom-left shows the model's name.",
        ],
        Inches(0.55),
        Inches(1.6),
        Inches(5.2),
        Inches(4.6),
        size=18,
    )
    rows = [
        ("", "No model", "With a model"),
        ("Sorting into Needs you / Worth knowing / Reference", "Yes (rules)", "Yes, with guard rails"),
        ("One-line summaries", "The email's key sentence", "Written by the model"),
        ("Tasks, due dates, fraud warnings, digest", "Yes", "Yes"),
        ("Ask CloseDesk", "Finds the emails", "Writes an answer with links"),
        ("Draft a reply", "Starter template", "Written draft"),
    ]
    top = Inches(1.6)
    for index, (a, b, c) in enumerate(rows):
        y = top + Inches(0.78) * index
        fill = NAVY if index == 0 else (LIGHT if index % 2 else SOFT)
        color = WHITE if index == 0 else INK
        _box(slide, Inches(6.0), y, Inches(6.8), Inches(0.74), fill)
        _text(slide, a, Inches(6.15), y + Inches(0.14), Inches(2.9), Inches(0.6), size=13, bold=index == 0, color=color)
        _text(slide, b, Inches(9.1), y + Inches(0.14), Inches(1.75), Inches(0.6), size=13, bold=index == 0, color=color)
        _text(slide, c, Inches(10.9), y + Inches(0.14), Inches(1.85), Inches(0.6), size=13, bold=index == 0, color=color)

    # 8 open in place
    slide = _page(prs, "Getting around", "Click an email and it opens right there.", 8)
    _rich(
        slide,
        [
            "Click any email, task, search result, or chat link: it opens in a panel on the right. **Esc** closes it.",
            "**Open in Outlook** opens the original message.",
            "**Draft a reply** writes a reply to copy or open in your mail app. Payment-change emails get “verify by phone” advice instead.",
            "Press **/** to search everything, including attachments.",
        ],
        Inches(0.55),
        Inches(1.5),
        Inches(4.0),
        Inches(5.4),
        size=15,
        gap=8,
    )
    _shot(slide, "preview.jpg")

    # 9 chat
    slide = _page(prs, "Getting around", "Ask CloseDesk, like a support chat.", 9)
    _rich(
        slide,
        [
            "The button at the **bottom-right** of every page.",
            "Try: **what's urgent today?** · **what needs a reply?** · **anything from Maya?** · **invoice 10482**",
            "Click a **[number]** to open that email.",
            "No model: it looks things up and lists the matching emails. With a model: it writes the answer.",
            "Nothing leaves the laptop.",
        ],
        Inches(0.55),
        Inches(1.5),
        Inches(4.0),
        Inches(5.4),
        size=15,
        gap=8,
    )
    _shot(slide, "chat.jpg", "With no model running, answers are a lookup with links to each email.")

    # 10 profile
    slide = _page(prs, "Make it yours", "Pick the kind of inbox on Setup.", 10)
    _rich(
        slide,
        [
            "**General** (default) suits anyone: replies, approvals, meetings, deadlines, notifications.",
            "**Finance & accounting** adds the month-end countdown and close sections.",
            "Invoices and payment-change warnings are caught either way.",
            "The Setup page also shows whether a local model is running.",
        ],
        Inches(0.55),
        Inches(1.5),
        Inches(4.0),
        Inches(5.4),
        size=16,
        gap=9,
    )
    _shot(slide, "setup.jpg")

    # 11 every morning
    slide = _page(prs, "Every morning", "Three steps, about five minutes.", 11)
    steps = [
        ("1", "Drag", "Yesterday's mail from Outlook into inbox/incoming."),
        ("2", "Double-click", "CloseDesk. It reads the mail and opens today's page."),
        ("3", "Work the list", "Top to bottom. Click Done as you go; past digests are kept."),
    ]
    for index, (num, title, body) in enumerate(steps):
        x = Inches(0.55 + index * 4.15)
        _box(slide, x, Inches(1.6), Inches(3.95), Inches(2.9), NAVY if index % 2 == 0 else NAVY_2)
        _text(slide, num, x + Inches(0.25), Inches(1.8), Inches(3.4), Inches(0.6), size=28, bold=True, color=GOLD, font="Georgia")
        _text(slide, title, x + Inches(0.25), Inches(2.5), Inches(3.4), Inches(0.5), size=22, bold=True, color=WHITE, font="Georgia")
        _text(slide, body, x + Inches(0.25), Inches(3.15), Inches(3.4), Inches(1.2), size=16, color=RGB_SOFT)
    _box(slide, Inches(0.55), Inches(4.8), Inches(12.25), Inches(1.9), DANGER)
    _text(slide, "A red “Do not process — verify by phone” box", Inches(0.85), Inches(5.0), Inches(11.7), Inches(0.5), size=22, bold=True, color=WHITE, font="Georgia")
    _text(
        slide,
        "means someone asked to change bank or payment details. Call a number you already have before doing anything. Don't pay or reply from that email.",
        Inches(0.85),
        Inches(5.6),
        Inches(11.7),
        Inches(0.9),
        size=17,
        color=WHITE,
    )

    # 12 ready + troubleshooting
    slide = _page(prs, "Check it works", "You're ready when…", 12)
    checks = [
        "Double-clicking CloseDesk opens the dashboard.",
        "An email dragged into inbox/incoming shows up after Process new mail.",
        "The header says Local AI: Off (fine), or the bottom-left shows your model.",
    ]
    for index, line in enumerate(checks):
        y = Inches(1.55 + index * 0.9)
        _box(slide, Inches(0.55), y, Inches(5.6), Inches(0.8), LIGHT)
        _text(slide, "✓", Inches(0.7), y + Inches(0.16), Inches(0.4), Inches(0.5), size=20, bold=True, color=GOLD)
        _text(slide, line, Inches(1.15), y + Inches(0.14), Inches(4.85), Inches(0.6), size=14, color=INK)
    _text(slide, "Full guide: docs/SETUP.md in the CloseDesk folder.", Inches(0.55), Inches(4.35), Inches(5.6), Inches(0.4), size=14, color=MUTED)
    fixes = [
        ("“Python 3.11 or newer is needed”", "Reinstall Python and tick Add python.exe to PATH."),
        ("The page doesn't open", "Keep the black window open; go to http://127.0.0.1:8765."),
        ("Local AI: Off", "Fine. Start LM Studio's server for written answers."),
        ("A file went to inbox/failed", "Read the .why.txt beside it; re-save from Outlook."),
        ("A category is wrong", "Open it → Wrong category? → one sentence why."),
    ]
    _box(slide, Inches(6.45), Inches(1.55), Inches(6.35), Inches(0.5), NAVY)
    _text(slide, "If something is off", Inches(6.65), Inches(1.6), Inches(6), Inches(0.4), size=15, bold=True, color=GOLD)
    for index, (seen, fix) in enumerate(fixes):
        y = Inches(2.1 + index * 0.9)
        _box(slide, Inches(6.45), y, Inches(6.35), Inches(0.84), LIGHT if index % 2 == 0 else SOFT)
        _text(slide, seen, Inches(6.65), y + Inches(0.06), Inches(6.0), Inches(0.4), size=14, bold=True, color=NAVY)
        _text(slide, fix, Inches(6.65), y + Inches(0.42), Inches(6.0), Inches(0.4), size=13, color=INK)

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)


if __name__ == "__main__":
    out = ROOT / "docs" / "CloseDesk-quick-setup.pptx"
    build(out)
    print(out)
