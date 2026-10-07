"""AP invoice cost coding against the user's own list of JDE cost codes.

The codes live in a workbook the user keeps (two columns: Description and Cost Code) in the
``AP cost codes`` folder. CloseDesk never ships codes of its own. For an AP invoice it looks
for those codes where AP writes them: stamped or typed on the invoice, or in the email.

E1 JDE codes are a business unit, an object and a subsidiary joined by periods
(``1100.6110.100``). A scan may read a period as a comma, an O for a zero, or drop the
business unit's leading zeros, so each code is matched with those slips allowed.

When no code is printed, the code you last confirmed for the same sender comes next, then
descriptions whose words are on the invoice. Every suggestion waits for Confirm or Revise.
"""

from __future__ import annotations

import hashlib
import re
import threading
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

from controller_inbox.models import DocumentType, EmailRecord

WORKBOOK = "AP cost codes.xlsx"
HEADERS = ("Description", "Cost Code")
SIGNATURE_KEY = "cost_codes_signature"

ON_INVOICE = "on the invoice"
SENDER_HISTORY = "you used it for this sender"
DESCRIPTION = "description matches"
YOU = "chosen by you"

# A digit a scan commonly reads as a letter, and back.
_LOOKALIKE = {"0": "[0O]", "1": "[1Il|]", "5": "[5S]", "8": "[8B]"}
_SEPARATOR = r"\s?[.,·]\s?"
_BEFORE = r"(?<![\w.,$€£])"
_AFTER = r"(?![\w]|[.,]\w)"
# Any code-shaped number when the workbook gives no shapes to go by: a business unit, an
# object of four to six digits, and an optional subsidiary.
_GENERIC_SHAPE = re.compile(_BEFORE + r"\d{2,12}\.\d{4,6}(?:\.[0-9A-Z]{1,8})?" + _AFTER)
_STOP = frozenset(
    "the and for with from per inc ltd llc corp company co all any other misc general various".split()
)


@dataclass(frozen=True)
class CostCode:
    code: str
    description: str
    # Typed into a number cell, so Excel may have dropped trailing zeros (1100.6110 -> 1100.611).
    loose: bool = False

    @property
    def key(self) -> str:
        return code_key(self.code)


@dataclass
class Codebook:
    codes: list[CostCode] = field(default_factory=list)
    files: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def signature(self) -> str:
        text = "\n".join(f"{code.code}\t{code.description}\t{code.loose}" for code in self.codes)
        return hashlib.sha256(text.encode()).hexdigest()[:16]

    def find(self, code: str) -> CostCode | None:
        key = code_key(code)
        return next((item for item in self.codes if item.key == key), None)


@dataclass
class Suggestion:
    codes: list[dict]
    others: list[dict]
    unlisted: list[str]

    @property
    def status(self) -> str:
        return "suggested" if self.codes else "unmatched"


def code_key(code: str) -> str:
    """One spelling per code: periods between segments, capitals, no leading zeros on the business unit."""
    segments = _segments(code)
    if not segments:
        return ""
    first = segments[0].lstrip("0") or "0"
    return ".".join([first, *segments[1:]]).upper()


def _segments(code: str) -> list[str]:
    return [part for part in re.split(r"[\s.,·]+", str(code).strip()) if part]


def is_ap_invoice(email: EmailRecord) -> bool:
    return email.category == DocumentType.AP_INVOICE or any(
        att.document_type == DocumentType.AP_INVOICE for att in email.attachments
    )


# The workbook -------------------------------------------------------------------------------


def workbook_path(settings) -> Path:
    return settings.cost_codes_folder / WORKBOOK


def _workbooks(folder: Path) -> list[Path]:
    if not folder.is_dir():
        return []
    return sorted(
        path
        for path in folder.iterdir()
        if path.suffix.lower() in {".xlsx", ".xlsm"} and not path.name.startswith(("~$", "."))
    )


def ensure_workbook(settings) -> Path:
    """The folder, with a blank workbook to fill in when there is none yet."""
    folder = settings.cost_codes_folder
    folder.mkdir(parents=True, exist_ok=True)
    path = workbook_path(settings)
    if not _workbooks(folder):
        _write_template(path)
    return path


def _write_template(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    book = Workbook()
    sheet = book.active
    sheet.title = "Cost codes"
    sheet.append(list(HEADERS))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(vertical="center")
    sheet.column_dimensions["A"].width = 52
    sheet.column_dimensions["B"].width = 24
    sheet.freeze_panes = "A2"
    # Text cells keep a code as typed. In a number cell Excel turns 1100.6110 into 1100.611.
    sheet.column_dimensions["B"].number_format = "@"
    for row in range(2, 1001):
        sheet.cell(row=row, column=2).number_format = "@"
    notes = book.create_sheet("How to use")
    notes.column_dimensions["A"].width = 110
    for line in (
        "One cost code per row on the Cost codes sheet: what it is for in Description, the JDE code in Cost Code.",
        "Write the code the way it is stamped on invoices, with its periods, for example 1100.6110.100.",
        "A good description names the spend and, where it helps, the vendor or site: Office supplies - head office.",
        "Save the file. CloseDesk reads it again by itself and checks every AP invoice against it.",
        "More rows, more sheets, or more workbooks in this folder are all read. A sheet needs the two headings.",
    ):
        notes.append([line])
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


_cache_lock = threading.Lock()
_cache: dict[tuple, Codebook] = {}


def load(settings) -> Codebook:
    """Every code in the folder's workbooks. Read again only when a workbook changes."""
    files = _workbooks(settings.cost_codes_folder)
    try:
        stamp = tuple((str(path), path.stat().st_mtime_ns, path.stat().st_size) for path in files)
    except OSError:
        stamp = ()
    with _cache_lock:
        if stamp in _cache:
            return _cache[stamp]
    book = Codebook(files=[path.name for path in files])
    seen: set[str] = set()
    for path in files:
        try:
            _read_workbook(path, book, seen)
        except Exception as exc:  # a workbook open in Excel or damaged: say so, keep the rest
            book.warnings.append(f"{path.name} couldn't be read ({exc.__class__.__name__}). Close it in Excel and check again.")
    with _cache_lock:
        _cache.clear()
        _cache[stamp] = book
    return book


def _read_workbook(path: Path, book: Codebook, seen: set[str]) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        numbers = 0
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            columns = None
            for _ in range(10):
                row = next(rows, None)
                if row is None:
                    break
                columns = _header_columns(row)
                if columns:
                    break
            if not columns:
                continue
            describe, coded = columns
            for row in rows:
                value = row[coded] if coded < len(row) else None
                code, loose = _code_text(value)
                if not code or not _segments(code):
                    continue
                key = code_key(code)
                if key in seen:
                    continue
                seen.add(key)
                numbers += loose
                description = str(row[describe]).strip() if describe < len(row) and row[describe] is not None else ""
                book.codes.append(CostCode(code, description, loose))
        if numbers:
            book.warnings.append(
                f"{numbers} code{'s' if numbers != 1 else ''} in {path.name} {'are' if numbers != 1 else 'is'} stored as "
                "numbers, so Excel may have dropped trailing zeros. Format the Cost Code column as Text and retype them."
            )
    finally:
        workbook.close()


def _header_columns(row: tuple) -> tuple[int, int] | None:
    labels = [str(value).strip().lower() if value is not None else "" for value in row]
    describe = next((i for i, label in enumerate(labels) if "description" in label or label in {"desc", "name"}), None)
    coded = next((i for i, label in enumerate(labels) if "code" in label or label in {"account", "gl account"}), None)
    if describe is None or coded is None or describe == coded:
        return None
    return describe, coded


def _code_text(value) -> tuple[str, bool]:
    if value is None or isinstance(value, bool):
        return "", False
    if isinstance(value, float):
        return (f"{value:.6f}".rstrip("0").rstrip(".") if not value.is_integer() else str(int(value))), True
    if isinstance(value, int):
        return str(value), True
    return str(value).strip(), False


# Matching -----------------------------------------------------------------------------------


def _pattern(code: CostCode) -> re.Pattern | None:
    segments = _segments(code.code)
    if len(segments) == 1 and len(segments[0]) < 6:
        # A bare short number is everywhere on an invoice; only its description can match.
        return None
    parts = []
    for index, segment in enumerate(segments):
        body = segment.lstrip("0") if index == 0 else segment
        piece = "".join(_LOOKALIKE.get(ch, re.escape(ch)) for ch in body.upper())
        if index == 0:
            piece = "[0O]*" + piece if body else "[0O]+"
        if index == len(segments) - 1 and code.loose:
            piece += "0*"
        parts.append(piece)
    return re.compile(_BEFORE + _SEPARATOR.join(parts) + _AFTER, re.IGNORECASE)


def _sources(email: EmailRecord) -> list[tuple[str, str]]:
    """Where a code could be written: the email, then each attachment."""
    out = [("the email", f"{email.subject}\n{email.body_text or ''}")]
    out += [(att.filename, att.extracted_text or "") for att in email.attachments if att.extracted_text]
    return out


def _snippet(text: str, start: int, end: int) -> str:
    left = text.rfind("\n", 0, start)
    right = text.find("\n", end)
    line = text[left + 1 : right if right >= 0 else len(text)].strip()
    if len(line) <= 140:
        return line
    begin = max(0, start - left - 1 - 60)
    return ("…" if begin else "") + line[begin : begin + 140].strip() + "…"


def codes_on(email: EmailRecord, book: Codebook) -> list[dict]:
    """Workbook codes written on the invoice or in the email, in the order they appear."""
    found: list[tuple[int, int, dict]] = []
    for code in book.codes:
        pattern = _pattern(code)
        if pattern is None:
            continue
        for order, (where, text) in enumerate(_sources(email)):
            match = pattern.search(text)
            if match:
                found.append(
                    (order, match.start(), {
                        "code": code.code,
                        "description": code.description,
                        "source": ON_INVOICE,
                        "where": where,
                        "evidence": _snippet(text, match.start(), match.end()),
                    })
                )
                break
    found.sort(key=lambda item: (item[0], item[1]))
    return [item for _order, _at, item in found]


def unlisted_codes(email: EmailRecord, book: Codebook, *, limit: int = 6) -> list[str]:
    """Code-shaped numbers on the invoice that the workbook doesn't have, so a typo or a new code shows up."""
    shapes = _shapes(book) or [_GENERIC_SHAPE]
    known = {code.key for code in book.codes}
    # A code kept in a number cell (1100.611) is the 1100.6110 on the invoice, which codes_on already found.
    loose = [pattern for pattern in (_pattern(code) for code in book.codes if code.loose) if pattern is not None]
    out: list[str] = []
    for _where, text in _sources(email):
        for shape in shapes:
            for match in shape.finditer(text):
                code = re.sub(r"\s", "", match.group(0)).replace(",", ".")
                listed = code_key(code) in known or any(pattern.fullmatch(code) for pattern in loose)
                if not listed and code not in out:
                    out.append(code)
                if len(out) >= limit:
                    return out
    return out


def _shapes(book: Codebook) -> list[re.Pattern]:
    """The shapes the user's own codes have (segment count and widths), as patterns."""
    seen: dict[tuple, re.Pattern] = {}
    for code in book.codes:
        segments = _segments(code.code)
        if len(segments) < 2:
            continue
        widths = tuple(len(segment) for segment in segments)
        if widths in seen:
            continue
        parts = [rf"\d{{1,{widths[0]}}}"] + [rf"[0-9A-Z]{{{width}}}" for width in widths[1:]]
        seen[widths] = re.compile(_BEFORE + r"\.".join(parts) + _AFTER)
    return list(seen.values())


def _words(text: str) -> set[str]:
    return {_stem(word) for word in re.findall(r"[a-z0-9]{3,}", text.lower()) if word not in _STOP}


def _stem(word: str) -> str:
    if len(word) > 4 and word.endswith("ies"):
        return word[:-3] + "y"
    if len(word) > 3 and word.endswith("s") and not word.endswith("ss"):
        return word[:-1]
    return word


def by_description(email: EmailRecord, book: Codebook, *, limit: int = 3) -> list[dict]:
    """Codes whose description is mostly words on the invoice, best first."""
    on_invoice = set()
    for _where, text in _sources(email):
        on_invoice |= _words(text[:60_000])
    ranked = []
    for code in book.codes:
        wanted = _words(code.description)
        if not wanted:
            continue
        hits = wanted & on_invoice
        share = len(hits) / len(wanted)
        enough = share >= 0.6 if len(wanted) > 1 else bool(hits) and len(next(iter(hits))) >= 4
        if enough:
            ranked.append((share, len(hits), sum(map(len, hits)), code, sorted(hits)))
    ranked.sort(key=lambda item: (item[0], item[1], item[2]), reverse=True)
    return [
        {
            "code": code.code,
            "description": code.description,
            "source": DESCRIPTION,
            "where": "",
            "evidence": "Matched: " + ", ".join(hits),
        }
        for _share, _n, _len, code, hits in ranked[:limit]
    ]


def suggest(store, email: EmailRecord, book: Codebook) -> Suggestion:
    unlisted = unlisted_codes(email, book)
    printed = codes_on(email, book)
    described = by_description(email, book)
    if printed:
        chosen = printed
    else:
        chosen = []
        for item in store.last_confirmed_codes(email.sender_email, exclude=email.id):
            listed = book.find(item["code"])
            chosen.append({
                "code": listed.code if listed else item["code"],
                "description": listed.description if listed else item.get("description", ""),
                "source": SENDER_HISTORY,
                "where": "",
                "evidence": "",
            })
        if not chosen and described:
            chosen = described[:1]
    picked = {code_key(item["code"]) for item in chosen}
    others = [item for item in described if code_key(item["code"]) not in picked]
    return Suggestion(chosen, others, unlisted)


def choices(book: Codebook, coding: dict | None) -> list[dict]:
    """The workbook for the Revise list: the current codes ticked first, then likely ones, then the rest."""
    current = {code_key(item["code"]) for item in (coding or {}).get("codes", [])}
    likely = {code_key(item["code"]) for item in (coding or {}).get("others", [])}

    def rank(code: CostCode) -> int:
        return 0 if code.key in current else 1 if code.key in likely else 2

    return [
        {"code": code.code, "description": code.description, "checked": code.key in current, "likely": code.key in likely}
        for code in sorted(book.codes, key=rank)
    ]


# Keeping the store up to date ---------------------------------------------------------------


def _stamp(book: Codebook, email: EmailRecord) -> str:
    text = "\n".join(text for _where, text in _sources(email))
    return hashlib.sha256(f"{book.signature}\n{email.sender_email}\n{text}".encode()).hexdigest()[:16]


def refresh(store, settings, *, email_ids: list[str] | None = None, force: bool = False) -> int:
    """Suggest codes for AP invoices that are new, changed, or checked against an older workbook.

    A coding you confirmed is never replaced. Mail that is no longer an AP invoice loses an
    unconfirmed suggestion. Returns how many invoices were checked.
    """
    book = load(settings)
    ids = store.ap_invoice_ids() if email_ids is None else email_ids
    checked = 0
    for email_id in ids:
        email = store.get_email(email_id)
        current = store.cost_coding(email_id)
        if email is None or not is_ap_invoice(email):
            if current and current["status"] != "confirmed":
                store.delete_cost_coding(email_id)
            continue
        if current and current["status"] == "confirmed":
            continue
        stamp = _stamp(book, email)
        if current and current["signature"] == stamp and not force:
            continue
        suggestion = suggest(store, email, book)
        store.save_cost_coding(
            email_id,
            status=suggestion.status,
            codes=suggestion.codes,
            others=suggestion.others,
            unlisted=suggestion.unlisted,
            signature=stamp,
        )
        checked += 1
    if email_ids is None:
        store.set_state(SIGNATURE_KEY, book.signature)
    return checked


def refresh_if_changed(store, settings) -> None:
    """Check every AP invoice again when the workbook has changed since the last check."""
    if store.get_state(SIGNATURE_KEY) != load(settings).signature:
        refresh(store, settings)


def confirm(store, settings, email_id: str) -> None:
    current = store.cost_coding(email_id)
    if not current or not current["codes"]:
        raise ValueError("There is no code to confirm yet. Use Revise to choose one.")
    store.save_cost_coding(
        email_id,
        status="confirmed",
        codes=current["codes"],
        others=current["others"],
        unlisted=current["unlisted"],
        signature=current["signature"],
        decided_at=_now(),
    )
    _recheck_sender(store, settings, email_id)


def revise(store, settings, email_id: str, codes: list[str]) -> None:
    """Replace the suggestion with codes picked from the workbook, and confirm them."""
    book = load(settings)
    chosen = []
    for code in dict.fromkeys(code.strip() for code in codes if code and code.strip()):
        listed = book.find(code)
        if listed is None:
            raise ValueError(f"{code} is not in the cost code workbook.")
        chosen.append({"code": listed.code, "description": listed.description, "source": YOU, "where": "", "evidence": ""})
    if not chosen:
        raise ValueError("Pick at least one code.")
    current = store.cost_coding(email_id) or {}
    store.save_cost_coding(
        email_id,
        status="confirmed",
        codes=chosen,
        others=current.get("others", []),
        unlisted=current.get("unlisted", []),
        signature=current.get("signature", ""),
        decided_at=_now(),
    )
    _recheck_sender(store, settings, email_id)


def _recheck_sender(store, settings, email_id: str) -> None:
    """Other open invoices from this sender may now follow what was just confirmed."""
    email = store.get_email(email_id)
    if email and email.sender_email:
        others = [other for other in store.email_ids_from(sender=email.sender_email) if other != email_id]
        refresh(store, settings, email_ids=others, force=True)


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")
