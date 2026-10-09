"""Read attachments into text a small model can navigate, and let it go back for more.

Every document becomes plain text with section markers:

    [page 3]                      PDF pages
    [sheet "Budget" A1:F40]       workbook sheets, one line per row, cells named A5, B5 …
    [slide 2]                     PowerPoint slides
    [part 4]                      long text without natural sections

Rows keep their cell references, and a formula is shown next to its value
(``C7: 45,000 (=SUM(C2:C6))``), so the model can follow a workbook's logic.
``split_parts`` cuts that text back into sections, ``search_parts`` finds the
ones that match a question, and ``read_cells`` / ``trace_cell`` open the
original workbook for exact ranges and formula precedents.
"""

from __future__ import annotations

import csv
import io
import logging
import math
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import ROUND_HALF_UP, Decimal

from controller_inbox import ocr, pdf_layout, tables

# pypdf and pdfminer warn on the small defects many real PDFs have ("EOF marker not found") and
# still read them; with no logging set up, those warnings land in the user's terminal.
logging.getLogger("pypdf").setLevel(logging.ERROR)
logging.getLogger("pdfminer").setLevel(logging.ERROR)

# Raise when attachments read differently, so text stored by an older reader is read again.
READER_VERSION = "18"
MAX_TEXT = 400_000
MAX_PDF_PAGES = 300
MAX_SHEETS = 40
MAX_ROWS = 1_000
MAX_COLUMNS = 60
PART_SIZE = 1_800

SPREADSHEET_SUFFIXES = (".xlsx", ".xlsm", ".xltx", ".xltm")
MARKER_RE = re.compile(r"^\[(page \d+|sheet \"[^\"\n]*\"[^\]\n]*|slide \d+|part \d+)\]\s*$", re.M)


@dataclass
class Part:
    label: str
    text: str


def extract_document(filename: str, content_type: str, data: bytes) -> str:
    """Structured text for one attachment, or '' when the file type has no text to read."""
    name = (filename or "").lower()
    ctype = (content_type or "").lower()
    if name.endswith(".pdf") or "pdf" in ctype:
        return _cap(pdf_text(data))
    office = name.endswith((".docx", ".pptx", *SPREADSHEET_SUFFIXES)) or "openxmlformats" in ctype
    if office and _password_protected(data):
        return _PROTECTED_NOTE
    if name.endswith(".docx") or "wordprocessingml" in ctype:
        return _cap(docx_text(data))
    if name.endswith(SPREADSHEET_SUFFIXES) or "spreadsheetml" in ctype:
        return _cap(xlsx_text(data))
    if name.endswith(".xls") or ctype == "application/vnd.ms-excel":
        return _cap(_excel_named_text(data, filename or "table.csv"))
    if name.endswith(".pptx") or "presentationml" in ctype:
        return _cap(pptx_text(data))
    if name.endswith((".csv", ".tsv")) or ctype in {"text/csv", "application/csv", "text/tab-separated-values"}:
        return _cap(csv_text(data, filename or "table.csv"))
    return ""


_PROTECTED_NOTE = "[This file is password-protected, so CloseDesk can't read it. Open it with its password to see what it says.]"


def _password_protected(data: bytes) -> bool:
    """A Word, Excel or PowerPoint file saved with a password: it is no longer a zip but an encrypted package
    inside an OLE file."""
    if not data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return False
    try:
        import olefile

        ole = olefile.OleFileIO(data)
        try:
            return bool(ole.exists("EncryptedPackage"))
        finally:
            ole.close()
    except Exception:
        return False


def _cap(text: str) -> str:
    if len(text) <= MAX_TEXT:
        return text
    return text[:MAX_TEXT].rsplit("\n", 1)[0] + "\n[CloseDesk stopped reading here: the file is very long.]"


# PDF -----------------------------------------------------------------------------------------


def pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    failure: Exception | None = None
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        # pypdf gives up on small defects (no "%%EOF" at the end) that pdfminer reads past.
        reader, failure = None, exc
    if reader is not None and reader.is_encrypted:
        try:
            if not reader.decrypt(""):
                return "[This PDF is password-protected, so CloseDesk can't read it.]"
        except Exception:
            return "[This PDF is password-protected, so CloseDesk can't read it.]"
    listed = None
    if reader is not None:
        try:
            total = len(reader.pages)
            listed = reader.pages[:MAX_PDF_PAGES]
        except Exception as exc:
            failure = exc
    # pdfminer unpacks a page's drawing instructions whole, with no limit (pypdf stops at its own). A file whose
    # pages unpack to far more than any page of text needs (a "PDF bomb") is read with pypdf alone.
    heavy = listed is not None and not _content_fits(listed)
    miner = None if heavy else _Miner.open(data)
    if listed is None:
        # pypdf can't read the file: pdfminer's pages alone, when it can.
        if miner is None or not miner.pages:
            raise failure or ValueError("not a readable PDF")
        total = len(miner.pages)
        pairs = [(None, mined) for mined in miner.pages]
    else:
        pairs = [(page, miner.page_for(page, index) if miner else None) for index, page in enumerate(listed)]
    if not total:
        return "[This PDF has no pages.]"
    pages = []
    words = 0
    scanned = []
    tried = 0
    previous: list[pdf_layout.Table] = []
    laid_pages: dict[int, str] = {}  # the pages read from their own text, by number
    for number, (page, mined) in enumerate(pairs, start=1):
        layout = miner.layout(mined) if miner is not None and mined is not None else None
        text = ""
        if layout is not None:
            laid_out = pdf_layout.page_text(pdf_layout.glyphs_of(layout), previous, pdf_layout.rules_of(layout))
            text, previous = laid_out.text, laid_out.tables
            if text.strip():
                laid_pages[number] = text.strip()
        if not text.strip() and page is not None:
            text = _pdf_page(page)
        if not text.strip() and page is not None and tried < MAX_OCR_PAGES and _has_images(page):
            tried += 1
            text = _ocr_page_images(page)
            if text.strip() or ocr.engine_name():
                scanned.append(number)
        words += len(text.split())
        pages.append(f"[page {number}]\n{text.strip() or '(no text on this page)'}")
    from controller_inbox import camelot_tables  # here: it imports vision, which imports this module

    if laid_pages and camelot_tables.available():
        # A second reading of the tables, where the page has figures: kept where it read more (camelot_tables.py).
        wanted = [
            number for number, body in laid_pages.items()
            if len(_FIGURES.findall(body)) >= camelot_tables.MIN_FIGURES or len(_LINE_ENDS_IN_NUMBER.findall(body)) >= camelot_tables.MIN_FIGURES
        ]
        found = camelot_tables.read(data, wanted) if wanted else {}
        for number, grids in found.items():
            if number in laid_pages:
                at = number - 1
                pages[at] = f"[page {number}]\n{camelot_tables.combine(laid_pages[number], grids)}"
    if total > MAX_PDF_PAGES:
        pages.append(f"[CloseDesk read the first {MAX_PDF_PAGES} of {total} pages.]")
    if heavy:
        pages.insert(0, _HEAVY_NOTE)
    elif total and words < 5 * len(pairs):
        pages.insert(0, _SCANNED_NOTE if ocr.engine_name() else _SCANNED_NOTE + _OCR_HINT)
    elif scanned:
        pages.insert(0, f"[Scanned {'page' if len(scanned) == 1 else 'pages'} {_page_list(scanned)} read with OCR: check figures against the file.]")
    return "\n\n".join(pages)


# A figure on a page: an amount with a thousands comma or decimals, or a number of three digits or more.
_FIGURES = re.compile(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+\.\d+|\d{3,}")
# A line ending in a number, as a table of contents' lines end in their page numbers.
_LINE_ENDS_IN_NUMBER = re.compile(r"(?m)\b[A-Z]?-?\d{1,4}\s*$")


def _page_list(numbers: list[int]) -> str:
    return ", ".join(map(str, numbers[:8])) + (f" and {len(numbers) - 8} more" if len(numbers) > 8 else "")


_SCANNED_NOTE = (
    "[This PDF looks scanned: its pages are pictures with little or no text layer. "
    "Text from it may be missing; open the file to read it.]"
)
_OCR_HINT = ' [To read scanned pages, install the OCR add-on: pip install -e ".[ocr]"]'
MAX_OCR_PAGES = 40
# A page of dense text or a detailed drawing is a few MB of drawing instructions; this is far past any real file.
MAX_PDF_CONTENT = 100_000_000
_HEAVY_NOTE = (
    "[This PDF unpacks to far more than its pages need, as a damaged or malicious file does, so CloseDesk read "
    "only the text it could get safely. Open the file to check it.]"
)


def _content_fits(pages) -> bool:
    """Whether the pages' drawing instructions (their content streams, and the forms they draw) unpack to less
    than ``MAX_PDF_CONTENT`` in all, with none past pypdf's own limit for one stream."""
    from pypdf.errors import LimitReachedError

    total = 0
    for page in pages:
        try:
            resources = page.get("/Resources")
            xobjects = resources.get_object().get("/XObject") if resources is not None else None
            forms = [item.get_object() for item in (xobjects.get_object().values() if xobjects is not None else [])]
            contents = page.get("/Contents")
            contents = contents.get_object() if contents is not None else None
            streams = list(contents) if isinstance(contents, list) else [contents] if contents is not None else []
            streams += [form for form in forms if form.get("/Subtype") == "/Form"]
            for stream in streams:
                total += len(stream.get_object().get_data())
        except LimitReachedError:
            return False
        except Exception:
            continue
        if total > MAX_PDF_CONTENT:
            return False
    return True


class _Miner:
    """pdfminer's pages of a PDF, each read on request into its characters' positions.

    pdfminer is asked for the characters only: its layout analysis (grouping them into text
    boxes) takes many times longer, and ``pdf_layout`` does that work itself.

    Pages are matched to pypdf's by object number, not by position. pdfminer's walk of the page
    tree passes over a page whose dictionary lacks ``/Type /Page`` and visits a page listed twice
    only once; matched by position, every page after it would get the wrong page's text.
    """

    def __init__(self, document, pages: list) -> None:
        from pdfminer.converter import PDFPageAggregator
        from pdfminer.pdfinterp import PDFPageInterpreter, PDFResourceManager

        self.document = document
        self.pages = pages
        self.by_id = {page.pageid: page for page in pages}
        manager = PDFResourceManager()
        self.device = PDFPageAggregator(manager, laparams=None)
        self.interpreter = PDFPageInterpreter(manager, self.device)

    @classmethod
    def open(cls, data: bytes) -> _Miner | None:
        """None when pdfminer can't open the file at all."""
        try:
            from pdfminer.pdfdocument import PDFDocument
            from pdfminer.pdfpage import PDFPage
            from pdfminer.pdfparser import PDFParser

            document = PDFDocument(PDFParser(io.BytesIO(data)))
        except Exception:
            return None
        pages: list = []
        try:
            for page in PDFPage.create_pages(document):
                pages.append(page)
                if len(pages) >= MAX_PDF_PAGES:
                    break
        except Exception:
            pass
        return cls(document, pages)

    def page_for(self, page, index: int):
        """pdfminer's page for one of pypdf's (the same object), built from the page's own dictionary
        when pdfminer's walk passed it over; by position when pypdf has no object number for it."""
        objid = getattr(getattr(page, "indirect_reference", None), "idnum", None)
        if objid is None:
            return self.pages[index] if index < len(self.pages) else None
        if objid not in self.by_id:
            self.by_id[objid] = self._build(objid)
        return self.by_id[objid]

    def _build(self, objid: int):
        from pdfminer.pdfpage import PDFPage
        from pdfminer.pdftypes import dict_value, resolve1

        try:
            attrs = dict(dict_value(self.document.getobj(objid)))
            parent, steps = attrs.get("Parent"), 0
            # The fonts and the page size can be set on the page tree above the page.
            while parent is not None and steps < 32:
                node = dict_value(resolve1(parent))
                for key in PDFPage.INHERITABLE_ATTRS:
                    if key not in attrs and key in node:
                        attrs[key] = node[key]
                parent, steps = node.get("Parent"), steps + 1
            return PDFPage(self.document, objid, attrs)
        except Exception:
            return None

    def layout(self, page):
        """The page's characters and drawn lines (pdfminer's ``LTPage``), or None when it can't be read."""
        try:
            self.interpreter.process_page(page)
            return self.device.get_result()
        except Exception:
            return None


def _pdf_page(page) -> str:
    """pypdf's text for a page, for the files pdfminer can't lay out."""
    try:
        text = page.extract_text(extraction_mode="layout")
    except Exception:
        text = ""
    if not text.strip():
        try:
            return page.extract_text() or ""
        except Exception:
            return ""
    lines = []
    for line in text.splitlines():
        line = re.sub(r" {3,}", " | ", line.rstrip()).strip(" |")
        lines.append(line)
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _has_images(page) -> bool:
    try:
        return len(page.images) > 0
    except Exception:
        return False


def _ocr_page_images(page) -> str:
    """OCR a scanned page's pictures when an OCR engine is installed; otherwise nothing."""
    texts = []
    try:
        for image in page.images[:4]:
            texts.append(ocr.image_text(image.data))
    except Exception:
        return ""
    return "\n".join(t for t in texts if t)


# Office files are zips of XML parts ---------------------------------------------------------------


def _office_refusal(data: bytes, *, workbook: bool = False) -> str:
    """A note saying why an Office file is not opened, or "" when it can be.

    A Word, Excel or PowerPoint file is a zip of XML parts, and a few hundred KB can unpack to gigabytes (a "zip
    bomb"); Word and PowerPoint files are read into memory whole. As with zip attachments, a file whose parts
    unpack to more than ``MAX_UNZIPPED`` (a workbook, read row by row, four times that), or a large part packed
    tighter than ``MAX_ZIP_RATIO``, is not opened.
    """
    import zipfile

    from controller_inbox.extract import MAX_UNZIPPED, MAX_ZIP_RATIO

    try:
        parts = zipfile.ZipFile(io.BytesIO(data)).infolist()
    except Exception:
        return ""  # not a zip: the reader says what is wrong with it
    limit, large = (4 * MAX_UNZIPPED, 25_000_000) if workbook else (MAX_UNZIPPED, 1_000_000)
    unpacked = sum(part.file_size for part in parts if part.filename.lower().endswith((".xml", ".rels", ".vml")))
    tight = any(part.file_size > large and part.file_size > MAX_ZIP_RATIO * max(1, part.compress_size) for part in parts)
    if unpacked <= limit and not tight:
        return ""
    return (
        f"[CloseDesk didn't open this file: it unpacks to {max(unpacked, 1_000_000) // 1_000_000:,} MB or more, "
        "far more than a real document holds, as a damaged or malicious file does. Open it only if you trust the sender.]"
    )


# Word ----------------------------------------------------------------------------------------

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def docx_text(data: bytes) -> str:
    """Paragraphs and tables in reading order, with headings, tracked insertions, and comments."""
    from docx import Document

    if refused := _office_refusal(data):
        return refused
    document = Document(io.BytesIO(data))
    lines: list[str] = []
    for block in _docx_blocks(document.element.body):
        if block.tag == f"{_W}p":
            line = _docx_paragraph(block)
            if line:
                lines.append(line)
        elif block.tag == f"{_W}tbl":
            lines.extend(_docx_table(block))
            lines.append("")
    comments = _docx_comments(document)
    if comments:
        lines += ["", "Comments in the document:"] + [f"- {item}" for item in comments]
    if any("[deleted:" in line for line in lines):
        lines.insert(0, "[This draft has tracked changes: inserted text is shown, deletions are marked [deleted: …].]")
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _docx_blocks(parent):
    """Paragraphs and tables in order, including those inside content controls (``w:sdt``), which
    templates use for cover pages, form fields and amounts."""
    for child in parent.iterchildren():
        if child.tag == f"{_W}sdt":
            content = child.find(f"{_W}sdtContent")
            if content is not None:
                yield from _docx_blocks(content)
        else:
            yield child


_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"


def _docx_runs(element, parts: list[str]) -> None:
    """The text under ``element``. A text box is stored twice (a drawing, then an older VML copy as the
    ``mc:Fallback``), so the copy is skipped; its paragraphs go on lines of their own. Text a tracked change
    moved is kept at its old place too (``w:moveFrom``, as plain ``w:t``); it is read where it went (``w:moveTo``)."""
    for node in element:
        if node.tag in (_MC_FALLBACK, f"{_W}pPr", f"{_W}moveFrom"):
            continue
        if node.tag == f"{_W}p":
            parts.append("\x00")
            _docx_runs(node, parts)
            parts.append("\x00")
        elif node.tag == f"{_W}t" and node.text:
            parts.append(node.text)
        elif node.tag == f"{_W}delText" and node.text:
            parts.append(f"[deleted: {node.text}]")
        elif node.tag == f"{_W}tab":
            parts.append("\t")
        elif node.tag in {f"{_W}br", f"{_W}cr"}:
            parts.append("\n")
        else:
            _docx_runs(node, parts)


def _docx_paragraph(element) -> str:
    parts: list[str] = []
    _docx_runs(element, parts)
    text = re.sub(r"\s*\x00[\x00\s]*", "\n", "".join(parts)).strip()
    if not text:
        return ""
    style = element.find(f"{_W}pPr/{_W}pStyle")
    name = (style.get(f"{_W}val") if style is not None else "") or ""
    level = re.match(r"(?i)heading\s*(\d)", name) or re.match(r"(?i)^title$", name)
    if level:
        depth = int(level.group(1)) if level.groups() else 1
        return "#" * min(depth, 4) + " " + text
    if element.find(f"{_W}pPr/{_W}numPr") is not None or re.match(r"(?i)list", name):
        return "- " + text
    return text


def _docx_table(element) -> list[str]:
    """Rows through ``tables``: a cell spanning columns is one cell, a vertically merged cell repeats
    the value above it, and the header row is the one Word marks as repeating, or a bold first row."""
    grid: list[list[str | None]] = []
    marked = False
    bold: list[bool] = []
    for number, row in enumerate(element.findall(f"{_W}tr")):
        if number == 0:
            marked = row.find(f"{_W}trPr/{_W}tblHeader") is not None
        # A row that starts further right (Word's "grid before") has no cells for the columns it skips.
        cells: list[str | None] = [""] * _docx_count(row.find(f"{_W}trPr/{_W}gridBefore"), 0)
        for cell in _docx_cells(row):
            text = " ".join(filter(None, (_docx_paragraph(p) for p in cell.iter(f"{_W}p") if _docx_own(p, cell)))).strip()
            properties = cell.find(f"{_W}tcPr")
            span = properties.find(f"{_W}gridSpan") if properties is not None else None
            merge = properties.find(f"{_W}vMerge") if properties is not None else None
            if merge is not None and merge.get(f"{_W}val", "continue") == "continue":
                above = grid[-1] if grid else []
                text = next((c for c in reversed(above[: len(cells) + 1]) if c is not None), "") if len(cells) < len(above) else ""
            cells.append(text)
            cells.extend([None] * (_docx_count(span, 1) - 1))
        grid.append(cells)
        if number == 0:
            runs = [r for r in row.iter(f"{_W}r") if "".join(t.text or "" for t in r.iter(f"{_W}t")).strip()]
            bold.append(bool(runs) and all(_docx_bold(r) for r in runs))
    header = tables.has_header(grid, marked=marked, bold_first=bool(bold and bold[0]))
    return tables.table_lines(grid, header=header)


def _docx_count(element, default: int) -> int:
    """A column count such as ``w:gridSpan``: the default when it is missing or unreadable, and no more
    columns than a Word table can have."""
    value = (element.get(f"{_W}val") or "") if element is not None else ""
    return min(int(value), 63) if value.isdigit() else default


def _docx_own(paragraph, cell) -> bool:
    """Whether a paragraph is the cell's own (or a nested table's), not one in a text box, which the
    paragraph holding the box already reads."""
    for parent in paragraph.iterancestors():
        if parent is cell:
            return True
        if parent.tag in (f"{_W}p", _MC_FALLBACK):
            return False
    return True


def _docx_cells(row):
    """A row's own cells, including those inside content controls, not those of a table nested in one."""
    for child in row:
        if child.tag == f"{_W}tc":
            yield child
        elif child.tag == f"{_W}sdt":
            content = child.find(f"{_W}sdtContent")
            if content is not None:
                yield from _docx_cells(content)


def _docx_bold(run) -> bool:
    mark = run.find(f"{_W}rPr/{_W}b")
    return mark is not None and mark.get(f"{_W}val", "true") not in ("0", "false", "off")


def _docx_comments(document) -> list[str]:
    try:
        comments = list(document.comments)
    except Exception:
        return []
    out = []
    for comment in comments[:50]:
        text = " ".join(p.text for p in comment.paragraphs).strip()
        if text:
            who = (getattr(comment, "author", "") or "").strip()
            out.append(f"{who}: {text}" if who else text)
    return out


# Workbooks -----------------------------------------------------------------------------------


def xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    if refused := _office_refusal(data, workbook=True):
        return refused
    values = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    lines: list[str] = []
    try:
        names = _defined_names(formulas)
        sheets = formulas.worksheets
        if len(sheets) > MAX_SHEETS:
            lines.append(f"[The workbook has {len(sheets)} sheets; CloseDesk read the first {MAX_SHEETS}.]")
        summary = ", ".join(f'"{ws.title}"' + (" (hidden)" if ws.sheet_state != "visible" else "") for ws in sheets)
        lines.append(f"Workbook with {len(sheets)} sheet{'s' if len(sheets) != 1 else ''}: {summary}")
        if names:
            lines.append("Named ranges: " + "; ".join(f"{k} = {v}" for k, v in names[:30]))
        for ws in sheets[:MAX_SHEETS]:
            lines.append("")
            lines.extend(_sheet_lines(ws, values[ws.title]))
    finally:
        values.close()
        formulas.close()
    return "\n".join(lines).strip()


def _defined_names(book) -> list[tuple[str, str]]:
    try:
        return [(name, str(item.attr_text)) for name, item in book.defined_names.items()]
    except Exception:
        return []


_HEADER_SEARCH = 10


@dataclass
class SheetRow:
    number: int
    cells: dict[int, str]
    bold: bool = False


def sheet_row_lines(rows: list[SheetRow], *, first_names_columns: bool = False) -> list[str]:
    """``A5: value | C5: value`` for each row. Under a header row each cell also names its column,
    ``C5 (Department): Finance``, and a blank between filled cells is written ``C5 (Department): not listed``,
    so a row about one person or account reads on its own. ``first_names_columns``: a CSV, whose first row
    is its header whenever it holds distinct labels, even over columns of names only."""
    from openpyxl.utils import get_column_letter

    header = _sheet_header(rows, first_names_columns=first_names_columns)
    labels = {c: v.replace("|", "/") for c, v in rows[header].cells.items()} if header is not None else {}
    out = []
    for index, row in enumerate(rows):
        named = header is not None and index > header
        columns = sorted(row.cells)
        if named and len(columns) >= 2:
            columns = sorted(set(columns) | {c for c in labels if columns[0] < c < columns[-1]})
        cells = []
        for column in columns:
            ref = f"{get_column_letter(column)}{row.number}"
            label = f" ({labels[column]})" if named and column in labels else ""
            cells.append(f"{ref}{label}: {row.cells.get(column) or tables.BLANK}")
        out.append(" | ".join(cells))
    return out


def _sheet_header(rows: list[SheetRow], *, first_names_columns: bool = False) -> int | None:
    """The index of the row naming the columns: one of the first rows, with two or more distinct labels,
    bold over rows that aren't, or over a column of figures."""
    for index, row in enumerate(rows[:_HEADER_SEARCH]):
        if len(row.cells) < 2:
            continue
        below = rows[index + 1 : index + 31]
        columns = range(min(row.cells), max(max(r.cells) for r in [row, *below]) + 1)
        grid = [[r.cells.get(c, "") for c in columns] for r in [row, *below]]
        bold = row.bold and not all(r.bold for r in below)
        if tables.has_header(grid, bold_first=bold, marked=first_names_columns and index == 0 and row.number == 1):
            return index
    return None


def _sheet_lines(formula_sheet, value_sheet) -> list[str]:
    from openpyxl.utils import get_column_letter, range_boundaries

    try:
        declared = formula_sheet.calculate_dimension()
    except Exception:
        declared = ""
    # Writers other than Excel often store too small a size (<dimension ref="A1"/>), and read-only mode
    # stops there; forget it, so every row in the file is read.
    formula_sheet.reset_dimensions()
    value_sheet.reset_dimensions()
    rows: list[SheetRow] = []
    more = 0
    formula_count = 0
    first_col = first_row = last_col = last_row = 0
    for index, (frow, vrow) in enumerate(
        zip(
            formula_sheet.iter_rows(max_col=MAX_COLUMNS),
            value_sheet.iter_rows(max_col=MAX_COLUMNS, values_only=True),
        )
    ):
        cells: dict[int, str] = {}
        bold = len(rows) < _HEADER_SEARCH
        for column, cell in enumerate(frow):
            raw = getattr(cell, "value", None)
            value = vrow[column] if column < len(vrow) else None
            if raw is None and value is None:
                continue
            first_col, last_col = min(first_col or column + 1, column + 1), max(last_col, column + 1)
            first_row, last_row = (first_row or index + 1), index + 1
            if formula := _formula(raw):
                formula_count += 1
                shown = f"{_formatted(value, cell)} ({formula})" if value is not None else formula
            else:
                shown = _formatted(value if value is not None else raw, cell)
            if shown:
                cells[column + 1] = shown
                bold = bold and bool(getattr(getattr(cell, "font", None), "b", False))
        if not cells:
            continue
        if len(rows) >= MAX_ROWS:
            more += 1
            continue
        number = next((cell.row for cell in frow if getattr(cell, "row", None)), index + 1)
        rows.append(SheetRow(number, cells, bold))
    _fill_merged_down(rows, _merged_ranges(formula_sheet))
    rows_text = sheet_row_lines(rows)
    try:
        min_col, min_row, max_col, max_row = range_boundaries(declared)
        fits = min_col <= first_col and min_row <= first_row and last_col <= max_col and last_row <= max_row
    except (ValueError, TypeError):
        fits = False
    extent = declared or "A1:A1"
    if last_row and not fits:
        extent = f"{get_column_letter(first_col)}{first_row}:{get_column_letter(last_col)}{last_row}"
    hidden = " hidden" if formula_sheet.sheet_state != "visible" else ""
    head = f'[sheet "{formula_sheet.title}" {extent}{hidden}]'
    note = f"({len(rows) + more} row{'s' if len(rows) + more != 1 else ''} with data" + (
        f", {formula_count} formula{'s' if formula_count != 1 else ''}" if formula_count else ""
    ) + ")"
    out = [head, note, *rows_text]
    if more:
        out.append(f"[{more} more rows not shown; ask for a range such as rows {MAX_ROWS + 1}–{MAX_ROWS + 200}.]")
    return out


_MERGE_CELL = re.compile(rb'<(?:\w+:)?mergeCell\b[^>]*?\bref="([A-Z]{1,3})(\d+):([A-Z]{1,3})(\d+)"')


def _merged_ranges(sheet) -> list[tuple[int, int, int, int]]:
    """The sheet's merged ranges as (first column, first row, last column, last row). Read-only mode doesn't
    load them; they are listed after the cells, so the sheet's file is scanned for them, not parsed again."""
    from openpyxl.utils import column_index_from_string

    found: list[tuple[int, int, int, int]] = []
    try:
        with sheet._get_source() as source:
            tail = b""
            while chunk := source.read(1 << 20):
                text = tail + chunk
                if b"mergeCell" in text:
                    ends = 0
                    for match in _MERGE_CELL.finditer(text):
                        first, top, last, bottom = match.groups()
                        found.append((column_index_from_string(first.decode()), int(top), column_index_from_string(last.decode()), int(bottom)))
                        ends = match.end()
                    tail = text[max(ends, len(text) - 200):]
                else:
                    tail = text[-200:]
    except Exception:
        return []
    return found


def _fill_merged_down(rows: list[SheetRow], merged: list[tuple[int, int, int, int]]) -> None:
    """A value merged down a column (a department over its vendors' rows) is held only by its top cell; the
    rows under it get it too, as Word's merged table cells do, so each row reads on its own. Only below the
    header row: a merge across a heading, or over a title, stays as it is."""
    tall = [area for area in merged if area[3] > area[1]]
    if not tall or not rows:
        return
    header = _sheet_header(rows)
    below = rows[header].number if header is not None else 0
    by_number = {row.number: row for row in rows}
    for column, top, _last, bottom in tall:
        start = by_number.get(top)
        value = start.cells.get(column) if start is not None and top > below else None
        if not value:
            continue
        for number in range(top + 1, min(bottom, top + MAX_ROWS) + 1):
            row = by_number.get(number)
            if row is not None and column not in row.cells:
                row.cells[column] = value


def _formula(raw) -> str:
    """A cell's formula as Excel writes it ("=SUM(B2:B3)"), or "" when the cell holds a value. An array formula
    (entered with Ctrl+Shift+Enter, or a dynamic array) and a what-if data table come from openpyxl as objects,
    not text."""
    from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

    if isinstance(raw, str):
        return raw if raw.startswith("=") else ""
    if isinstance(raw, ArrayFormula):
        text = str(raw.text or "")
        return text if text.startswith("=") else f"={text}"
    if isinstance(raw, DataTableFormula):
        return f"=TABLE({raw.r1 or ''},{raw.r2 or ''})"
    return ""


def xls_text(data: bytes) -> str:
    import xlrd
    from openpyxl.utils import get_column_letter

    book = xlrd.open_workbook(file_contents=data)
    lines = [f"Workbook with {book.nsheets} sheet{'s' if book.nsheets != 1 else ''}: " + ", ".join(f'"{s.name}"' for s in book.sheets())]
    for sheet in book.sheets()[:MAX_SHEETS]:
        lines.append("")
        last = get_column_letter(max(1, min(sheet.ncols, MAX_COLUMNS)))
        lines.append(f'[sheet "{sheet.name}" A1:{last}{max(sheet.nrows, 1)}]')
        rows: list[SheetRow] = []
        for r in range(sheet.nrows):
            cells: dict[int, str] = {}
            for c in range(min(sheet.ncols, MAX_COLUMNS)):
                cell = sheet.cell(r, c)
                if cell.value in ("", None):
                    continue
                shown = _fmt(_xls_value(cell, book.datemode))
                if shown:
                    cells[c + 1] = shown
            if cells:
                if len(rows) >= MAX_ROWS:
                    rows_left = sheet.nrows - r
                    break
                rows.append(SheetRow(r + 1, cells))
        else:
            rows_left = 0
        lines.extend(sheet_row_lines(rows))
        if rows_left:
            lines.append(f"[{rows_left} more rows not shown.]")
    return "\n".join(lines).strip()


def _xls_value(cell, datemode: int):
    """An Excel 97-2003 cell's value as the sheet shows it. xlrd gives an error such as #N/A as its code
    (42) and TRUE as 1, which would read as figures."""
    import xlrd

    if cell.ctype == xlrd.XL_CELL_ERROR:
        return xlrd.error_text_from_code.get(cell.value, "#ERROR")
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate.xldate_as_datetime(cell.value, datemode)
        except Exception:
            pass
    return cell.value


def _excel_named_text(data: bytes, filename: str) -> str:
    """A file called a workbook, by its name or by the type Windows gives every .csv, read as what it
    is: an Excel 97-2003 workbook, an .xlsx renamed, a web page table saved as .xls, or delimited text."""
    if data[:4] == b"\xd0\xcf\x11\xe0":
        return xls_text(data)
    if data[:4] == b"PK\x03\x04":
        return xlsx_text(data)
    text = decode_text(data[:4096]).lstrip("﻿ \t\r\n")
    if re.match(r"(?is)(?:<!--.*?-->\s*)*<(?:!doctype\s+html|html|head|body|meta|table|\?xml)\b", text):
        return markup_table_text(data, filename)
    return csv_text(data, filename)


def markup_table_text(data: bytes, filename: str) -> str:
    """The rows of the tables in a web page, as a sheet. "Export to Excel" in many web apps sends one
    named .xls; an Excel 2003 XML workbook has the same shape (``Row`` and ``Cell``). Each table, or each
    worksheet of an XML workbook, is a sheet of its own: its rows are under its own headers."""
    from html.parser import HTMLParser

    # (sheet name, rows) for each table; the rows of the one being read are ``state["rows"]``.
    sections: list[tuple[str, list[list[str]]]] = []
    loose: list[str] = []
    cell: list[str] = []
    state: dict = {"in_cell": False, "span": 1, "down": 0, "start": 0, "skip": 0, "rows": [], "carry": {}}

    def new_section(name: str = "") -> None:
        # A table nested in a cell, or the Table inside a Worksheet, goes on in the section it opens in until
        # that one has rows.
        if sections and not any(any(row) for row in state["rows"]):
            sections[-1] = (name or sections[-1][0], state["rows"])
            return
        state["rows"], state["carry"] = [], {}
        sections.append((name, state["rows"]))

    def pad_to(width: int) -> None:
        """Fill the row up to ``width`` cells: with the value of a cell merged down from a row above (rowspan,
        ss:MergeDown), which the rows under it leave out, else blank."""
        row, carry = state["rows"][-1], state["carry"]
        while len(row) < width:
            held = carry.get(len(row))
            row.append(held[1] if held else "")
            if held:
                held[0] -= 1
                if held[0] <= 0:
                    del carry[len(row) - 1]

    def merged_here() -> None:
        while len(state["rows"][-1]) in state["carry"]:
            pad_to(len(state["rows"][-1]) + 1)

    def end_row() -> None:
        if state["rows"] and state["carry"]:
            ahead = [column for column in state["carry"] if column >= len(state["rows"][-1])]
            if ahead:
                pad_to(max(ahead) + 1)

    def end_cell() -> None:
        if state["in_cell"]:
            row = state["rows"][-1]
            text = re.sub(r"\s+", " ", "".join(cell)).strip()
            row.append(text)
            row.extend([""] * (state["span"] - 1))
            if state["down"]:
                for column in range(state["start"], state["start"] + state["span"]):
                    state["carry"][column] = [state["down"], text if column == state["start"] else ""]
            cell.clear()
            state["in_cell"] = False

    class Reader(HTMLParser):
        def handle_starttag(self, tag, attrs):
            tag = tag.rsplit(":", 1)[-1]
            found = {name.rsplit(":", 1)[-1]: value or "" for name, value in attrs}
            if tag in ("script", "style"):
                state["skip"] += 1
            elif tag in ("table", "worksheet"):
                end_cell()
                end_row()
                new_section(found.get("name", "") if tag == "worksheet" else "")
            elif tag in ("tr", "row"):
                end_cell()
                end_row()
                if not sections:
                    new_section()
                state["rows"].append([])
            elif tag in ("td", "th", "cell"):
                end_cell()
                if not sections:
                    new_section()
                if not state["rows"]:
                    state["rows"].append([])
                merged_here()
                # An XML workbook leaves out empty cells and gives the next one's column.
                if found.get("index", "").isdigit():
                    pad_to(min(int(found["index"]), MAX_COLUMNS) - 1)
                    merged_here()
                span, across = found.get("colspan", ""), found.get("mergeacross", "")
                span = int(span) if span.isdigit() else int(across) + 1 if across.isdigit() else 1
                state["span"] = max(1, min(span, MAX_COLUMNS))
                down, below = found.get("rowspan", ""), found.get("mergedown", "")
                down = int(down) - 1 if down.isdigit() and int(down) > 0 else int(below) if below.isdigit() else 0
                state["down"] = max(0, min(down, MAX_ROWS))
                state["start"] = len(state["rows"][-1])
                state["in_cell"] = True
            elif tag in ("br", "p", "div") and state["in_cell"]:
                cell.append(" ")

        def handle_endtag(self, tag):
            tag = tag.rsplit(":", 1)[-1]
            if tag in ("script", "style"):
                state["skip"] = max(0, state["skip"] - 1)
            elif tag in ("td", "th", "cell"):
                end_cell()
            elif tag in ("tr", "row", "table"):
                end_cell()
                end_row()

        def handle_data(self, text):
            if state["skip"]:
                return
            (cell if state["in_cell"] else loose).append(text)

    reader = Reader(convert_charrefs=True)
    reader.feed(decode_text(data))
    reader.close()
    end_cell()
    filled = [(name, rows) for name, rows in sections if any(any(row) for row in rows)]
    if not filled:
        return re.sub(r"\s+", " ", " ".join(loose)).strip()
    if len(filled) == 1:
        return _rows_text(filled[0][1], filename)
    names = [(name or f"{filename} table {number}").replace('"', "'") for number, (name, _rows) in enumerate(filled, start=1)]
    return "\n\n".join(_rows_text(rows, name) for name, (_name, rows) in zip(names, filled))


def csv_text(data: bytes, filename: str) -> str:
    text = decode_text(data).replace("\x00", "")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        # The sniffer turns doubled quotes off when its sample has none, and a quote written twice further down
        # ('"Pipe 3"" PVC, schedule 40"') would then end the cell at its comma. Every spreadsheet writes them.
        dialect.doublequote = True
    except csv.Error:
        dialect = csv.excel_tab if filename.lower().endswith(".tsv") else csv.excel
    quote = getattr(dialect, "quotechar", None) or '"'
    try:
        rows = list(csv.reader(io.StringIO(text), dialect))
        broken = _swallowed(rows, text, dialect)
    except csv.Error:  # a cell over 128 KB, as one opened by a stray quote can be
        broken = True
    if broken:
        rows = [_csv_line(line, dialect, quote) for line in text.splitlines()]
    return _rows_text(rows, filename)


def _swallowed(rows: list[list[str]], text: str, dialect) -> bool:
    """Whether a stray quote ran one cell over the rows below it: it never closed, or it closed only at a
    quote further on that isn't the end of a cell ('Desk,"1"'). A cell whose quotes pair up properly around
    lines of its own (a remit-to address: "PO Box 1200 / Suite 4, Building B") is one cell, unless lines shaped
    like the file's rows in it clearly outnumber the rows read."""
    cells = [cell for row in rows for cell in row if "\n" in cell]
    if not cells:
        return False
    try:
        # A quote inside a quoted cell is written twice: a lone one ends the cell, so a delimiter must follow it.
        for _row in csv.reader(io.StringIO(text), dialect, strict=True, doublequote=True):
            pass
    except csv.Error:
        return True
    width = len(rows[0]) - 1
    shaped = max(sum(line.count(dialect.delimiter) == width for line in cell.split("\n")[1:]) for cell in cells)
    return width >= 1 and shaped >= max(3, 2 * len(rows))


def _csv_line(line: str, dialect, quote: str) -> list[str]:
    """One line of a file whose quotes don't pair up. Quoted cells on a line with whole pairs still read as
    one cell each; a line with a stray quote is split at every delimiter, the quote kept as a character."""
    try:
        if line.count(quote) % 2 == 0:
            return next(csv.reader([line], dialect), [])
        return next(csv.reader([line], dialect, quoting=csv.QUOTE_NONE), [])
    except csv.Error:
        return line.split(dialect.delimiter)


def _rows_text(rows: list[list[str]], filename: str) -> str:
    """Rows of text cells as one sheet named after the file; the first row is its header when it names the columns."""
    from openpyxl.utils import get_column_letter

    width = max((len(r) for r in rows), default=1)
    lines = [f'[sheet "{filename}" A1:{get_column_letter(max(1, min(width, MAX_COLUMNS)))}{max(len(rows), 1)}]']
    sheet = [
        SheetRow(number, {c + 1: _fmt(v) for c, v in enumerate(row[:MAX_COLUMNS]) if v.strip()})
        for number, row in enumerate(rows[:MAX_ROWS], start=1)
    ]
    lines.extend(sheet_row_lines([row for row in sheet if row.cells], first_names_columns=True))
    if len(rows) > MAX_ROWS:
        lines.append(f"[{len(rows) - MAX_ROWS} more rows not shown.]")
    return "\n".join(lines)


def decode_text(data: bytes) -> str:
    """Text from a file of unknown encoding: UTF-16 with a byte-order mark (or the NUL pattern
    of one without), UTF-8, then Windows-1252."""
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16", errors="replace")
    head = data[:4096]
    if len(head) >= 4 and head.count(b"\x00") >= len(head) // 3:
        odd, even = head[1::2].count(b"\x00"), head[0::2].count(b"\x00")
        return data.decode("utf-16-le" if odd >= even else "utf-16-be", errors="replace")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


_MONTHS = ("January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December")
_DATE_PART = re.compile(r'(?i)"([^"]*)"|\\(.)|mmmm|mmm|mm|m|yyyy|yy|[^a-z]')


def _formatted(value, cell) -> str:
    """A workbook value as the sheet shows it where that reads differently from the value itself: a
    percentage (0.15 formatted 0.0% is 15.0%) and a month (a date formatted mmm-yy is Mar-26, not the
    first of the month), and a number whose format has no thousands separator ("0": invoice 100235, year 2026,
    not 100,235). Other numbers, and dates with a day, keep the plain form ``_fmt`` gives them."""
    number_format = getattr(cell, "number_format", None)
    if not isinstance(number_format, str) or number_format == "General":
        return _fmt(value)
    # The format for positive values, without its colour or locale, padding and fill; then without its literal text.
    section = re.sub(r"\[[^\]]*\]|[_*].", "", number_format.split(";")[0])
    bare = re.sub(r'"[^"]*"|\\.', "", section)
    if isinstance(value, (int, float)) and not isinstance(value, bool) and "%" in bare and math.isfinite(value):
        decimals = re.search(r"\.([0#?]+)", bare)
        # Rounded half up, as Excel shows it (0.125 as 0% is 13%), from the value as written, not its binary form.
        places = len(decimals.group(1)) if decimals else 0
        shown = (Decimal(str(value)) * 100).quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
        return f"{shown:,.{places}f}%"
    if isinstance(value, (datetime, date)) and "m" in bare.lower() and "y" in bare.lower():
        parts = list(_DATE_PART.finditer(section))
        if "".join(part.group(0) for part in parts) == section:
            month, year = _MONTHS[value.month - 1], value.year
            words = {"mmmm": month, "mmm": month[:3], "mm": f"{value.month:02d}", "m": str(value.month), "yyyy": f"{year:04d}", "yy": f"{year % 100:02d}"}
            shown = []
            for part in parts:
                literal = part.group(1) if part.group(1) is not None else part.group(2)
                shown.append(literal if literal is not None else words.get(part.group(0).lower(), part.group(0)))
            return "".join(shown).strip()
    if isinstance(value, (int, float)) and not isinstance(value, bool) and "," not in bare:
        return _fmt(value, grouped=False)
    return _fmt(value)


def _fmt(value, *, grouped: bool = True) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return str(value)
        if value.is_integer():
            return f"{int(value):,}" if grouped else str(int(value))
        return (f"{value:,.4f}" if grouped else f"{value:.4f}").rstrip("0").rstrip(".")
    if isinstance(value, int):
        return f"{value:,}" if grouped else str(value)
    return re.sub(r"\s+", " ", str(value)).strip()


# PowerPoint ----------------------------------------------------------------------------------


def pptx_text(data: bytes) -> str:
    from pptx import Presentation

    if refused := _office_refusal(data):
        return refused
    deck = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for index, slide in enumerate(deck.slides, start=1):
        if index > 150:
            parts.append("[More slides not shown.]")
            break
        lines = [f"[slide {index}]"]
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        if title:
            lines.append(f"# {title}")
        for shape in _pptx_shapes(slide.shapes):
            if shape == slide.shapes.title:
                continue
            if getattr(shape, "has_table", False) and shape.has_table:
                lines.extend(["", *_pptx_table(shape.table), ""])
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    lines.append(text)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                lines.append(f"Speaker notes: {notes}")
        parts.append(re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip())
    return "\n\n".join(parts)


def _pptx_shapes(shapes):
    """Every shape on a slide, including those inside groups, in drawing order."""
    for shape in shapes:
        inner = getattr(shape, "shapes", None) if getattr(shape, "shape_type", None) == 6 else None  # MSO_SHAPE_TYPE.GROUP
        if inner is not None:
            yield from _pptx_shapes(inner)
        else:
            yield shape


def _pptx_table(table) -> list[str]:
    """Rows through ``tables``; the header is the first row when the table is styled with one."""
    grid: list[list[str | None]] = []
    for row in table.rows:
        cells: list[str | None] = []
        for column, cell in enumerate(row.cells):
            if cell._tc.get("hMerge") in ("1", "true"):
                cells.append(None)
            elif cell._tc.get("vMerge") in ("1", "true"):
                cells.append(grid[-1][column] if grid and column < len(grid[-1]) else "")
            else:
                cells.append(cell.text.strip())
        grid.append(cells)
    marked = bool(getattr(table, "first_row", False))
    return tables.table_lines(grid, header=tables.has_header(grid, marked=marked))


# Navigating the text -------------------------------------------------------------------------


def split_parts(text: str, *, size: int = PART_SIZE) -> list[Part]:
    """The document's sections, each no longer than ``size`` characters."""
    text = text or ""
    marks = list(MARKER_RE.finditer(text))
    sections: list[tuple[str, str]] = []
    if not marks:
        sections.append(("", text))
    else:
        head = text[: marks[0].start()].strip()
        if head:
            sections.append(("start", head))
        for i, mark in enumerate(marks):
            end = marks[i + 1].start() if i + 1 < len(marks) else len(text)
            sections.append((mark.group(1), text[mark.end() : end].strip()))
    parts: list[Part] = []
    for label, body in sections:
        pieces = _pieces(body, size)
        for n, piece in enumerate(pieces, start=1):
            name = label or f"part {len(parts) + 1}"
            if len(pieces) > 1:
                name = _sub_label(label, piece, n) if label else f"part {len(parts) + 1}"
            parts.append(Part(name, piece))
    return [p for p in parts if p.text.strip()] or [Part("part 1", "")]


def _pieces(body: str, size: int) -> list[str]:
    if len(body) <= size:
        return [body]
    out, current = [], []
    length = 0
    for line in body.split("\n"):
        while len(line) > size:
            if current:
                out.append("\n".join(current))
                current, length = [], 0
            cut = line[:size].rfind(" ")
            cut = cut if cut > size // 2 else size
            out.append(line[:cut])
            line = line[cut:].lstrip()
        if length + len(line) + 1 > size and current:
            out.append("\n".join(current))
            current, length = [], 0
        current.append(line)
        length += len(line) + 1
    if current:
        out.append("\n".join(current))
    return out


def _sub_label(label: str, piece: str, n: int) -> str:
    if label.startswith("sheet"):
        name = re.match(r'sheet "([^"]*)"', label)
        refs = re.findall(r"(?m)^[A-Z]{1,3}(\d+)(?::| \()", piece)
        if name and refs:
            return f'sheet "{name.group(1)}" rows {refs[0]}–{refs[-1]}'
    return f"{label} (cont. {n})"


def outline(text: str, *, limit: int = 40) -> list[str]:
    """One line per section: its label and how it starts."""
    rows = []
    for part in split_parts(text):
        first = next((line.strip() for line in part.text.splitlines() if line.strip() and not line.startswith("(")), "")
        rows.append(f"{part.label}: {first[:90]}")
    if len(rows) > limit:
        rows = rows[:limit] + [f"… {len(rows) - limit} more sections"]
    return rows


_TERM = re.compile(r"[A-Za-z0-9][A-Za-z0-9.$%&'-]*[A-Za-z0-9%]|[A-Za-z0-9]")


def terms_of(query: str, stop: set[str] | frozenset[str] = frozenset()) -> list[str]:
    words = []
    for match in _TERM.findall(query or ""):
        word = match.lower().strip(".'-")
        if (len(word) < 3 and not word.isdigit()) or word in stop:
            continue
        words.append(word)
    return list(dict.fromkeys(words))


def search_parts(text: str, query: str, *, limit: int = 3, stop: set[str] | frozenset[str] = frozenset()) -> list[Part]:
    """The sections that best match ``query`` (rarer words count more)."""
    parts = split_parts(text)
    words = terms_of(query, stop)
    if not words:
        return []
    lowered = [p.text.lower() + "\n" + p.label.lower() for p in parts]
    scored = []
    for index, body in enumerate(lowered):
        score = 0.0
        for word in words:
            hits = body.count(word)
            if hits:
                spread = sum(1 for other in lowered if word in other)
                score += (1 + math.log(hits)) * math.log(1 + len(parts) / spread)
        if score:
            scored.append((score, -index, parts[index]))
    scored.sort(key=lambda item: (item[0], item[1]), reverse=True)
    return [item[2] for item in scored[:limit]]


def excerpt(text: str, query: str, size: int) -> str:
    """At most ``size`` characters: the whole text if it fits, else its section list, its opening, and the matching sections."""
    text = (text or "").strip()
    if len(text) <= size:
        return text
    parts = split_parts(text)
    lines = []
    if len(parts) > 1:
        labels = [part.label for part in parts[:12]] + ([f"… {len(parts) - 12} more"] if len(parts) > 12 else [])
        lines.append("Sections: " + " / ".join(labels))
    room = size - sum(len(line) + 1 for line in lines)
    matches = [part for part in search_parts(text, query, limit=3) if part.label != parts[0].label]
    opening_cap = room // 3 if matches else room
    shown: dict[str, str] = {}
    for part in matches:
        shown[part.label] = _clip(part, (room - opening_cap) // len(matches))
    used = sum(len(block) + 1 for block in shown.values())
    shown[parts[0].label] = _clip(parts[0], room - used)
    order = {part.label: index for index, part in enumerate(parts)}
    lines += [block for label, block in sorted(shown.items(), key=lambda item: order[item[0]]) if block]
    return "\n".join(lines)


def skim(parts: list[Part], room: int, *, tag: str = "", line_size: int = 240) -> str:
    """The lines that stand out from ``parts``, labelled by section, within ``room`` characters.

    Lines whose shape repeats (page headers, boilerplate, "Section 3.2: no exceptions…") are dropped, so what is
    left is what differs from page to page: findings, totals, names, dates. Sections take turns so a long file is
    covered end to end instead of only its first pages.
    """

    def shape(line: str) -> str:
        return re.sub(r"\d+", "#", line.lower()).strip()

    counts: dict[str, int] = {}
    for part in parts:
        for line in part.text.splitlines():
            counts[shape(line)] = counts.get(shape(line), 0) + 1
    standout = [
        [line.strip()[:line_size] for line in part.text.splitlines() if len(line.split()) >= 3 and counts[shape(line)] < 3]
        for part in parts
    ]
    chosen: list[list[str]] = [[] for _ in parts]
    used = 0
    for depth in range(max((len(lines) for lines in standout), default=0)):
        for index, lines in enumerate(standout):
            if depth >= len(lines):
                continue
            cost = len(lines[depth]) + 1 + (0 if chosen[index] else len(tag) + len(parts[index].label) + 3)
            if used + cost > room:
                return _skim_text(parts, chosen, tag)
            chosen[index].append(lines[depth])
            used += cost
    return _skim_text(parts, chosen, tag)


def _skim_text(parts: list[Part], chosen: list[list[str]], tag: str) -> str:
    return "\n".join(f"[{tag}{part.label}]\n" + "\n".join(lines) for part, lines in zip(parts, chosen) if lines)


def _clip(part: Part, room: int) -> str:
    head = f"[{part.label}]\n"
    if room - len(head) < 100:
        return ""
    if len(head) + len(part.text) <= room:
        return head + part.text
    return head + part.text[: room - len(head) - 2].rsplit(" ", 1)[0] + " …"


def locate(parts: list[Part], at: str) -> tuple[int, str] | None:
    """Where a citation points: the index of the section for "page 3", "slide 2", 'sheet "Staff"', a cell
    such as "Staff!C5" or "C5", and the line holding that cell ("" when the whole section is meant)."""
    at = re.sub(r"\s+", " ", (at or "").strip())[:120]
    if not at:
        return None
    cell = re.fullmatch(r"(?:['\"]?([^'\"!]+?)['\"]?!)?\$?([A-Za-z]{1,3})\$?(\d{1,6})", at)
    if cell and not re.fullmatch(r"(?i)(?:page|slide|part)\d+", at):
        sheet, ref = cell.group(1), f"{cell.group(2).upper()}{cell.group(3)}"
        holds = re.compile(rf"(?:^|\| ){ref}(?: \(|:)")
        for index, part in enumerate(parts):
            if sheet and not part.label.lower().startswith(f'sheet "{sheet.lower()}"'):
                continue
            line = next((line for line in part.text.splitlines() if holds.search(line)), None)
            if line is not None:
                return index, line
    wanted = at.lower()
    number = re.fullmatch(r"(page|slide|part) ?(\d+)", wanted)
    for index, part in enumerate(parts):
        label = part.label.lower()
        if label == wanted or (number and re.match(rf"{number.group(1)} {number.group(2)}\b", label)):
            return index, ""
    for index, part in enumerate(parts):
        if wanted in part.label.lower():
            return index, ""
    return None


def read_part(text: str, label: str) -> Part | None:
    """A section by its label ("page 3", "slide 2", "Budget", "part 4"), matched loosely."""
    parts = split_parts(text)
    wanted = re.sub(r"\s+", " ", (label or "").strip().lower())
    if len(wanted) > 1 and wanted[0] == wanted[-1] == '"':
        wanted = wanted[1:-1]
    if not wanted:
        return parts[0] if parts else None
    for part in parts:
        if part.label.lower() == wanted:
            return part
    number = re.fullmatch(r"(?:(page|slide|part)\s*)?(\d+)", wanted)
    if number:
        kind = number.group(1)
        for part in parts:
            if re.match(rf"(?:{kind or 'page|slide|part'}) {number.group(2)}\b", part.label.lower()):
                return part
        index = int(number.group(2)) - 1
        if not kind and 0 <= index < len(parts):
            return parts[index]
    # "Budget" or 'sheet "Budget"' is that sheet, not "Budget 2025" that merely starts the same way.
    for part in parts:
        name = part.label.lower()
        if name.startswith((wanted + " ", f'sheet "{wanted}"', f'{wanted}"')):
            return part
    for part in parts:
        if wanted in part.label.lower():
            return part
    return None


# Exact cells from the original workbook ------------------------------------------------------

# Not part of a longer name (LOG10, ATAN2, 1.5E10) or of a link to another workbook ([1]Sheet1!B2),
# and not a function's name (LOG10( ).
_FORMULA_REF = re.compile(
    r"(?<![\w.\]!$'])(?:'((?:[^']|'')+)'!|([A-Za-z0-9_.]+)!)?\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?(?![\w(!])"
)
_FORMULA_TEXT = re.compile(r'"(?:[^"]|"")*"')
_FORMULA_LINK = re.compile(r"'?\[[^\[\]]+\][^\[\]!(),;+\-*/^&=<>]*!\$?[A-Za-z_]*\$?\d*(?::\$?[A-Z]{1,3}\$?\d+)?")


def read_cells(data: bytes, sheet: str, cells: str, *, limit: int = 200) -> str:
    """Values and formulas for a range such as ``B2:D20`` on ``sheet`` (name or 1-based number)."""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter, range_boundaries

    if refused := _office_refusal(data, workbook=True):
        return refused
    values = load_workbook(io.BytesIO(data), data_only=True)
    formulas = load_workbook(io.BytesIO(data), data_only=False)
    try:
        ws = _find_sheet(formulas, sheet)
        if ws is None:
            return f"No sheet called {sheet!r}. Sheets: " + ", ".join(f'"{s}"' for s in formulas.sheetnames)
        vs = values[ws.title]
        try:
            min_col, min_row, max_col, max_row = range_boundaries((cells or "A1:J40").replace("$", "").upper())
        except (ValueError, TypeError):
            return f"{cells!r} isn't a cell range. Use something like A1:F20."
        min_col, min_row = min_col or 1, min_row or 1
        max_col = min(max_col or ws.max_column, min_col + MAX_COLUMNS - 1)
        max_row = min(max_row or ws.max_row, ws.max_row)
        lines = [f'Sheet "{ws.title}" {get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}']
        shown = 0
        held = _held(ws, min_col, min_row, max_col, max_row)
        held_values = _held(vs, min_col, min_row, max_col, max_row)
        for r in sorted({r for r, _c in held} | {r for r, _c in held_values}):
            row = []
            for c in range(min_col, max_col + 1):
                raw = getattr(held.get((r, c)), "value", None)
                value = getattr(held_values.get((r, c)), "value", None)
                if raw is None and value is None:
                    continue
                ref = f"{get_column_letter(c)}{r}"
                if formula := _formula(raw):
                    row.append(f"{ref}: {_fmt(value)} ({formula})" if value is not None else f"{ref}: {formula}")
                else:
                    row.append(f"{ref}: {_fmt(value if value is not None else raw)}")
            if row:
                lines.append(" | ".join(row))
                shown += 1
                if shown >= limit:
                    lines.append(f"[Stopped at row {r}; ask for the next range.]")
                    break
        if shown == 0:
            lines.append("(these cells are empty)")
        return "\n".join(lines)
    finally:
        values.close()
        formulas.close()


def _held(ws, min_col: int, min_row: int, max_col: int, max_row: int) -> dict:
    """The cells of a range the sheet holds, by (row, column) in order, without making a cell for every
    empty address in it as ``ws.cell`` does: a sheet formatted down to row 1,048,576 would take minutes."""
    return {
        key: cell
        for key, cell in sorted(ws._cells.items())
        if min_row <= key[0] <= max_row and min_col <= key[1] <= max_col
    }


def trace_cell(data: bytes, sheet: str, cell: str, *, depth: int = 2) -> str:
    """A cell's formula and the cells it depends on, ``depth`` levels down."""
    from openpyxl import load_workbook

    if refused := _office_refusal(data, workbook=True):
        return refused
    values = load_workbook(io.BytesIO(data), data_only=True)
    formulas = load_workbook(io.BytesIO(data), data_only=False)
    try:
        ws = _find_sheet(formulas, sheet)
        if ws is None:
            return f"No sheet called {sheet!r}. Sheets: " + ", ".join(f'"{s}"' for s in formulas.sheetnames)
        ref = (cell or "").replace("$", "").upper().strip()
        if not re.fullmatch(r"[A-Z]{1,3}\d+", ref):
            return f"{cell!r} isn't a single cell. Use something like C12."
        lines: list[str] = []
        _trace(formulas, values, ws.title, ref, depth, 0, lines, set())
        return "\n".join(lines)
    finally:
        values.close()
        formulas.close()


def compare_columns(data: bytes, sheet: str, first: str, second: str, *, limit: int = 40) -> str:
    """How each row changes from column ``first`` to ``second`` (letters or header text), biggest increase first."""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    if refused := _office_refusal(data, workbook=True):
        return refused
    book = load_workbook(io.BytesIO(data), data_only=True)
    try:
        ws = _find_sheet(book, sheet)
        if ws is None:
            return f"No sheet called {sheet!r}. Sheets: " + ", ".join(f'"{s}"' for s in book.sheetnames)
        found = _columns(ws, first, second)
        if isinstance(found, str):
            return found
        header, a, b = found
        rows, totals = [], []
        for r in range(header + 1, ws.max_row + 1):
            old, new = ws.cell(r, a).value, ws.cell(r, b).value
            if not all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in (old, new)):
                continue
            words = [str(ws.cell(r, c).value).strip() for c in range(1, min(a, b)) if isinstance(ws.cell(r, c).value, str)]
            label = " · ".join(word for word in words if word) or f"row {r}"
            (totals if re.search(r"\btotal\b", label, re.I) else rows).append((new - old, r, label, old, new))
        if not rows and not totals:
            return (
                f"No rows on \"{ws.title}\" have numbers in both {get_column_letter(a)} and {get_column_letter(b)}. "
                "If they are formulas, the file may have been saved without values: open and save it in Excel."
            )
        def name(col: int) -> str:
            head = ws.cell(header, col).value if header else None
            return f"{get_column_letter(col)} “{_fmt(head)}”" if head is not None else get_column_letter(col)

        rows.sort(key=lambda row: row[0], reverse=True)
        lines = [f'Sheet "{ws.title}": {name(a)} → {name(b)}, {len(rows)} rows, biggest increase first']
        lines += [_change_line(*row) for row in rows[:limit]]
        if len(rows) > limit:
            lines.append(f"… {len(rows) - limit} more rows")
        if rows:
            up = f"Biggest increase: {rows[0][2]}." if rows[0][0] > 0 else "Nothing went up."
            down = f"Biggest decrease: {rows[-1][2]}." if rows[-1][0] < 0 else "Nothing went down."
            lines.append(f"{up} {down}")
        if totals:
            lines += ["Total rows:"] + [_change_line(*row) for row in totals]
        return "\n".join(lines)
    finally:
        book.close()


def _columns(ws, first: str, second: str):
    """(header row, first column, second column) from letters or header text, or a message saying what's there."""
    from openpyxl.utils import column_index_from_string

    headers: dict[str, tuple[int, int]] = {}
    for r in range(1, min(ws.max_row, 10) + 1):
        for c in range(1, min(ws.max_column, 60) + 1):
            value = ws.cell(r, c).value
            if isinstance(value, str) and value.strip():
                headers.setdefault(value.strip().lower(), (r, c))
    picked = []
    for wanted in (first, second):
        text = (wanted or "").strip().strip("'\"")
        # A header that is exactly this text wins over reading it as column letters ("Jan", "Qty", "Net").
        hit = headers.get(text.lower())
        if hit is None and re.fullmatch(r"[A-Za-z]{1,3}", text):
            try:
                picked.append((None, column_index_from_string(text.upper())))
                continue
            except ValueError:
                pass
        hit = hit or next((spot for head, spot in headers.items() if text and text.lower() in head), None)
        if hit is None:
            shown = ", ".join(f'"{head}"' for head in list(headers)[:20])
            return f"No column called {wanted!r}. Give column letters like C and D, or one of these headers: {shown}"
        picked.append(hit)
    rows = [r for r, _c in picked if r is not None]
    if rows:
        header = max(rows)
    else:
        header = next((r for r in range(1, min(ws.max_row, 10) + 1) if all(isinstance(ws.cell(r, c).value, str) for _r, c in picked)), 0)
    return header, picked[0][1], picked[1][1]


def _change_line(diff, row: int, label: str, old, new) -> str:
    sign = "+" if diff >= 0 else "−"
    pct = f" ({sign}{abs(diff / old * 100):.1f}%)" if old else ""
    return f"{label} (row {row}): {_fmt(old)} → {_fmt(new)}, {sign}{_fmt(abs(diff))}{pct}"


def _trace(formulas, values, sheet: str, ref: str, depth: int, level: int, lines: list[str], seen: set) -> None:
    from openpyxl.utils import range_boundaries, get_column_letter

    key = (sheet, ref)
    if key in seen or len(lines) > 80:
        return
    seen.add(key)
    raw = _formula(formulas[sheet][ref].value) or formulas[sheet][ref].value
    value = values[sheet][ref].value
    indent = "  " * level
    where = f"{sheet}!{ref}"
    if isinstance(raw, str) and raw.startswith("="):
        shown = f" = {_fmt(value)}" if value is not None else ""
        lines.append(f"{indent}{where}{shown} ← {raw}")
        if level >= depth:
            return
        formula = _FORMULA_TEXT.sub('""', raw[1:])
        for link in dict.fromkeys(_FORMULA_LINK.findall(formula)):
            lines.append(f"{indent}  {link}: in another workbook, not in this file")
        for match in _FORMULA_REF.finditer(formula):
            target = (match.group(1) or match.group(2) or sheet).replace("''", "'")
            if target not in formulas.sheetnames:
                continue
            start = f"{match.group(3)}{match.group(4)}"
            if match.group(5):
                end = f"{match.group(5)}{match.group(6)}"
                min_col, min_row, max_col, max_row = range_boundaries(f"{start}:{end}")
                # Only the part the sheet uses: an old lookup's A1:Z65536 is a few hundred cells, not 1.7 million.
                used = values[target]
                max_col, max_row = min(max_col, used.max_column), min(max_row, used.max_row)
                count = max(0, max_col - min_col + 1) * max(0, max_row - min_row + 1)
                if count > 12:
                    filled = [c.value for c in _held(used, min_col, min_row, max_col, max_row).values() if c.value is not None]
                    numbers = [v for v in filled if isinstance(v, (int, float)) and not isinstance(v, bool)]
                    lines.append(
                        f"{indent}  {target}!{start}:{end}: {len(filled)} filled cells"
                        + (f", {len(numbers)} numbers adding to {_fmt(sum(numbers))}" if numbers else "")
                    )
                    continue
                for r in range(min_row, max_row + 1):
                    for c in range(min_col, max_col + 1):
                        _trace(formulas, values, target, f"{get_column_letter(c)}{r}", depth, level + 1, lines, seen)
            else:
                _trace(formulas, values, target, start, depth, level + 1, lines, seen)
    else:
        lines.append(f"{indent}{where} = {_fmt(value if value is not None else raw) or '(empty)'}")


def _find_sheet(book, sheet: str):
    name = (sheet or "").strip().strip("'\"")
    if not name:
        return book.worksheets[0]
    for ws in book.worksheets:
        if ws.title.lower() == name.lower():
            return ws
    if name.isdigit() and 1 <= int(name) <= len(book.worksheets):
        return book.worksheets[int(name) - 1]
    for ws in book.worksheets:
        if name.lower() in ws.title.lower():
            return ws
    return None
