from __future__ import annotations

import hashlib
import io
import re
from datetime import date, datetime
from typing import Iterable

from dateutil import parser as date_parser

from controller_inbox.models import ExtractedFields


INVOICE_RE = re.compile(
    r"\b(?:invoice|inv\.?|bill)[\s#:No.-]*([A-Z]{1,6}[-_]?\d{2,12}|\d{3,12})",
    re.IGNORECASE,
)
INVOICE_BARE_RE = re.compile(r"\b(INV[-_]?\d{3,8}|IN[-_]?\d{4,8})\b", re.IGNORECASE)
PO_RE = re.compile(
    r"\b(?:purchase\s+order|p\.?o\.?)[\s#:No.-]*([A-Z]{0,4}-?\d{3,10})\b",
    re.IGNORECASE,
)
AMOUNT_RE = re.compile(
    r"(?<!\w)(?:USD|US\$|\$)\s*([0-9]{1,3}(?:,[0-9]{3})*(?:\.[0-9]{2})|[0-9]+\.[0-9]{2})(?!\w)"
)
AMOUNT_WORDS_RE = re.compile(
    r"\b(?:amount(?:\s+due)?|total(?:\s+due)?|balance(?:\s+due)?|grand\s+total)\s*[:\-]?\s*\$?\s*"
    r"([0-9]{1,3}(?:,[0-9]{3})*(?:\.[0-9]{2})|[0-9]+\.[0-9]{2})",
    re.IGNORECASE,
)
DUE_RE = re.compile(
    r"\b(?:due(?:\s+date)?|payment\s+due|remit\s+by|pay\s+by|respond\s+by|needed\s+by|"
    r"please\s+(?:complete|provide|respond|approve)\s+by|deadline|by)\s*[:\-]?\s*"
    r"([A-Za-z]{3,9}\.?\s+\d{1,2},?\s+\d{2,4}|\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{4}-\d{2}-\d{2}|"
    r"EOD|COB|today|tomorrow|Monday|Tuesday|Wednesday|Thursday|Friday)",
    re.IGNORECASE,
)
ACCOUNT_RE = re.compile(
    r"\b(?:account(?:\s+number)?|acct\.?|a/c|iban)[\s#:]*([A-Z0-9]{6,34})\b",
    re.IGNORECASE,
)
ROUTING_RE = re.compile(
    r"\b(?:routing(?:\s+number)?|aba|sort\s+code)[\s#:]*(\d{6,9})\b",
    re.IGNORECASE,
)
VENDOR_RE = re.compile(
    r"\b(?:from|vendor|supplier|bill\s+from|sold\s+by|remit\s+to)\s*[:\-]\s*([A-Z][A-Za-z0-9&.,' \-]{2,60})",
)
ATTACHMENT_MENTION_RE = re.compile(
    r"\b(attached|attachment|enclosed|please\s+see\s+attached|see\s+the\s+attached)\b",
    re.IGNORECASE,
)
BANK_SECRET_RE = re.compile(
    r"\b(?:routing(?:\s+number)?|aba)[\s#:]*\d{6,9}\b|"
    r"\b(?:account(?:\s+number)?|acct\.?)[\s#:]*[A-Z0-9]{6,34}\b|"
    r"\b(?:iban)[\s#:]*[A-Z]{2}\d{2}[A-Z0-9]{10,30}\b",
    re.IGNORECASE,
)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def html_to_text(html: str) -> str:
    text = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", html)
    text = re.sub(r"(?i)<br\s*/?>", "\n", text)
    text = re.sub(r"(?i)</p>", "\n", text)
    text = re.sub(r"(?i)</div>", "\n", text)
    text = re.sub(r"(?s)<[^>]+>", " ", text)
    text = re.sub(r"&nbsp;", " ", text)
    text = re.sub(r"&amp;", "&", text)
    text = re.sub(r"&lt;", "<", text)
    text = re.sub(r"&gt;", ">", text)
    text = re.sub(r"&#39;", "'", text)
    text = re.sub(r"&quot;", '"', text)
    return collapse_ws(text)


def collapse_ws(text: str) -> str:
    text = text.replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def redact_financial_secrets(text: str) -> str:
    def _mask(match: re.Match[str]) -> str:
        raw = match.group(0)
        digits = re.sub(r"\W", "", raw)
        last4 = digits[-4:] if len(digits) >= 4 else "****"
        if re.search(r"routing|aba|sort", raw, re.I):
            return f"routing ****{last4}"
        if re.search(r"iban", raw, re.I):
            return f"IBAN ****{last4}"
        return f"account ****{last4}"

    return BANK_SECRET_RE.sub(_mask, text)


def extract_text_from_bytes(filename: str, content_type: str, data: bytes) -> str:
    name = (filename or "").lower()
    ctype = (content_type or "").lower()
    if not data:
        return ""
    try:
        if name.endswith(".pdf") or "pdf" in ctype:
            return _pdf_text(data)
        if name.endswith(".docx") or "wordprocessingml" in ctype:
            return _docx_text(data)
        if name.endswith((".xlsx", ".xlsm", ".xltx")) or "spreadsheetml" in ctype:
            return _xlsx_text(data)
        if name.endswith(".xls") or ctype == "application/vnd.ms-excel":
            return _xls_text(data)
        if name.endswith(".pptx") or "presentationml" in ctype:
            return _pptx_text(data)
        if name.endswith(".csv") or ctype in {"text/csv", "application/csv"}:
            return data.decode("utf-8", errors="replace")[:50_000]
        if name.endswith((".txt", ".md", ".tsv")) or ctype.startswith("text/plain"):
            return data.decode("utf-8", errors="replace")[:50_000]
        if name.endswith(".rtf") or "rtf" in ctype:
            return _rtf_text(data)
        if "html" in ctype or name.endswith((".html", ".htm")):
            return html_to_text(data.decode("utf-8", errors="replace"))
        if ctype.startswith("image/") or name.endswith((".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp")):
            return _image_text(data)
    except Exception as exc:  # extraction should never fail the pipeline
        return f"[extraction error: {exc}]"
    # Last resort: if it looks like text, keep a sample.
    sample = data[:2000]
    if b"\x00" not in sample:
        try:
            decoded = sample.decode("utf-8")
            if decoded.isprintable() or "\n" in decoded:
                return decoded
        except UnicodeDecodeError:
            return ""
    return ""


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages[:50]:
        pages.append(page.extract_text() or "")
    return collapse_ws("\n".join(pages))


def _docx_text(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return collapse_ws("\n".join(parts))


def _xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets[:6]:
        parts.append(f"[sheet:{sheet.title}]")
        for i, row in enumerate(sheet.iter_rows(max_row=80, values_only=True)):
            values = [str(cell) for cell in row if cell is not None]
            if values:
                parts.append(" | ".join(values))
            if i >= 80:
                break
    wb.close()
    return collapse_ws("\n".join(parts))


def _xls_text(data: bytes) -> str:
    import xlrd

    book = xlrd.open_workbook(file_contents=data)
    parts: list[str] = []
    for sheet in book.sheets()[:6]:
        parts.append(f"[sheet:{sheet.name}]")
        for i in range(min(sheet.nrows, 80)):
            values = [str(cell) for cell in sheet.row_values(i) if cell not in ("", None)]
            if values:
                parts.append(" | ".join(values))
    return collapse_ws("\n".join(parts))


def _pptx_text(data: bytes) -> str:
    from pptx import Presentation

    deck = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for index, slide in enumerate(deck.slides, start=1):
        if index > 40:
            break
        parts.append(f"[slide:{index}]")
        for shape in slide.shapes:
            text = getattr(shape, "text", "") or ""
            if text.strip():
                parts.append(text)
    return collapse_ws("\n".join(parts))


def _rtf_text(data: bytes) -> str:
    raw = data.decode("latin-1", errors="replace")
    raw = re.sub(r"\\'[0-9a-fA-F]{2}", " ", raw)
    raw = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", raw)
    raw = raw.replace("\\", " ")
    raw = re.sub(r"[{}]", " ", raw)
    return collapse_ws(raw)


def _image_text(data: bytes) -> str:
    """OCR when a local Tesseract install is present. Otherwise the file stays an image scan."""
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        return ""
    try:
        image = Image.open(io.BytesIO(data))
        return collapse_ws(pytesseract.image_to_string(image))[:20_000]
    except Exception:
        return ""


def explode_archives(items: list, *, limit: int = 40) -> list:
    """Unpack zip attachments into the files inside, so a zipped workbook is still classified."""
    from controller_inbox.models import RawAttachment

    exploded: list[RawAttachment] = []
    for item in items:
        name = (item.filename or "").lower()
        if not name.endswith(".zip") and "zip" not in (item.content_type or "").lower():
            exploded.append(item)
            continue
        inner = _unzip(item, limit=limit)
        exploded.extend(inner or [item])
    return exploded


def _unzip(item, *, limit: int) -> list:
    import zipfile

    from controller_inbox.models import RawAttachment

    if not item.content:
        return []
    try:
        archive = zipfile.ZipFile(io.BytesIO(item.content))
    except zipfile.BadZipFile:
        return []
    out = []
    for info in archive.infolist():
        if info.is_dir() or len(out) >= limit:
            continue
        if info.file_size > 30_000_000 or info.filename.startswith("__MACOSX"):
            continue
        filename = _basename(info.filename)
        if not filename or filename.startswith("."):
            continue
        try:
            payload = archive.read(info)
        except Exception:
            continue
        out.append(
            RawAttachment(
                id=f"{item.id}:{filename}",
                filename=filename,
                content_type="application/octet-stream",
                size_bytes=len(payload),
                content=payload,
            )
        )
    return out


def _basename(filename: str) -> str:
    return filename.replace("\\", "/").split("/")[-1]


def parse_amount(raw: str) -> float | None:
    try:
        return round(float(raw.replace(",", "")), 2)
    except ValueError:
        return None


def parse_due_date(raw: str, *, as_of: date) -> str | None:
    token = raw.strip()
    low = token.lower()
    if low in {"eod", "cob", "today"}:
        return as_of.isoformat()
    if low == "tomorrow":
        return date.fromordinal(as_of.toordinal() + 1).isoformat()
    weekdays = {
        "monday": 0,
        "tuesday": 1,
        "wednesday": 2,
        "thursday": 3,
        "friday": 4,
    }
    if low in weekdays:
        delta = (weekdays[low] - as_of.weekday()) % 7
        if delta == 0:
            delta = 7
        return date.fromordinal(as_of.toordinal() + delta).isoformat()
    try:
        parsed = date_parser.parse(token, default=datetime(as_of.year, as_of.month, as_of.day), fuzzy=False)
        return parsed.date().isoformat()
    except (ValueError, OverflowError, TypeError):
        return None


def extract_fields(text: str, *, as_of: date, extra_vendor: str | None = None) -> ExtractedFields:
    invoices = _unique(_normalize_id(m.group(1)) for m in INVOICE_RE.finditer(text))
    invoices += [v for v in _unique(_normalize_id(m.group(1)) for m in INVOICE_BARE_RE.finditer(text)) if v not in invoices]
    pos = _unique(_normalize_id(m.group(1)) for m in PO_RE.finditer(text))
    amounts: list[float] = []
    for match in list(AMOUNT_RE.finditer(text)) + list(AMOUNT_WORDS_RE.finditer(text)):
        amount = parse_amount(match.group(1))
        if amount is not None and 0.01 <= amount <= 10_000_000:
            amounts.append(amount)
    amounts = _unique(amounts)
    due_dates: list[str] = []
    for match in DUE_RE.finditer(text):
        parsed = parse_due_date(match.group(1), as_of=as_of)
        if parsed:
            due_dates.append(parsed)
    due_dates = _unique(due_dates)
    vendors = _unique(m.group(1).strip(" \t-,.") for m in VENDOR_RE.finditer(text))
    if extra_vendor:
        vendors = _unique([extra_vendor, *vendors])
    last4: list[str] = []
    mentions_account = False
    for match in list(ACCOUNT_RE.finditer(text)) + list(ROUTING_RE.finditer(text)):
        mentions_account = True
        digits = re.sub(r"\W", "", match.group(1))
        if len(digits) >= 4:
            last4.append(digits[-4:])
    return ExtractedFields(
        invoice_numbers=invoices[:8],
        po_numbers=pos[:8],
        amounts=amounts[:8],
        due_dates=due_dates[:6],
        vendor_candidates=vendors[:6],
        account_last4=_unique(last4)[:6],
        mentions_routing_or_account=mentions_account,
        mentions_attachment=bool(ATTACHMENT_MENTION_RE.search(text)),
    )


def _normalize_id(value: str) -> str:
    return re.sub(r"\s+", "", value).strip(".,;:").upper()


def _unique(items: Iterable) -> list:
    seen: set[str] = set()
    out = []
    for item in items:
        if item is None or item == "":
            continue
        key = str(item)
        if key in seen:
            continue
        seen.add(key)
        out.append(item)
    return out


def build_search_blob(*parts: str) -> str:
    return collapse_ws("\n".join(p for p in parts if p))
