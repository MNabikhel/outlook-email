"""A workbook, CSV or Word document as it looks, for the Preview tab: every sheet with its cells shown as Excel
shows them (number formats, merged cells, bold, fills, column widths), or a Word document's headings, paragraphs
and tables in order. Read from the file as it arrived; nothing here runs a macro or a formula.

The JSON is small on purpose: a sheet sends at most ``MAX_ROWS`` rows of ``MAX_COLS`` columns, and says how many
more there are.
"""

from __future__ import annotations

import csv
import io
import math
import re
from datetime import date, datetime, time, timedelta
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

MAX_ROWS = 1000
MAX_COLS = 52
MAX_SHEETS = 40
MAX_CELL = 2000  # characters shown in one cell
MAX_DOC_BLOCKS = 3000
MAX_BYTES = 25_000_000  # a larger workbook is downloaded, not drawn
MAX_UNZIPPED = 150_000_000  # nor one that unzips to more than this: openpyxl would need gigabytes to load it
TOO_BIG = "This file is too big to show here. Download it to open it."

SHEET_SUFFIXES = {".xlsx", ".xlsm", ".xltx", ".xltm", ".xls", ".csv", ".tsv"}
DOC_SUFFIXES = {".docx"}


def kind(filename: str) -> str:
    """"sheet", "doc" or "" for a file this module shows."""
    suffix = Path(filename or "").suffix.lower()
    if suffix in SHEET_SUFFIXES:
        return "sheet"
    if suffix in DOC_SUFFIXES:
        return "doc"
    return ""


def preview(data: bytes, filename: str, *, rows: int = MAX_ROWS, cols: int = MAX_COLS) -> dict:
    """{"kind": "sheet", "sheets": [...]} or {"kind": "doc", "blocks": [...]}. A workbook that can't be drawn says
    why in a sentence instead: {"kind": "sheet", "sheets": [], "note": "..."}."""
    which = kind(filename)
    if which == "doc":
        if data[:2] == b"PK" and _unzipped_size(data) > MAX_UNZIPPED:
            return {"kind": "doc", "blocks": [], "note": TOO_BIG}
        return {"kind": "doc", "blocks": docx_blocks(data)}
    if which != "sheet":
        raise ValueError("This kind of file has no preview.")
    rows, cols = max(1, min(rows, MAX_ROWS)), max(1, min(cols, MAX_COLS))
    suffix = Path(filename).suffix.lower()
    if len(data) > MAX_BYTES or (data[:2] == b"PK" and _unzipped_size(data) > MAX_UNZIPPED):
        return {"kind": "sheet", "sheets": [], "note": TOO_BIG}
    if data[:2] == b"PK":
        sheets = _xlsx(data, rows, cols)
    elif data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        sheets = _xls(data, rows, cols)
    elif suffix == ".xls" and _MARKUP.match(data[:512]):
        # Many reporting tools save a web page or an XML spreadsheet and call it .xls; Excel opens it, this doesn't.
        note = "This .xls file is really a web page or XML file saved by another program, so it isn't drawn here. Download it to open it in Excel, or see its Text tab."
        return {"kind": "sheet", "sheets": [], "note": note}
    else:
        sheets = [_csv(data, filename, rows, cols, tab=suffix == ".tsv")]
    return {"kind": "sheet", "sheets": sheets}


_MARKUP = re.compile(rb"^(?:\xef\xbb\xbf|\xff\xfe|\xfe\xff)?[\s\x00]*<")


def _unzipped_size(data: bytes) -> int:
    """What the workbook's parts add up to unzipped, as the zip says, or 0 when it isn't a zip that opens."""
    import zipfile

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            return sum(info.file_size for info in archive.infolist())
    except (zipfile.BadZipFile, OSError, ValueError):
        return 0


# Excel workbooks ---------------------------------------------------------------------------------------


def _xlsx(data: bytes, max_rows: int, max_cols: int) -> list[dict]:
    from openpyxl import load_workbook

    values = load_workbook(io.BytesIO(data), data_only=True)
    formulas = load_workbook(io.BytesIO(data), data_only=False)
    sheets = []
    try:
        for ws in formulas.worksheets[:MAX_SHEETS]:
            sheets.append(_sheet(ws, values[ws.title], max_rows, max_cols))
    finally:
        values.close()
        formulas.close()
    return sheets


def _sheet(ws, vs, max_rows: int, max_cols: int) -> dict:
    from openpyxl.utils import get_column_letter

    # Only the cells the sheet holds are looked at: asking openpyxl for any other cell makes it, and a sheet
    # formatted down to row 1,048,576 would make a million of them.
    held = {key: cell for key, cell in ws._cells.items() if cell.value is not None}
    values = vs._cells
    merged: dict[tuple[int, int], tuple[int, int]] = {}
    covered: set[tuple[int, int]] = set()
    for rng in ws.merged_cells.ranges:
        merged[(rng.min_row, rng.min_col)] = (rng.max_row, rng.max_col)
        for r in range(rng.min_row, min(rng.max_row, rng.min_row + max_rows) + 1):
            for c in range(rng.min_col, min(rng.max_col, max_cols) + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    covered.add((r, c))
    used = set(held) | set(merged)
    last_row = max((r for r, _c in used), default=0)
    last_col = max((c for _r, c in used), default=0)
    shown_cols = min(last_col, max_cols)
    data_rows = sorted({r for r, _c in used})
    hidden_rows = {r for r, dim in ws.row_dimensions.items() if dim.hidden}
    columns = _columns(ws, shown_cols)
    hidden_cols = {c for c, (_width, hidden) in columns.items() if hidden}
    out_rows = []
    shown_last = 0
    for r in range(1, last_row + 1):
        if len(out_rows) >= max_rows:
            break
        if r in hidden_rows:
            continue
        shown_last = r
        cells = []
        for c in range(1, shown_cols + 1):
            if (r, c) in covered:
                continue
            span = merged.get((r, c))
            cell = held.get((r, c)) or (ws._cells.get((r, c)) if span else None)
            if cell is None:
                continue
            raw = cell.value
            value_cell = values.get((r, c))
            value = value_cell.value if value_cell is not None else None
            formula = raw if isinstance(raw, str) and raw.startswith("=") else ""
            if not formula and type(raw).__name__ == "ArrayFormula":
                formula = str(getattr(raw, "text", "") or "")
            shown_value = value if formula else raw
            text, red = display(shown_value, cell.number_format)
            entry: dict = {"c": c, "v": text[:MAX_CELL]}
            if formula:
                entry["f"] = formula[:500]
                if shown_value is None:
                    # Saved without its result (a file written by a program, not Excel): the formula, marked as such.
                    entry["v"], entry["nocalc"] = formula[:MAX_CELL], True
            if isinstance(shown_value, (int, float, Decimal, datetime, date, time, timedelta)) and not isinstance(shown_value, bool):
                entry["num"] = True  # numbers and dates sit at the right of the cell, as in Excel
            if red:
                entry["red"] = True
            if span:
                # The grid leaves hidden rows and columns out, so a merge spans only the ones it shows.
                entry["rs"] = _shown(r, span[0], hidden_rows, max_rows)
                entry["cs"] = _shown(c, min(span[1], shown_cols), hidden_cols, shown_cols)
            entry.update(_style(cell))
            cells.append(entry)
        # An empty row still takes its place in the grid, so row numbers read as in Excel.
        out_rows.append({"n": r, "cells": cells})
    while out_rows and not out_rows[-1]["cells"]:
        out_rows.pop()
    widths = []
    for c in range(1, shown_cols + 1):
        width, hidden = columns.get(c, (None, False))
        widths.append({"n": c, "letter": get_column_letter(c), "w": _px(width), "hidden": hidden})
    freeze = ws.freeze_panes if isinstance(ws.freeze_panes, str) else ""
    return {
        "name": ws.title,
        "hidden": ws.sheet_state != "visible",
        "cols": widths,
        "rows": out_rows,
        "more_rows": sum(1 for r in data_rows if r > shown_last),
        "more_cols": max(0, last_col - shown_cols),
        "freeze": freeze,
    }


def _columns(ws, shown_cols: int) -> dict[int, tuple[float | None, bool]]:
    """Each column's width and whether it's hidden. One <col> can cover a range (Excel's "Hide" on B:D, or a
    collapsed group), and openpyxl files it under the range's first letter only, so every column in it is set."""
    from openpyxl.utils import column_index_from_string

    out: dict[int, tuple[float | None, bool]] = {}
    for letter, dim in ws.column_dimensions.items():
        try:
            first = int(dim.min or column_index_from_string(letter))
            last = int(dim.max or first)
        except (TypeError, ValueError):
            continue
        for c in range(max(1, first), min(last, shown_cols) + 1):
            out[c] = (dim.width or None, bool(dim.hidden))
    return out


def _shown(start: int, end: int, hidden: set[int], limit: int) -> int:
    """How many of the rows (or columns) from ``start`` to ``end`` the grid shows, counting at most ``limit``."""
    count = 0
    for n in range(start, end + 1):
        if n not in hidden:
            count += 1
            if count >= limit:
                break
    return max(1, count)


def _px(width) -> int:
    """Excel's column width (characters of the default font) in pixels, as Excel draws it."""
    if not width:
        return 64
    return max(16, min(480, int(round(float(width) * 7 + 5))))


_THEMELESS = re.compile(r"^[0-9A-Fa-f]{8}$")


def _rgb(color) -> str:
    try:
        if color is None or getattr(color, "type", "") != "rgb":
            return ""
        value = str(color.rgb or "")
    except Exception:
        return ""
    if not _THEMELESS.match(value):
        return ""
    return "#" + value[2:].lower()


def _style(cell) -> dict:
    out: dict = {}
    font = getattr(cell, "font", None)
    if font is not None:
        if font.b:
            out["b"] = True
        if font.i:
            out["i"] = True
        if font.u and font.u != "none":
            out["u"] = True
        fg = _rgb(font.color)
        if fg and fg not in ("#000000",):
            out["fg"] = fg
    fill = getattr(cell, "fill", None)
    if fill is not None and getattr(fill, "fill_type", None) == "solid":
        bg = _rgb(fill.fgColor) or _rgb(fill.start_color)
        if bg and bg not in ("#ffffff",):
            out["bg"] = bg
    alignment = getattr(cell, "alignment", None)
    if alignment is not None:
        if alignment.horizontal in ("center", "right", "left", "centerContinuous"):
            out["al"] = "center" if alignment.horizontal == "centerContinuous" else alignment.horizontal
        if alignment.wrap_text:
            out["wrap"] = True
        if alignment.indent:
            out["ind"] = int(alignment.indent)
    border = getattr(cell, "border", None)
    if border is not None:
        lines = "".join(side[0] for side in ("top", "bottom", "left", "right") if getattr(getattr(border, side, None), "style", None))
        if lines:
            out["bd"] = lines
    return out


# Excel 97-2003 ----------------------------------------------------------------------------------------


def _xls(data: bytes, max_rows: int, max_cols: int) -> list[dict]:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=data, formatting_info=True, on_demand=True)
    except Exception:  # noqa: BLE001 - read again without formats
        # xlrd reads formats only from some workbooks (not an Excel 2.x one, say): those show as General.
        book = xlrd.open_workbook(file_contents=data, formatting_info=False, on_demand=True)
    sheets = []
    try:
        for index in range(min(book.nsheets, MAX_SHEETS)):
            sheet = book.sheet_by_index(index)
            shown_cols = min(sheet.ncols, max_cols)
            out_rows = []
            for r in range(min(sheet.nrows, max_rows)):
                cells = []
                for c in range(shown_cols):
                    cell = sheet.cell(r, c)
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    value = cell.value
                    fmt = _xls_format(book, cell)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        if not _is_date_format(_sections(fmt)[0]):
                            # The format wasn't read: a time of day, or a date, as Excel shows one by default.
                            fmt = "h:mm:ss AM/PM" if 0 <= value < 1 else "m/d/yyyy" if float(value).is_integer() else "m/d/yyyy h:mm"
                        if getattr(book, "datemode", 0) == 1 and not re.search(r"\[[hms]", fmt.lower()):
                            value += 1462  # the 1904 date system counts from 1/1/1904
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = xlrd.error_text_from_code.get(value, "#ERROR")
                    text, red = display(value, fmt)
                    entry = {"c": c + 1, "v": text[:MAX_CELL]}
                    if isinstance(value, float) and cell.ctype in (xlrd.XL_CELL_NUMBER, xlrd.XL_CELL_DATE):
                        entry["num"] = True
                    if red:
                        entry["red"] = True
                    cells.append(entry)
                out_rows.append({"n": r + 1, "cells": cells})
            while out_rows and not out_rows[-1]["cells"]:
                out_rows.pop()
            sheets.append(_plain(sheet.name, out_rows, shown_cols, max(0, sheet.nrows - max_rows), max(0, sheet.ncols - shown_cols), hidden=bool(getattr(sheet, "visibility", 0))))
    finally:
        book.release_resources()
    return sheets


def _xls_format(book, cell) -> str:
    """The cell's number format from its XF record, or "General" when the workbook was read without formats."""
    try:
        return book.format_map[book.xf_list[cell.xf_index].format_key].format_str or "General"
    except (AttributeError, IndexError, KeyError, TypeError):
        return "General"


def _plain(name: str, rows: list[dict], ncols: int, more_rows: int, more_cols: int, *, hidden: bool = False) -> dict:
    from openpyxl.utils import get_column_letter

    return {
        "name": name,
        "hidden": hidden,
        "cols": [{"n": c, "letter": get_column_letter(c), "w": 96, "hidden": False} for c in range(1, ncols + 1)],
        "rows": rows,
        "more_rows": more_rows,
        "more_cols": more_cols,
        "freeze": "",
    }


# CSV --------------------------------------------------------------------------------------------------


def _csv(data: bytes, filename: str, max_rows: int, max_cols: int, *, tab: bool = False) -> dict:
    from controller_inbox.documents import decode_text

    text = decode_text(data)
    if tab:
        dialect: type[csv.Dialect] | csv.Dialect = csv.excel_tab
    else:
        try:
            dialect = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
    # Excel's CSV import quotes only with '"' and keeps the spaces after a comma; the sniffer can guess "'" (any
    # text with apostrophes) and merge fields that way.
    reader = csv.reader(io.StringIO(text), dialect, quotechar='"', doublequote=True, skipinitialspace=False)
    rows = []
    widest = 0
    more = 0
    for number, record in enumerate(reader, start=1):
        if number > max_rows:
            more += 1
            continue
        widest = max(widest, len(record))
        cells = []
        for c, value in enumerate(record[:max_cols], start=1):
            if value == "":
                continue
            entry = {"c": c, "v": value[:MAX_CELL]}
            if _looks_numeric(value):
                entry["num"] = True
            cells.append(entry)
        rows.append({"n": number, "cells": cells})
    shown = min(widest, max_cols)
    return _plain(Path(filename).stem or "Sheet1", rows, shown, more, max(0, widest - shown))


_NUMERIC = re.compile(r"^\s*[-+(]?\s*[$€£¥]?\s*[\d,]*\.?\d+\s*%?\)?\s*$")


def _looks_numeric(value: str) -> bool:
    return bool(_NUMERIC.match(value))


# How Excel shows a value ------------------------------------------------------------------------------

_COLOR = re.compile(r"\[(black|blue|cyan|green|magenta|red|white|yellow|color\s*\d+)\]", re.IGNORECASE)
_CURRENCY = re.compile(r"\[\$([^\]\-]*)(?:-[^\]]*)?\]")
_CONDITION = re.compile(r"\[[<>=][^\]]*\]")
_DATE_TOKENS = re.compile(r'(?i)"[^"]*"|\\.|_.|\*.|\[h+\]|\[m+\]|\[s+\]|yyyy|yy|mmmmm|mmmm|mmm|mm|m|dddd|ddd|dd|d|hh|h|ss|s|am/pm|a/p|\.0+|[^"\\]')
_MONTHS = ("January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December")
_DAYS = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday")
_EPOCH = datetime(1899, 12, 30)


def display(value, number_format: str | None) -> tuple[str, bool]:
    """``value`` as Excel shows it with ``number_format``, and whether Excel shows it in red."""
    fmt = number_format if isinstance(number_format, str) else "General"
    if value is None:
        return "", False
    if isinstance(value, bool):
        return ("TRUE" if value else "FALSE"), False
    if isinstance(value, str):
        sections = _sections(fmt)
        # Text uses the fourth section, or the only one when it has an "@" in it.
        text_section = sections[3] if len(sections) >= 4 else sections[0] if len(sections) == 1 else ""
        if any(kind == "code" and ch == "@" for kind, ch in _tokens(text_section)):
            return "".join(value if kind == "code" and ch == "@" else ch for kind, ch in _tokens(text_section)), False
        return value, False
    if isinstance(value, time):
        value = datetime.combine(date(1899, 12, 30), value)
    if isinstance(value, (datetime, date)):
        if not isinstance(value, datetime):
            value = datetime.combine(value, time(0))
        return _date(value, fmt), False
    if isinstance(value, timedelta):
        value = value.total_seconds() / 86400
    if isinstance(value, Decimal):
        value = float(value)
    if not isinstance(value, (int, float)):
        return str(value), False
    if isinstance(value, float) and not math.isfinite(value):
        return "#NUM!", False
    if fmt.strip().lower() in ("general", ""):
        return _general(value), False
    sections = _sections(fmt)
    picked = _pick(value, sections)
    if picked is None:
        # No section takes the number (it fails every condition): Excel fills the cell with hashes.
        return "####", False
    section, signed = picked
    if _is_date_format(section):
        try:
            return _date(_EPOCH + timedelta(days=float(value)), section), False
        except (OverflowError, ValueError):
            return _general(value), False
    colour = _COLOR.search(section)
    red = bool(colour) and "red" in colour.group(1).lower()
    text = _number(abs(value), section)
    if signed and value < 0 and text and any(ch.isdigit() and ch != "0" for ch in text):
        text = "-" + text
    return text, red


_CONDITION_PARTS = re.compile(r"\[(<=|>=|<>|<|>|=)\s*([-+]?(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?)\]")


def _condition(section: str):
    match = _CONDITION_PARTS.search(section)
    return (match.group(1), float(match.group(2))) if match else None


def _meets(value: float, condition) -> bool:
    op, limit = condition
    return {"<": value < limit, "<=": value <= limit, ">": value > limit, ">=": value >= limit, "=": value == limit, "<>": value != limit}[op]


def _pick(value: float, sections: list[str]) -> tuple[str, bool] | None:
    """The section Excel formats ``value`` with, and whether the number keeps its minus sign there. Without
    conditions: positive;negative;zero;text. With a condition on the first (and maybe the second) section, the
    first section whose condition holds, else the next section without one."""
    first = _condition(sections[0])
    second = _condition(sections[1]) if len(sections) > 1 else None
    if first or second:
        if first and _meets(value, first):
            return sections[0], True
        if second and _meets(value, second):
            return sections[1], True
        if not first:
            return sections[0], True
        if len(sections) > 1 and not second:
            return sections[1], True
        if len(sections) > 2:
            return sections[2], True
        return None
    if value < 0 and len(sections) >= 2:
        return sections[1], False
    if value == 0 and len(sections) >= 3:
        return sections[2], False
    return sections[0], True


def _sections(fmt: str) -> list[str]:
    out, current, quoted, escaped = [], [], False, False
    for ch in fmt:
        if escaped:
            current.append(ch)
            escaped = False
        elif ch == "\\":
            current.append(ch)
            escaped = True
        elif ch == '"':
            current.append(ch)
            quoted = not quoted
        elif ch == ";" and not quoted:
            out.append("".join(current))
            current = []
        else:
            current.append(ch)
    out.append("".join(current))
    return out


def _general(value) -> str:
    """Excel's General format: whole numbers as they are (up to 15 digits, as in a column wide enough for them),
    others in at most 11 characters, and very small or large ones in scientific form with an uppercase E
    (1.234E-05)."""
    if value == 0:
        return "0"
    if isinstance(value, int) and abs(value) < 10**15 or isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
        return str(int(value))
    sign, value = ("-" if value < 0 else ""), abs(float(value))
    power = math.floor(math.log10(value))
    if -5 < power < 11:
        text = f"{value:.{max(0, 10 - max(1, power + 1))}f}"
        if "." in text:
            text = text.rstrip("0").rstrip(".")
        if len(text) <= 11:
            return sign + text
    mantissa, _, exponent = f"{value:.5E}".partition("E")
    mantissa = mantissa.rstrip("0").rstrip(".") if "." in mantissa else mantissa
    return f"{sign}{mantissa}E{exponent}"


def _tokens(section: str) -> list[tuple[str, str]]:
    """The section one character at a time, each marked "code" (a placeholder, a separator or a sign Excel
    reads) or "lit" (quoted or escaped text, and the space a "_x" leaves), with colours, conditions and fill
    characters taken out and a currency like [$€-407] turned into its symbol."""
    section = _CURRENCY.sub(lambda m: '"' + m.group(1).replace('"', "") + '"', section)
    section = _CONDITION.sub("", _COLOR.sub("", section))
    out: list[tuple[str, str]] = []
    i = 0
    while i < len(section):
        ch = section[i]
        if ch == '"':
            end = section.find('"', i + 1)
            end = len(section) if end < 0 else end
            out.extend(("lit", c) for c in section[i + 1 : end])
            i = end + 1
        elif ch == "\\" and i + 1 < len(section):
            out.append(("lit", section[i + 1]))
            i += 2
        elif ch == "_" and i + 1 < len(section):
            out.append(("lit", " "))
            i += 2
        elif ch == "*" and i + 1 < len(section):
            i += 2
        elif ch == "[" and "]" in section[i:]:
            i = section.index("]", i) + 1  # a locale or other bracketed code ([$-409], [DBNum1])
        else:
            out.append(("code", ch))
            i += 1
    return out


_DIGITS = "0#?"


def _fill_whole(digits: str, positions: list[int], code: str, fills: dict[int, str]) -> None:
    """Puts the digits of a whole number into its placeholders from the right, as Excel does across literal text
    ("000-00-0000"); a placeholder left over shows 0, a space or nothing, and digits left over go in the first."""
    left = list(digits)
    for pos in reversed(positions):
        if left:
            fills[pos] = left.pop()
        else:
            fills[pos] = {"0": "0", "?": " "}.get(code[pos], "")
    if left and positions:
        fills[positions[0]] = "".join(left) + fills[positions[0]]


def _fill_decimals(digits: str, positions: list[int], code: str, fills: dict[int, str]) -> None:
    """The digits after the point from the left; past the last one that isn't 0, "#" shows nothing and "?" a
    space."""
    last = max((k for k, d in enumerate(digits) if d != "0"), default=-1)
    for k, pos in enumerate(positions):
        digit = digits[k] if k < len(digits) else "0"
        fills[pos] = digit if k <= last or code[pos] == "0" else (" " if code[pos] == "?" else "")


def _number(value: float, section: str) -> str:
    """A number that isn't negative, laid out in one section of a number format."""
    tokens = _tokens(section)
    # One character per token; text that was quoted or escaped is masked so nothing in it is read as a placeholder.
    code = "".join(ch if kind == "code" else "\0" for kind, ch in tokens)
    general = re.search("(?i)general", code)
    places = [i for i, ch in enumerate(code) if ch in _DIGITS]
    if general or not places:
        # No digits of its own: "General" or "@" show the number as General does; otherwise only the text shows.
        shown = _general(value * 100 ** code.count("%")) if (general or "@" in code) else ""
        out = []
        for i, (_kind, ch) in enumerate(tokens):
            if general and general.start() <= i < general.end():
                out.append(shown if i == general.start() else "")
            else:
                out.append(shown if code[i] == "@" else ch)
        return "".join(out).strip()
    value *= 100 ** code.count("%")
    fills: dict[int, str] = {}
    slash = _fraction_slash(code, places)
    exponent = next((i for i in range(1, len(code) - 1) if code[i] in "Ee" and code[i + 1] in "+-" and any(p < i for p in places)), -1)
    if slash >= 0:
        _fraction(value, code, places, slash, fills)
    elif exponent >= 0:
        _scientific(value, code, places, exponent, fills)
    else:
        _decimal(value, code, places, fills)
    return "".join(fills.get(i, ch) for i, (_kind, ch) in enumerate(tokens)).strip()


def _decimal(value: float, code: str, places: list[int], fills: dict[int, str]) -> None:
    dot = code.find(".")
    dot = len(code) if dot < 0 else dot
    whole_at = [p for p in places if p < dot]
    decimal_at = [p for p in places if p > dot]
    grouped, scale = False, 0
    for i, ch in enumerate(code):
        if ch != ",":
            continue
        if whole_at and whole_at[0] < i < whole_at[-1]:
            grouped, fills[i] = True, ""
            continue
        # Commas after the last digit placeholder scale by a thousand each ("#,##0," shows thousands).
        j = i - 1
        while j >= 0 and code[j] == ",":
            j -= 1
        if j >= 0 and (code[j] in _DIGITS or code[j] == "."):
            scale, fills[i] = scale + 1, ""
    rounded = Decimal(repr(float(value))).scaleb(-3 * scale).quantize(Decimal(1).scaleb(-len(decimal_at)), rounding=ROUND_HALF_UP)
    whole, _, fraction = f"{rounded:f}".partition(".")
    whole = "" if whole.lstrip("0") == "" else whole.lstrip("0")
    if whole_at:
        _fill_whole(whole, whole_at, code, fills)
        if grouped:
            joined = "".join(fills[p] for p in whole_at)
            digits = joined.lstrip(" ")
            groups = [digits[max(0, k - 3) : k] for k in range(len(digits), 0, -3)][::-1]
            fills[whole_at[0]] = " " * (len(joined) - len(digits)) + ",".join(groups)
            for p in whole_at[1:]:
                fills[p] = ""
    elif whole and dot < len(code):
        fills[dot] = whole + "."
    _fill_decimals(fraction, decimal_at, code, fills)


def _scientific(value: float, code: str, places: list[int], at: int, fills: dict[int, str]) -> None:
    dot = code.find(".", 0, at)
    dot = at if dot < 0 else dot
    whole_at = [p for p in places if p < dot]
    decimal_at = [p for p in places if dot < p < at]
    power_at = [p for p in places if p > at + 1]
    width = len(whole_at)  # with none (".00E+00") the mantissa is below 1
    # "##0.0E+0" is engineering notation: the exponent is a multiple of the number of whole placeholders.
    engineering = width > 1 and "#" in "".join(code[p] for p in whole_at)
    number = Decimal(repr(float(value)))
    step = width if engineering else 1
    power = number.adjusted() if number else 0
    if number:
        power = power - power % width if engineering else power - (width - 1)
    quantum = Decimal(1).scaleb(-len(decimal_at))
    mantissa = number.scaleb(-power).quantize(quantum, rounding=ROUND_HALF_UP)
    if number and mantissa >= 10**width:
        power += step
        mantissa = number.scaleb(-power).quantize(quantum, rounding=ROUND_HALF_UP)
    whole, _, fraction = f"{mantissa:f}".partition(".")
    _fill_whole("" if whole.lstrip("0") == "" else whole.lstrip("0"), whole_at, code, fills)
    _fill_decimals(fraction, decimal_at, code, fills)
    fills[at] = "E"
    fills[at + 1] = "-" if power < 0 else ("+" if code[at + 1] == "+" else "")
    _fill_whole(str(abs(power)) if power else "", power_at, code, fills)


def _fraction_slash(code: str, places: list[int]) -> int:
    """Where a fraction's "/" is ("# ?/?", "?/8"), or -1: a slash right after a placeholder and before one or a
    digit."""
    for i, ch in enumerate(code):
        if ch == "/" and i > 0 and code[i - 1] in _DIGITS and i + 1 < len(code) and (code[i + 1] in _DIGITS or code[i + 1].isdigit()):
            return i
    return -1


def _fraction(value: float, code: str, places: list[int], slash: int, fills: dict[int, str]) -> None:
    from fractions import Fraction

    start = slash
    while start > 0 and code[start - 1] in _DIGITS:
        start -= 1
    numerator_at = list(range(start, slash))
    end = slash + 1
    while end < len(code) and (code[end] in _DIGITS or code[end].isdigit()):
        end += 1
    denominator_at = list(range(slash + 1, end))
    whole_at = [p for p in places if p < start]
    fixed = "".join(code[p] for p in denominator_at)
    if whole_at:
        whole, part = int(value), value - int(value)
    else:
        whole, part = 0, value
    if any(ch in "123456789" for ch in fixed) and fixed.isdigit():
        bottom = int(fixed)
        top = int(Decimal(repr(part * bottom)).quantize(Decimal(1), rounding=ROUND_HALF_UP))
    else:
        best = Fraction(repr(part)).limit_denominator(10 ** len(denominator_at) - 1)
        top, bottom = best.numerator, best.denominator
    if whole_at and top and top >= bottom:
        whole, top = whole + top // bottom, top % bottom
    if whole_at:
        _fill_whole(str(whole) if whole else ("" if top else "0"), whole_at, code, fills)
    if whole_at and top == 0:
        # A whole number: the fraction's room is left blank.
        for p in numerator_at + [slash] + denominator_at:
            fills[p] = " "
        return
    _fill_whole(str(top), numerator_at, code, fills)
    if not fixed.isdigit() or not any(ch in "123456789" for ch in fixed):
        text = str(bottom)
        pad = " " if "?" in fixed else ""
        fills[denominator_at[0]] = text + pad * max(0, len(denominator_at) - len(text))
        for p in denominator_at[1:]:
            fills[p] = ""


def _is_date_format(section: str) -> bool:
    bare = _CONDITION.sub("", _COLOR.sub("", _CURRENCY.sub("", section)))
    bare = re.sub(r'"[^"]*"|\\.|_.|\*.', "", bare).lower()
    # Bracketed codes other than elapsed time ([h], [mm], [ss]) say nothing about dates; ".00" is fractions of a second.
    bare = re.sub(r"\.0+", "", re.sub(r"\[(?![hms]+\])[^\]]*\]", "", bare))
    return bool(re.search(r"[dy]|m{1,5}|h|ss", bare)) and not re.search(r"[0#?]", bare)


def _date(value: datetime, fmt: str) -> str:
    section = _sections(fmt)[0] if fmt else ""
    if not _is_date_format(section):
        # A date in a General or number cell: shown as Excel does by default.
        if value.time() == time(0):
            return f"{value.month}/{value.day}/{value.year}"
        return f"{value.month}/{value.day}/{value.year} {value.hour}:{value.minute:02d}"
    section = _CONDITION.sub("", _COLOR.sub("", _CURRENCY.sub(lambda m: m.group(1), section)))
    tokens = [m.group(0) for m in _DATE_TOKENS.finditer(section)]
    if not any(t.startswith(".0") for t in tokens):
        # Without fractions of a second Excel rounds to the nearest second (29:59:59.6 shows as 30:00:00).
        try:
            value = (value + timedelta(microseconds=500_000)).replace(microsecond=0)
        except OverflowError:
            pass
    twelve = any(t.lower() in ("am/pm", "a/p") for t in tokens)
    # Elapsed time ([h]:mm, [mm]:ss, [ss]) counts from Excel's day 0 instead of wrapping at a day or an hour.
    elapsed = int((value - _EPOCH).total_seconds())
    out = []
    for index, token in enumerate(tokens):
        low = token.lower()
        prev = next((t.lower().lstrip("[") for t in reversed(tokens[:index]) if t.strip() and t[0] not in '"\\_*' and t.lower() not in ("/", "-", ".", ",", ":", " ")), "")
        nxt = next((t.lower().lstrip("[") for t in tokens[index + 1 :] if t.strip() and t[0] not in '"\\_*' and t.lower() not in ("/", "-", ".", ",", ":", " ")), "")
        minute = low in ("m", "mm") and (prev.startswith("h") or nxt.startswith("s"))
        hour = value.hour % 12 or 12 if twelve else value.hour
        if low.startswith("[") and len(low) > 2:
            unit = {"h": 3600, "m": 60, "s": 1}[low[1]]
            out.append(f"{elapsed // unit:0{len(low) - 2}d}")
        elif token.startswith("_"):
            out.append(" ")
        elif token.startswith("*"):
            continue
        elif token.startswith('"'):
            out.append(token[1:-1])
        elif token.startswith("\\"):
            out.append(token[1:])
        elif low == "yyyy":
            out.append(f"{value.year:04d}")
        elif low == "yy":
            out.append(f"{value.year % 100:02d}")
        elif minute:
            out.append(f"{value.minute:02d}" if low == "mm" else str(value.minute))
        elif low == "mmmmm":
            out.append(_MONTHS[value.month - 1][0])
        elif low == "mmmm":
            out.append(_MONTHS[value.month - 1])
        elif low == "mmm":
            out.append(_MONTHS[value.month - 1][:3])
        elif low == "mm":
            out.append(f"{value.month:02d}")
        elif low == "m":
            out.append(str(value.month))
        elif low == "dddd":
            out.append(_DAYS[value.weekday()])
        elif low == "ddd":
            out.append(_DAYS[value.weekday()][:3])
        elif low == "dd":
            out.append(f"{value.day:02d}")
        elif low == "d":
            out.append(str(value.day))
        elif low == "hh":
            out.append(f"{hour:02d}")
        elif low == "h":
            out.append(str(hour))
        elif low == "ss":
            out.append(f"{value.second:02d}")
        elif low == "s":
            out.append(str(value.second))
        elif low == "am/pm":
            out.append("AM" if value.hour < 12 else "PM")
        elif low == "a/p":
            out.append("A" if value.hour < 12 else "P")
        elif low.startswith(".0"):
            out.append("." + f"{value.microsecond:06d}"[: len(token) - 1])
        elif token in ("_", "*"):
            continue
        else:
            out.append(token)
    return "".join(out).strip()


# Word -------------------------------------------------------------------------------------------------


def docx_blocks(data: bytes) -> list[dict]:
    """The document's headings, paragraphs, list items and tables in order: {"t": "h1".."h4" | "p" | "li" |
    "table", ...}. Tracked changes are shown accepted (inserted and moved-to text in, deleted and moved-from text
    out); text and table cells only, no pictures. A table cell is its text, or {"v": text, "cs": n} when it spans
    n columns; a numbered list item says its number ("num": "1.")."""
    from docx import Document
    from docx.text.paragraph import Paragraph

    document = Document(io.BytesIO(data))
    numbers = _Numbering(document)
    blocks: list[dict] = []
    for child in _block_elements(document.element.body):
        if len(blocks) >= MAX_DOC_BLOCKS:
            blocks.append({"t": "p", "text": "[The rest of the document isn't shown here; it's in the Text tab.]"})
            break
        if _tag(child) == "p":
            paragraph = Paragraph(child, document)
            text = _paragraph_text(child)
            if not text.strip():
                continue
            style = (paragraph.style.name if paragraph.style is not None else "") or ""
            level = re.match(r"(?i)heading (\d)", style)
            if style.lower() == "title":
                blocks.append({"t": "h1", "text": text})
            elif level:
                blocks.append({"t": f"h{min(4, int(level.group(1)) + 1)}", "text": text})
            elif "list" in style.lower() or child.find(".//{*}numPr") is not None:
                label = numbers.label(child, paragraph.style)
                blocks.append({"t": "li", "text": text, **({"num": label} if label else {})})
            else:
                runs = paragraph.runs
                bold = bool(runs) and all(run.bold for run in runs if run.text.strip())
                blocks.append({"t": "p", "text": text, **({"b": True} if bold else {})})
        else:
            grid = _table_rows(child)
            if grid:
                blocks.append({"t": "table", "rows": grid})
    return blocks


def _tag(node) -> str:
    """The element's name without its namespace, or "" for a comment or processing instruction."""
    return node.tag.rsplit("}", 1)[-1] if isinstance(node.tag, str) else ""


def _children(parent, *names: str):
    """``parent``'s child elements called one of ``names``, looking through content controls (w:sdt) and custom
    XML, which wrap paragraphs, tables, rows and cells without changing what they read."""
    for child in parent.iterchildren():
        tag = _tag(child)
        if tag in names:
            yield child
        elif tag == "sdt":
            for content in child.iterchildren():
                if _tag(content) == "sdtContent":
                    yield from _children(content, *names)
        elif tag == "customXml":
            yield from _children(child, *names)


def _block_elements(parent):
    """The paragraphs and tables in a body or a table cell, in order."""
    return _children(parent, "p", "tbl")


def _val(node, path: str):
    """The w:val of the first element at ``path`` under ``node``, or None."""
    found = node.find(path) if node is not None else None
    if found is None:
        return None
    return next((v for k, v in found.attrib.items() if k.rsplit("}", 1)[-1] == "val"), None)


def _table_rows(tbl) -> list[list]:
    grid = []
    for tr in list(_children(tbl, "tr"))[:500]:
        cells: list = []
        before = _val(tr, "{*}trPr/{*}gridBefore")
        if before and before.isdigit() and int(before) > 0:
            # The row starts further right than the others: an empty cell holds its place.
            cells.append({"v": "", "cs": int(before)} if int(before) > 1 else "")
        for tc in list(_children(tr, "tc"))[:40]:
            # A cell merged with the one above it is empty in Word's file; the text is in the top cell.
            merge = tc.find("{*}tcPr/{*}vMerge")
            text = "" if merge is not None and _val(tc, "{*}tcPr/{*}vMerge") in (None, "continue") else _cell_text(tc)
            span = _val(tc, "{*}tcPr/{*}gridSpan")
            span = int(span) if span and span.isdigit() else 1
            cells.append({"v": text, "cs": span} if span > 1 else text)
        grid.append(cells)
    return grid


def _cell_text(tc) -> str:
    """A table cell's paragraphs, one a line, with a table inside it read row by row (its cells split by tabs)."""
    lines = []
    for block in _block_elements(tc):
        if _tag(block) == "p":
            lines.append(_paragraph_text(block))
        else:
            for row in _table_rows(block):
                lines.append("\t".join(cell if isinstance(cell, str) else cell["v"] for cell in row))
    return "\n".join(lines).strip()[:MAX_CELL]


# Parts of a paragraph that aren't read: deleted and moved-away text (tracked changes, accepted), the second copy
# of a text box Word saves for older readers (mc:Fallback), and paragraph and run properties (whose w:tab elements
# are tab stops, not tabs).
_UNREAD = {"del", "moveFrom", "Fallback", "pPr", "rPr"}


def _paragraph_text(p) -> str:
    """The paragraph as it reads with tracked changes accepted: inserted text in, deleted text out."""
    parts: list[str] = []

    def walk(node) -> None:
        for child in node.iterchildren():
            tag = _tag(child)
            if not tag or tag in _UNREAD:
                continue
            if tag == "t":
                parts.append(child.text or "")
            elif tag == "tab":
                parts.append("\t")
            elif tag in ("br", "cr"):
                parts.append("\n")
            else:
                walk(child)

    walk(p)
    return "".join(parts)


_ROMAN = ((1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"), (90, "xc"), (50, "l"), (40, "xl"), (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i"))


def _count_text(n: int, kind: str) -> str:
    if kind in ("lowerLetter", "upperLetter"):
        text = chr(ord("a") + (n - 1) % 26) * ((n - 1) // 26 + 1)
        return text.upper() if kind == "upperLetter" else text
    if kind in ("lowerRoman", "upperRoman") and 0 < n < 4000:
        out = ""
        for value, letters in _ROMAN:
            while n >= value:
                out, n = out + letters, n - value
        return out.upper() if kind == "upperRoman" else out
    if kind == "decimalZero":
        return f"{n:02d}"
    return str(n)


class _Numbering:
    """The number Word shows before each numbered paragraph ("1.", "a)", "2.1"), worked out from the document's
    numbering definitions and counted in document order. Bullets and anything unreadable give no number."""

    def __init__(self, document) -> None:
        self.levels: dict[str, dict[int, tuple[str, str, int]]] = {}
        self.counts: dict[str, dict[int, int]] = {}
        try:
            numbering = document.part.numbering_part.element
        except Exception:  # noqa: BLE001 - a document without readable numbering shows bullets
            return
        abstract = {}
        for node in numbering.findall("{*}abstractNum"):
            key = next((v for k, v in node.attrib.items() if k.endswith("}abstractNumId")), None)
            abstract[key] = node
        for num in numbering.findall("{*}num"):
            num_id = next((v for k, v in num.attrib.items() if k.endswith("}numId")), None)
            definition = abstract.get(_val(num, "{*}abstractNumId"))
            if num_id is None or definition is None:
                continue
            levels = {}
            for lvl in definition.findall("{*}lvl"):
                ilvl = next((v for k, v in lvl.attrib.items() if k.endswith("}ilvl")), "0")
                start = _val(lvl, "{*}start")
                levels[int(ilvl) if ilvl.isdigit() else 0] = (_val(lvl, "{*}numFmt") or "decimal", _val(lvl, "{*}lvlText") or "", int(start) if start and start.isdigit() else 1)
            for override in num.findall("{*}lvlOverride"):
                ilvl = next((v for k, v in override.attrib.items() if k.endswith("}ilvl")), "0")
                start = _val(override, "{*}startOverride")
                if ilvl.isdigit() and int(ilvl) in levels and start and start.isdigit():
                    kind, text, _start = levels[int(ilvl)]
                    levels[int(ilvl)] = (kind, text, int(start))
            self.levels[num_id] = levels

    def label(self, p, style) -> str:
        num_pr = p.find("{*}pPr/{*}numPr")
        if num_pr is None and style is not None:
            num_pr = style.element.find("{*}pPr/{*}numPr")  # a list style ("List Number") numbers its paragraphs
        num_id = _val(num_pr, "{*}numId")
        ilvl = _val(num_pr, "{*}ilvl")
        ilvl = int(ilvl) if ilvl and ilvl.isdigit() else 0
        levels = self.levels.get(num_id or "")
        if not levels or ilvl not in levels or levels[ilvl][0] in ("bullet", "none"):
            return ""
        counts = self.counts.setdefault(num_id, {})
        counts[ilvl] = counts.get(ilvl, levels[ilvl][2] - 1) + 1
        for deeper in [k for k in counts if k > ilvl]:
            del counts[deeper]  # a new item restarts the levels under it

        def number(match: re.Match) -> str:
            level = int(match.group(1)) - 1
            kind, _text, start = levels.get(level, ("decimal", "", 1))
            return _count_text(counts.get(level, start), kind)

        return re.sub(r"%([1-9])", number, levels[ilvl][1]).strip()
