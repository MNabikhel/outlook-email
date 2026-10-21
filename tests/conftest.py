from __future__ import annotations

import io
import os
from datetime import datetime, timezone
from pathlib import Path

import pytest

from openpyxl import Workbook

from controller_inbox.config import Settings
from controller_inbox.demo import make_pdf
from controller_inbox.folder_mail import ingest_folder
from controller_inbox.pipeline import ingest_demo
from controller_inbox.store import Store


@pytest.fixture(autouse=True)
def _isolated_env(monkeypatch: pytest.MonkeyPatch) -> None:
    for key in list(os.environ):
        if key.startswith("CONTROLLER_INBOX_"):
            monkeypatch.delenv(key, raising=False)
    monkeypatch.setenv("CONTROLLER_INBOX_LLM", "false")


@pytest.fixture(autouse=True)
def _fresh_model_memory():
    """What a run learned about the loaded model, and the results kept to save time, must not leak between tests."""
    from controller_inbox import cost_codes, local_llm, memo, semantic, sheet_preview

    caches = (
        local_llm._status_cache,
        local_llm._reasoning_seen,
        local_llm._effort_rejected,
        local_llm._tools_rejected,
        local_llm._context_raised,
        local_llm._reader_failed,
        local_llm._pending_loads,
        local_llm._in_use,
        semantic._model_cache,
        semantic._vector_cache,
        semantic._question_cache,
        sheet_preview._cache,
        cost_codes._cache,
    )
    for cache in caches:
        cache.clear()
    memo.clear_all()
    yield
    for cache in caches:
        cache.clear()
    memo.clear_all()


@pytest.fixture(autouse=True)
def _no_table_finder_download(monkeypatch):
    """Tests never fetch the layout model from the internet: without it, tables are found as before."""
    from controller_inbox import table_finder

    def refuse(_target):
        raise OSError("tests don't download the table finder")

    monkeypatch.setattr(table_finder, "_download", refuse)
    monkeypatch.setattr(table_finder, "_session", None)
    monkeypatch.setattr(table_finder, "_last_try", 0.0)


AS_OF_NOW = datetime(2026, 9, 22, 12, 0, tzinfo=timezone.utc)


@pytest.fixture
def as_of_now() -> datetime:
    return AS_OF_NOW


def _test_settings(root: Path) -> Settings:
    # A model server nobody listens on: a test that forgets to fake a model call fails fast, here and in CI.
    return Settings(
        data_dir=root, inbox_dir=root / "inbox", timezone="America/New_York", llm_base_url="http://127.0.0.1:9/v1", _env_file=None
    )


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    return _test_settings(tmp_path)


@pytest.fixture
def store(settings: Settings) -> Store:
    return Store(settings.db_path)


# Where settings differ between tests without changing what reading the sample gives.
_PER_TEST = {"data_dir", "inbox_dir"}


@pytest.fixture(scope="session")
def _sample(tmp_path_factory) -> tuple[Settings, dict]:
    """The sample mailbox read once, for `loaded` to copy into each test.

    A session fixture is set up before a test's own fixtures, so nothing a test patches can reach it; the
    environment and the table finder are kept as the autouse fixtures above keep them for a test.
    """
    from controller_inbox import memo, table_finder

    root = tmp_path_factory.mktemp("sample")
    with pytest.MonkeyPatch.context() as patch:
        for key in list(os.environ):
            if key.startswith("CONTROLLER_INBOX_"):
                patch.delenv(key, raising=False)
        patch.setenv("CONTROLLER_INBOX_LLM", "false")

        def refuse(_target):
            raise OSError("tests don't download the table finder")

        patch.setattr(table_finder, "_download", refuse)
        patch.setattr(table_finder, "_session", None)
        patch.setattr(table_finder, "_last_try", 0.0)
        settings = _test_settings(root)
        ingest_demo(Store(settings.db_path), settings, now=AS_OF_NOW)
        memo.clear_all()
    return settings, settings.model_dump(exclude=_PER_TEST)


def _only_a_new_database(settings: Settings) -> bool:
    """Nothing but the empty database the store fixture made: no file dropped, no cost codes, no folder set up."""
    roots = [settings.data_dir, settings.inbox_dir, settings.cost_codes_folder]
    found = {p.name for root in roots if root.exists() for p in root.rglob("*")}
    return found <= {settings.db_path.name, settings.db_path.name + "-wal", settings.db_path.name + "-shm", "inbox"}


@pytest.fixture
def loaded(_sample, store: Store, settings: Settings, as_of_now: datetime) -> Store:
    """The sample mailbox read as of `as_of_now`. Copied from one reading per session when this test's settings
    and folders are what that reading had (it gives the same mail, ids and files); read afresh otherwise."""
    import shutil
    import sqlite3

    template, wanted = _sample
    if as_of_now != AS_OF_NOW or settings.model_dump(exclude=_PER_TEST) != wanted or not _only_a_new_database(settings):
        ingest_demo(store, settings, now=as_of_now)
        return store
    with sqlite3.connect(template.db_path) as source, sqlite3.connect(store.path) as target:
        source.backup(target)
    source.close()
    target.close()
    shutil.copytree(template.inbox_dir, settings.inbox_dir, dirs_exist_ok=True)
    return store


@pytest.fixture
def client(settings: Settings, loaded: Store):
    """The app on the sample mailbox, in this process. Modules that need another mailbox define their own."""
    from fastapi.testclient import TestClient

    from controller_inbox.web import create_app

    return TestClient(create_app(settings, loaded))


def _budget_xlsx() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Budget"
    sheet.append(["Line", "Q3", "Q4", "Change"])
    sheet.append(["Ads", 1000, 1500, "=C2-B2"])
    sheet.append(["Travel", 250, 300, "=C3-B3"])
    sheet.append(["Total", "=SUM(B2:B3)", "=SUM(C2:C3)", "=D2+D3"])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _memo_docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("Offsite plan (draft)", 1)
    doc.add_paragraph("The offsite moves to Lisbon on 14 November. Budget is capped at $42,000.")
    doc.add_paragraph("Open item: confirm the venue deposit with Finance.")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


@pytest.fixture
def mail(store, settings):
    """Three real .msg files: a budget from a colleague, a forwarded quote, and a bank-change scam with a file."""
    # Imported here, not at the top: building .msg files needs aspose-email-foss (the dev extra), and only the tests
    # that do should skip without it.
    from msgfactory import DOCX, PDF, XLSX, build_message, write_msg

    settings.trusted_domains = "taz.com"
    settings.ensure_data_dir()
    write_msg(
        settings.inbox_incoming / "budget.msg",
        "Q4 budget draft",
        "Hi, the Q4 budget draft is attached. Can you check the totals before Friday?",
        sender_name="Maya Chen",
        sender_email="maya@taz.com",
        attachments=[
            ("Q4 budget.xlsx", _budget_xlsx(), XLSX),
            ("Offsite memo.docx", _memo_docx(), DOCX),
        ],
    )
    quote = build_message(
        "Acme quote",
        "Our quote is attached.",
        sender_name="Acme Sales",
        sender_email="sales@acme.com",
        attachments=[("quote.pdf", make_pdf([["Acme quote Q-881", ["Item", "Amount"], ["Support plan", "$9,600.00"]]]), PDF)],
    )
    write_msg(
        settings.inbox_incoming / "fw quote.msg",
        "FW: Acme quote",
        "Is this in line with last year?",
        sender_name="Priya Raman",
        sender_email="priya@taz.com",
        forwarded=[("Acme quote.msg", quote)],
    )
    write_msg(
        settings.inbox_incoming / "scam.msg",
        "Updated remittance details",
        "Our bank details have changed. Please use the following account for all payments from today.",
        sender_name="Acme Billing",
        sender_email="billing@acme-payments.net",
        attachments=[("new bank letter.pdf", make_pdf([["Remit to account 5566778899 routing 021000021"]]), PDF)],
    )
    records = {email.subject: email for email in ingest_folder(store, settings)}
    return records


# ---------- A real browser (Playwright's Chromium), for the tests marked "browser" ----------


@pytest.hookimpl(tryfirst=True)
def pytest_collection_modifyitems(items):
    """Every test that uses the browser fixture is a browser test, so `-m browser` / `-m "not browser"` split them."""
    for item in items:
        if "browser" in getattr(item, "fixturenames", ()):
            item.add_marker(pytest.mark.browser)


@pytest.fixture(scope="module")
def browser():
    pytest.importorskip("playwright")
    from playwright.sync_api import sync_playwright

    from liveserver import chromium_path

    path = chromium_path()
    if path is None:
        pytest.skip("no Chromium for Playwright on this machine")
    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path=path)
        yield browser
        browser.close()


@pytest.fixture
def site(settings, loaded):
    """The sample mailbox served by a real CloseDesk on a free local port."""
    from controller_inbox import web
    from liveserver import serving

    with serving(web.create_app(settings, loaded)) as base:
        yield base


@pytest.fixture
def page(browser):
    """A fresh browser tab; the test fails if the page threw an error."""
    context = browser.new_context()
    context.set_default_timeout(15000)
    context.errors = []
    context.on("weberror", lambda error: context.errors.append(str(error.error)))
    yield context.new_page()
    assert context.errors == []
    context.close()


# ---------- OCR faked as reading both scans from Harbor Steel (ar@taz.com) ----------

# What OCR reads there: the name with a zero for its "o", and a table of three lines no outline is found for (only
# two columns line up).
HARBOR_LINES = [
    (60, 40, 330, 70, "Harb0r Steel LLC"),
    (60, 120, 260, 140, "Description"),
    (900, 120, 1040, 140, "Amount"),
    (60, 160, 300, 180, "Steel beams"),
    (900, 160, 1040, 180, "1,250.00"),
    (60, 200, 300, 220, "Bolts M12"),
    (900, 200, 1040, 220, "310.50"),
    (60, 240, 300, 260, "Delivery"),
    (900, 240, 1040, 260, "185.00"),
    (60, 400, 700, 420, "Thank you for your business."),
]


@pytest.fixture
def harbor_ocr(monkeypatch):
    """RapidOCR as the engine, reading HARBOR_LINES on every page (tests ask for it with usefixtures)."""
    from controller_inbox import ocr

    monkeypatch.setattr(ocr, "engine_name", lambda: "RapidOCR")
    monkeypatch.setattr(
        ocr,
        "line_boxes",
        lambda _png: [{"left": float(l), "top": float(t), "right": float(r), "bottom": float(b), "text": text, "score": 0.9}
                      for l, t, r, b, text in HARBOR_LINES],
    )
