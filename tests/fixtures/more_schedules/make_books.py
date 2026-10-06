"""Six more schedules the way a controller builds them in Excel, saved as PDF with LibreOffice.

Layouts that differ from the seven in ../excel_schedules:

    1. AR aging, Total before the buckets, two-row header, a credit balance in parentheses
    2. Accrued liabilities rollforward, sections with subtotals, negatives in parentheses
    3. Bank reconciliations, three accounts side by side, row labels as keys, plus a detail table
    4. Monthly revenue by customer, too wide for one page, print-title columns repeated on page 2
    5. Payroll register, department subtotals, dates like "Oct 15, 2026"
    6. Intercompany due to / due from matrix with row and column totals and a net-position block

Every figure is held in Python as a Decimal. Formulas are written where an Excel user would write them;
LibreOffice computes them when it saves the PDF, and the script checks the PDF text for the computed
figures. It also writes questions.py: questions about the schedules with their exact answers, as
regexes, used to measure how well CloseDesk answers questions about layouts it was not built on.

The PDFs beside this file are its output, used by tests/test_more_schedules.py and
tests/test_table_query.py. To make them again (needs openpyxl, LibreOffice's soffice and poppler's
pdftotext and pdfinfo on the PATH):

    python tests/fixtures/more_schedules/make_books.py tests/fixtures/more_schedules
"""

from __future__ import annotations

import os
import random
import re
import subprocess
import sys
import tempfile
from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(sys.argv[1] if len(sys.argv) > 1 else "books").resolve()
PDF = OUT
XLSX = OUT / "xlsx"
PROFILE = Path(tempfile.mkdtemp(prefix="lo-profile-"))
CENT = Decimal("0.01")
ZERO = Decimal("0.00")
COMPANY = "Cobalt Ridge Industries, Inc."

ACCT = '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)'
COMMA = '_(* #,##0.00_);_(* (#,##0.00);_(* "-"??_);_(@_)'
NUM = "#,##0.00_);(#,##0.00)"
BOLD = Font(bold=True)
THIN = Side(style="thin")
DOUBLE = Side(style="double")
MEDIUM = Side(style="medium")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
GRAY = PatternFill("solid", fgColor="D9D9D9")
LIGHT = PatternFill("solid", fgColor="F2F2F2")


def D(x) -> Decimal:
    """Exact Decimal for a literal written with two decimals (floats round-trip through str)."""
    return Decimal(str(x)).quantize(CENT)


def r2(x: Decimal) -> Decimal:
    """Excel ROUND(x, 2) for the values used here: half away from zero."""
    return x.quantize(CENT, rounding=ROUND_HALF_UP)


def show(v: Decimal) -> str:
    """How a comma/accounting format displays the figure (without the $)."""
    v = r2(v)
    if v == 0:
        return "-"
    return f"({-v:,.2f})" if v < 0 else f"{v:,.2f}"


def plain(v: Decimal) -> str:
    v = r2(v)
    return f"-{-v:,.2f}" if v < 0 else f"{v:,.2f}"


# ------------------------------------------------------------------------------ answer regexes
def _core(v: Decimal) -> str:
    v = r2(abs(v))
    whole = int(v)
    cents = int((v - whole) * 100)
    digits = f"{whole:,}".replace(",", ",?")
    tail = r"(?:\.00)?" if cents == 0 else rf"\.{cents:02d}"
    return rf"(?<![\d.,]){digits}{tail}(?!\d|[.,]\d)"


def money(v: Decimal) -> str:
    """The figure, with or without $ and thousands separators (and without .00 when whole)."""
    return r"\$?\s?" + _core(v)


def neg_money(v: Decimal, words: tuple[str, ...] = ("credit",)) -> str:
    """A negative figure: -1,234.56, (1,234.56), $(1,234.56), 1,234.56 CR, or a word like 'credit' next to it."""
    n = _core(v)
    alts = [
        rf"[-−]\s?\$?\s?{n}",
        rf"\(\s?\$?\s?{n}\s?\)",
        rf"\$\s?[-−]\s?{n}",
        rf"{n}\s?(?:CR|Cr|cr)\b",
    ]
    for w in words:
        alts.append(rf"(?i:{w})[^\d\n]{{0,40}}\$?\s?{n}")
        alts.append(rf"{n}[^\d\n]{{0,25}}(?i:{w})")
    return "(?:" + "|".join(alts) + ")"


def number(v: Decimal) -> str:
    """A plain quantity like hours: 232.50 may also be written 232.5."""
    core = _core(v)
    return re.sub(r"\\\.(\d)0\(", r"\\.\1(?:0)?(", core)


WORDS = {0: "zero|none|no", 1: "one", 2: "two", 3: "three", 4: "four", 5: "five", 6: "six", 7: "seven",
         8: "eight", 9: "nine", 10: "ten", 11: "eleven", 12: "twelve"}


def count(n: int) -> str:
    return rf"(?:(?<![\d.,$]){n}(?!\d|[.,]\d)|\b(?i:{WORDS[n]})\b)"


QUESTIONS: list[tuple[str, str, list[str], list[str], str]] = []
EXPECT: dict[str, list[str]] = {}  # pdf name -> figures that must appear in pdftotext output


def ask(pdf: str, question: str, must: list[str], must_not: list[str], note: str) -> None:
    QUESTIONS.append((pdf, question, must, must_not, note))


def unanswerable(pdf: str, question: str, why: str) -> None:
    QUESTIONS.append((pdf, question, [], [], "UNANSWERABLE: " + why))


# ------------------------------------------------------------------------------ sheet helpers
def titles(ws, rows: list[str], width: int, *, merge: bool = True, size: int = 14) -> int:
    for r, text in enumerate(rows, start=1):
        cell = ws.cell(r, 1, text)
        cell.font = Font(bold=True, size=size) if r == 1 else BOLD
        if merge:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
            cell.alignment = Alignment(horizontal="center")
    return len(rows) + 2


def page(ws, *, landscape: bool, fit_wide: bool = True, footer_left: str = "&F", footer_right: str = "") -> None:
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    if fit_wide:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    ws.oddFooter.left.text = footer_left
    ws.oddFooter.center.text = "Page &P of &N"
    if footer_right:
        ws.oddFooter.right.text = footer_right
    ws.page_margins.left = ws.page_margins.right = 0.4


def widths(ws, ws_widths: list[float]) -> None:
    for c, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w


def total_border(cell) -> None:
    cell.border = Border(top=THIN, bottom=DOUBLE)


# ============================================================================== 1. AR aging
def ar_aging() -> str:
    name = "AR Aging 9-30-26"
    pdf = name + ".pdf"
    wb = Workbook()
    ws = wb.active
    ws.title = "AR Aging"
    # customer, terms, current, 1-30, 31-60, 61-90, over 90
    rows = [
        ("Alderbrook Hospitality Group", "Net 30", 18_442.10, 6_210.00, 0, 0, 0),
        ("Brightline Transit Authority", "Net 60", 42_785.33, 15_600.00, 9_875.40, 0, 0),
        ("Cascadia Food Co-op", "2% 10, Net 30", 3_912.75, 0, 0, 0, 0),
        ("Dunmore Precision Parts", "Net 45", 0, 0, 4_418.90, 2_960.15, 11_304.62),
        ("Eastgate School District", "Net 30", 27_150.00, 0, 0, 0, 0),
        ("Foxhollow Veterinary Clinics", "Net 30", 1_286.44, 1_286.44, 1_286.44, 0, 0),
        ("Granite Peak Outfitters", "Net 30", 9_034.68, 2_117.25, 0, 0, 3_480.00),
        ("Halcyon Biotech LLC", "Net 60", 56_300.00, 0, 0, 0, 0),
        ("Ironwood Construction", "Net 45", 12_775.91, 8_940.37, 6_233.18, 5_105.50, 0),
        ("Juniper Ridge Apartments", "Due on receipt", 0, 0, 0, 0, 7_615.29),
        ("Kestrel Aviation Services", "Net 30", 21_488.06, 4_302.88, 0, 0, 0),
        ("Lakeshore Medical Group *", "Net 30", -2_150.00, 0, 0, 0, 0),
        ("Meridian Freight Lines", "Net 45", 7_640.52, 7_640.52, 0, 1_875.00, 0),
        ("Northgate Pharmacy", "Net 30", 4_229.87, 0, 0, 0, 0),
        ("Oakhurst Property Mgmt", "Net 60", 0, 0, 0, 0, 19_962.40),
    ]
    data = [(c, t, [D(x) for x in b]) for c, t, *b in rows]
    width = 8
    top = titles(ws, [COMPANY, "Accounts Receivable Aging Summary", "As of September 30, 2026",
                      "Aged by days past invoice due date"], width)
    # two-row header: Customer / Terms / Total Balance / Current span both rows; Days Past Due spans E:H
    for c, label in enumerate(["Customer", "Terms", "Total Balance", "Current"], start=1):
        ws.cell(top, c, label)
        ws.merge_cells(start_row=top, start_column=c, end_row=top + 1, end_column=c)
    ws.cell(top, 5, "Days Past Due")
    ws.merge_cells(start_row=top, start_column=5, end_row=top, end_column=8)
    for c, label in enumerate(["1 - 30", "31 - 60", "61 - 90", "Over 90"], start=5):
        ws.cell(top + 1, c, label)
    fill = PatternFill("solid", fgColor="305496")
    for rr in (top, top + 1):
        for c in range(1, width + 1):
            cell = ws.cell(rr, c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill
            cell.alignment = CENTER
            cell.border = BOX
    first = r = top + 2
    for cust, terms, buckets in data:
        ws.cell(r, 1, cust)
        ws.cell(r, 2, terms).alignment = Alignment(horizontal="center")
        fmt = ACCT if r == first else COMMA
        ws.cell(r, 3, f"=SUM(D{r}:H{r})").number_format = fmt
        ws.cell(r, 3).font = BOLD
        for c, v in enumerate(buckets, start=4):
            ws.cell(r, c, float(v)).number_format = fmt
        r += 1
    last = r - 1
    ws.cell(r, 1, "Total Accounts Receivable").font = BOLD
    for c in range(3, width + 1):
        col = get_column_letter(c)
        cell = ws.cell(r, c, f"=SUM({col}{first}:{col}{last})")
        cell.number_format = ACCT
        cell.font = BOLD
        total_border(cell)
    ws.cell(r + 2, 1, "* Lakeshore Medical Group credit balance is unapplied credit memo CM-2291 dated 9/22/2026; "
                      "to be refunded or applied in October.").font = Font(italic=True, size=9)
    widths(ws, [30, 14, 15, 14, 13, 13, 13, 13])
    page(ws, landscape=False, footer_right="AR Subledger tie-out: GL 1200")
    wb.save(XLSX / f"{name}.xlsx")

    # ---- truth
    def tot(b):
        return sum(b, ZERO)
    by = {c.rstrip(" *"): (t, b, tot(b)) for c, t, b in data}
    col_tot = [sum((b[i] for _, _, b in data), ZERO) for i in range(5)]
    grand = sum(col_tot, ZERO)
    EXPECT[pdf] = [show(v) for _, (_, _, v) in by.items()] + [show(v) for v in col_tot] + [show(grand)]

    iw = by["Ironwood Construction"][2]
    ask(pdf, "how much does Ironwood Construction owe us in total", [money(iw)], [],
        f"Ironwood total balance = {plain(iw)}")
    v = by["Brightline Transit Authority"][1][2]
    ask(pdf, "whats brightline sitting at in the 31-60 bucket", [money(v)], [],
        f"Brightline 31-60 = {plain(v)}")
    over = [c for c, (_, b, _) in by.items() if b[4] != 0]
    assert over == ["Dunmore Precision Parts", "Granite Peak Outfitters", "Juniper Ridge Apartments", "Oakhurst Property Mgmt"]
    ask(pdf, "which customers have anything over 90 days?",
        [r"Dunmore", r"Granite Peak", r"Juniper", r"Oakhurst"], [r"Ironwood", r"Meridian", r"Brightline"],
        "Over 90: " + "; ".join(f"{c} {plain(by[c][1][4])}" for c in over))
    past_due = sum(col_tot[1:], ZERO)
    ask(pdf, "total past due AR (everything not current)?", [money(past_due)], [],
        f"Past due = 1-30 + 31-60 + 61-90 + over 90 = {plain(past_due)}")
    n6190 = [c for c, (_, b, _) in by.items() if b[3] != 0]
    assert len(n6190) == 3
    ask(pdf, "how many customers have a balance in 61-90", [count(len(n6190))], [],
        f"{len(n6190)} customers: {', '.join(n6190)}")
    ask(pdf, "who has the biggest current balance", [r"Halcyon", money(by["Halcyon Biotech LLC"][1][0])], [r"Brightline"],
        f"Halcyon Biotech LLC, current {plain(by['Halcyon Biotech LLC'][1][0])}")
    lk = by["Lakeshore Medical Group"][2]
    ask(pdf, "what's the balance on Lakeshore Medical", [neg_money(lk)], [],
        f"Lakeshore Medical Group = ({-lk:,.2f}), a credit balance (unapplied credit memo CM-2291)")
    k, a = by["Kestrel Aviation Services"][2], by["Alderbrook Hospitality Group"][2]
    ask(pdf, "difference between Kestrel's total balance and Alderbrook's", [money(k - a)], [],
        f"Kestrel {plain(k)} - Alderbrook {plain(a)} = {plain(k - a)}")
    groups: dict[str, Decimal] = {}
    for c, (t, b, total) in by.items():
        groups[t] = groups.get(t, ZERO) + total
    ask(pdf, "can you break total AR out by payment terms", [money(v) for v in groups.values()], [],
        "By terms: " + "; ".join(f"{t} {plain(v)}" for t, v in groups.items()))
    net60_pd = [c for c, (t, b, _) in by.items() if t == "Net 60" and sum(b[1:], ZERO) != 0]
    assert net60_pd == ["Brightline Transit Authority", "Oakhurst Property Mgmt"]
    ask(pdf, "any Net 60 customers past due? who", [r"Brightline", r"Oakhurst"], [],
        "Net 60 with past-due balances: Brightline Transit Authority, Oakhurst Property Mgmt (Halcyon is all current)")
    unanswerable(pdf, "what's Dunmore's credit limit", "the aging has no credit limit column")
    unanswerable(pdf, "how much did Granite Peak pay us in october", "the aging is as of 9/30/26 and shows no October collections")
    return pdf


# ============================================================================== 2. accrued liabilities rollforward
def accruals() -> str:
    name = "Accrued Liabilities Rollforward Q3 2026"
    pdf = name + ".pdf"
    wb = Workbook()
    ws = wb.active
    ws.title = "Accruals RF"
    sections = {
        "Payroll": [
            ("Accrued wages", 184_220.15, 1_912_446.80, 1_889_105.62, 0),
            ("Accrued bonus", 412_500.00, 137_500.00, 0, 25_000.00),
            ("Accrued PTO", 96_318.42, 41_775.60, 38_902.17, 0),
            ("Accrued payroll taxes", 22_804.91, 146_302.18, 151_617.44, 0),
            ("Accrued 401(k) match", 4_115.30, 57_218.66, 62_146.36, 0),
        ],
        "Professional Fees": [
            ("Audit fees", 85_000.00, 42_500.00, 97_750.00, 0),
            ("Legal fees", 31_640.25, 58_212.90, 49_377.15, 6_500.00),
            ("Tax compliance", 18_000.00, -3_250.00, 9_800.00, 0),
            ("Consulting", 12_480.00, 27_915.75, 22_036.50, 4_200.00),
        ],
        "Other": [
            ("Utilities", 14_206.77, 43_118.29, 41_902.63, 0),
            ("Freight", 9_873.60, 31_447.18, 33_019.92, 1_150.00),
            ("Property taxes", 46_500.00, 23_250.00, 0, 0),
            ("Sales & use tax", 7_915.44, 19_602.31, 21_377.08, 0),
            ("Customer rebates", 28_750.00, 16_385.20, 30_100.00, 0),
            ("Accrued interest", -1_320.50, 18_750.00, 16_875.00, 0),
        ],
    }
    truth: dict[str, dict[str, tuple[Decimal, ...]]] = {}
    for sec, lines in sections.items():
        truth[sec] = {}
        for acct, b, a, p, rv in lines:
            b, a, p, rv = D(b), D(a), D(p), D(rv)
            truth[sec][acct] = (b, a, p, rv, b + a - p - rv)
    width = 6
    top = titles(ws, [COMPANY, "Accrued Liabilities Rollforward", "Quarter Ended September 30, 2026"], width)
    labels = ["Account", "Beginning Balance\n6/30/2026", "Additions", "Payments", "Reversals", "Ending Balance\n9/30/2026"]
    for c, label in enumerate(labels, start=1):
        cell = ws.cell(top, c, label)
        cell.font = BOLD
        cell.alignment = CENTER
        cell.border = Border(top=MEDIUM, bottom=MEDIUM)
    ws.row_dimensions[top].height = 32
    r = top + 1
    subtotal_rows = []
    for sec, accts in truth.items():
        ws.cell(r, 1, sec).font = Font(bold=True, underline="single")
        r += 1
        start = r
        for acct, (b, a, p, rv, e) in accts.items():
            ws.cell(r, 1, acct).alignment = Alignment(indent=2)
            for c, v in enumerate((b, a, p, rv), start=2):
                ws.cell(r, c, float(v)).number_format = NUM
            ws.cell(r, 6, f"=B{r}+C{r}-D{r}-E{r}").number_format = NUM
            r += 1
        ws.cell(r, 1, f"Total {sec}").font = BOLD
        for c in range(2, 7):
            col = get_column_letter(c)
            cell = ws.cell(r, c, f"=SUBTOTAL(9,{col}{start}:{col}{r - 1})")
            cell.number_format = NUM
            cell.font = BOLD
            cell.border = Border(top=THIN)
        subtotal_rows.append(r)
        r += 2
    ws.cell(r, 1, "Total Accrued Liabilities").font = BOLD
    for c in range(2, 7):
        col = get_column_letter(c)
        cell = ws.cell(r, c, "=" + "+".join(f"{col}{s}" for s in subtotal_rows))
        cell.number_format = ACCT
        cell.font = BOLD
        total_border(cell)
    grand_row = r
    all_lines = [(sec, acct, *v) for sec, accts in truth.items() for acct, v in accts.items()]
    end_total = sum((x[6] for x in all_lines), ZERO)
    r += 2
    ws.cell(r, 1, "Balance per trial balance, GL 2300-2399")
    ws.cell(r, 6, float(end_total)).number_format = ACCT
    ws.cell(r + 1, 1, "Difference")
    ws.cell(r + 1, 6, f"=F{grand_row}-F{r}").number_format = ACCT
    ws.cell(r + 3, 1, "Negative amounts in parentheses. Accrued 401(k) match is in a debit position at 9/30 "
                      "(Q3 true-up deposit exceeded the accrual).").font = Font(italic=True, size=9)
    widths(ws, [30, 17, 16, 16, 14, 17])
    page(ws, landscape=False, footer_right="Prepared by: G. Kim")
    wb.save(XLSX / f"{name}.xlsx")

    # ---- truth
    flat = {acct: v for accts in truth.values() for acct, v in accts.items()}
    sec_tot = {sec: tuple(sum((v[i] for v in accts.values()), ZERO) for i in range(5)) for sec, accts in truth.items()}
    grand = tuple(sum((s[i] for s in sec_tot.values()), ZERO) for i in range(5))
    EXPECT[pdf] = [show(v[4]) for v in flat.values()] + [show(x) for s in sec_tot.values() for x in s] + [show(x) for x in grand]

    ask(pdf, "ending balance for accrued bonus?", [money(flat["Accrued bonus"][4])], [],
        f"Accrued bonus ending 9/30/26 = {plain(flat['Accrued bonus'][4])}")
    ask(pdf, "total payroll accruals at 9/30", [money(sec_tot["Payroll"][4])], [],
        f"Total Payroll ending = {plain(sec_tot['Payroll'][4])}")
    rev = [a for a, v in flat.items() if v[3] != 0]
    assert rev == ["Accrued bonus", "Legal fees", "Consulting", "Freight"]
    ask(pdf, "which accounts had reversals this qtr", [r"(?i)bonus", r"(?i)legal", r"(?i)consulting", r"(?i)freight"],
        [r"(?i)audit", r"(?i)utilities", r"(?i)property tax"],
        "Reversals: " + "; ".join(f"{a} {plain(flat[a][3])}" for a in rev))
    k = flat["Accrued 401(k) match"][4]
    debit = [a for a, v in flat.items() if v[4] < 0]
    assert debit == ["Accrued 401(k) match"]
    ask(pdf, "any accruals in a debit balance at quarter end?",
        [r"401\s?\(?k\)?", neg_money(k, ("debit", "negative", "credit"))], [r"(?i)accrued interest"],
        f"Accrued 401(k) match, ending ({-k:,.2f}) debit balance")
    change = grand[4] - grand[0]
    ask(pdf, "how much did total accrued liabilities change during the quarter", [money(change)], [],
        f"{plain(grand[0])} -> {plain(grand[4])}, increase of {plain(change)}")
    inc = {a: v[4] - v[0] for a, v in flat.items()}
    best = max(inc, key=inc.get)
    assert best == "Accrued bonus" and sorted(inc.values())[-2] < inc[best] - 50_000
    ask(pdf, "which accrual went up the most over the quarter", [r"(?i)bonus", money(inc[best])], [],
        f"Accrued bonus, up {plain(inc[best])} ({plain(flat[best][0])} -> {plain(flat[best][4])})")
    ask(pdf, "total additions in professional fees?", [money(sec_tot["Professional Fees"][1])],
        [money(sec_tot["Professional Fees"][1] + 2 * D(3250))],
        f"Professional Fees additions = {plain(sec_tot['Professional Fees'][1])} (includes the (3,250.00) tax true-up)")
    up_other = [a for a, v in truth["Other"].items() if v[4] > v[0]]
    assert up_other == ["Utilities", "Property taxes", "Accrued interest"]
    ask(pdf, "how many of the Other accruals ended the quarter higher than they started", [count(len(up_other))], [],
        f"{len(up_other)}: " + ", ".join(up_other))
    tc = flat["Tax compliance"][1]
    ask(pdf, "what were Q3 additions to the tax compliance accrual", [neg_money(tc, ("true-up", "reduc", "negative", "decrease"))], [],
        f"Tax compliance additions = ({-tc:,.2f}), a negative true-up")
    ask(pdf, "payments by section pls", [money(s[2]) for s in sec_tot.values()], [],
        "Payments: " + "; ".join(f"{sec} {plain(s[2])}" for sec, s in sec_tot.items()))
    unanswerable(pdf, "what's the GL account number for accrued PTO", "accounts are listed by name only; no GL numbers per line")
    unanswerable(pdf, "what was accrued wages at 12/31/25", "the rollforward starts at 6/30/2026")
    return pdf


# ============================================================================== 3. bank reconciliations side by side
def bank_recs() -> str:
    name = "Bank Reconciliations 9-30-26"
    pdf = name + ".pdf"
    wb = Workbook()
    ws = wb.active
    ws.title = "Bank Recs"
    accounts = ["Operating", "Payroll", "Lockbox"]
    heads = ["Operating\nFirst Lakes Bank\nAcct x4471", "Payroll\nHarborview Bank\nAcct x0932",
             "Lockbox\nFirst Lakes Bank\nAcct x2205", "Total"]
    checks = [
        ("Tidewater Metals", 20418, date(2026, 9, 14), 52_300.00),
        ("Summit Ridge Electric", 20425, date(2026, 9, 17), 18_744.65),
        ("Pinecrest Fleet Services", 20433, date(2026, 9, 22), 6_215.80),
        ("Oakline Office Supply", 20437, date(2026, 9, 24), 1_088.37),
        ("Mesa Verde Chemical", 20440, date(2026, 9, 25), 27_931.12),
        ("County of Lake Treasurer", 20444, date(2026, 9, 29), 23_250.00),
        ("Brennan & Holt LLP", 20446, date(2026, 9, 30), 14_357.25),
    ]
    checks = [(p, n, d, D(a)) for p, n, d, a in checks]
    op_os = sum((c[3] for c in checks), ZERO)
    assert op_os == D(143_887.19)
    bank = [D(1_284_615.42), D(48_210.07), D(362_904.88)]
    dit = [D(86_412.50), ZERO, D(14_775.30)]
    os_ = [op_os, D(31_402.66), ZERO]
    other = [ZERO, ZERO, D(-1_215.00)]
    adj_bank = [bank[i] + dit[i] - os_[i] + other[i] for i in range(3)]
    fees = [D(385.00), D(62.50), D(140.00)]
    interest = [D(1_122.84), ZERO, ZERO]
    target_var = [ZERO, ZERO, D(18.40)]
    adj_gl = [adj_bank[i] - target_var[i] for i in range(3)]
    gl = [adj_gl[i] + fees[i] - interest[i] for i in range(3)]
    var = [adj_bank[i] - adj_gl[i] for i in range(3)]

    width = 5
    top = titles(ws, [COMPANY, "Bank Reconciliations", "September 30, 2026"], width)
    ws.cell(4, 1, "Prepared by: L. Wei  10/2/2026").font = Font(size=9)
    ws.cell(5, 1, "Reviewed by: J. Alvarez  10/5/2026").font = Font(size=9)
    top = 7
    for c, label in enumerate([""] + heads, start=1):
        cell = ws.cell(top, c, label)
        cell.font = BOLD
        cell.alignment = CENTER
        cell.border = Border(bottom=MEDIUM)
    ws.row_dimensions[top].height = 46
    layout = [
        ("Balance per bank statement, 9/30/2026", bank, "v"),
        ("Add: Deposits in transit", dit, "v"),
        ("Less: Outstanding checks", [-x for x in os_], "v"),
        ("Add/(Less): Other reconciling items", other, "v"),
        ("Adjusted bank balance", adj_bank, "sum_bank"),
        None,
        ("Balance per general ledger, 9/30/2026", gl, "v"),
        ("Less: Bank service charges not recorded", [-x for x in fees], "v"),
        ("Add: Interest earned not recorded", interest, "v"),
        ("Adjusted GL balance", adj_gl, "sum_gl"),
        None,
        ("Variance (adjusted bank less adjusted GL)", var, "var"),
    ]
    r = top + 1
    rows_at: dict[str, int] = {}
    n_rows = len(layout) + 1
    detail_top = top + 1 + n_rows + 3  # header row of the check detail
    detail_total = detail_top + len(checks) + 1
    for item in layout:
        if item is None:
            r += 1
            continue
        label, values, kind = item
        ws.cell(r, 1, label)
        rows_at[kind if kind != "v" else label] = r
        for i, v in enumerate(values):
            col = get_column_letter(2 + i)
            if kind == "sum_bank":
                f = f"=SUM({col}{r - 4}:{col}{r - 1})"
            elif kind == "sum_gl":
                f = f"=SUM({col}{r - 3}:{col}{r - 1})"
            elif kind == "var":
                f = f"={col}{rows_at['sum_bank']}-{col}{rows_at['sum_gl']}"
            elif label.startswith("Less: Outstanding") and i == 0:
                f = f"=-D{detail_total}"
            else:
                f = float(v)
            cell = ws.cell(r, 2 + i, f)
            cell.number_format = ACCT if kind != "v" or label.startswith("Balance per") else COMMA
        ws.cell(r, 5, f"=SUM(B{r}:D{r})").number_format = ACCT if kind != "v" or label.startswith("Balance per") else COMMA
        if kind != "v":
            for c in range(1, 6):
                ws.cell(r, c).font = BOLD
            for c in range(2, 6):
                ws.cell(r, c).border = Border(top=THIN, bottom=DOUBLE if kind == "var" else None)
        if kind == "var":
            for c in range(1, 6):
                ws.cell(r, c).fill = LIGHT
        r += 1
    ws.cell(r, 1, "Status")
    for i, v in enumerate(var):
        cell = ws.cell(r, 2 + i, "Reconciled" if v == 0 else "Investigate")
        cell.alignment = Alignment(horizontal="center")
        cell.font = Font(italic=True, color="000000" if v == 0 else "C00000")
    # check detail
    ws.cell(detail_top - 1, 1, "Outstanding Checks - Operating x4471").font = Font(bold=True, underline="single")
    for c, label in enumerate(["Payee", "Check #", "Issue Date", "Amount"], start=1):
        cell = ws.cell(detail_top, c, label)
        cell.font = BOLD
        cell.fill = GRAY
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center")
    rr = detail_top + 1
    for payee, num, when, amt in checks:
        ws.cell(rr, 1, payee).border = BOX
        ws.cell(rr, 2, num).border = BOX
        ws.cell(rr, 2).alignment = Alignment(horizontal="center")
        cell = ws.cell(rr, 3, when)
        cell.number_format = "m/d/yyyy"
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center")
        cell = ws.cell(rr, 4, float(amt))
        cell.number_format = COMMA
        cell.border = BOX
        rr += 1
    assert rr == detail_total
    ws.cell(rr, 1, "Total outstanding checks").font = BOLD
    cell = ws.cell(rr, 4, f"=SUM(D{detail_top + 1}:D{rr - 1})")
    cell.number_format = ACCT
    cell.font = BOLD
    total_border(cell)
    widths(ws, [44, 18, 18, 18, 18])
    page(ws, landscape=False, footer_right="Cash - GL 1010 / 1020 / 1030")
    wb.save(XLSX / f"{name}.xlsx")

    # ---- truth
    EXPECT[pdf] = [show(x) for x in adj_bank + adj_gl + gl + var] + [show(sum(adj_bank, ZERO)), show(-op_os), show(op_os),
                                                                     show(-sum(os_, ZERO)), show(sum(var, ZERO))]
    ask(pdf, "adjusted bank balance for payroll acct?", [money(adj_bank[1])], [],
        f"Payroll adjusted bank balance = {plain(adj_bank[1])}")
    ask(pdf, "what's the variance on lockbox", [money(var[2])], [], f"Lockbox variance = {plain(var[2])} (Investigate)")
    ask(pdf, "deposits in transit on the operating account", [money(dit[0])], [], f"Operating DIT = {plain(dit[0])}")
    tos = sum(os_, ZERO)
    ask(pdf, "total outstanding checks across all three accounts", [money(tos)], [],
        f"Total outstanding checks = {plain(tos)} (shown as ({tos:,.2f}))")
    ask(pdf, "which accounts don't reconcile?", [r"(?i)lockbox|x?2205"], [],
        f"Only Lockbox (x2205), variance {plain(var[2])}; Operating and Payroll are reconciled")
    ask(pdf, "GL balance for operating before adjustments", [money(gl[0])], [],
        f"Operating balance per GL = {plain(gl[0])}")
    d = bank[0] - gl[0]
    ask(pdf, "difference between the bank balance and the GL balance on operating", [money(d)], [],
        f"{plain(bank[0])} - {plain(gl[0])} = {plain(d)}")
    big = [c for c in checks if c[3] > 10_000]
    assert len(big) == 5
    ask(pdf, "how many outstanding checks on operating are over 10k", [count(len(big))], [],
        f"{len(big)}: " + ", ".join(f"#{c[1]} {plain(c[3])}" for c in big))
    top_check = max(checks, key=lambda c: c[3])
    ask(pdf, "biggest outstanding check on operating - who's it to and how much", [r"Tidewater", money(top_check[3])], [],
        f"Check #{top_check[1]} to {top_check[0]}, {plain(top_check[3])}")
    late = [c for c in checks if c[2] > date(2026, 9, 20)]
    late_sum = sum((c[3] for c in late), ZERO)
    ask(pdf, "total of the outstanding checks written after Sep 20", [money(late_sum)], [],
        f"{len(late)} checks dated 9/22-9/30 = {plain(late_sum)}")
    unanswerable(pdf, "when did check 20418 clear the bank", "the rec lists outstanding checks as of 9/30; no clearing dates")
    unanswerable(pdf, "what was the operating balance per bank at 8/31", "only the 9/30/2026 reconciliation is shown")
    return pdf


# ============================================================================== 4. monthly revenue, two pages wide
def revenue() -> str:
    name = "Revenue by Customer FY2026"
    pdf = name + ".pdf"
    wb = Workbook()
    ws = wb.active
    ws.title = "Rev by Customer"
    customers = [
        ("Apex Dental Partners", "Northeast", 38_000), ("Bayfront Hotels", "South", 61_500),
        ("Cedar Valley Foods", "Midwest", 24_800), ("Driftwood Marine", "West", 9_700),
        ("Elmstead Logistics", "Midwest", 47_200), ("Fairhaven Senior Living", "Northeast", 18_300),
        ("Glacier Point Software", "West", 72_900), ("Hollis & Grant Retail", "South", 14_600),
        ("Inland Empire Builders", "West", 33_400), ("Jasper County Schools", "South", 21_100),
        ("Keystone Auto Group", "Northeast", 27_750), ("Larkspur Pharma", "Northeast", 55_000),
        ("Monarch Stadium Group", "Midwest", 12_900), ("Nimbus Data Centers", "West", 84_000),
        ("Orchard Lane Grocers", "Midwest", 30_600), ("Pioneer Rail Services", "South", 43_800),
        ("Quarry Hill Brewing", "Northeast", 6_450), ("Redwood Health System", "West", 66_200),
        ("Silverline Telecom", "South", 39_900), ("Thornbury Packaging", "Midwest", 16_750),
    ]
    season = [0.88, 0.91, 1.02, 0.97, 1.04, 1.08, 0.95, 0.99, 1.06, 1.03, 0.98, 1.12]
    starts = {"Larkspur Pharma": 3, "Nimbus Data Centers": 4}
    ends = {"Quarry Hill Brewing": 7}
    rng = random.Random(20261006)
    data: dict[str, tuple[str, list[Decimal]]] = {}
    for cust, region, base in customers:
        vals = []
        for m in range(1, 13):
            x = base * season[m - 1] * rng.uniform(0.86, 1.14)
            if m < starts.get(cust, 1) or m > ends.get(cust, 12):
                x = 0
            vals.append(D(round(x, 2)))
        data[cust] = (region, vals)
    data["Hollis & Grant Retail"][1][4] = D(-1_842.60)  # May: credit memo larger than billings
    months = [date(2026, m, 1) for m in range(1, 13)]
    width = 2 + 12 + 1
    small = Font(size=8)
    small_bold = Font(size=8, bold=True)
    ws.cell(1, 1, COMPANY).font = Font(bold=True, size=11)
    ws.cell(2, 1, "Monthly Revenue by Customer - Fiscal Year 2026").font = Font(bold=True, size=10)
    ws.cell(3, 1, "Jan-Sep actual; Oct-Dec forecast. USD.").font = Font(italic=True, size=8)
    top = 5
    for c, label in enumerate(["Customer", "Region"] + months + ["FY 2026 Total"], start=1):
        cell = ws.cell(top, c, label)
        if isinstance(label, date):
            cell.number_format = "mmm-yy"
        cell.font = Font(size=8, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="375623")
        cell.alignment = CENTER
        cell.border = BOX
    r = top + 1
    for cust, (region, vals) in data.items():
        ws.cell(r, 1, cust).font = small
        ws.cell(r, 2, region).font = small
        for k, v in enumerate(vals):
            cell = ws.cell(r, 3 + k, float(v))
            cell.number_format = COMMA
            cell.font = small
        cell = ws.cell(r, width, f"=SUM(C{r}:N{r})")
        cell.number_format = COMMA
        cell.font = small_bold
        if (r - top) % 2 == 0:
            for c in range(1, width + 1):
                ws.cell(r, c).fill = LIGHT
        r += 1
    ws.cell(r, 1, "Total Revenue").font = small_bold
    for c in range(3, width + 1):
        col = get_column_letter(c)
        cell = ws.cell(r, c, f"=SUM({col}{top + 1}:{col}{r - 1})")
        cell.number_format = ACCT
        cell.font = small_bold
        total_border(cell)
    ws.cell(r + 2, 1, "Hollis & Grant May reflects credit memo CM-0518 for returned fixtures.").font = Font(italic=True, size=8)
    widths(ws, [24, 10] + [11] * 12 + [13])
    page(ws, landscape=True, fit_wide=False, footer_right="Source: NetSuite saved search 'Rev by Cust by Month'")
    ws.print_title_rows = f"{top}:{top}"
    ws.print_title_cols = "A:B"
    ws.page_setup.pageOrder = "overThenDown"
    wb.save(XLSX / f"{name}.xlsx")

    # ---- truth
    fy = {c: sum(v, ZERO) for c, (_, v) in data.items()}
    month_tot = [sum((v[k] for _, v in data.values()), ZERO) for k in range(12)]
    grand = sum(month_tot, ZERO)
    EXPECT[pdf] = [show(x) for x in fy.values()] + [show(x) for x in month_tot] + [show(grand), "(1,842.60)"]
    for c, (reg, v) in data.items() if os.environ.get("SHOW_DATA") else ():
        print(f"  {c:26} {reg:9} FY {plain(fy[c]):>14}  " + " ".join(plain(x) for x in v), file=sys.stderr)

    cv = data["Cedar Valley Foods"][1]
    q2 = cv[3] + cv[4] + cv[5]
    ask(pdf, "Q2 revenue for Cedar Valley Foods?", [money(q2)], [money(cv[0] + cv[1] + cv[2])],
        f"Apr {plain(cv[3])} + May {plain(cv[4])} + Jun {plain(cv[5])} = {plain(q2)}")
    q1 = month_tot[0] + month_tot[1] + month_tot[2]
    ask(pdf, "total revenue jan through mar, all customers", [money(q1)], [],
        f"Jan {plain(month_tot[0])} + Feb {plain(month_tot[1])} + Mar {plain(month_tot[2])} = {plain(q1)}")
    ask(pdf, "FY total for Glacier Point", [money(fy["Glacier Point Software"])], [],
        f"Glacier Point Software FY2026 = {plain(fy['Glacier Point Software'])}")
    dec = {c: v[11] for c, (_, v) in data.items()}
    top_dec = max(dec, key=dec.get)
    runner = sorted(dec.values())[-2]
    assert dec[top_dec] - runner > 1_000
    ask(pdf, "who was our biggest customer in December", [re.escape(top_dec.split()[0]), money(dec[top_dec])], [],
        f"{top_dec}, Dec-26 = {plain(dec[top_dec])}")
    regions: dict[str, Decimal] = {}
    for c, (reg, _) in data.items():
        regions[reg] = regions.get(reg, ZERO) + fy[c]
    ask(pdf, "full-year revenue by region?", [money(v) for v in regions.values()], [],
        "By region: " + "; ".join(f"{k} {plain(v)}" for k, v in regions.items()))
    zero_jan = [c for c, (_, v) in data.items() if v[0] == 0]
    assert zero_jan == ["Larkspur Pharma", "Nimbus Data Centers"]
    ask(pdf, "how many customers had no revenue in January", [count(len(zero_jan))], [],
        f"{len(zero_jan)}: Larkspur Pharma (starts Mar), Nimbus Data Centers (starts Apr)")
    hg = data["Hollis & Grant Retail"][1][4]
    ask(pdf, "Hollis & Grant May revenue?", [neg_money(hg, ("credit", "negative"))], [],
        f"Hollis & Grant Retail May-26 = ({-hg:,.2f}) (credit memo)")
    dj = month_tot[6] - month_tot[5]
    ask(pdf, "how much did total revenue change from June to July", [money(dj)], [],
        f"Jul {plain(month_tot[6])} - Jun {plain(month_tot[5])} = {plain(dj)}")
    sv = data["Silverline Telecom"][1]
    h2 = sum(sv[6:], ZERO)
    ask(pdf, "H2 revenue for Silverline", [money(h2)], [money(sum(sv[:6], ZERO))],
        f"Silverline Telecom Jul-Dec = {plain(h2)}")
    avg = r2(fy["Keystone Auto Group"] / 12)
    ask(pdf, "average monthly revenue for Keystone Auto", [money(avg)], [],
        f"Keystone FY {plain(fy['Keystone Auto Group'])} / 12 = {plain(avg)}")
    big = [c for c in data if fy[c] > 600_000]
    for c in data:
        assert abs(fy[c] - 600_000) > 15_000, c
    near = [c for c in data if 400_000 < fy[c] <= 600_000]
    ask(pdf, "which customers did more than $600k for the year",
        [re.escape(c.split()[0]) for c in big], [re.escape(c.split()[0]) for c in near],
        "Over 600k: " + "; ".join(f"{c} {plain(fy[c])}" for c in big))
    unanswerable(pdf, "what was Apex Dental's 2025 revenue", "the schedule only covers FY2026")
    return pdf


# ============================================================================== 5. payroll register
def payroll() -> str:
    name = "Payroll Register 10-15-26"
    pdf = name + ".pdf"
    wb = Workbook()
    ws = wb.active
    ws.title = "Register"
    # employee, dept, hours, rate, fed %, state %
    emps = [
        ("Ava Thornton", "Finance", 80.00, 52.40, 0.162, 0.0495),
        ("Marcus Bell", "Finance", 80.00, 38.75, 0.118, 0.0495),
        ("Nina Patel", "Finance", 72.50, 34.10, 0.104, 0.0315),
        ("Diego Ramirez", "Operations", 86.50, 27.80, 0.087, 0.0495),
        ("Hannah Cole", "Operations", 80.00, 24.65, 0.079, 0.0315),
        ("Isaac Moreno", "Operations", 91.25, 26.35, 0.093, 0.0495),
        ("Keisha Grant", "Operations", 64.00, 22.90, 0.071, 0.0495),
        ("Liam O'Connor", "Sales", 80.00, 45.20, 0.141, 0.0495),
        ("Sofia Lindqvist", "Sales", 80.00, 41.85, 0.126, 0.0315),
        ("Trevor Huang", "Sales", 76.00, 39.60, 0.122, 0.0495),
        ("Uma Krishnan", "Engineering", 80.00, 68.25, 0.188, 0.0495),
        ("Victor Osei", "Engineering", 80.00, 61.40, 0.171, 0.0495),
        ("Wen Zhao", "Engineering", 78.50, 57.90, 0.166, 0.0315),
        ("Zachary Reed", "Engineering", 80.00, 49.10, 0.149, 0.0495),
    ]
    people = []
    for emp, dept, hours, rate, fp, sp in emps:
        h, rt = Decimal(str(hours)), Decimal(str(rate))
        exact = h * rt
        assert (exact * 1000) % 10 != 5 or exact == exact.quantize(CENT), (emp, exact)  # no ROUND ties
        gross = r2(exact)
        fica_exact = gross * Decimal("0.0765")
        assert (fica_exact * 1000) % 10 != 5, (emp, fica_exact)
        fica = r2(fica_exact)
        fed = r2(gross * Decimal(str(fp)))
        state = r2(gross * Decimal(str(sp)))
        net = gross - fed - state - fica
        people.append((emp, dept, h, rt, gross, fed, state, fica, net))
    width = 9
    ws.cell(1, 1, COMPANY).font = Font(bold=True, size=13)
    ws.cell(2, 1, "Payroll Register - Biweekly").font = BOLD
    facts = [("Period Begin:", date(2026, 9, 27)), ("Period End:", date(2026, 10, 10)), ("Pay Date:", date(2026, 10, 15)),
             ("Pay Group:", "Hourly & Salaried - IL/IN")]
    for i, fact in enumerate(facts, start=4):
        ws.cell(i, 1, fact[0]).font = BOLD
        for j, v in enumerate(fact[1:], start=2):
            cell = ws.cell(i, j, v)
            if isinstance(v, date):
                cell.number_format = "mmm d, yyyy"
            cell.alignment = Alignment(horizontal="left")
    top = 9
    labels = ["Employee", "Dept", "Hours", "Rate", "Gross", "Federal W/H", "State W/H", "FICA", "Net Pay"]
    for c, label in enumerate(labels, start=1):
        cell = ws.cell(top, c, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="404040")
        cell.alignment = CENTER
    r = top + 1
    sub_rows = []
    depts = list(dict.fromkeys(p[1] for p in people))
    for dept in depts:
        start = r
        for emp, d, h, rt, gross, fed, state, fica, net in [p for p in people if p[1] == dept]:
            ws.cell(r, 1, emp)
            ws.cell(r, 2, d)
            ws.cell(r, 3, float(h)).number_format = "0.00"
            ws.cell(r, 4, float(rt)).number_format = "0.00"
            ws.cell(r, 5, f"=ROUND(C{r}*D{r},2)").number_format = NUM
            ws.cell(r, 6, float(fed)).number_format = NUM
            ws.cell(r, 7, float(state)).number_format = NUM
            ws.cell(r, 8, f"=ROUND(E{r}*0.0765,2)").number_format = NUM
            ws.cell(r, 9, f"=E{r}-F{r}-G{r}-H{r}").number_format = NUM
            r += 1
        ws.cell(r, 1, f"{dept} Subtotal")
        for c in (3, 5, 6, 7, 8, 9):
            col = get_column_letter(c)
            cell = ws.cell(r, c, f"=SUM({col}{start}:{col}{r - 1})")
            cell.number_format = "#,##0.00" if c == 3 else NUM
        for c in range(1, width + 1):
            ws.cell(r, c).font = BOLD
            ws.cell(r, c).fill = GRAY
            ws.cell(r, c).border = Border(top=THIN)
        sub_rows.append(r)
        r += 2
    ws.cell(r, 1, "Company Total").font = BOLD
    for c in (3, 5, 6, 7, 8, 9):
        col = get_column_letter(c)
        cell = ws.cell(r, c, "=" + "+".join(f"{col}{s}" for s in sub_rows))
        cell.number_format = "#,##0.00" if c == 3 else ACCT
        cell.font = BOLD
        total_border(cell)
    ws.cell(r + 1, 1, f"Employees paid: {len(people)}")
    ws.cell(r + 3, 1, "Direct deposit file transmitted to Harborview Bank on Oct 13, 2026. "
                      "Employer FICA match not shown.").font = Font(italic=True, size=9)
    widths(ws, [20, 13, 9, 9, 13, 13, 12, 12, 13])
    page(ws, landscape=False, footer_right="CONFIDENTIAL - Payroll")
    wb.save(XLSX / f"{name}.xlsx")

    # ---- truth
    by = {p[0]: p for p in people}
    sums = {d: [sum((p[i] for p in people if p[1] == d), Decimal(0)) for i in range(2, 9)] for d in depts}
    tot = [sum((p[i] for p in people), Decimal(0)) for i in range(2, 9)]
    EXPECT[pdf] = ([show(p[4]) for p in people] + [show(p[7]) for p in people] + [show(p[8]) for p in people]
                   + [show(x) for s in sums.values() for x in s[2:]] + [show(x) for x in tot[2:]] + ["Oct 15, 2026"])

    w = by["Wen Zhao"]
    ask(pdf, "what's Wen Zhao's net pay this check", [money(w[8])], [], f"Wen Zhao net pay = {plain(w[8])}")
    ot = [p[0] for p in people if p[2] > 80]
    assert ot == ["Diego Ramirez", "Isaac Moreno"]
    ask(pdf, "who worked more than 80 hours?", [r"Diego", r"Isaac"], [r"Ava", r"Wen Zhao", r"Zachary"],
        "Diego Ramirez 86.50, Isaac Moreno 91.25")
    ask(pdf, "total federal withholding on this payroll", [money(tot[3])], [], f"Federal W/H total = {plain(tot[3])}")
    hi = max(people, key=lambda p: p[3])
    ask(pdf, "highest hourly rate and who", [r"Uma", money(hi[3])], [], f"{hi[0]} at {hi[3]:.2f}/hr")
    n_sales = sum(1 for p in people if p[1] == "Sales")
    ask(pdf, "how many people in Sales got paid", [count(n_sales)], [], f"{n_sales}: Liam O'Connor, Sofia Lindqvist, Trevor Huang")
    eng = [p for p in people if p[1] == "Engineering"]
    avg = r2(sum((p[4] for p in eng), Decimal(0)) / len(eng))
    ask(pdf, "avg gross pay in engineering", [money(avg)], [],
        f"Engineering gross {plain(sums['Engineering'][2])} / {len(eng)} = {plain(avg)}")
    ask(pdf, "net pay by department", [money(s[6]) for s in sums.values()], [],
        "Net pay: " + "; ".join(f"{d} {plain(s[6])}" for d, s in sums.items()))
    ask(pdf, "what's the pay date on this register",
        [r"(?i)(?:Oct(?:ober)?\.?\s+15(?:th)?,?\s+2026|10/15/(?:20)?26|2026-10-15)"], [],
        "Pay date Oct 15, 2026 (period Sep 27 - Oct 10, 2026)")
    lo, tr = by["Liam O'Connor"], by["Trevor Huang"]
    ask(pdf, "how much more did Liam O'Connor gross than Trevor Huang", [money(lo[4] - tr[4])], [],
        f"{plain(lo[4])} - {plain(tr[4])} = {plain(lo[4] - tr[4])}")
    ask(pdf, "total hours for finance", [number(sums["Finance"][0])], [], f"Finance hours = {sums['Finance'][0]:,.2f}")
    unanswerable(pdf, "what's Ava Thornton's YTD gross", "register covers one pay period; no YTD columns")
    unanswerable(pdf, "how much was taken out for Marcus Bell's 401k", "no 401(k) deduction column on the register")
    return pdf


# ============================================================================== 6. intercompany matrix
def intercompany() -> str:
    name = "Intercompany Matrix 9-30-26"
    pdf = name + ".pdf"
    wb = Workbook()
    ws = wb.active
    ws.title = "IC Matrix"
    ents = [("US01", "Cobalt Ridge Industries, Inc."), ("CA02", "Cobalt Ridge Canada ULC"),
            ("UK03", "Cobalt Ridge UK Ltd."), ("MX04", "Cobalt Ridge de México S. de R.L."),
            ("DE05", "Cobalt Ridge GmbH"), ("SG06", "Cobalt Ridge Asia Pte. Ltd.")]
    codes = [e[0] for e in ents]
    # grid[i][j] = amount due TO entity i FROM entity j (i's receivable, j's payable)
    raw = {
        "US01": {"CA02": 1_245_800.00, "UK03": 892_416.55, "MX04": 2_318_740.12, "DE05": 467_290.38, "SG06": 0},
        "CA02": {"US01": 318_205.47, "UK03": 0, "MX04": 74_118.90, "DE05": 0, "SG06": 12_640.00},
        "UK03": {"US01": 156_930.21, "CA02": 0, "MX04": 0, "DE05": 288_415.66, "SG06": 41_207.73},
        "MX04": {"US01": 905_112.84, "CA02": 22_780.50, "UK03": 0, "DE05": 0, "SG06": 0},
        "DE05": {"US01": 0, "CA02": 9_315.08, "UK03": 133_672.40, "MX04": 0, "SG06": 58_904.17},
        "SG06": {"US01": 214_560.00, "CA02": 0, "UK03": 19_488.62, "MX04": 6_027.35, "DE05": 0},
    }
    grid = {i: {j: D(raw[i][j]) for j in codes if j != i} for i in codes}
    n = len(codes)
    width = n + 2
    top = titles(ws, [COMPANY, "Intercompany Balances Matrix (USD)", "As of September 30, 2026"], width)
    # header: A merged two rows; B..G "Due From" over entity codes; H "Total Due To" merged two rows
    ws.cell(top, 1, "Due To (receivable entity)")
    ws.merge_cells(start_row=top, start_column=1, end_row=top + 1, end_column=1)
    ws.cell(top, 2, "Due From (payable entity)")
    ws.merge_cells(start_row=top, start_column=2, end_row=top, end_column=n + 1)
    ws.cell(top, width, "Total Due To")
    ws.merge_cells(start_row=top, start_column=width, end_row=top + 1, end_column=width)
    for j, code in enumerate(codes, start=2):
        ws.cell(top + 1, j, code)
    for rr in (top, top + 1):
        for c in range(1, width + 1):
            cell = ws.cell(rr, c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="7030A0")
            cell.alignment = CENTER
            cell.border = BOX
    first = r = top + 2
    for i, (code, legal) in enumerate(ents):
        ws.cell(r, 1, f"{code}  {legal}").border = BOX
        for j, other in enumerate(codes, start=2):
            cell = ws.cell(r, j)
            cell.border = BOX
            if other == code:
                cell.fill = GRAY
                continue
            cell.value = float(grid[code][other])
            cell.number_format = COMMA
        cell = ws.cell(r, width, f"=SUM(B{r}:{get_column_letter(n + 1)}{r})")
        cell.number_format = COMMA
        cell.font = BOLD
        cell.border = BOX
        r += 1
    last = r - 1
    ws.cell(r, 1, "Total Due From").font = BOLD
    for c in range(2, width + 1):
        col = get_column_letter(c)
        cell = ws.cell(r, c, f"=SUM({col}{first}:{col}{last})")
        cell.number_format = ACCT
        cell.font = BOLD
        total_border(cell)
    tot_row = r
    r += 3
    ws.cell(r - 1, 1, "Net Intercompany Position").font = Font(bold=True, underline="single")
    for c, label in enumerate(["Entity", "Receivable\n(Due To)", "Payable\n(Due From)", "Net Receivable\n/(Payable)"], start=1):
        cell = ws.cell(r, c, label)
        cell.font = BOLD
        cell.alignment = CENTER
        cell.border = Border(bottom=THIN)
    ws.row_dimensions[r].height = 30
    net_top = r + 1
    for k, code in enumerate(codes):
        rr = net_top + k
        col = get_column_letter(2 + k)
        ws.cell(rr, 1, code)
        ws.cell(rr, 2, f"={get_column_letter(width)}{first + k}").number_format = COMMA
        ws.cell(rr, 3, f"={col}{tot_row}").number_format = COMMA
        ws.cell(rr, 4, f"=B{rr}-C{rr}").number_format = COMMA
    rr = net_top + n
    ws.cell(rr, 1, "Total").font = BOLD
    for c in (2, 3, 4):
        col = get_column_letter(c)
        cell = ws.cell(rr, c, f"=ROUND(SUM({col}{net_top}:{col}{rr - 1}),2)")
        cell.number_format = ACCT
        cell.font = BOLD
        total_border(cell)
    ws.cell(rr + 2, 1, "Read across: each cell is the amount owed TO the row entity BY the column entity. "
                       "Balances translated at 9/30/2026 spot rates.").font = Font(italic=True, size=9)
    widths(ws, [36] + [15] * n + [16])
    page(ws, landscape=True, footer_right="Consolidation - IC elimination support")
    wb.save(XLSX / f"{name}.xlsx")

    # ---- truth
    due_to = {i: sum(grid[i].values(), ZERO) for i in codes}            # row totals: receivable
    due_from = {j: sum((grid[i][j] for i in codes if i != j), ZERO) for j in codes}  # column totals: payable
    net = {k: due_to[k] - due_from[k] for k in codes}
    grand = sum(due_to.values(), ZERO)
    assert grand == sum(due_from.values(), ZERO) and sum(net.values(), ZERO) == 0
    EXPECT[pdf] = [show(x) for x in due_to.values()] + [show(x) for x in due_from.values()] + [show(x) for x in net.values()] + [show(grand)]

    ask(pdf, "how much does MX04 owe US01", [money(grid["US01"]["MX04"])], [money(grid["MX04"]["US01"])],
        f"Due to US01 from MX04 = {plain(grid['US01']['MX04'])}")
    ask(pdf, "what does the US entity owe Mexico", [money(grid["MX04"]["US01"])], [money(grid["US01"]["MX04"])],
        f"Due to MX04 from US01 = {plain(grid['MX04']['US01'])}")
    ask(pdf, "total due to UK03", [money(due_to["UK03"])], [money(due_from["UK03"])], f"UK03 row total = {plain(due_to['UK03'])}")
    ask(pdf, "how much does DE05 owe the other entities in total", [money(due_from["DE05"])], [money(due_to["DE05"])],
        f"DE05 column total (Total Due From) = {plain(due_from['DE05'])}")
    ca = net["CA02"]
    assert ca < 0
    ask(pdf, "net IC position for Canada?", [neg_money(ca, ("payable", "owes", "net payable"))], [],
        f"CA02 receivable {plain(due_to['CA02'])} - payable {plain(due_from['CA02'])} = ({-ca:,.2f}) net payable")
    owe_sg = [j for j, v in grid["SG06"].items() if v != 0]
    assert owe_sg == ["US01", "UK03", "MX04"]
    ask(pdf, "how many entities owe SG06 anything", [count(len(owe_sg))], [], f"{len(owe_sg)}: US01, UK03, MX04")
    big = max(due_to, key=due_to.get)
    ask(pdf, "which entity has the largest total intercompany receivable", [r"US01|Cobalt Ridge Industries", money(due_to[big])], [],
        f"{big}, total due to {plain(due_to[big])}")
    ask(pdf, "grand total of all IC balances", [money(grand)], [], f"Total = {plain(grand)}")
    de = [j for j, v in grid["DE05"].items() if v > 50_000]
    assert de == ["UK03", "SG06"]
    ask(pdf, "which entities owe DE05 more than 50k", [r"UK03|UK Ltd", r"SG06|Asia"], [],
        f"UK03 {plain(grid['DE05']['UK03'])}, SG06 {plain(grid['DE05']['SG06'])} (CA02 only {plain(grid['DE05']['CA02'])})")
    a, b = grid["DE05"]["UK03"], grid["UK03"]["DE05"]
    ask(pdf, "difference between what UK owes Germany and what Germany owes UK", [money(b - a)], [],
        f"Germany owes UK {plain(b)}; UK owes Germany {plain(a)}; difference {plain(b - a)}")
    unanswerable(pdf, "what FX rate did we use for the GBP balances", "the note says 9/30 spot rates but no rates are shown")
    unanswerable(pdf, "what were IC balances at 6/30", "the matrix is as of 9/30/2026 only")
    return pdf


# ============================================================================== output
def _raw(p: str) -> str:
    return f'r"{p}"' if '"' not in p and not p.endswith("\\") else repr(p)


def write_questions(path: Path) -> None:
    out = ['"""Questions about the schedules beside this file, written by make_books.py (do not edit by hand).',
           "",
           "Each entry: (pdf filename, question, must_match regexes, must_not_match regexes, note).",
           "Money regexes accept the figure with or without $ and thousands separators (and without .00 when",
           "whole). Negative figures accept -1,234.56, (1,234.56) or a word like 'credit' next to the figure.",
           'Unanswerable questions have empty lists and a note starting "UNANSWERABLE:".',
           '"""', "", "NEW = ["]
    current = None
    for pdf, question, must, must_not, note in QUESTIONS:
        if pdf != current:
            out.append(f"    # ---------------------------------------------------------------- {pdf}")
            current = pdf
        out.append(f"    ({pdf!r}, {question!r},")
        out.append(f"     [{', '.join(_raw(p) for p in must)}],")
        out.append(f"     [{', '.join(_raw(p) for p in must_not)}],")
        out.append(f"     {note!r}),")
    out.append("]")
    path.write_text("\n".join(out) + "\n")


def self_check(path: Path) -> None:
    ns: dict = {}
    exec(path.read_text(), ns)
    assert ns["NEW"] == QUESTIONS, "questions file does not round-trip"
    for pdf, question, must, must_not, note in QUESTIONS:
        assert (PDF / pdf).exists(), pdf
        if note.startswith("UNANSWERABLE:"):
            assert not must and not must_not
            continue
        assert must, question
        for p in must:
            assert re.search(p, note), (question, p, note)
        for p in must_not:
            re.compile(p)
            assert not re.search(p, note), (question, p, note)
    # the money regex accepts the usual renderings and rejects near misses
    m = money(D(12_345.67))
    for ok in ["$12,345.67", "12345.67", "$ 12,345.67", "is 12,345.67."]:
        assert re.search(m, ok), ok
    for bad in ["112,345.67", "12,345.678", "12,345.6", "2,345.67"]:
        assert not re.search(m, bad), bad
    w = money(D(1_250))
    for ok in ["$1,250", "1250.00", "$1,250.00."]:
        assert re.search(w, ok), ok
    for bad in ["1,250.50", "11,250", "1,2500"]:
        assert not re.search(w, bad), bad
    n = neg_money(D(-2_150))
    for ok in ["-$2,150.00", "($2,150.00)", "$(2,150.00)", "-2150", "a credit balance of $2,150.00", "2,150.00 CR"]:
        assert re.search(n, ok), ok
    assert not re.search(n, "owes $2,150.00")


def convert() -> None:
    PROFILE.mkdir(exist_ok=True)
    books = sorted(XLSX.glob("*.xlsx"))
    for old in PDF.glob("*.pdf"):
        old.unlink()
    subprocess.run(["soffice", f"-env:UserInstallation=file://{PROFILE}", "--headless", "--calc",
                    "--convert-to", "pdf", "--outdir", str(PDF), *map(str, books)],
                   check=True, capture_output=True, timeout=300)


def verify_pdfs() -> None:
    for pdf, figures in EXPECT.items():
        path = PDF / pdf
        text = subprocess.run(["pdftotext", "-layout", str(path), "-"], check=True, capture_output=True, text=True).stdout
        pages = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", str(path)], check=True, capture_output=True,
                                                                  text=True).stdout).group(1))
        missing = sorted({f for f in figures if f != "-" and f not in text})
        print(f"{pdf}: {pages} page(s), {len(figures)} computed figures checked" + (f", MISSING {missing}" if missing else ""))
        assert not missing, (pdf, missing)


def main() -> None:
    XLSX.mkdir(exist_ok=True)
    PDF.mkdir(exist_ok=True)
    for build in (ar_aging, accruals, bank_recs, revenue, payroll, intercompany):
        build()
    convert()
    verify_pdfs()
    path = OUT / "questions.py"
    write_questions(path)
    self_check(path)
    per = {}
    for q in QUESTIONS:
        per[q[0]] = per.get(q[0], 0) + 1
    for pdf, k in per.items():
        print(f"  {k:2d} questions  {pdf}")
    print(f"{len(QUESTIONS)} questions -> {path}")


if __name__ == "__main__":
    main()
