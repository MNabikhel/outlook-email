"""Two more schedules, built the way a controller builds them in Excel and saved as PDF with LibreOffice.
Layouts unlike the others in ../excel_schedules and beside this file:

    1. Sales tax collected by state: states down the side, Jul-26..Sep-26 under a merged "Tax Collected"
       heading, a Q3 total column, then two text/date columns at the far right (Filing Frequency,
       Next Return Due); prepared-by / source lines above the table; a negative month in parentheses;
       lettered footnotes under the table.
    2. Headcount & salary budget: departments down the side in three sections with subtotals; a
       three-row header (merged "Location" band, then merged Chicago | Austin | Remote groups, each split
       into HC and Salary, plus a Total group); blank location cells shown as "-"; prepared-by,
       as-of and version lines above the table; whole-dollar accounting format.

The data lives at module level (Decimals), so questions about the schedules can be answered from it. The
PDFs beside this file are its output, used by tests/test_more_schedules.py. To make them again (needs
openpyxl, LibreOffice's soffice and poppler's pdftotext on the PATH):

    python tests/fixtures/more_schedules/make_more_books.py tests/fixtures/more_schedules
"""

from __future__ import annotations

import shutil
import subprocess
import sys
import tempfile
from datetime import date
from decimal import Decimal
from pathlib import Path

OUT = Path(sys.argv[1] if len(sys.argv) > 1 else "books").resolve()
PDF = OUT
XLSX = OUT / "xlsx"

ACCT = '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)'
COMMA = '_(* #,##0.00_);_(* (#,##0.00);_(* "-"??_);_(@_)'
ACCT0 = '_($* #,##0_);_($* (#,##0);_($* "-"??_);_(@_)'
COMMA0 = '_(* #,##0_);_(* (#,##0);_(* "-"??_);_(@_)'
HC_FMT = '_(* #,##0_);_(* (#,##0);_(* "-"_);_(@_)'


def D(x) -> Decimal:
    return Decimal(str(x)).quantize(Decimal("0.01"))


# ============================================================================== 1. sales tax by state
TAX_PDF = "Sales Tax Collected Q3 2026.pdf"
TAX_COMPANY = "Harborline Supply Co."
TAX_PREPARED = ("R. Delgado", date(2026, 10, 12))
TAX_SOURCE = "Avalara liability report (collected basis)"
TAX_MONTHS = [date(2026, 7, 1), date(2026, 8, 1), date(2026, 9, 1)]
# state: (Jul, Aug, Sep, filing frequency, next return due)
TAX_ROWS = {
    "Arizona": (4_812.37, 5_106.90, 4_977.15, "Monthly", date(2026, 10, 20)),
    "California": (18_442.80, 19_915.32, 21_077.64, "Quarterly", date(2026, 10, 31)),
    "Colorado": (2_214.06, 2_388.71, 2_150.49, "Monthly", date(2026, 10, 20)),
    "Florida": (9_731.25, 8_964.10, 10_218.77, "Monthly", date(2026, 10, 20)),
    "Georgia": (3_560.44, 3_912.08, 3_705.91, "Monthly", date(2026, 10, 20)),
    "Illinois": (11_284.63, 12_047.19, 11_630.02, "Monthly", date(2026, 10, 20)),
    "Minnesota": (1_945.70, 2_031.36, 1_876.58, "Quarterly", date(2026, 10, 20)),
    "Nevada": (1_102.18, 987.45, 1_240.66, "Quarterly", date(2026, 10, 31)),
    "New Jersey": (3_318.27, 3_522.90, -412.30, "Quarterly", date(2026, 10, 20)),
    "New York": (7_903.55, 8_416.02, 8_129.88, "Quarterly", date(2026, 10, 20)),
    "Ohio": (640.12, 702.55, 615.80, "Semi-annual", date(2027, 1, 23)),
    "Texas": (0, 6_284.40, 7_015.93, "Monthly", date(2026, 10, 20)),
    "Wyoming": (212.40, 198.75, 236.10, "Annual", date(2027, 1, 31)),
}
TAX = {s: ([D(j), D(a), D(p)], f, due) for s, (j, a, p, f, due) in TAX_ROWS.items()}
TAX_NOTES = [
    "(a) New Jersey September is net of a $4,180.00 refund of tax on returned order SO-88213.",
    "(b) Texas registration effective 8/1/2026; no July collections.",
]


def sales_tax(ws) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin, double = Side(style="thin"), Side(style="double")
    box = Border(top=thin, bottom=thin, left=thin, right=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    fill = PatternFill("solid", fgColor="833C0B")
    width = 7
    for r, (text, size) in enumerate([(TAX_COMPANY, 14), ("Sales Tax Collected by State", 12),
                                      ("Quarter Ended September 30, 2026", 11)], start=1):
        cell = ws.cell(r, 1, text)
        cell.font = Font(bold=True, size=size)
        cell.alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
    ws.cell(5, 1, "Prepared by:").font = Font(bold=True)
    ws.cell(5, 2, TAX_PREPARED[0])
    ws.cell(5, 5, "Date:").font = Font(bold=True)
    ws.cell(5, 6, TAX_PREPARED[1]).number_format = "m/d/yyyy"
    ws.cell(5, 6).alignment = Alignment(horizontal="left")
    ws.cell(6, 1, "Source:").font = Font(bold=True)
    ws.cell(6, 2, TAX_SOURCE)
    top = 8
    # two-row header: State / Q3 Total / Filing Frequency / Next Return Due span both rows,
    # "Tax Collected" spans the three month columns
    ws.cell(top, 1, "State")
    ws.merge_cells(start_row=top, start_column=1, end_row=top + 1, end_column=1)
    ws.cell(top, 2, "Tax Collected")
    ws.merge_cells(start_row=top, start_column=2, end_row=top, end_column=4)
    for c, m in enumerate(TAX_MONTHS, start=2):
        ws.cell(top + 1, c, m).number_format = "mmm-yy"
    for c, label in ((5, "Q3 2026 Total"), (6, "Filing Frequency"), (7, "Next Return Due")):
        ws.cell(top, c, label)
        ws.merge_cells(start_row=top, start_column=c, end_row=top + 1, end_column=c)
    for rr in (top, top + 1):
        for c in range(1, width + 1):
            cell = ws.cell(rr, c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill
            cell.alignment = center
            cell.border = box
    first = r = top + 2
    for state, (vals, freq, due) in TAX.items():
        label = state + (" (a)" if state == "New Jersey" else " (b)" if state == "Texas" else "")
        ws.cell(r, 1, label).border = box
        fmt = ACCT if r == first else COMMA
        for c, v in enumerate(vals, start=2):
            cell = ws.cell(r, c, float(v))
            cell.number_format = fmt
            cell.border = box
        cell = ws.cell(r, 5, f"=SUM(B{r}:D{r})")
        cell.number_format = fmt
        cell.font = Font(bold=True)
        cell.border = box
        cell = ws.cell(r, 6, freq)
        cell.alignment = Alignment(horizontal="center")
        cell.border = box
        cell = ws.cell(r, 7, due)
        cell.number_format = "m/d/yyyy"
        cell.alignment = Alignment(horizontal="center")
        cell.border = box
        r += 1
    last = r - 1
    ws.cell(r, 1, "Total Sales Tax Collected").font = Font(bold=True)
    for c in range(2, 6):
        col = "ABCDE"[c - 1]
        cell = ws.cell(r, c, f"=SUM({col}{first}:{col}{last})")
        cell.number_format = ACCT
        cell.font = Font(bold=True)
        cell.border = Border(top=thin, bottom=double)
    for k, note in enumerate(TAX_NOTES):
        ws.cell(r + 2 + k, 1, note).font = Font(italic=True, size=9)
    for c, w in enumerate([24, 14, 14, 14, 16, 14, 14], start=1):
        ws.column_dimensions["ABCDEFG"[c - 1]].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.left.text = "&F"
    ws.oddFooter.center.text = "Page &P of &N"
    ws.oddFooter.right.text = "Sales tax payable - GL 2210"


# ============================================================================== 2. headcount & salary by dept x location
HC_PDF = "Headcount Salary Budget FY2027.pdf"
HC_COMPANY = "Bramblewood Software, Inc."
HC_PREPARED = "D. Okafor, FP&A"
HC_ASOF = date(2026, 8, 31)
HC_VERSION = "Board draft v3"
HC_LOCATIONS = ["Chicago", "Austin", "Remote"]
# section -> dept -> {location: (HC, annual base salary $)}
HC_DATA = {
    "General & Administrative": {
        "Finance": {"Chicago": (9, 1_026_500), "Austin": (2, 182_000), "Remote": (1, 118_500)},
        "Human Resources": {"Chicago": (4, 398_250), "Remote": (2, 171_000)},
        "Legal": {"Chicago": (3, 552_000)},
        "IT": {"Chicago": (3, 309_400), "Austin": (4, 386_000)},
    },
    "Research & Development": {
        "Engineering": {"Chicago": (12, 1_812_000), "Austin": (21, 3_024_750), "Remote": (14, 2_051_300)},
        "Product Management": {"Chicago": (4, 664_000), "Austin": (3, 471_500), "Remote": (2, 318_000)},
        "Data Science": {"Austin": (5, 790_000), "Remote": (3, 462_600)},
    },
    "Go-to-Market": {
        "Sales": {"Chicago": (11, 1_243_000), "Austin": (6, 642_000), "Remote": (8, 904_800)},
        "Marketing": {"Chicago": (5, 585_250), "Austin": (2, 214_000), "Remote": (3, 327_000)},
        "Customer Success": {"Chicago": (4, 352_000), "Austin": (7, 581_700), "Remote": (5, 410_000)},
    },
}
HC_SHORT = {"General & Administrative": "G&A", "Research & Development": "R&D", "Go-to-Market": "GTM"}


def headcount(ws) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin, double, medium = Side(style="thin"), Side(style="double"), Side(style="medium")
    box = Border(top=thin, bottom=thin, left=thin, right=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    fill = PatternFill("solid", fgColor="2F5597")
    band = PatternFill("solid", fgColor="D6DCE4")
    width = 9  # A dept, B-G three location pairs, H-I total pair
    for r, (text, size) in enumerate([(HC_COMPANY, 14), ("FY2027 Headcount & Salary Budget", 12),
                                      ("by Department and Location", 11)], start=1):
        cell = ws.cell(r, 1, text)
        cell.font = Font(bold=True, size=size)
        cell.alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
    facts = [("Prepared by:", HC_PREPARED), ("As of:", HC_ASOF), ("Version:", HC_VERSION)]
    for i, (k, v) in enumerate(facts, start=5):
        ws.cell(i, 1, k).font = Font(bold=True)
        cell = ws.cell(i, 2, v)
        if isinstance(v, date):
            cell.number_format = "mmmm d, yyyy"
        cell.alignment = Alignment(horizontal="left")
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=4)
    ws.cell(8, 1, "HC = budgeted FTEs at fiscal year end. Salary = annual base salary only "
                  "(excludes benefits, bonus and payroll taxes).").font = Font(italic=True, size=9)
    top = 10
    ws.cell(top, 1, "Department")
    ws.merge_cells(start_row=top, start_column=1, end_row=top + 2, end_column=1)
    ws.cell(top, 2, "Location")
    ws.merge_cells(start_row=top, start_column=2, end_row=top, end_column=7)
    ws.cell(top, 8, "Total")
    ws.merge_cells(start_row=top, start_column=8, end_row=top + 1, end_column=9)
    for k, loc in enumerate(HC_LOCATIONS):
        c = 2 + 2 * k
        ws.cell(top + 1, c, loc)
        ws.merge_cells(start_row=top + 1, start_column=c, end_row=top + 1, end_column=c + 1)
    for c in range(2, 10, 2):
        ws.cell(top + 2, c, "HC")
        ws.cell(top + 2, c + 1, "Salary")
    for rr in (top, top + 1, top + 2):
        for c in range(1, width + 1):
            cell = ws.cell(rr, c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill
            cell.alignment = center
            cell.border = box
    r = top + 3
    sub_rows = []
    first_data = True
    for section, depts in HC_DATA.items():
        ws.cell(r, 1, section).font = Font(bold=True, italic=True)
        for c in range(1, width + 1):
            ws.cell(r, c).fill = band
        r += 1
        start = r
        for dept, locs in depts.items():
            ws.cell(r, 1, dept).alignment = Alignment(indent=1)
            for k, loc in enumerate(HC_LOCATIONS):
                c = 2 + 2 * k
                hc, sal = locs.get(loc, (0, 0))
                ws.cell(r, c, hc).number_format = HC_FMT
                ws.cell(r, c + 1, sal).number_format = ACCT0 if first_data else COMMA0
            ws.cell(r, 8, f"=B{r}+D{r}+F{r}").number_format = HC_FMT
            ws.cell(r, 9, f"=C{r}+E{r}+G{r}").number_format = ACCT0 if first_data else COMMA0
            for c in (8, 9):
                ws.cell(r, c).font = Font(bold=True)
            first_data = False
            r += 1
        ws.cell(r, 1, f"Total {HC_SHORT[section]}").font = Font(bold=True)
        for c in range(2, width + 1):
            col = get_column_letter(c)
            cell = ws.cell(r, c, f"=SUM({col}{start}:{col}{r - 1})")
            cell.number_format = HC_FMT if c % 2 == 0 else COMMA0
            cell.font = Font(bold=True)
            cell.border = Border(top=thin)
        sub_rows.append(r)
        r += 2
    ws.cell(r, 1, "Total Company").font = Font(bold=True)
    for c in range(2, width + 1):
        col = get_column_letter(c)
        cell = ws.cell(r, c, "=" + "+".join(f"{col}{s}" for s in sub_rows))
        cell.number_format = HC_FMT if c % 2 == 0 else ACCT0
        cell.font = Font(bold=True)
        cell.border = Border(top=medium, bottom=double)
    ws.cell(r + 2, 1, "Remote = employees not assigned to an office; salaries budgeted at the national band midpoint."
            ).font = Font(italic=True, size=9)
    ws.column_dimensions["A"].width = 26
    for c in range(2, width + 1):
        ws.column_dimensions[get_column_letter(c)].width = 7 if c % 2 == 0 else 14
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1  # one page: keep the Total Company row with the table
    ws.oddFooter.left.text = "&F"
    ws.oddFooter.center.text = "Page &P of &N"
    ws.oddFooter.right.text = "CONFIDENTIAL - FP&&A"  # "&&" prints a literal "&" ("&A" is the sheet name)


# ============================================================================== figures expected in the PDF text
def _show(v: Decimal, decimals: int = 2) -> str:
    if v == 0:
        return "-"
    s = f"{abs(v):,.{decimals}f}"
    return f"({s})" if v < 0 else s


def expected() -> dict[str, list[str]]:
    zero = Decimal("0.00")
    tax_tot = {s: sum(v, zero) for s, (v, _, _) in TAX.items()}
    month_tot = [sum((v[k] for v, _, _ in TAX.values()), zero) for k in range(3)]
    tax_fig = ([_show(x) for v, _, _ in TAX.values() for x in v] + [_show(x) for x in tax_tot.values()]
               + [_show(x) for x in month_tot] + [_show(sum(month_tot, zero)), "R. Delgado", "10/12/2026",
                                                    "Semi-annual", "1/31/2027", "1/23/2027"])
    hc_fig = ["D. Okafor, FP&A", "August 31, 2026", "Board draft v3"]
    grand = [0] * 8
    for depts in HC_DATA.values():
        sub = [0] * 8
        for locs in depts.values():
            row = []
            for loc in HC_LOCATIONS:
                row += list(locs.get(loc, (0, 0)))
            row += [sum(row[0::2]), sum(row[1::2])]
            hc_fig += [f"{x:,}" for x in row if x]
            sub = [a + b for a, b in zip(sub, row)]
        hc_fig += [f"{x:,}" for x in sub if x]
        grand = [a + b for a, b in zip(grand, sub)]
    hc_fig += [f"{x:,}" for x in grand]
    return {TAX_PDF: tax_fig, HC_PDF: hc_fig}


def build() -> None:
    from openpyxl import Workbook

    XLSX.mkdir(parents=True, exist_ok=True)
    PDF.mkdir(parents=True, exist_ok=True)
    books = []
    for pdf, fill in ((TAX_PDF, sales_tax), (HC_PDF, headcount)):
        wb = Workbook()
        ws = wb.active
        ws.title = "Sales Tax Q3" if fill is sales_tax else "HC Budget"
        fill(ws)
        path = XLSX / (pdf[:-4] + ".xlsx")
        wb.save(path)
        books.append(path)
    profile = Path(tempfile.mkdtemp(prefix="lo-fresh-"))
    try:
        subprocess.run(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--calc",
                        "--convert-to", "pdf", "--outdir", str(PDF), *map(str, books)],
                       check=True, capture_output=True, timeout=300)
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    for pdf, figures in expected().items():
        text = subprocess.run(["pdftotext", "-layout", str(PDF / pdf), "-"], check=True, capture_output=True,
                              text=True).stdout
        missing = sorted({f for f in figures if f != "-" and f not in text})
        print(f"{pdf}: {len(figures)} figures checked" + (f", MISSING {missing}" if missing else ""))
        assert not missing, (pdf, missing)


if __name__ == "__main__":
    build()
