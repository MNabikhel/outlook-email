"""AP invoice cost coding against the user's own workbook of JDE codes."""

from __future__ import annotations

from datetime import datetime, timezone

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from controller_inbox import cost_codes, web
from controller_inbox.demo import make_pdf
from controller_inbox.folder_mail import ingest_folder
from controller_inbox.models import AttachmentRecord, DocumentType, EmailRecord, Importance
from msgfactory import PDF, write_msg

CODES = [
    ("Office supplies - head office", "1100.6110.100"),
    ("Office supplies - warehouse", "1200.6110.100"),
    ("Freight and delivery", "1100.6420"),
    ("Legal fees", "00100.7250.10"),
    ("Software subscriptions", "1100.6510.200"),
]


def _write_codes(settings, rows, *, title_row: bool = False, numbers: tuple = ()) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Cost codes"
    if title_row:
        sheet.append(["JDE cost codes, Kingston"])
    sheet.append(["Description", "Cost Code"])
    for description, code in rows:
        sheet.append([description, code])
    for description, number in numbers:
        sheet.append([description, number])
    path = cost_codes.workbook_path(settings)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def _invoice(email_id: str, *, body: str = "", text: str = "", sender: str = "billing@staples.example", category=DocumentType.AP_INVOICE) -> EmailRecord:
    attachments = []
    if text:
        attachments.append(
            AttachmentRecord(
                id=f"{email_id}-1",
                email_id=email_id,
                filename="invoice.pdf",
                content_type="application/pdf",
                size_bytes=len(text),
                sha256=email_id,
                extracted_text=text,
                document_type=category,
                document_confidence=0.9,
            )
        )
    return EmailRecord(
        id=email_id,
        subject=f"Invoice {email_id}",
        sender_name="Billing",
        sender_email=sender,
        received_at=datetime(2026, 9, 21, 9, tzinfo=timezone.utc).isoformat(),
        body_text=body,
        body_preview=body[:80],
        has_attachments=bool(attachments),
        outlook_importance="normal",
        is_read=False,
        category=category,
        category_confidence=0.9,
        importance=Importance.MEDIUM,
        importance_score=40,
        attachments=attachments,
    )


def test_a_blank_workbook_is_ready_to_fill_in(settings):
    settings.ensure_data_dir()
    path = cost_codes.workbook_path(settings)
    assert path.exists() and path.parent == settings.inbox_dir.parent / "AP cost codes"
    sheet = load_workbook(path)["Cost codes"]
    assert [cell.value for cell in sheet[1]] == ["Description", "Cost Code"]
    assert sheet["B2"].number_format == "@", "codes must stay text so 1100.6110 keeps its zero"
    assert cost_codes.load(settings).codes == []


def test_codes_are_read_under_their_headings_and_numbers_are_flagged(settings):
    _write_codes(settings, CODES, title_row=True, numbers=(("Rent", 1100.611),))
    book = cost_codes.load(settings)
    assert [code.code for code in book.codes][:5] == [code for _description, code in CODES]
    rent = book.find("1100.611")
    assert rent and rent.loose and rent.description == "Rent"
    assert any("stored as numbers" in warning for warning in book.warnings)


def test_a_code_stamped_on_a_scan_is_found_despite_reading_slips(settings, store):
    _write_codes(settings, CODES, numbers=(("Rent", 1100.611),))
    book = cost_codes.load(settings)
    # A comma for a period, O for zero, the business unit's leading zeros dropped, a code cut short by Excel.
    scan = "INVOICE 4471\nAmount due $1,240.00\nAP coding: 11OO,6110.1OO\nGL 100.7250.10\nRent 1100.6110 per lease"
    found = cost_codes.codes_on(_invoice("a", text=scan), book)
    assert [item["code"] for item in found] == ["1100.6110.100", "00100.7250.10", "1100.611"]
    assert found[0]["where"] == "invoice.pdf" and "11OO,6110.1OO" in found[0]["evidence"]


def test_numbers_that_only_look_like_codes_do_not_match(settings):
    _write_codes(settings, CODES)
    book = cost_codes.load(settings)
    text = "Total $1,100.64\nRef 1100.6110.1001\nPhone 1100.6420.55 ext\nPO 21100.6420\nVersion 1.1006420"
    assert cost_codes.codes_on(_invoice("b", text=text), book) == []


def test_split_coding_keeps_every_code_in_the_order_printed(settings, store):
    _write_codes(settings, CODES)
    email = _invoice("c", text="Coding\n1100.6420 $80.00\n1100.6110.100 $420.00")
    store.upsert_email(email)
    cost_codes.refresh(store, settings)
    coding = store.cost_coding("c")
    assert coding["status"] == "suggested"
    assert [item["code"] for item in coding["codes"]] == ["1100.6420", "1100.6110.100"]


def test_a_code_shaped_number_not_in_the_workbook_is_pointed_out(settings, store):
    _write_codes(settings, CODES)
    store.upsert_email(_invoice("d", text="Coded to 1100.6999.100 and 1100.6110.100"))
    cost_codes.refresh(store, settings)
    coding = store.cost_coding("d")
    assert [item["code"] for item in coding["codes"]] == ["1100.6110.100"]
    assert coding["unlisted"] == ["1100.6999.100"]


def test_without_a_printed_code_the_description_decides(settings, store):
    _write_codes(settings, CODES)
    store.upsert_email(_invoice("e", text="Staples\nPrinter paper, toner and other office supplies\nShip to: Warehouse, Spanish Town"))
    cost_codes.refresh(store, settings)
    coding = store.cost_coding("e")
    assert coding["codes"][0]["code"] == "1200.6110.100"
    assert coding["codes"][0]["source"] == cost_codes.DESCRIPTION
    assert [item["code"] for item in coding["others"]] == ["1100.6110.100"]


def test_a_confirmed_code_follows_the_sender_and_is_never_overwritten(settings, store):
    _write_codes(settings, CODES)
    store.upsert_email(_invoice("f1", text="Monthly licence renewal"))
    store.upsert_email(_invoice("f2", text="Monthly licence renewal"))
    cost_codes.refresh(store, settings)
    assert store.cost_coding("f2")["status"] == "unmatched"

    cost_codes.revise(store, settings, "f1", ["1100.6510.200"])
    first = store.cost_coding("f1")
    assert first["status"] == "confirmed" and first["codes"][0]["source"] == cost_codes.YOU
    second = store.cost_coding("f2")
    assert second["status"] == "suggested"
    assert second["codes"][0]["code"] == "1100.6510.200" and second["codes"][0]["source"] == cost_codes.SENDER_HISTORY

    _write_codes(settings, CODES[:2])
    cost_codes.refresh_if_changed(store, settings)
    assert store.cost_coding("f1")["codes"][0]["code"] == "1100.6510.200"
    with pytest.raises(ValueError):
        cost_codes.revise(store, settings, "f2", ["9999.9999"])


def test_mail_that_stops_being_an_invoice_loses_its_suggestion(settings, store):
    _write_codes(settings, CODES)
    email = _invoice("g", text="1100.6420")
    store.upsert_email(email)
    cost_codes.refresh(store, settings)
    assert store.cost_coding("g")
    store.upsert_email(_invoice("g", text="1100.6420", category=DocumentType.NOTIFICATION))
    cost_codes.refresh(store, settings)
    assert store.cost_coding("g") is None


def test_search_finds_an_invoice_by_its_code_or_description(settings, store):
    _write_codes(settings, CODES)
    store.upsert_email(_invoice("h", text="Your order of office supplies for the warehouse"))
    store.upsert_email(_invoice("i", text="Unrelated", sender="other@vendor.example"))
    cost_codes.refresh(store, settings)
    assert [email.id for email in store.list_emails(q="1200.6110")] == ["h"]
    assert [email.id for email in store.list_emails(q="warehouse 6110")] == ["h"]
    assert [row["email_id"] for row in store.cost_codings(q="1200.6110")] == ["h"]


def test_email_page_confirms_and_revises_the_code(settings, store):
    _write_codes(settings, CODES)
    store.upsert_email(_invoice("j", text="AP stamp: 1100.6420 approved"))
    store.upsert_email(_invoice("k", body="Lunch on Friday?", category=DocumentType.REPLY_NEEDED))
    client = TestClient(web.create_app(settings, store))

    page = client.get("/inbox/j")
    assert "AP coding" in page.text and "Suggested · check and confirm" in page.text
    assert "1100.6420" in page.text and "Freight and delivery" in page.text
    assert 'action="/inbox/j/coding/confirm"' in page.text and "Revise" in page.text
    assert "AP coding</h3>" not in client.get("/inbox/k").text
    assert "AP coding" in client.get("/inbox/j/preview?next=/coding").text

    done = client.post("/inbox/j/coding/confirm", data={"next": "/inbox/j"}, follow_redirects=False)
    assert done.status_code == 303 and done.headers["location"] == "/inbox/j?notice=coding-confirmed#coding"
    assert store.cost_coding("j")["status"] == "confirmed"

    revised = client.post("/inbox/j/coding/revise", data={"codes": ["1100.6110.100", "1100.6420"], "next": "/coding"}, follow_redirects=False)
    assert revised.headers["location"] == "/coding?notice=coding-revised#coding"
    assert [item["code"] for item in store.cost_coding("j")["codes"]] == ["1100.6110.100", "1100.6420"]
    assert client.post("/inbox/j/coding/revise", data={"codes": ["0000.0000"]}).status_code == 400

    listing = client.get("/coding?status=confirmed")
    assert "Office supplies - head office" in listing.text and "Confirmed (1)" in listing.text
    assert "AP coding" in client.get("/settings").text


def test_an_invoice_dropped_in_the_inbox_is_coded_when_processed(settings, store):
    _write_codes(settings, CODES)
    settings.ensure_data_dir()
    pdf = make_pdf([["INVOICE INV-2231", ["Item", "Amount"], ["Freight Kingston to Montego Bay", "$640.00"], ["Amount due", "$640.00"], ["Coding 1100.6420"]]])
    write_msg(
        settings.inbox_incoming / "freight.msg",
        "Invoice INV-2231 for freight",
        "Please find our invoice attached. Payment due in 30 days.",
        sender_name="Island Freight",
        sender_email="ar@islandfreight.example",
        attachments=[("INV-2231.pdf", pdf, PDF)],
    )
    [email] = ingest_folder(store, settings)
    assert cost_codes.is_ap_invoice(email)
    coding = store.cost_coding(email.id)
    assert [item["code"] for item in coding["codes"]] == ["1100.6420"]
    assert coding["codes"][0]["where"] == "INV-2231.pdf"


def test_a_scan_spacing_the_periods_finds_only_the_longer_code(settings):
    _write_codes(settings, [("Office supplies", "1100.6110"), ("Office supplies - head office", "1100.6110.100")])
    book = cost_codes.load(settings)
    for scan in ("Coded 1100.6110. 100", "Coded 1100. 6110. 100", "Coded 1100.6110 .100"):
        assert [item["code"] for item in cost_codes.codes_on(_invoice("s", text=scan), book)] == ["1100.6110.100"], scan
    # The shorter code still counts where it is written on its own, in the order printed.
    both = cost_codes.codes_on(_invoice("t", text="Coded 1100.6110. 100\nFreight 1100.6110"), book)
    assert [item["code"] for item in both] == ["1100.6110.100", "1100.6110"]


def test_a_reading_slip_of_a_listed_code_is_not_called_unlisted(settings):
    _write_codes(settings, CODES)
    book = cost_codes.load(settings)
    email = _invoice("u", text="AP coding: 1100.6110.1OO\nFreight 1100.642O\nNew code 1100.6999.100")
    assert [item["code"] for item in cost_codes.codes_on(email, book)] == ["1100.6110.100", "1100.6420"]
    assert cost_codes.unlisted_codes(email, book) == ["1100.6999.100"]


def test_a_cost_code_description_heading_is_not_taken_for_the_code_column(settings):
    # "Cost Code Description" | "Cost Code": both headings say "code", so the sheet was skipped.
    book = Workbook()
    sheet = book.active
    sheet.append(["Cost Code Description", "Cost Code"])
    sheet.append(["Office supplies - head office", "1100.6110.100"])
    path = cost_codes.workbook_path(settings)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    codes = cost_codes.load(settings).codes
    assert [(c.code, c.description) for c in codes] == [("1100.6110.100", "Office supplies - head office")]


def test_quantities_and_unit_prices_are_not_unlisted_codes():
    # With a 4.4 code shape, "Qty 2.0000  Unit price 1250.0000" was reported as two unlisted codes.
    book = cost_codes.Codebook(codes=[cost_codes.CostCode("1100.6420", "Freight"), cost_codes.CostCode("1100.6110.100", "Office supplies")])
    email = _invoice("q", body="Line 1: Toner  Qty 2.0000  Unit price 1250.0000  Total 2,500.00\nCode 1100.6420")
    assert cost_codes.unlisted_codes(email, book) == []


def test_a_tax_code_column_is_not_taken_for_the_account_column():
    # "Account | Description | Tax Code": the account is the code; the tax code is not.
    assert cost_codes._header_columns(("Account", "Description", "Tax Code")) == (1, 0)
