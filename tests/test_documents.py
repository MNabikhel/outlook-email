"""Attachments come out as text a model can navigate: pages, sheets with cell references, slides."""

from __future__ import annotations

import io

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from openpyxl import Workbook

from controller_inbox.demo import make_pdf
from controller_inbox.documents import (
    Part,
    compare_columns,
    outline,
    pdf_text,
    read_cells,
    read_part,
    search_parts,
    skim,
    split_parts,
    trace_cell,
    xlsx_text,
)
from controller_inbox.extract import extract_text_from_bytes


def _budget() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Budget"
    sheet.append(["Line", "Q3", "Q4", "Change"])
    sheet.append(["Ads", 1000, 1500, "=C2-B2"])
    sheet.append(["Travel", 250.5, 300, "=C3-B3"])
    sheet.append(["Total", "=SUM(B2:B3)", "=SUM(C2:C3)", "=D2+D3"])
    notes = book.create_sheet("Assumptions")
    notes["A1"] = "Growth"
    notes["B1"] = 0.1
    notes.sheet_state = "hidden"
    sheet["E4"] = "=D4*(1+Assumptions!B1)"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_workbook_keeps_cell_references_formulas_and_hidden_sheets():
    text = extract_text_from_bytes("budget.xlsx", "", _budget())
    assert 'Workbook with 2 sheets: "Budget", "Assumptions" (hidden)' in text
    assert '[sheet "Budget" A1:E4]' in text
    assert "A1: Line | B1: Q3 | C1: Q4 | D1: Change" in text
    assert "A2 (Line): Ads | B2 (Q3): 1,000 | C2 (Q4): 1,500 | D2 (Change): =C2-B2" in text
    assert "B4 (Q3): =SUM(B2:B3)" in text
    assert "(4 rows with data, 6 formulas)" in text
    assert '[sheet "Assumptions" A1:B1 hidden]' in text


def test_a_formula_can_be_traced_back_to_its_inputs_across_sheets():
    trace = trace_cell(_budget(), "budget", "E4")
    assert trace.splitlines()[0] == "Budget!E4 ← =D4*(1+Assumptions!B1)"
    assert "  Budget!D4 ← =D2+D3" in trace
    assert "  Assumptions!B1 = 0.1" in trace
    assert "    Budget!D2 ← =C2-B2" in trace


def test_a_trace_skips_other_workbooks_function_names_and_text_and_reads_only_cells_in_use(monkeypatch):
    from openpyxl.styles import Font
    from openpyxl.worksheet.worksheet import Worksheet

    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet["B2"] = 999  # shares its address with the cell in the other workbook
    sheet["A1"] = "=[1]Sheet1!B2*2"
    sheet["A2"] = '=IF(C1>0,"Q1 plan","Q2 plan")'
    sheet["A3"] = "=LOG10(C1)+ATAN2(C1,C2)"
    sheet["A4"] = "=VLOOKUP(5,Data!$A$1:$Z$65536,2,FALSE)"
    sheet["C1"], sheet["C2"] = 5, 7
    data = book.create_sheet("Data")
    for row in range(1, 50):
        data.append([row, row * 10])
    data.cell(1048576, 1).font = Font(bold=True)  # formatted to the bottom of the sheet
    out = io.BytesIO()
    book.save(out)
    made = []
    real = Worksheet.cell
    monkeypatch.setattr(Worksheet, "cell", lambda self, *args, **kwargs: made.append(1) or real(self, *args, **kwargs))
    assert trace_cell(out.getvalue(), "Sheet1", "A1").splitlines() == ["Sheet1!A1 ← =[1]Sheet1!B2*2", "  [1]Sheet1!B2: in another workbook, not in this file"]
    assert trace_cell(out.getvalue(), "Sheet1", "A2").splitlines()[1:] == ["  Sheet1!C1 = 5"]
    assert trace_cell(out.getvalue(), "Sheet1", "A3").splitlines()[1:] == ["  Sheet1!C1 = 5", "  Sheet1!C2 = 7"]
    assert trace_cell(out.getvalue(), "Sheet1", "A4").splitlines()[1] == "  Data!A1:Z65536: 98 filled cells, 98 numbers adding to 13,475"
    assert read_cells(out.getvalue(), "Data", "A48:B1048576").splitlines()[1:] == ["A48: 48 | B48: 480", "A49: 49 | B49: 490"]
    assert len(made) < 1000, "a cell was made for every empty address in the range"


def test_reading_a_range_shows_values_and_formulas():
    cells = read_cells(_budget(), "Budget", "A3:D4")
    assert cells.splitlines() == [
        'Sheet "Budget" A3:D4',
        "A3: Travel | B3: 250.5 | C3: 300 | D3: =C3-B3",
        "A4: Total | B4: =SUM(B2:B3) | C4: =SUM(C2:C3) | D4: =D2+D3",
    ]
    assert "No sheet called 'Forecast'" in read_cells(_budget(), "Forecast", "A1:B2")
    assert "isn't a cell range" in read_cells(_budget(), "Budget", "rows please")


def test_pdf_pages_and_table_columns():
    data = make_pdf(
        [
            ["Northwind Traders - Invoice INV-2231", ["Item", "Qty", "Amount"], ["Consulting", "10", "$4,000.00"], ["Total", "", "$4,350.00"]],
            ["Terms: net 30. Remit to Northwind."],
        ]
    )
    text = extract_text_from_bytes("inv.pdf", "application/pdf", data)
    assert text.startswith("[page 1]\n") and "Northwind Traders - Invoice INV-2231" in text
    assert "Item | Qty | Amount\nItem: Consulting | Qty: 10 | Amount: $4,000.00\nItem: Total | Qty: not listed | Amount: $4,350.00" in text
    assert "[page 2]\nTerms: net 30" in text
    assert read_part(text, "page 2").text.startswith("Terms: net 30")


def test_a_scanned_pdf_says_so_instead_of_looking_empty():
    text = extract_text_from_bytes("scan.pdf", "application/pdf", make_pdf([[], []]))
    assert text.startswith("[This PDF looks scanned")
    assert "[page 2]\n(no text on this page)" in text


def test_word_draft_in_order_with_tracked_changes_and_comments():
    document = Document()
    document.add_heading("Q4 Marketing Plan (DRAFT)", 1)
    paragraph = document.add_paragraph("We propose raising the ads budget to ")
    inserted = OxmlElement("w:ins")
    inserted.set(qn("w:author"), "Maya")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "$18,000"
    run.append(text)
    inserted.append(run)
    paragraph._p.append(inserted)
    deleted = OxmlElement("w:del")
    deleted.set(qn("w:author"), "Maya")
    old = OxmlElement("w:r")
    old_text = OxmlElement("w:delText")
    old_text.text = "$12,000"
    old.append(old_text)
    deleted.append(old)
    paragraph._p.append(deleted)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Item", "Cost"
    table.cell(1, 0).text, table.cell(1, 1).text = "Agency", "$6,000"
    document.add_paragraph("Next steps come after the table.")
    document.add_comment(document.paragraphs[1].runs, text="Has finance approved this?", author="Priya")
    out = io.BytesIO()
    document.save(out)
    result = extract_text_from_bytes("plan.docx", "", out.getvalue())
    lines = result.splitlines()
    assert lines[0].startswith("[This draft has tracked changes")
    assert "# Q4 Marketing Plan (DRAFT)" in lines
    assert "We propose raising the ads budget to $18,000[deleted: $12,000]" in lines
    assert lines.index("Agency | $6,000") < lines.index("Next steps come after the table.")
    assert "- Priya: Has finance approved this?" in lines


def test_csv_reads_like_a_sheet():
    text = extract_text_from_bytes("vendors.csv", "text/csv", "Vendor;Amount\nAcme;100\nGlobex;250\n".encode())
    assert text.splitlines() == [
        '[sheet "vendors.csv" A1:B3]', "A1: Vendor | B1: Amount", "A2 (Vendor): Acme | B2 (Amount): 100", "A3 (Vendor): Globex | B3 (Amount): 250"
    ]
    pairs = extract_text_from_bytes("terms.csv", "text/csv", b"Amount due,12480.00\nPaid,0\n").splitlines()
    assert pairs[1:] == ["A1: Amount due | B1: 12480.00", "A2: Paid | B2: 0"], "a list of values isn't a header"


def test_a_file_called_a_workbook_is_read_as_what_it_is():
    # Windows sends every .csv as application/vnd.ms-excel; only a real 97-2003 workbook goes to xlrd.
    data = b'Vendor,Invoice,Amount\nAcme,INV-1001,"1,250.00"\nGlobex,INV-1002,980.00\n'
    lines = extract_text_from_bytes("payments.csv", "application/vnd.ms-excel", data).splitlines()
    assert lines[0] == '[sheet "payments.csv" A1:C3]'
    assert lines[2] == "A2 (Vendor): Acme | B2 (Invoice): INV-1001 | C2 (Amount): 1,250.00"
    page = (
        b'<html xmlns:o="urn:schemas-microsoft-com:office:office"><head><style>td {color: red}</style></head><body>'
        b"<table><tr><th>Vendor</th><th>Invoice</th><th>Amount</th></tr>"
        b"<tr><td>Acme &amp; Sons</td><td>INV-1001</td><td>1,250.00</td></tr>"
        b"<tr><td colspan=2>Total</td><td>1,250.00</td></tr></table></body></html>"
    )
    lines = extract_text_from_bytes("export.xls", "application/vnd.ms-excel", page).splitlines()
    assert lines[2:] == [
        "A2 (Vendor): Acme & Sons | B2 (Invoice): INV-1001 | C2 (Amount): 1,250.00",
        "A3 (Vendor): Total | B3 (Invoice): not listed | C3 (Amount): 1,250.00",
    ]
    xml = (
        b'<?xml version="1.0"?><Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" '
        b'xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet"><Worksheet ss:Name="AP"><Table>'
        b'<Row><Cell><Data ss:Type="String">Vendor</Data></Cell><Cell><Data ss:Type="String">Invoice</Data></Cell>'
        b'<Cell><Data ss:Type="String">Amount</Data></Cell></Row>'
        b'<Row><Cell><Data ss:Type="String">Acme</Data></Cell><Cell ss:Index="3"><Data ss:Type="Number">1250</Data></Cell></Row>'
        b"</Table></Worksheet></Workbook>"
    )
    assert "A2 (Vendor): Acme | B2 (Invoice): not listed | C2 (Amount): 1250" in extract_text_from_bytes("export.xls", "", xml)
    renamed = extract_text_from_bytes("budget.xls", "application/vnd.ms-excel", _budget())
    assert "A2 (Line): Ads | B2 (Q3): 1,000" in renamed


def test_excel_97_errors_and_yes_no_cells_are_not_read_as_figures():
    import xlrd
    from xlrd.sheet import Cell

    from controller_inbox.documents import _fmt, _xls_value

    def shown(ctype: int, value) -> str:
        return _fmt(_xls_value(Cell(ctype, value), 0))

    assert shown(xlrd.XL_CELL_ERROR, 0x2A) == "#N/A" and shown(xlrd.XL_CELL_ERROR, 0x07) == "#DIV/0!"
    assert shown(xlrd.XL_CELL_BOOLEAN, 1) == "TRUE" and shown(xlrd.XL_CELL_BOOLEAN, 0) == "FALSE"
    assert shown(xlrd.XL_CELL_NUMBER, 42.0) == "42" and shown(xlrd.XL_CELL_DATE, 46295.0) == "2026-09-30"


def test_a_csv_quote_that_never_closes_does_not_swallow_the_file():
    from controller_inbox.documents import csv_text

    for repeat in (2, 12_000):  # past 128 KB the csv module raises instead
        data = 'Item,Qty,Price\n"Monitor 27,1,249.99\n' + "Cable,5,9.99\n" * repeat + 'Desk,"1",120.00\n'
        lines = csv_text(data.encode(), "order.csv").splitlines()
        assert lines[0] == f'[sheet "order.csv" A1:C{repeat + 3}]'
        assert lines[2:4] == ['A2 (Item): "Monitor 27 | B2 (Qty): 1 | C2 (Price): 249.99', "A3 (Item): Cable | B3 (Qty): 5 | C3 (Price): 9.99"]
        assert lines[-1] == ("A5 (Item): Desk | B5 (Qty): 1 | C5 (Price): 120.00" if repeat == 2 else "[11003 more rows not shown.]")
    closed = csv_text(b'Item,Note,Price\nDesk,"two\nlines",120.00\nLamp,"one, with a comma",30.00\n', "order.csv")
    assert "B2 (Note): two lines" in closed and "B3 (Note): one, with a comma" in closed
    # The stray quote closes further on at a quote that ends a cell ('TV 55"'): still a stray quote.
    late = 'Item,Qty,Price\n"Monitor 27,1,249.99\n' + "Cable,5,9.99\n" * 8 + 'TV 55",1,499.00\nDesk,1,120.00\n'
    assert "A12 (Item): Desk | B12 (Qty): 1 | C12 (Price): 120.00" in csv_text(late.encode(), "order.csv")


def test_a_csv_address_on_several_lines_in_its_quotes_is_one_cell():
    # A vendor list's remit-to addresses: each line of an address has a comma, as many as the file's rows do.
    from controller_inbox.documents import csv_text

    data = (
        b'Vendor,Remit To\nHarbor Steel LLC,"PO Box 1200\nSuite 4, Building B\nSpringfield, IL 62701"\n'
        b'Acme Supply,"18 Main St\nFloor 2, Unit 9\nDayton, OH 45402"\n'
    )
    assert csv_text(data, "vendors.csv").splitlines() == [
        '[sheet "vendors.csv" A1:B3]',
        "A1: Vendor | B1: Remit To",
        "A2 (Vendor): Harbor Steel LLC | B2 (Remit To): PO Box 1200 Suite 4, Building B Springfield, IL 62701",
        "A3 (Vendor): Acme Supply | B3 (Remit To): 18 Main St Floor 2, Unit 9 Dayton, OH 45402",
    ]


def test_a_csv_of_names_only_still_names_its_columns():
    data = "Employee,Department,Manager\nJonathan Reyes,,Priya Raman\nLi Wei,Finance,Dana Cole\n"
    lines = extract_text_from_bytes("staff.csv", "text/csv", data.encode()).splitlines()
    assert lines[2] == "A2 (Employee): Jonathan Reyes | B2 (Department): not listed | C2 (Manager): Priya Raman"


def test_csv_rows_under_a_header_name_their_columns_and_blanks():
    data = "Employee,Department,Manager,Salary\nMaya Chen,Finance,Priya Raman,98500\nJonathan Alvarez,,Priya Raman,112000\n"
    lines = extract_text_from_bytes("staff.csv", "text/csv", data.encode()).splitlines()
    assert lines[1] == "A1: Employee | B1: Department | C1: Manager | D1: Salary"
    assert lines[3] == "A3 (Employee): Jonathan Alvarez | B3 (Department): not listed | C3 (Manager): Priya Raman | D3 (Salary): 112000"


def _roster_book() -> bytes:
    from openpyxl.styles import Font

    book = Workbook()
    sheet = book.active
    sheet.title = "Staff"
    sheet["A1"] = "Staff roster, September"
    for column, label in enumerate(["Employee", "ID", "Department", "Manager", "Notes"], start=1):
        sheet.cell(3, column, label).font = Font(bold=True)
    sheet.append(["Maya Chen", "E-1001", "Finance", "Priya Raman", "Hybrid"])
    sheet.append(["Jonathan Alvarez", "E-1002", None, "Priya Raman"])
    sheet.append(["Sofia Rossi", "E-1003", "Operations"])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_workbook_rows_under_a_bold_header_name_their_columns():
    lines = extract_text_from_bytes("roster.xlsx", "", _roster_book()).splitlines()
    assert "A1: Staff roster, September" in lines
    assert "A3: Employee | B3: ID | C3: Department | D3: Manager | E3: Notes" in lines
    assert "A5 (Employee): Jonathan Alvarez | B5 (ID): E-1002 | C5 (Department): not listed | D5 (Manager): Priya Raman" in lines
    assert "A6 (Employee): Sofia Rossi | B6 (ID): E-1003 | C6 (Department): Operations" in lines


def test_labelled_sheet_parts_and_citations_still_find_their_cells():
    from controller_inbox.answer_check import cells_of

    text = extract_text_from_bytes("roster.xlsx", "", _roster_book())
    cells = cells_of(text)
    assert cells[("Staff", "C4")] == "Finance" and cells[("Staff", "C5")] == "not listed"
    rows = "\n".join(f"A{r} (Employee): Person {r} | B{r} (Notes): " + "x" * 40 for r in range(2, 400))
    parts = split_parts('[sheet "Staff" A1:B400]\nA1: Employee | B1: Notes\n' + rows, size=2000)
    assert parts[1].label.startswith('sheet "Staff" rows ')


def test_a_citation_finds_its_section_and_cell():
    from controller_inbox.documents import locate

    parts = split_parts(extract_text_from_bytes("roster.xlsx", "", _roster_book()) + "\n")
    assert locate(parts, "Staff!C5") == (1, "A5 (Employee): Jonathan Alvarez | B5 (ID): E-1002 | C5 (Department): not listed | D5 (Manager): Priya Raman")
    assert locate(parts, "c5")[0] == 1
    assert locate(parts, 'sheet "Staff"') == (1, "")
    pages = split_parts("[page 1]\nIntro\n[page 2]\nTerms\n[page 10]\nAppendix")
    assert locate(pages, "page 1") == (0, "") and locate(pages, "Page 10") == (2, "")
    assert locate(pages, "page 4") is None and locate(pages, "") is None


def _text_box(*lines: str):
    """A run holding a text box the way Word 2010 and later save it: as a drawing, then again as VML for older Word."""
    from docx.oxml import parse_xml

    box = "".join(f"<w:p><w:r><w:t>{line}</w:t></w:r></w:p>" for line in lines)
    return parse_xml(
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        'xmlns:v="urn:schemas-microsoft-com:vml"><mc:AlternateContent>'
        '<mc:Choice Requires="wps"><w:drawing><wp:anchor><a:graphic><a:graphicData><wps:wsp><wps:txbx>'
        f"<w:txbxContent>{box}</w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>"
        f"<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent>{box}</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback>"
        "</mc:AlternateContent></w:r>"
    )


def test_a_word_text_box_is_read_once_a_line_at_a_time():
    document = Document()
    document.add_paragraph("Invoice INV-2001")
    document.add_paragraph("See box").runs[0]._r.addnext(_text_box("Bill To:", "Acme Corp", "Amount due: $4,500.00"))
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Remit to"
    table.cell(0, 1).paragraphs[0]._p.append(_text_box("First Bank", "Account ending 4471"))
    out = io.BytesIO()
    document.save(out)
    text = extract_text_from_bytes("invoice.docx", "", out.getvalue())
    assert "See box\nBill To:\nAcme Corp\nAmount due: $4,500.00\n" in text
    assert text.count("Acme Corp") == 1 and text.count("First Bank") == 1
    assert "Remit to | First Bank Account ending 4471" in text


def test_word_table_with_a_header_row_merged_cells_and_blanks():
    document = Document()
    table = document.add_table(rows=6, cols=4)
    for column, label in enumerate(["Employee", "Department", "Manager", "Start date"]):
        table.cell(0, column).paragraphs[0].add_run(label).bold = True
    table.cell(1, 0).text, table.cell(1, 1).text, table.cell(1, 2).text, table.cell(1, 3).text = "Maya Chen", "Finance", "Priya Raman", "2021-03-01"
    table.cell(2, 0).text, table.cell(2, 3).text = "Jonathan Alvarez", "2019-07-15"
    table.cell(3, 0).text, table.cell(3, 1).text = "Sofia Rossi", "Operations"
    table.cell(2, 2).merge(table.cell(3, 2)).paragraphs[0].text = "Priya Raman"
    table.cell(4, 0).merge(table.cell(4, 3)).text = "Contractors"
    table.cell(5, 0).text, table.cell(5, 1).text, table.cell(5, 3).text = "Omar Haddad", "Finance", "2026-01-05"
    document.add_paragraph("After the table.")
    out = io.BytesIO()
    document.save(out)
    lines = extract_text_from_bytes("staff.docx", "", out.getvalue()).splitlines()
    assert lines[0] == "Employee | Department | Manager | Start date"
    assert "Employee: Jonathan Alvarez | Department: not listed | Manager: Priya Raman | Start date: 2019-07-15" in lines
    assert "Employee: Sofia Rossi | Department: Operations | Manager: Priya Raman | Start date: not listed" in lines
    assert lines.index("Contractors") == lines.index("Employee: Omar Haddad | Department: Finance | Manager: not listed | Start date: 2026-01-05") - 1
    assert lines[-1] == "After the table."


def test_a_word_row_that_starts_further_right_keeps_its_figures_in_their_columns(monkeypatch):
    from controller_inbox import tables

    document = Document()
    table = document.add_table(rows=0, cols=4)
    for values in (["Description", "Q1", "Q2", "Q3"], ["Rent", "3,000", "3,100", "3,200"], ["Total", "3,450", "3,570", "3,720"], ["Notes", "", "", ""]):
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    total = table.rows[2]._tr
    total.remove(total.tc_lst[0])
    skipped = OxmlElement("w:gridBefore")
    skipped.set(qn("w:val"), "1")
    total.get_or_add_trPr().append(skipped)
    # A span with no width is one column; one wider than Word allows is cut to Word's 63.
    table.rows[1]._tr.tc_lst[0].get_or_add_tcPr().append(OxmlElement("w:gridSpan"))
    notes = table.rows[3]._tr
    for cell in notes.tc_lst[1:]:
        notes.remove(cell)
    wide = OxmlElement("w:gridSpan")
    wide.set(qn("w:val"), "99999")
    notes.tc_lst[0].get_or_add_tcPr().append(wide)
    widths = []
    real = tables.table_lines
    monkeypatch.setattr(tables, "table_lines", lambda grid, **kw: widths.append(max(map(len, grid))) or real(grid, **kw))
    out = io.BytesIO()
    document.save(out)
    lines = extract_text_from_bytes("budget.docx", "", out.getvalue()).splitlines()
    assert "Description: Rent | Q1: 3,000 | Q2: 3,100 | Q3: 3,200" in lines
    assert "Description: not listed | Q1: 3,450 | Q2: 3,570 | Q3: 3,720" in lines
    assert "Notes" in lines and widths == [63]


def test_a_lone_figure_in_a_table_row_keeps_its_column():
    from controller_inbox.tables import labelled_row

    document = Document()
    table = document.add_table(rows=0, cols=4)
    for values in (["Description", "Q1", "Q2", "Q3"], ["Rent", "3,000", "3,000", "3,000"], ["Utilities", "450", "", "520"], ["", "", "", "12,970"]):
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    out = io.BytesIO()
    document.save(out)
    lines = extract_text_from_bytes("costs.docx", "", out.getvalue()).splitlines()
    assert lines[-1] == "Description: not listed | Q1: not listed | Q2: not listed | Q3: 12,970"
    assert labelled_row(["Employee", "Department", "Manager"], ["Contractors", "", ""]) == "Contractors"


def test_powerpoint_table_rows_name_their_columns():
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Headcount"
    shape = slide.shapes.add_table(3, 3, Inches(1), Inches(2), Inches(6), Inches(2))
    for (r, c), value in {(0, 0): "Team", (0, 1): "Lead", (0, 2): "Open roles", (1, 0): "Finance", (1, 1): "Priya Raman", (1, 2): "2", (2, 0): "Sales", (2, 2): "1"}.items():
        shape.table.cell(r, c).text = value
    out = io.BytesIO()
    deck.save(out)
    text = extract_text_from_bytes("headcount.pptx", "", out.getvalue())
    assert "Team | Lead | Open roles\nTeam: Finance | Lead: Priya Raman | Open roles: 2\nTeam: Sales | Lead: not listed | Open roles: 1" in text


def test_long_documents_split_into_labelled_parts_and_search_finds_the_right_one():
    pages = [[f"Section {n} filler text about the quarterly close." for _ in range(3)] for n in range(1, 6)]
    pages[3].append("The accrual for Globex consulting is $48,200 and needs sign-off.")
    text = extract_text_from_bytes("pack.pdf", "application/pdf", make_pdf(pages))
    assert [part.label for part in split_parts(text)] == [f"page {n}" for n in range(1, 6)]
    assert [part.label for part in search_parts(text, "Globex accrual")] == ["page 4"]
    assert outline(text)[0] == "page 1: Section 1 filler text about the quarterly close."


def test_a_skim_keeps_what_differs_from_page_to_page():
    pages = [[f"Audit pack page {n}"] + [f"Section {n}.{i}: area reviewed, no exceptions noted in the sample." for i in range(8)] for n in range(1, 21)]
    pages[16][3] = "FINDING 4 (HIGH): bank-detail changes approved without a call-back. Owner: AP lead."
    pages[19][2] = "Total sampled: $4,812,300; exceptions: $61,450."
    parts = split_parts(extract_text_from_bytes("audit.pdf", "application/pdf", make_pdf(pages)))
    text = skim(parts, 600, tag="audit.pdf · ")
    assert text == (
        "[audit.pdf · page 17]\nFINDING 4 (HIGH): bank-detail changes approved without a call-back. Owner: AP lead.\n"
        "[audit.pdf · page 20]\nTotal sampled: $4,812,300; exceptions: $61,450."
    )
    varied = [Part(f"page {n}", f"Heading {chr(64 + n)} topic\nDetail line for {chr(64 + n)} alone") for n in range(1, 9)]
    covered = skim(varied, 180)
    assert "[page 7]\nHeading G topic" in covered and "[page 8]" not in covered and "Detail line" not in covered
    assert len(covered) <= 180


def test_compare_columns_ranks_each_row_and_keeps_totals_apart():
    book = Workbook()
    sheet = book.active
    sheet.title = "Detail"
    sheet.append(["Dept", "Line", "Q3 actual", "Q4 budget"])
    for row in (["Mkt", "Trade show", 18000, 36500], ["Sales", "Travel", 31000, 29500], ["Ops", "Lease", 45000, 45000], ["", "Total", 94000, 111000]):
        sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    by_header = compare_columns(out.getvalue(), "detail", "Q3", "Q4 budget")
    assert by_header.splitlines() == [
        'Sheet "Detail": C “Q3 actual” → D “Q4 budget”, 3 rows, biggest increase first',
        "Mkt · Trade show (row 2): 18,000 → 36,500, +18,500 (+102.8%)",
        "Ops · Lease (row 4): 45,000 → 45,000, +0 (+0.0%)",
        "Sales · Travel (row 3): 31,000 → 29,500, −1,500 (−4.8%)",
        "Biggest increase: Mkt · Trade show. Biggest decrease: Sales · Travel.",
        "Total rows:",
        "Total (row 5): 94,000 → 111,000, +17,000 (+18.1%)",
    ]
    assert compare_columns(out.getvalue(), "", "C", "D") == by_header
    assert compare_columns(out.getvalue(), "", "Q2", "Q4").startswith("No column called 'Q2'. Give column letters")
    assert "saved without values" in compare_columns(_budget(), "Budget", "Change", "D")


def test_percentages_and_months_read_as_the_sheet_shows_them():
    from datetime import datetime

    from openpyxl.styles import Font

    book = Workbook()
    sheet = book.active
    sheet.title = "Fees"
    sheet.append(["Line", datetime(2026, 3, 1), datetime(2026, 4, 1), "Rate", "Paid"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet["B1"].number_format, sheet["C1"].number_format = "mmm-yy", '[$-409]mmm\\-yy;@'
    sheet.append(["Rent", 3000, 3100, 0.15, datetime(2026, 3, 15)])
    sheet.append(["Fees", 400, 410, 0.0525, datetime(2026, 4, 2)])
    sheet.append(["Period", datetime(2026, 9, 1)])
    sheet["D2"].number_format, sheet["D3"].number_format, sheet["E2"].number_format = "0%", '0.00%;[Red]-0.00%', "d-mmm-yy"
    sheet["B4"].number_format = 'mmmm" "yyyy'
    out = io.BytesIO()
    book.save(out)
    lines = extract_text_from_bytes("fees.xlsx", "", out.getvalue()).splitlines()
    assert "A1: Line | B1: Mar-26 | C1: Apr-26 | D1: Rate | E1: Paid" in lines
    assert "A2 (Line): Rent | B2 (Mar-26): 3,000 | C2 (Apr-26): 3,100 | D2 (Rate): 15% | E2 (Paid): 2026-03-15" in lines
    assert "A3 (Line): Fees | B3 (Mar-26): 400 | C3 (Apr-26): 410 | D3 (Rate): 5.25% | E3 (Paid): 2026-04-02" in lines
    assert "A4 (Line): Period | B4 (Mar-26): September 2026" in lines


def test_big_sheets_are_cut_into_row_ranges():
    book = Workbook()
    sheet = book.active
    sheet.title = "Ledger"
    for n in range(1, 301):
        sheet.append([f"Line {n}", n * 10, f"note {n}"])
    out = io.BytesIO()
    book.save(out)
    text = extract_text_from_bytes("ledger.xlsx", "", out.getvalue())
    labels = [part.label for part in split_parts(text)]
    assert labels[0] == "start"
    assert labels[1].startswith('sheet "Ledger" rows 1–') and len(labels) > 3
    last = split_parts(text)[-1]
    assert "A300: Line 300" in last.text
    assert read_part(text, "Ledger").label == labels[1]


def _rewritten(data: bytes, member: str, change) -> bytes:
    """An Office file (a zip) with one part's XML changed."""
    import zipfile

    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(out, "w") as target:
        for item in source.infolist():
            body = source.read(item)
            target.writestr(item, change(body) if item.filename == member else body)
    return out.getvalue()


def test_a_sheet_that_understates_its_size_is_read_to_its_last_row():
    import re

    book = Workbook()
    sheet = book.active
    sheet.title = "AP"
    sheet.append(["Vendor", "Invoice", "Amount"])
    for n in range(5):
        sheet.append([f"Vendor {n}", f"INV-{100 + n}", 1000.5 + n])
    out = io.BytesIO()
    book.save(out)
    # Writers other than Excel often store <dimension ref="A1"/> whatever the sheet holds.
    data = _rewritten(out.getvalue(), "xl/worksheets/sheet1.xml", lambda xml: re.sub(rb'<dimension ref="[^"]*"/>', b'<dimension ref="A1"/>', xml))
    text = extract_text_from_bytes("ap.xlsx", "", data)
    assert '[sheet "AP" A1:C6]\n(6 rows with data)' in text
    assert "A6 (Vendor): Vendor 4 | B6 (Invoice): INV-104 | C6 (Amount): 1,004.5" in text
    unsized = _rewritten(out.getvalue(), "xl/worksheets/sheet1.xml", lambda xml: re.sub(rb'<dimension ref="[^"]*"/>', b"", xml))
    assert '[sheet "AP" A1:C6]' in extract_text_from_bytes("ap.xlsx", "", unsized)


def _scanned_pdf(lines: list[str]) -> bytes:
    from PIL import Image, ImageDraw, ImageFont

    page = Image.new("L", (1240, 1754), 255)
    draw = ImageDraw.Draw(page)
    # Pillow's own font on every OS: it is set tight, like the scans that make OCR drop spaces.
    font = ImageFont.load_default(size=34)
    for n, line in enumerate(lines):
        draw.text((100, 150 + n * 70), line, fill=0, font=font)
    out = io.BytesIO()
    page.save(out, "PDF", resolution=150)
    return out.getvalue()


def test_scanned_pdf_pages_are_read_with_ocr():
    from controller_inbox import ocr

    if not ocr.engine_name():
        pytest.skip("no OCR engine installed")
    text = pdf_text(_scanned_pdf(["INVOICE 4471", "Northwind Supply Co.", "Amount due: $12,480.00", "Due date: 15 October 2026"]))
    assert text.startswith("[Scanned page 1 read with OCR")
    assert "$12,480.00" in text and "Northwind Supply Co." in text and "15 October 2026" in text


def test_scanned_pdf_without_ocr_says_how_to_add_it(monkeypatch):
    from controller_inbox import ocr

    monkeypatch.setattr(ocr, "engine_name", lambda: "")
    text = pdf_text(_scanned_pdf(["Amount due: $12,480.00"]))
    assert "looks scanned" in text and 'pip install -e ".[ocr]"' in text and "12,480" not in text


def test_a_landscape_scan_keeps_its_ocr_when_a_strip_cant_be_read_again():
    """RapidOCR fits a picture to 2,000 pixels; a thin full-width strip of a landscape page (read again for a missed
    heading) then shrinks to nothing and raises. That used to throw away the whole page's reading."""
    from PIL import Image

    from controller_inbox import ocr

    def refuses(*_a, **_k):
        raise ValueError("ResizeImgError")

    assert ocr._engine_lines(refuses, Image.new("RGB", (3302, 24), "white")) == []
    if ocr.engine_name() != "RapidOCR":
        pytest.skip("RapidOCR not installed")
    from rapidocr_onnxruntime import RapidOCR

    wide = Image.new("RGB", (1651, 1275), "white")
    assert ocr._read_gap(RapidOCR(), wide, 100, 108, wide.height) == []


def test_cells_on_one_baseline_stay_on_one_line():
    from controller_inbox.ocr import _rows

    # Tops differ by a few pixels, the way a scan reads a table cell and its label.
    hits = [
        (97, 697, "2023"),
        (103, 61, "（in millions except per share data）"),
        (113, 56, "GAAP basis(1):"),
        (631, 56, "Book value per share (3)"),
        (636, 722, "259.34"),
        (636, 1029, "259.34"),
        (638, 659, "$"),
        (660, 657, "$"),
        (661, 742, "5.00"),
        (661, 900, "4.88"),
        (661, 1041, "15.00"),
        (661, 1200, "14.64"),
        (662, 60, "Cash dividends declared and paid per share"),
    ]
    lines = _rows(hits)
    book = next(line for line in lines if "Book value" in line)
    dividends = next(line for line in lines if "Cash dividends" in line)
    assert book.index("Book value") < book.index("259.34")
    assert "5.00" not in book and "4.88" not in book
    assert dividends.index("Cash dividends") < dividends.index("5.00")
    assert "5.00" in dividends and "4.88" in dividends and "15.00" in dividends and "14.64" in dividends
    assert all("GAAP" not in line or "in millions" not in line for line in lines)


def test_wide_gaps_and_glued_words_become_readable():
    from controller_inbox.ocr import _open_gaps, _polish

    def boxes(chars: str, gaps: list[float]) -> list:
        x = 0.0
        made = []
        for index, _char in enumerate(chars):
            if index:
                x += 8 + gaps[index - 1]
            made.append([[x, 0], [x + 8, 0], [x + 8, 14], [x, 14]])
        return made

    assert _open_gaps(list("YourCityAZ12345"), boxes("YourCityAZ12345", [1, 1, 1, 9, 1, 1, 1, 10, 1, 10, 1, 1, 1, 1])) == "Your City AZ 12345"
    assert _open_gaps(list("ACC #12341234"), boxes("ACC #12341234", [0, 0, 0, 0, 5, 0, 0, 1, 5, 0, 1, 0])) == "ACC # 1234 1234"
    assert _open_gaps(list("OrderNumber"), boxes("OrderNumber", [0, 1, 0, 0, 5, 5, 0, 1, 0, 1])) == "OrderNumber"
    assert _polish("OrderNumber") == "Order Number"
    assert _polish("Totalrevenue") == "Total revenue"
    assert _polish("Thisisasampledescription..") == "This is a sample description..."
    assert _polish("Cashdividendsdeclaredandpaidpershare") == "Cash dividends declared and paid per share"
    assert _polish("INV-3337") == "INV-3337"
    assert _polish("within 30days from date of invoice.Late payment of5%") == "within 30 days from date of invoice. Late payment of 5%"
    assert _polish("income(expense),less net") == "income (expense), less net"
    assert _polish("EXECUTIVESUMMARY") == "EXECUTIVE SUMMARY"
    assert _polish("September 30,2023") == "September 30, 2023"
    from controller_inbox.ocr import _already
    assert not _already(
        "Three Months Ended September 30, 2023 Compared with Three Months Ended September 30, 2022",
        ["threemonthsended"],
    )
    assert _already("Total revenue", ["totalrevenue"])
    assert _polish("$12,480.00") == "$12,480.00"
    assert _polish("（in millions except per share data）") == "(in millions except per share data)"
    assert _polish("28.412,414") == "28,412,414"
    assert _polish("(1.097.978)") == "(1,097,978)"
    assert _polish("7.792") == "7,792"
    assert _polish("0.61") == "0.61"
    assert _polish("$12,480.00") == "$12,480.00"
    assert _polish("JS'O00") == "JS'000"
    # Names are not in any word list: a capital marks where one starts, and two unknown capitalised pieces stay one name.
    assert _polish("NetincomeattributabletoNorthwind") == "Net income attributable to Northwind"
    assert _polish("TotalBlackRockstockholders'equity") == "Total BlackRock stockholders'equity"
    assert _polish("NetProfitattributableto:") == "Net Profit attributable to:"
    assert _polish("PayPal") == "PayPal" and _polish("McDonald") == "McDonald" and _polish("Chartwell") == "Chartwell"
    from controller_inbox.ocr import _rows
    assert _rows([(10, 1, "From continuing"), (12, 700, "0.61"), (11, 800, "S"), (12, 880, "0.78")]) == [
        "From continuing 0.61 $ 0.78"
    ]
    from controller_inbox.ocr import _footnote_marker

    assert _footnote_marker(["(1)", "(2) As adjusted items are described in more detail"]) == "(1)"
    assert _footnote_marker(["(2) As adjusted items are described"]) == "(2)"
    invoices = list("InvoicesI ")
    gaps = [0] * (len(invoices) - 1)
    gaps[7] = 7  # the gap before the bar that was read as I
    assert "Invoices |" in _open_gaps(invoices, boxes(invoices, gaps))
    word = list(" Invoices")
    assert "I" in _open_gaps(word, boxes(word, [0] * (len(word) - 1)))


def test_a_missing_cell_stays_in_its_own_column():
    """One amount on a row must not slide into the other year's column."""
    from controller_inbox.ocr import _rows

    def cell(top, left, right, text):
        return (top, left, right, text)

    hits = [cell(30, 150, 200, "2018"), cell(30, 270, 320, "2019")]
    rows = [
        ("Revenue", "24,544,049", "28,412,414"),
        ("Cost of sales", "(15,421,144)", "(17,878,208)"),
        ("Gross Profit", "9,122,905", "10,534,206"),
        ("Share of results of associates", "", "7,792"),
        ("Profit from discontinued operations", "41,555", ""),
    ]
    for index, (label, left, right) in enumerate(rows):
        top = 80 + index * 24
        hits.append(cell(top, 10, 140, label))
        if left:
            hits.append(cell(top, 150, 200, left))
        if right:
            hits.append(cell(top, 270, 320, right))
    # A different shape: four columns, and a sentence that is not a row of the table.
    hits.append(cell(400, 10, 500, "The audit found no exceptions in the sample."))
    lines = _rows(hits)
    text = "\n".join(lines)
    assert "2018: 24,544,049" in text and "2019: 28,412,414" in text
    assert "2018: not listed | 2019: 7,792" in text
    assert "2018: 41,555 | 2019: not listed" in text
    assert "The audit found no exceptions in the sample." in text

    wide = []
    for index, values in enumerate(
        (
            ("4,522", "4,311", "13,228", "13,536"),
            ("2,885", "2,785", "8,538", "8,578"),
            ("1,637", "1,526", "4,690", "4,958"),
            ("", "", "15.00", "14.64"),
        )
    ):
        top = 80 + index * 22
        wide.append(cell(top, 10, 80, f"Row {index}"))
        for number, value in enumerate(values):
            if value:
                left = 200 + number * 120
                wide.append(cell(top, left, left + 50, value))
    short = next(line for line in _rows(wide) if "Row 3" in line)
    assert "15.00" in short and "14.64" in short and "not listed" in short
    assert short.index("not listed") < short.index("15.00")


def test_scanned_invoice_lines_and_statement_dashes_keep_their_columns():
    from controller_inbox.ocr import _rows

    def cell(top, left, right, text):
        return (top, left, right, text)

    invoice = [cell(20, 10, 120, "Description"), cell(20, 300, 330, "Qty"), cell(20, 400, 480, "Unit Price"), cell(20, 560, 620, "Amount")]
    for index, (item, qty, price, amount) in enumerate(
        [("Paper A4", "10", "4.50", "45.00"), ("Toner", "2", "80.00", "160.00"), ("Shipping", "1", "12.00", "12.00"), ("Discount", "", "", "(5.00)")]
    ):
        top = 60 + index * 24
        invoice.append(cell(top, 10, 150, item))
        if qty:
            invoice.append(cell(top, 315, 330, qty))
        if price:
            invoice.append(cell(top, 440, 480, price))
        invoice.append(cell(top, 570, 620, amount))
    text = "\n".join(_rows(invoice))
    assert "Toner | Qty: 2 | Unit Price: 80.00 | Amount: 160.00" in text
    assert "Discount | Qty: not listed | Unit Price: not listed | Amount: (5.00)" in text

    # A dash or N/A printed in a column is that column's entry, and a reference beside the date stays on the row.
    statement = [cell(10, 240, 300, "Debit"), cell(10, 370, 420, "Credit")]
    entries = [
        ("01/03/2024", "INV-1001", "1,200.00", "-"),
        ("01/09/2024", "PMT-77", "-", "800.00"),
        ("01/20/2024", "INV-1002", "450.00", "-"),
        ("01/28/2024", "CRN-4", "N/A", "50.00"),
        ("02/02/2024", "INV-1003", "75.00", "20.00"),
    ]
    for index, (date, reference, debit, credit) in enumerate(entries):
        top = 40 + index * 22
        statement += [cell(top, 10, 90, date), cell(top, 120, 190, reference)]
        statement += [cell(top, 300 - 6 * len(debit), 300, debit), cell(top, 420 - 6 * len(credit), 420, credit)]
    text = "\n".join(_rows(statement))
    assert "01/03/2024 INV-1001 | Debit: 1,200.00 | Credit: -" in text
    assert "01/09/2024 PMT-77 | Debit: - | Credit: 800.00" in text
    assert "CRN-4 | Debit: N/A | Credit: 50.00" in text


def test_scanned_negatives_keep_their_column_and_rates_keep_their_decimals():
    from controller_inbox.ocr import _polish, _rows

    hits = [(100, 60, 140, "Line item"), (100, 400, 440, "2024"), (100, 600, 640, "2023")]
    for top, (label, this_year, last_year) in zip(
        (130, 160, 190, 220),
        [("Revenue", "12,500", "11,900"), ("Cost of sales", "7,100", "6,800"), ("Other income", "-1,234", "450"), ("Net income", "4,166", "5,550")],
    ):
        hits += [(top, 60, 200, label), (top, 385, 440, this_year), (top, 590, 640, last_year)]
    assert "Other income | 2024: -1,234 | 2023: 450" in _rows(hits)
    assert _polish("Interest rate 5.125%") == "Interest rate 5.125%"
    assert _polish("Mileage rate $0.655 per mile") == "Mileage rate $0.655 per mile"
    assert _polish("FX rate 0.125, fee 1.250") == "FX rate 0.125, fee 1,250"


def test_ocr_puts_back_dropped_spaces_without_splitting_codes_or_times():
    from controller_inbox.ocr import _spaced

    assert _spaced("Duedate:15October2026") == "Duedate: 15 October 2026"
    assert _spaced("Due date: 150ctober 2026") == "Due date: 15 October 2026"
    assert _spaced("Order 200ct2026, Due10ctober") == "Order 20 Oct 2026, Due 1 October"
    for kept in ("INVOICE4471", "INV-4471 Q4 FY26", "Meeting at 10:30", "Total:$12,480.00", "Box of 500ct", "Paid $100ct 2026"):
        assert _spaced(kept) == kept


# A small Office file or PDF that unpacks to far more than any real document is not read into memory -----------


def _inflated(data: bytes, part: str, filler: bytes, before: bytes) -> bytes:
    """``data`` (an Office zip) with ``filler`` written into ``part`` just before ``before``, packed tight."""
    import zipfile

    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for info in source.infolist():
            body = source.read(info)
            if info.filename == part:
                head, tail = body.split(before, 1)
                body = head + filler + before + tail
            target.writestr(info, body, compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)
    return out.getvalue()


def test_a_word_file_that_unpacks_to_far_more_than_it_holds_is_not_opened():
    plain = io.BytesIO()
    document = Document()
    document.add_paragraph("Invoice INV-9 total $10.00")
    document.save(plain)
    bomb = _inflated(plain.getvalue(), "word/document.xml", b"<w:p><w:r><w:t>0000000000</w:t></w:r></w:p>" * 120_000, b"<w:sectPr")
    assert len(bomb) < 100_000
    text = extract_text_from_bytes("invoice.docx", "", bomb)
    assert text.startswith("[CloseDesk didn't open this file") and "MB" in text
    assert "INV-9" in extract_text_from_bytes("invoice.docx", "", plain.getvalue()), "a normal file is still read"


def test_a_workbook_part_packed_like_a_zip_bomb_is_not_opened():
    book = Workbook()
    book.active["A1"] = "Total"
    plain = io.BytesIO()
    book.save(plain)
    filler = b'<row r="2"><c r="A2" t="inlineStr"><is><t>0</t></is></c></row>' * 450_000
    bomb = _inflated(plain.getvalue(), "xl/worksheets/sheet1.xml", filler, b"</sheetData>")
    assert extract_text_from_bytes("ledger.xlsx", "", bomb).startswith("[CloseDesk didn't open this file")
    assert read_cells(bomb, "", "A1:A2").startswith("[CloseDesk didn't open this file")
    assert "Total" in extract_text_from_bytes("ledger.xlsx", "", plain.getvalue())


def test_a_pdf_whose_pages_unpack_to_too_much_is_read_without_pdfminer(monkeypatch):
    from controller_inbox import documents

    data = make_pdf([["Invoice INV-77", "Amount due $500.00 by October 15, 2026"]])
    assert "unpacks to far more" not in pdf_text(data)
    monkeypatch.setattr(documents, "MAX_PDF_CONTENT", 10)
    monkeypatch.setattr(documents._Miner, "open", classmethod(lambda cls, data: pytest.fail("pdfminer was asked")))
    text = pdf_text(data)
    assert text.startswith("[This PDF unpacks to far more") and "INV-77" in text and "$500.00" in text


def test_a_picture_far_larger_than_any_scan_is_not_read_with_ocr(monkeypatch):
    from PIL import Image

    from controller_inbox import ocr

    picture = io.BytesIO()
    Image.new("L", (400, 300), 255).save(picture, format="PNG")
    monkeypatch.setattr(ocr, "engine_name", lambda: "RapidOCR")
    monkeypatch.setattr(ocr, "_rapid", lambda data: "Invoice INV-5")
    assert ocr.image_text(picture.getvalue()) == "Invoice INV-5"
    monkeypatch.setattr(ocr, "MAX_PIXELS", 400 * 300 - 1)
    assert ocr.image_text(picture.getvalue()) == ""


def test_a_password_protected_office_file_says_so():
    ooxml = pytest.importorskip("msoffcrypto.format.ooxml")
    book = Workbook()
    book.active["A1"] = "Payroll"
    plain = io.BytesIO()
    book.save(plain)
    locked = io.BytesIO()
    ooxml.OOXMLFile(io.BytesIO(plain.getvalue())).encrypt("secret", locked)
    for name in ("payroll.xlsx", "payroll.docx"):
        text = extract_text_from_bytes(name, "", locked.getvalue())
        assert text.startswith("[This file is password-protected") and "zip" not in text
    assert "Payroll" in extract_text_from_bytes("payroll.xlsx", "", plain.getvalue())


def test_a_two_column_sheet_of_names_and_amounts_names_its_columns():
    # Budget on its own sheet, plain headings (not bold): each row still says which department and budget it is,
    # so a question across the two sheets can be worked out by joining them.
    book = Workbook()
    sheet = book.active
    sheet.title = "Q3 Budget"
    sheet.append(["Department", "Q3 Budget"])
    for name, amount in [("Finance", 88000), ("Marketing", 128800), ("Operations", 142800), ("IT", 142500)]:
        sheet.append([name, amount])
    out = io.BytesIO()
    book.save(out)
    text = xlsx_text(out.getvalue())
    assert "A4 (Department): Operations | B4 (Q3 Budget): 142,800" in text
    # A short list of labels and values keeps its first line as a row of its own.
    book = Workbook()
    for row in [["Customer", "Northwind"], ["Subtotal", 1200], ["Tax", 96], ["Total", 1296]]:
        book.active.append(row)
    out = io.BytesIO()
    book.save(out)
    text = xlsx_text(out.getvalue())
    assert "A2: Subtotal | B2: 1,200" in text and "(Customer)" not in text
