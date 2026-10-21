"""Reading attachments right: fixes kept undoable, struck-out and form figures, sideways photos, a fixed figure
reaching the email, large workbooks streamed, grand totals on schedules."""

from __future__ import annotations

import pytest

from controller_inbox import fixes, ocr  # noqa: F401
from test_page_fixes import LINES, _fixes_path, _index, _post, fake_ocr  # noqa: F401
from test_page_view import PAGE, _get, _n, _regions, client, files  # noqa: F401


def _box(region: dict, pad: float = 0.004) -> dict:
    return {"x": region["x"] - pad, "y": region["y"] - pad, "w": region["w"] + 2 * pad, "h": region["h"] + 2 * pad}


def test_undoing_a_figure_fix_after_the_email_is_filed_again_puts_the_figure_back(client, store, files):
    """A fix is applied to the text when the email is loaded; filing the email again (teaching a field) saves that
    shown text as the stored text, so the fix can never be taken back."""
    email = files["Invoice INV-1042"]
    note = _n(email, "delivery note.png")
    att = email.attachments[note - 1]
    store.set_attachment_text(att.id, "Harb0r Steel LLC\nDescription Amount\nSteel beams 1,250.00\nBolts M12 310.50")
    view = _regions(client, email, "delivery note.png")
    reply = _post(client, _fixes_path(email, "delivery note.png"), {"kind": "text", "region": _index(view, "310.50"), "now": "301.50"})
    # Teach a field on the invoice: the email is filed again.
    pdf = _n(email, "invoice.pdf")
    page = _get(client, f"/api/mail/{email.id}/files/{pdf}/pages/1/regions").json()
    total = max((r for r in page["regions"] if r["text"] == "1,560.50"), key=lambda r: r["y"])
    taught = client.post(f"/api/mail/{email.id}/files/{pdf}/pages/1/fixes", headers=PAGE, json={"kind": "field", "field": "total", "box": _box(total)})
    assert taught.status_code == 200, taught.text
    _post(client, f"{_fixes_path(email, 'delivery note.png')}/{reply['id']}/undo", {})
    text = store.get_email(email.id).attachments[note - 1].extracted_text
    assert "Bolts M12 310.50" in text, text


# Word: tracked changes ------------------------------------------------------------------------------


def _tracked_docx() -> bytes:
    """A vendor's revised invoice sent with Track Changes on: the total was cut from $12,500.00 to $10,250.00."""
    import io

    from docx import Document
    from docx.oxml import parse_xml

    W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    doc = Document()
    doc.add_paragraph("Invoice INV-7781")
    doc.add_paragraph("Northwind Services")
    para = parse_xml(
        f'<w:p {W}><w:r><w:t xml:space="preserve">Total due: </w:t></w:r>'
        f'<w:del w:id="1" w:author="AR" w:date="2026-10-01T00:00:00Z"><w:r><w:delText>$12,500.00</w:delText></w:r></w:del>'
        f'<w:ins w:id="2" w:author="AR" w:date="2026-10-01T00:00:00Z"><w:r><w:t>$10,250.00</w:t></w:r></w:ins></w:p>'
    )
    doc.element.body.insert(len(doc.element.body) - 1, para)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_a_figure_deleted_in_tracked_changes_isnt_read_as_the_invoice_amount():
    from datetime import date

    from controller_inbox.documents import docx_text
    from controller_inbox.extract import extract_fields

    text = docx_text(_tracked_docx())
    assert "$10,250.00" in text
    fields = extract_fields(text, as_of=date(2026, 10, 20))
    assert 12500.0 not in fields.amounts, (fields.amounts, text)
    assert fields.primary_amount == 10250.0


# Pictures: a phone photo's orientation ---------------------------------------------------------------


def test_a_phone_photo_turned_by_its_exif_tag_is_read_upright():
    """A phone stores a photo's pixels as the sensor saw them and says in its EXIF tag how to turn them. The Page tab
    turns them (vision.render), but the text read as the email arrives (ocr.image_text) is read from the pixels
    sideways, so the line items lose their figures."""
    import io

    from PIL import Image

    from controller_inbox import ocr, page_view, vision
    from test_page_view import _invoice

    if ocr.engine_name() != "RapidOCR":
        pytest.skip("RapidOCR isn't installed")
    upright = Image.open(io.BytesIO(vision.render(_invoice(), "invoice.pdf", 1, reader=page_view.VIEW))).convert("RGB")
    exif = Image.Exif()
    exif[0x0112] = 6  # turn 90 degrees clockwise to view
    out = io.BytesIO()
    upright.transpose(Image.Transpose.ROTATE_90).save(out, "JPEG", quality=92, exif=exif.tobytes())
    text = ocr.image_text(out.getvalue())
    assert "Bolts M12 200 310.50" in text and "Total 1,560.50" in text, text


# Front end: Undo from the note after going back to the email -----------------------------------------


from test_workspace import browser, page  # noqa: E402,F401 - fixtures used by name


def test_undoing_a_taught_total_from_its_note_on_the_email_page_redraws_the_email(settings, store, page):
    """Teach the total on the Page tab, go back to the email (its facts show the taught total), then click Undo on
    the note that is still showing: the fix is taken back, but the email page keeps showing the taught total."""
    from controller_inbox import web
    from controller_inbox.folder_mail import ingest_folder
    from liveserver import serving
    from test_teach_and_preview import _drop, _invoice

    settings.trusted_domains = "keystonepaper.com"
    settings.ensure_data_dir()
    _drop(settings, "a", "Your Keystone invoice", _invoice("KP-55120", "$2,118.40", "11/15/2026"), "KP-55120.pdf")
    [email] = ingest_folder(store, settings)
    with serving(web.create_app(settings, store)) as base:
        page.set_viewport_size({"width": 1440, "height": 900})
        page.goto(f"{base}/app/mail/{email.id}/file/1?tab=page")
        page.wait_for_function("() => { const i = document.querySelector('.pv-sheet img'); return i && i.naturalWidth > 0; }")
        page.select_option(".pv-teach-pick", "total")
        boxes = page.query_selector_all(".pv-box[aria-label='$2,118.40']")
        lowest = max(boxes, key=lambda b: b.bounding_box()["y"])
        lowest.scroll_into_view_if_needed()
        box = lowest.bounding_box()
        page.mouse.move(box["x"] - 3, box["y"] - 3)
        page.mouse.down()
        page.mouse.move(box["x"] + box["width"] + 3, box["y"] + box["height"] + 3, steps=4)
        with page.expect_response(lambda r: r.url.endswith("/fixes") and r.request.method == "POST"):
            page.mouse.up()
        page.wait_for_selector(".toast-action:has-text('Undo')")
        page.click(".rd-bar a.btn-quiet")  # back to the email
        page.wait_for_selector(".rd-facts dt:has-text('Total')")
        assert "$2,118.40" in page.text_content(".rd-facts")
        with page.expect_response(lambda r: r.url.endswith("/undo")):
            page.click(".toast-action:has-text('Undo')")
        assert "total" not in store.get_email(email.id).extracted.taught
        page.wait_for_function("() => ![...document.querySelectorAll('.rd-facts dt')].some((dt) => dt.textContent === 'Total')", timeout=4000)


# A figure fixed on a scan never reaches the email's amount ------------------------------------------------


def test_a_total_fixed_on_a_scan_is_the_emails_amount(store, settings, monkeypatch):
    """OCR misreads a scanned invoice's total ($1,860.50 for $1,560.50). The user fixes the box on the Page tab; the
    file's text has the fix, but the email's amount (and its task) stay the misread figure."""
    from fastapi.testclient import TestClient

    from controller_inbox import ocr, page_view, vision, web
    from controller_inbox.folder_mail import ingest_folder
    from msgfactory import write_msg
    from test_page_view import _invoice

    read = [(60, 40, 330, 70, "Harbor Steel LLC"), (60, 120, 330, 140, "Invoice HS-2201"), (60, 400, 300, 420, "Total due"), (700, 400, 900, 420, "$1,860.50")]
    monkeypatch.setattr(ocr, "engine_name", lambda: "RapidOCR")
    from controller_inbox import extract

    misread = lambda _data: "Harbor Steel LLC\nInvoice HS-2201\nTotal due $1,860.50"  # noqa: E731
    monkeypatch.setattr(ocr, "image_text", misread)
    monkeypatch.setattr(extract, "image_text", misread)
    monkeypatch.setattr(
        ocr, "line_boxes",
        lambda _png: [{"left": float(l), "top": float(t), "right": float(r), "bottom": float(b), "text": x, "score": 0.9} for l, t, r, b, x in read],
    )
    settings.trusted_domains = "taz.com"
    settings.ensure_data_dir()
    scan = vision.render(_invoice(), "invoice.pdf", 1, reader=page_view.VIEW)
    write_msg(settings.inbox_incoming / "scan.msg", "Invoice HS-2201", "Scan attached.", sender_name="Harbor Steel",
              sender_email="ar@taz.com", attachments=[("HS-2201.png", scan, "image/png")])
    [email] = ingest_folder(store, settings)
    assert email.extracted.primary_amount == 1860.50
    client = TestClient(web.create_app(settings, store))
    view = client.get(f"/api/mail/{email.id}/files/1/pages/1/regions", headers=PAGE).json()
    index = next(i for i, r in enumerate(view["regions"]) if r["text"] == "$1,860.50")
    reply = client.post(f"/api/mail/{email.id}/files/1/pages/1/fixes", headers=PAGE, json={"kind": "text", "region": index, "now": "$1,560.50"})
    assert reply.status_code == 200, reply.text
    stored = store.get_email(email.id)
    assert "$1,560.50" in stored.attachments[0].extracted_text
    assert stored.extracted.primary_amount == 1560.50, stored.extracted.amounts


# Preview: a big export --------------------------------------------------------------------------------


def test_the_thumbnail_of_a_5_mb_ledger_export_doesnt_load_the_whole_workbook_twice(monkeypatch):
    """A 5 MB general-ledger export (200,000 rows) is well inside the preview's limits (25 MB, 150 MB unzipped).
    The email's thumbnail asks for 8 rows: the workbook is streamed (read-only), never loaded in full, which took
    about 25 seconds and 1 GB of memory twice over (values and formulas)."""
    import io

    import openpyxl
    from openpyxl import Workbook

    from controller_inbox import sheet_preview

    book = Workbook(write_only=True)
    ws = book.create_sheet("GL")
    ws.append(["Date", "Account", "Description", "Debit", "Credit"])
    for i in range(200_000):
        ws.append(["2026-09-30", 4000 + i % 50, f"Journal line {i}", i * 1.01, 0])
    out = io.BytesIO()
    book.save(out)
    data = out.getvalue()
    assert len(data) < sheet_preview.MAX_BYTES and sheet_preview._unzipped_size(data) < sheet_preview.MAX_UNZIPPED
    opened = []
    real = openpyxl.load_workbook

    def load_workbook(*args, **kwargs):
        opened.append(bool(kwargs.get("read_only")))
        return real(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", load_workbook)
    got = sheet_preview.preview(data, "GL detail.xlsx", rows=8, cols=12)
    assert len(got["sheets"][0]["rows"]) == 8
    assert opened and all(opened), "the whole workbook was loaded, not streamed"


# PDF forms -------------------------------------------------------------------------------------------


def _filled_form_pdf() -> bytes:
    """An invoice filled in as a PDF form: the amount due is a text field's value (with its appearance, as Acrobat
    and every form filler save it), not text on the page."""
    content = (
        b"BT /F1 12 Tf 72 720 Td (Harbor Steel LLC - Invoice INV-5521) Tj ET\n"
        b"BT /F1 11 Tf 72 680 Td (Amount due:) Tj ET\n"
        b"BT /F1 11 Tf 72 660 Td (Please remit by the due date shown on this form.) Tj ET"
    )
    ap = b"/Tx BMC q BT /F1 11 Tf 2 5 Td ($4,210.00) Tj ET Q EMC"
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R /AcroForm << /Fields [6 0 R] /DR << /Font << /F1 5 0 R >> >> /DA (/F1 11 Tf 0 g) >> >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> /Annots [6 0 R] >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
        b"<< /Type /Annot /Subtype /Widget /FT /Tx /T (amount_due) /V ($4,210.00) /Rect [160 674 300 692] /P 3 0 R /F 4 "
        b"/DA (/F1 11 Tf 0 g) /AP << /N 7 0 R >> >>",
        b"<< /Type /XObject /Subtype /Form /BBox [0 0 140 18] /Resources << /Font << /F1 5 0 R >> >> /Length %d >>\nstream\n" % len(ap)
        + ap + b"\nendstream",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1) + b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer << /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return out


def test_a_filled_pdf_forms_values_are_read():
    from controller_inbox.documents import pdf_text

    text = pdf_text(_filled_form_pdf())
    assert "4,210.00" in text, text


def test_a_filled_pdf_forms_values_are_drawn_on_the_page():
    import io

    from PIL import Image

    from controller_inbox import page_view, vision

    image = Image.open(io.BytesIO(vision.render(_filled_form_pdf(), "form.pdf", 1, reader=page_view.VIEW))).convert("L")
    scale = image.width / 612
    field = image.crop((int(160 * scale), int((792 - 692) * scale), int(300 * scale), int((792 - 674) * scale)))
    assert field.getextrema()[0] < 128, "the field's value $4,210.00 isn't drawn: the page shows an empty box"


# A page fix taken back after the model read the email (from the chat sweep; same root cause as the first test) ---


def _reader(settings, handler):
    import httpx

    from controller_inbox.local_llm import LocalReader

    return LocalReader(settings, model="tiny-3b", client=httpx.Client(transport=httpx.MockTransport(handler)))


def _reply(obj: dict):
    import json

    import httpx

    return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps(obj)}, "finish_reason": "stop"}]})


def test_a_page_fix_taken_back_after_the_model_read_the_email_is_really_taken_back(loaded, settings):
    from controller_inbox.overnight import read_queue

    email = loaded.get_email("demo-inv-10482")
    att = email.attachments[0]
    assert "Horizon Goods LLC" in att.extracted_text
    fix_id = loaded.add_page_fix({
        "attachment_id": att.id, "sha256": att.sha256, "page": 1, "kind": "text",
        "x": 0.1, "y": 0.1, "w": 0.2, "h": 0.05, "was": "Horizon Goods LLC", "now": "Horizon Foods LLC", "sender": email.sender_email,
    })
    assert "Horizon Foods LLC" in loaded.get_email(email.id).attachments[0].extracted_text

    reading = {"category": "ap_invoice", "folder": "important", "importance": "high",
               "summary": "Invoice INV-10482 to enter.", "why": "Vendor bill.", "actions": []}
    read_queue(loaded, settings, limit=50, reader=_reader(settings, lambda _r: _reply(reading)))
    assert loaded.get_email(email.id).model_status == "bionic"
    # The stored (first) reading of the file must stay as it was read: the fix is laid over it when shown.
    assert "Horizon Goods LLC" in loaded.stored_text(att.id), "the user's fix was written into the file's stored text"

    assert loaded.remove_page_fix(att.id, fix_id)  # Undo on the Page tab
    shown = loaded.get_email(email.id).attachments[0].extracted_text
    assert "Horizon Goods LLC" in shown, "the fix was taken back but the chat and summaries still read the fixed text"


# Undoing a figure fixed on a scan puts the email's amount back ----------------------------------------------------


def test_undoing_a_total_fixed_on_a_scan_puts_the_emails_amount_back(store, settings, monkeypatch):
    from fastapi.testclient import TestClient

    from controller_inbox import extract, page_view, vision, web
    from controller_inbox.folder_mail import ingest_folder
    from msgfactory import write_msg
    from test_page_view import _invoice

    read = [(60, 40, 330, 70, "Harbor Steel LLC"), (60, 120, 330, 140, "Invoice HS-2201"), (60, 400, 300, 420, "Total due"), (700, 400, 900, 420, "$1,860.50")]
    misread = lambda _data: "Harbor Steel LLC\nInvoice HS-2201\nTotal due $1,860.50"  # noqa: E731
    monkeypatch.setattr(ocr, "engine_name", lambda: "RapidOCR")
    monkeypatch.setattr(ocr, "image_text", misread)
    monkeypatch.setattr(extract, "image_text", misread)
    monkeypatch.setattr(ocr, "line_boxes", lambda _png: [{"left": float(l), "top": float(t), "right": float(r), "bottom": float(b), "text": x, "score": 0.9} for l, t, r, b, x in read])
    settings.trusted_domains = "taz.com"
    settings.ensure_data_dir()
    write_msg(settings.inbox_incoming / "scan.msg", "Invoice HS-2201", "Scan attached.", sender_name="Harbor Steel", sender_email="ar@taz.com",
              attachments=[("HS-2201.png", vision.render(_invoice(), "invoice.pdf", 1, reader=page_view.VIEW), "image/png")])
    [email] = ingest_folder(store, settings)
    client = TestClient(web.create_app(settings, store))
    view = client.get(f"/api/mail/{email.id}/files/1/pages/1/regions", headers=PAGE).json()
    index = next(i for i, r in enumerate(view["regions"]) if r["text"] == "$1,860.50")
    reply = client.post(f"/api/mail/{email.id}/files/1/pages/1/fixes", headers=PAGE, json={"kind": "text", "region": index, "now": "$1,560.50"}).json()
    assert "$1,560.50" in reply["message"], reply
    assert store.get_email(email.id).extracted.primary_amount == 1560.50
    undone = client.post(f"/api/mail/{email.id}/files/1/pages/1/fixes/{reply['id']}/undo", headers=PAGE)
    assert undone.status_code == 200, undone.text
    assert store.get_email(email.id).extracted.primary_amount == 1860.50


# A field taught in a box doesn't replace another field taught in the same box --------------------------------------


def test_two_fields_taught_in_one_box_are_both_kept(store):
    row = {"attachment_id": "e:a", "sha256": "", "page": 1, "kind": "field", "x": 0.5, "y": 0.1, "w": 0.2, "h": 0.03, "was": "10/01/2026"}
    first = store.add_page_fix({**row, "now": "invoice_date"})
    second = store.add_page_fix({**row, "now": "due_date"})
    assert {fix["id"] for fix in store.page_fixes("e:a")} == {first, second}
    third = store.add_page_fix({**row, "now": "due_date", "was": "10/31/2026"})
    assert {fix["id"] for fix in store.page_fixes("e:a")} == {first, third}, "the same field taught again replaces it"


# The Page tab's total on a schedule whose total row isn't a total due ----------------------------------------------


def test_a_schedules_total_due_is_never_its_total_cost_nor_a_row_of_joined_figures(settings):
    from pathlib import Path

    from test_remaining_fixes import _details

    fixtures = Path(__file__).parent / "fixtures"
    prepaid = _details(settings, (fixtures / "excel_schedules" / "Prepaid Amortization FY2026.pdf").read_bytes())
    assert "462,120.00" not in prepaid.get("Total due", ""), "the Total Cost column is what was paid, not what is due"
    matrix = _details(settings, (fixtures / "more_schedules" / "Intercompany Matrix 9-30-26.pdf").read_bytes())
    total = matrix.get("Total due")
    assert total is None or total.count(".") <= 1, total
    aging = _details(settings, (fixtures / "excel_schedules" / "AP Aging 9-30-26.pdf").read_bytes())
    assert aging["Total due"].replace(" ", "") == "$267,159.31"


# Preview: a merge whose top-left cell is hidden ---------------------------------------------------------------------


def test_a_merge_whose_first_row_is_hidden_still_shows():
    import io

    from openpyxl import Workbook

    from controller_inbox import sheet_preview

    book = Workbook()
    ws = book.active
    ws["A1"] = "Q3 accruals"
    ws.merge_cells("A1:C2")
    ws.row_dimensions[1].hidden = True
    ws["A3"], ws["B3"], ws["C3"] = "Rent", 4200, 4200
    out = io.BytesIO()
    book.save(out)
    sheet = sheet_preview.preview(out.getvalue(), "accruals.xlsx")["sheets"][0]
    cells = {cell["v"]: (row["n"], cell) for row in sheet["rows"] for cell in row["cells"]}
    assert "Q3 accruals" in cells, sheet["rows"]
    n, cell = cells["Q3 accruals"]
    assert n == 2 and cell["c"] == 1 and cell.get("cs") == 3 and cell.get("rs", 1) == 1


def test_a_large_workbook_is_streamed_with_its_figures_as_excel_shows_them(monkeypatch):
    import io

    from openpyxl import Workbook

    from controller_inbox import sheet_preview

    book = Workbook()
    ws = book.active
    ws.append(["Account", "Debit", "Credit"])
    ws.append(["Cash", 1250.5, None])
    ws["B2"].number_format = "#,##0.00"
    ws["C3"] = "=B2"
    ws["C3"].number_format = "#,##0.00"
    out = io.BytesIO()
    book.save(out)
    data = out.getvalue()
    full = sheet_preview.preview(data, "gl.xlsx")["sheets"][0]
    monkeypatch.setattr(sheet_preview, "FULL_UNZIPPED", 0)
    streamed = sheet_preview.preview(data, "gl2.xlsx")["sheets"][0]
    assert streamed.get("note")
    assert [[(c["c"], c["v"], c.get("f")) for c in r["cells"]] for r in streamed["rows"]] == [[(c["c"], c["v"], c.get("f")) for c in r["cells"]] for r in full["rows"]]
    assert streamed["rows"][1]["cells"][1]["v"] == "1,250.50"


# Preview: a TIFF's pages are counted without waiting on page 1's OCR ------------------------------------------------


def test_a_tiffs_page_count_comes_with_the_file(store, settings, monkeypatch):
    from fastapi.testclient import TestClient

    from controller_inbox import web
    from controller_inbox.folder_mail import ingest_folder
    from msgfactory import write_msg
    from test_page_view import _tiff

    monkeypatch.setattr(ocr, "image_text", lambda _data: "")
    settings.ensure_data_dir()
    write_msg(settings.inbox_incoming / "fax.msg", "Faxed invoice", "Fax attached.", sender_name="Harbor Steel", sender_email="ar@taz.com",
              attachments=[("fax.tif", _tiff("Invoice 5521 page one", "Total due 9,876.54"), "image/tiff")])
    [email] = ingest_folder(store, settings)
    got = TestClient(web.create_app(settings, store)).get(f"/api/mail/{email.id}/files/1", headers=PAGE).json()
    assert got["file"]["show"] == "pages" and got["file"]["pages"] == 2, got["file"]
