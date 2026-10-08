"""Regressions for the ingest and document-reading fixes (one or more tests per bug)."""

from __future__ import annotations

import io
import os
import time
import zipfile
from datetime import date, datetime, timezone
from email.message import EmailMessage
from pathlib import Path

import pytest

from controller_inbox import documents, folder_mail
from controller_inbox.config import Settings
from controller_inbox.extract import extract_fields, extract_text_from_bytes, parse_due_date, redact_financial_secrets
from controller_inbox.folder_mail import collect_messages, ingest_folder, reread_attachments
from controller_inbox.graph import GraphMailbox
from controller_inbox.models import RawAttachment, RawMessage
from controller_inbox.pipeline import ingest_demo, ingest_mailbox, process_message, rescore_stored
from controller_inbox.store import Store

AS_OF = date(2026, 10, 6)
NOW = datetime(2026, 10, 6, 14, 0, tzinfo=timezone.utc)


def _eml(settings: Settings, name: str, body: str = "See attached.", *, message_id: str = "<m1@x>", attach=(), date_header="Mon, 05 Oct 2026 10:00:00 -0400") -> Path:
    msg = EmailMessage()
    msg["From"] = "Vendor AP <ap@vendor.com>"
    msg["To"] = "me@co.com"
    msg["Subject"] = name
    msg["Message-ID"] = message_id
    msg["Date"] = date_header
    msg.set_content(body)
    for filename, data in attach:
        msg.add_attachment(data, maintype="text", subtype="plain", filename=filename)
    path = settings.inbox_incoming / f"{name}.eml"
    path.write_bytes(bytes(msg))
    _age(path)
    return path


def _age(path: Path) -> None:
    """Files written just now look like they are still being copied; date them a minute back."""
    old = time.time() - 60
    os.utime(path, (old, old))


def _raw(body: str, *, received: datetime, attachments=None, preview: str = "") -> RawMessage:
    return RawMessage(
        id="raw-1",
        subject="Docs",
        sender_name="A Vendor",
        sender_email="a@vendor.com",
        received_at=received,
        body_text=body,
        body_preview=preview,
        has_attachments=bool(attachments),
        source="folder",
        attachments=attachments or [],
    )


# 1. Account masking needs digits ---------------------------------------------------------------


def test_masking_leaves_ordinary_words_alone_and_masks_real_numbers():
    text = "Review the account statement, ask the account manager. Accountability. account number 123456789012"
    masked = redact_financial_secrets(text)
    assert "account statement" in masked and "account manager" in masked and "Accountability" in masked
    assert "123456789012" not in masked and "account ****9012" in masked
    assert "acct no. 55512345" not in redact_financial_secrets("acct no. 55512345")


def test_account_last4_comes_only_from_digits():
    words = extract_fields("Your account balance is overdue", as_of=AS_OF)
    assert words.account_last4 == [] and not words.mentions_routing_or_account
    real = extract_fields("Account: AB12345678", as_of=AS_OF)
    assert real.account_last4 == ["5678"] and real.mentions_routing_or_account


# 2. Whole-dollar amounts ------------------------------------------------------------------------


def test_whole_dollar_amounts_are_read():
    fields = extract_fields("Please wire $48,000 today. Fee USD 500. Total: 5,000", as_of=AS_OF)
    assert {48000.0, 500.0, 5000.0} <= set(fields.amounts)
    assert extract_fields("Pay $1,2345 now", as_of=AS_OF).amounts == []
    assert extract_fields("total 3 items", as_of=AS_OF).amounts == []


# 3. Repeated attachment names -------------------------------------------------------------------


def test_two_attachments_with_one_name_both_ingest_with_their_own_text(settings: Settings, store: Store):
    settings.ensure_data_dir()
    _eml(settings, "Two invoices", attach=[("invoice.csv", b"Item,Amount\nFirst,100\n"), ("invoice.csv", b"Item,Amount\nSecond,200\n")])
    report: dict = {}
    [record] = ingest_folder(store, settings, report=report)
    assert report["failed"] == []
    by_name = {att.filename: att for att in record.attachments}
    assert sorted(by_name) == ["invoice (2).csv", "invoice.csv"]
    assert by_name["invoice.csv"].id == f"{record.id}:invoice.csv"  # an unrepeated name keeps its old id
    folder = settings.inbox_extracted / record.id
    assert b"First" in (folder / "invoice.csv").read_bytes() and b"Second" in (folder / "invoice (2).csv").read_bytes()

    store.set_state(folder_mail.READER_KEY, "old")
    reread_attachments(store, settings)
    texts = {att.filename: att.extracted_text for att in store.get_email(record.id).attachments}
    assert "First" in texts["invoice.csv"] and "Second" not in texts["invoice.csv"]
    assert "Second" in texts["invoice (2).csv"]


def test_repeated_names_are_numbered_the_same_way_every_time(settings: Settings):
    settings.ensure_data_dir()
    _eml(settings, "Repeat", attach=[("a.txt", b"1"), ("A.TXT", b"2"), ("a.txt", b"3")])
    first = [att.filename for att in collect_messages(settings)[0][0].attachments]
    second = [att.filename for att in collect_messages(settings)[0][0].attachments]
    assert first == second == ["a.txt", "A (2).TXT", "a (3).txt"]


def test_zip_entries_with_one_name_get_their_own_ids(store: Store, settings: Settings):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("jan/report.txt", "January report")
        archive.writestr("feb/report.txt", "February report")
    zipped = RawAttachment(id="pack.zip", filename="pack.zip", content_type="application/zip", size_bytes=1, content=buffer.getvalue())
    record = process_message(_raw("files", received=NOW, attachments=[zipped]), store, settings, now=NOW)
    ids = [att.id for att in record.attachments]
    assert len(ids) == 2 and len(set(ids)) == 2


# 4. Preview is masked ---------------------------------------------------------------------------


def test_preview_does_not_show_account_numbers(store: Store, settings: Settings):
    raw = _raw("Please update our account number 123456789012 routing 021000021", received=NOW)
    record = process_message(raw, store, settings, now=NOW)
    stored = store.get_email(record.id)
    assert "123456789012" not in stored.body_preview and "021000021" not in stored.body_preview
    assert "123456789012" not in " ".join(a.title + a.detail for a in stored.actions)
    assert stored.extracted.mentions_routing_or_account


# 5. Folder imports don't move the Graph cursor -------------------------------------------------


def test_folder_import_and_sample_leave_the_graph_cursor_alone(settings: Settings, store: Store):
    settings.ensure_data_dir()
    ingest_demo(store, settings, now=NOW)
    assert store.get_state("last_sync_at") is None
    _eml(settings, "Cursor")
    assert ingest_folder(store, settings, now=NOW)
    assert store.get_state("last_sync_at") is None
    assert store.get_state("last_folder_ingest")


# 6. Graph item attachments and a bad message ----------------------------------------------------


class _FakeClient:
    def __init__(self, mime: bytes = b""):
        self.mime = mime
        self.params: list[dict | None] = []

    def _user_root(self) -> str:
        return "/me"

    def get_json(self, path, params=None):
        self.params.append(params)
        if path.endswith("/attachments"):
            return {"value": [{"@odata.type": "#microsoft.graph.itemAttachment", "id": "ITEM1", "name": "Old thread"}]}
        return {"value": []}

    def get_bytes(self, path):
        assert path.endswith("/attachments/ITEM1/$value")
        return self.mime


def test_graph_item_attachment_is_read_as_mime():
    inner = EmailMessage()
    inner["From"] = "Supplier <s@supplier.com>"
    inner["Subject"] = "Old thread"
    inner.set_content("Our bank details changed.")
    inner.add_attachment(b"Invoice text", maintype="text", subtype="plain", filename="inv.txt")
    found = GraphMailbox(_FakeClient(bytes(inner))).get_attachments("MSG1")
    assert [att.filename for att in found] == ["Old thread.txt", "Old thread › inv.txt"]
    assert b"Our bank details changed." in found[0].content
    assert len({att.id for att in found}) == 2


class _Mailbox:
    def __init__(self, raws):
        self.raws = raws

    def list_messages(self, received_after=None):
        return iter(self.raws)

    def get_attachments(self, message_id):
        if message_id == "bad":
            raise ValueError("Graph sent something odd")
        return []

    def apply_categories(self, message_id, categories, flag):
        return "skipped"


def test_one_bad_graph_message_does_not_stop_the_sync(store: Store, settings: Settings):
    good = _raw("hello", received=NOW)
    good.id, good.source = "good", "graph"
    bad = _raw("broken", received=NOW, attachments=[])
    bad.id, bad.source, bad.has_attachments = "bad", "graph", True
    report: dict = {}
    records = ingest_mailbox(_Mailbox([bad, good]), store, settings, now=NOW, report=report)
    assert [r.id for r in records] == ["good"]
    assert report["failed"][0]["id"] == "bad"
    assert store.get_state("last_sync_at") == NOW.isoformat()


# 7. Relative dates read against the day the mail was sent --------------------------------------


def test_relative_dates_anchor_to_the_received_date(store: Store, settings: Settings):
    received = datetime(2026, 9, 28, 14, 0, tzinfo=timezone.utc)
    record = process_message(_raw("Please send the W-9 by tomorrow.", received=received), store, settings, now=NOW)
    assert record.extracted.due_dates == ["2026-09-29"]
    task = next(a for a in record.actions if "W-9" in a.title)
    assert task.due_date == "2026-09-29"

    class Check:
        level = "none"

    again = rescore_stored(store, settings, record, Check(), now=NOW)
    assert again.extracted.due_dates == ["2026-09-29"]
    assert next(a for a in again.actions if "W-9" in a.title).due_date == "2026-09-29"


# 8. A message named like the drop folder --------------------------------------------------------


def test_a_message_named_incoming_does_not_take_the_loose_files(settings: Settings):
    settings.ensure_data_dir()
    (settings.inbox_incoming / "incoming.eml").write_bytes(b"From: a@b.com\r\nSubject: x\r\n\r\nhi\r\n")
    (settings.inbox_incoming / "unrelated.pdf").write_bytes(b"%PDF-1.4")
    batches = dict((p.name, [s.name for s in sidecars]) for p, sidecars in folder_mail.collect_batches(settings))
    assert batches["incoming.eml"] == [] and "unrelated.pdf" in batches


def test_a_folder_named_after_the_message_still_holds_its_files(settings: Settings):
    settings.ensure_data_dir()
    (settings.inbox_incoming / "invoice.eml").write_bytes(b"From: a@b.com\r\nSubject: x\r\n\r\nhi\r\n")
    (settings.inbox_attachments / "invoice").mkdir()
    (settings.inbox_attachments / "invoice" / "scan.pdf").write_bytes(b"%PDF-1.4")
    batches = dict((p.name, [s.name for s in sidecars]) for p, sidecars in folder_mail.collect_batches(settings))
    assert batches["invoice.eml"] == ["scan.pdf"]


# 9. A forwarded email inside an .eml ------------------------------------------------------------


def test_forwarded_eml_is_an_attachment_not_part_of_the_body(tmp_path: Path):
    inner = EmailMessage()
    inner["From"] = "fraud@evil.com"
    inner["Subject"] = "old thread"
    inner.set_content("INNER BODY: new bank details")
    inner.add_attachment(b"letter", maintype="text", subtype="plain", filename="letter.txt")
    outer = EmailMessage()
    outer["From"] = "boss@co.com"
    outer["Subject"] = "fwd"
    outer.set_content("OUTER BODY: see forwarded")
    outer.add_attachment(inner)
    path = tmp_path / "o.eml"
    path.write_bytes(bytes(outer))
    raw = folder_mail._parse_eml(path)
    assert "INNER BODY" not in raw.body_text and "OUTER BODY" in raw.body_text
    assert [att.filename for att in raw.attachments] == ["old thread.txt", "old thread › letter.txt"]
    assert b"INNER BODY" in raw.attachments[0].content and b"fraud@evil.com" in raw.attachments[0].content


# 10. Graph filter order -------------------------------------------------------------------------


@pytest.mark.parametrize("after", [None, datetime(2026, 10, 1, tzinfo=timezone.utc)])
def test_graph_filter_names_the_sort_property_first(after):
    client = _FakeClient()
    list(GraphMailbox(client).list_messages(received_after=after))
    params = client.params[0]
    assert params["$orderby"].startswith("receivedDateTime")
    assert params["$filter"].startswith("receivedDateTime ge ")


# 11. compare_columns and three-letter headers ---------------------------------------------------


def test_compare_columns_prefers_a_header_over_column_letters():
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "P"
    for row in (["Account", "Jan", "Feb"], ["Rent", 100, 120], ["Fees", 50, 40]):
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    out = documents.compare_columns(buffer.getvalue(), "P", "Jan", "Feb")
    assert "Rent (row 2): 100 → 120" in out
    assert "Rent (row 2): 100 → 120" in documents.compare_columns(buffer.getvalue(), "P", "B", "C")


# 12. Content the readers used to drop -----------------------------------------------------------


def test_docx_reads_text_inside_content_controls():
    from docx import Document
    from docx.oxml import parse_xml

    doc = Document()
    doc.add_paragraph("Visible")
    doc.element.body.insert(0, parse_xml(
        '<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:sdtContent>'
        "<w:p><w:r><w:t>Amount due $9,999.00</w:t></w:r></w:p></w:sdtContent></w:sdt>"
    ))
    buffer = io.BytesIO()
    doc.save(buffer)
    assert documents.docx_text(buffer.getvalue()) == "Amount due $9,999.00\nVisible"


def test_pptx_reads_grouped_shapes():
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = "Budget 1.2M in a group"
    buffer = io.BytesIO()
    deck.save(buffer)
    assert "Budget 1.2M in a group" in documents.pptx_text(buffer.getvalue())


def test_utf16_csv_from_excel_is_read():
    data = "Name\tAmount\r\nAcme\t100\r\n".encode("utf-16")
    text = documents.csv_text(data, "export.csv")
    assert "Acme" in text and "100" in text


def test_last_resort_sample_does_not_split_a_utf8_character():
    data = ("a" * 1999 + "é" + " rest of the log").encode("utf-8")
    assert extract_text_from_bytes("notes.log", "application/octet-stream", data).startswith("aaaa")


# 13. Invoice numbers and year-less due dates ----------------------------------------------------


def test_invoice_numbers_keep_every_segment():
    fields = extract_fields("Invoice No: INV-2024-0457. Again INV-2024-0457. Invoice 12345-due", as_of=AS_OF)
    assert fields.invoice_numbers == ["INV-2024-0457", "12345"]


def test_due_dates_without_a_year():
    assert extract_fields("Payment due October 15.", as_of=AS_OF).due_dates == ["2026-10-15"]
    assert parse_due_date("October 15th", as_of=AS_OF) == "2026-10-15"
    assert parse_due_date("January 5", as_of=date(2026, 12, 20)) == "2027-01-05"
    assert parse_due_date("Oct 15, 24", as_of=AS_OF) == "2024-10-15"


# 14. Drop folder: files still copying, locked files, the sample ---------------------------------


def test_a_file_still_being_copied_waits_for_the_next_run(settings: Settings, store: Store, monkeypatch):
    settings.ensure_data_dir()
    path = settings.inbox_incoming / "growing.txt"
    path.write_bytes(b"first half")

    def still_writing(_seconds):
        with path.open("ab") as handle:
            handle.write(b" second half")

    monkeypatch.setattr(folder_mail.time, "sleep", still_writing)
    report: dict = {}
    assert ingest_folder(store, settings, report=report) == []
    assert report["waiting"] == ["growing.txt"] and path.exists()

    _age(path)
    assert len(ingest_folder(store, settings)) == 1 and not path.exists()


def test_a_freshly_written_finished_file_is_read_now(settings: Settings, store: Store):
    settings.ensure_data_dir()
    (settings.inbox_incoming / "done.txt").write_bytes(b"complete")
    assert len(ingest_folder(store, settings)) == 1


def test_a_locked_file_is_left_and_the_batch_goes_on(settings: Settings, store: Store, monkeypatch):
    settings.ensure_data_dir()
    _eml(settings, "A locked", message_id="<a@x>")
    _eml(settings, "B fine", message_id="<b@x>")
    real = folder_mail._parse_message

    def parse(path):
        if path.name.startswith("A"):
            raise PermissionError("in use by another process")
        return real(path)

    monkeypatch.setattr(folder_mail, "_parse_message", parse)
    report: dict = {}
    records = ingest_folder(store, settings, report=report)
    assert [r.subject for r in records] == ["B fine"]
    assert report["waiting"] == ["A locked.eml"] and report["failed"] == []
    assert (settings.inbox_incoming / "A locked.eml").exists()


def test_a_failed_move_to_quarantine_does_not_stop_the_batch(settings: Settings, store: Store, monkeypatch):
    settings.ensure_data_dir()
    (settings.inbox_incoming / "a-broken.msg").write_bytes(b"not an outlook file")
    _age(settings.inbox_incoming / "a-broken.msg")
    _eml(settings, "b-fine", message_id="<b@x>")
    real_move = folder_mail.shutil.move

    def move(src, dst):
        if "broken" in src:
            raise PermissionError("locked")
        return real_move(src, dst)

    monkeypatch.setattr(folder_mail.shutil, "move", move)
    report: dict = {}
    records = ingest_folder(store, settings, report=report)
    assert [r.subject for r in records] == ["b-fine"]
    assert len(report["failed"]) == 1


def test_a_message_that_fails_after_parsing_keeps_the_sample(settings: Settings, loaded: Store, monkeypatch):
    settings.ensure_data_dir()
    before = loaded.counts()["emails"]
    _eml(settings, "Breaks later")

    def boom(*_args, **_kwargs):
        raise RuntimeError("classifier fell over")

    monkeypatch.setattr(folder_mail, "process_message", boom)
    report: dict = {}
    ingest_folder(loaded, settings, report=report)
    assert len(report["failed"]) == 1 and "sample_cleared" not in report
    assert loaded.counts()["emails"] == before and not loaded.real_mail_count()


def test_an_address_in_the_display_name_is_not_the_sender(store: Store, settings: Settings):
    # "Acme Billing <billing@acme.com>" is only the name: the mail comes from acme-payments.net.
    settings.ensure_data_dir()
    settings.trusted_domains = "acme.com"
    path = settings.inbox_incoming / "spoof.eml"
    path.write_bytes(
        b'From: "Acme Billing <billing@acme.com>" <billing@acme-payments.net>\r\nTo: ap@co.com\r\n'
        b"Subject: Updated remittance details\r\nMessage-ID: <sp1@x>\r\nDate: Mon, 05 Oct 2026 10:00:00 -0400\r\n\r\n"
        b"Our bank details have changed. Please use the new account for all payments from today.\r\n"
    )
    _age(path)
    record = ingest_folder(store, settings)[0]
    assert record.sender_email == "billing@acme-payments.net"
    assert store.fraud_check(record.id)["level"] == "high" and "do_not_process" in record.flags
    # Headers the strict parser gives up on still give the address in the brackets.
    assert folder_mail._split_address("Chen, Maya <maya@x.com>") == ("Chen, Maya", "maya@x.com")
    assert folder_mail._split_address("Maya Chen") == ("Maya Chen", "")


# 15. The overnight reading never overwrites what the user did while the model was reading -------


def test_a_reading_is_not_saved_over_a_correction_or_verdict_made_meanwhile(store: Store, settings: Settings):
    from controller_inbox.fraud import record_fraud_verdict
    from controller_inbox.learn import record_correction
    from controller_inbox.overnight import read_queue
    from test_bionic import AgreeingReader

    ingest_demo(store, settings, now=NOW)
    first, corrected, flagged = (e.id for e in store.list_emails(model_status="script_draft", order="queue", limit=3))

    class SlowReader(AgreeingReader):
        def read(self, packet):
            if packet["email_id"] == first:
                # While the model reads the first email, the user corrects the second and calls the third fraud.
                record_correction(store, settings, email_id=corrected, corrected_category="newsletter", reason="it is a newsletter")
                record_fraud_verdict(store, settings, flagged, verdict="fraud", note="phoned the vendor; it is fake")
            return super().read(packet)

    result = read_queue(store, settings, limit=3, now=NOW, reader=SlowReader())
    assert result["read_ids"] == [first]
    kept = store.get_email(corrected)
    assert kept.category.value == "newsletter" and kept.model_status == "corrected"
    verdict = store.get_email(flagged)
    assert {"fraud_risk", "fraud_confirmed"} <= set(verdict.flags)
    assert verdict.model_status == "script_draft", "read again on the next run, as it is now"


def test_a_limit_of_zero_reads_nothing(loaded: Store, settings: Settings):
    from controller_inbox.overnight import read_queue
    from test_bionic import AgreeingReader

    assert read_queue(loaded, settings, limit=0, reader=AgreeingReader())["read_ids"] == []
    assert loaded.counts()["waiting_on_bionic"] == 19


# 16. Another copy of a stored message keeps the files that came with the first ---------------


def _sidecar(settings: Settings, name: str, data: bytes) -> Path:
    path = settings.inbox_incoming / name
    path.write_bytes(data)
    _age(path)
    return path


def test_a_copy_without_the_files_keeps_the_ones_dropped_with_the_first(settings: Settings, store: Store):
    settings.ensure_data_dir()
    _eml(settings, "Invoice 4410", "Invoice attached, due Oct 30.", message_id="<inv4410@vendor.com>")
    _sidecar(settings, "Invoice 4410.txt", b"Invoice INV-4410 Freight $1,250.00")
    [record] = ingest_folder(store, settings)
    assert [att.filename for att in record.attachments] == ["Invoice 4410.txt"]

    # Exported again later (a bulk "save as" of the folder), this time with no file beside it.
    _eml(settings, "Invoice 4410", "Invoice attached, due Oct 30.", message_id="<inv4410@vendor.com>")
    report: dict = {}
    ingest_folder(store, settings, report=report)
    assert report["already_read"] == 1 and report["failed"] == []
    again = store.get_email(record.id)
    assert [att.filename for att in again.attachments] == ["Invoice 4410.txt"]
    assert 1250.0 in again.extracted.amounts and again.has_attachments
    assert "missing_attachment" not in again.flags


def test_a_second_copy_in_one_drop_still_brings_its_files(settings: Settings, store: Store):
    # "Invoice 4410 (2).eml" sorts first and has no file beside it; the copy that does must not be skipped.
    settings.ensure_data_dir()
    _eml(settings, "Invoice 4410", message_id="<inv4410@vendor.com>")
    _eml(settings, "Invoice 4410", message_id="<inv4410@vendor.com>").rename(settings.inbox_incoming / "Invoice 4410 (2).eml")
    _eml(settings, "Invoice 4410", message_id="<inv4410@vendor.com>")
    _sidecar(settings, "Invoice 4410.txt", b"Invoice INV-4410 Freight $1,250.00")
    report: dict = {}
    records = ingest_folder(store, settings, report=report)
    assert (report["read"], report["already_read"]) == (1, 1)
    assert len(records) == 1 and [att.filename for att in records[0].attachments] == ["Invoice 4410.txt"]
    assert "INV-4410" in records[0].attachments[0].extracted_text
    assert not list(settings.inbox_incoming.iterdir())


def test_a_copy_of_a_read_message_adds_its_new_files_and_keeps_the_reading(settings: Settings, store: Store):
    from controller_inbox.reading import apply_bionic_reading

    settings.ensure_data_dir()
    _eml(settings, "Invoice 4410", message_id="<inv4410@vendor.com>")
    _sidecar(settings, "Invoice 4410.txt", b"Invoice INV-4410 Freight $1,250.00")
    [record] = ingest_folder(store, settings)
    reading = {"category": "ap_invoice", "folder": "important", "importance": "high", "summary": "Freight invoice.", "actions": [], "why": "x"}
    apply_bionic_reading(store, record.id, reading)

    # The same message again, with a different file under the same name: both files are kept.
    _eml(settings, "Invoice 4410", message_id="<inv4410@vendor.com>")
    _sidecar(settings, "Invoice 4410.txt", b"Revised invoice INV-4410 Freight $1,300.00")
    ingest_folder(store, settings)
    kept = store.get_email(record.id)
    assert kept.model_status == "bionic" and kept.summary == "Freight invoice."
    assert sorted(att.filename for att in kept.attachments) == ["Invoice 4410 (2).txt", "Invoice 4410.txt"]
    folder = settings.inbox_extracted / record.id
    assert b"$1,250.00" in (folder / "Invoice 4410.txt").read_bytes()
    assert b"$1,300.00" in (folder / "Invoice 4410 (2).txt").read_bytes()

    # A third copy with a file already stored adds nothing.
    _eml(settings, "Invoice 4410", message_id="<inv4410@vendor.com>")
    _sidecar(settings, "Invoice 4410.txt", b"Revised invoice INV-4410 Freight $1,300.00")
    report: dict = {}
    assert ingest_folder(store, settings, report=report) == [] and report["already_read"] == 1
    assert len(store.get_email(record.id).attachments) == 2


# 17. Different emails that share a Message-ID stay apart ------------------------------------


def test_different_emails_with_one_message_id_are_kept_apart(settings: Settings, store: Store):
    # A scan-to-email copier gives every scan the same Message-ID.
    settings.ensure_data_dir()
    scan = "<scan@copier.local>"
    _eml(settings, "Scan 0001", "Scanned: vendor invoice", message_id=scan, attach=[("Scan 0001.txt", b"Invoice INV-7001 total $4,200.00")])
    second = _eml(settings, "Scan 0002", "Scanned: signed contract", message_id=scan, attach=[("Scan 0002.txt", b"Contract signed")],
                  date_header="Mon, 05 Oct 2026 11:00:00 -0400").read_bytes()
    report: dict = {}
    ingest_folder(store, settings, report=report)
    assert (report["read"], report["already_read"]) == (2, 0)

    # The next day another scan, while the first two are still script drafts.
    _eml(settings, "Scan 0003", "Scanned: bank statement", message_id=scan, attach=[("Scan 0003.txt", b"Statement Sept 2026")],
         date_header="Tue, 06 Oct 2026 09:00:00 -0400")
    ingest_folder(store, settings)
    stored = {email.subject: [att.filename for att in email.attachments] for email in store.list_emails()}
    assert stored == {"Scan 0001": ["Scan 0001.txt"], "Scan 0002": ["Scan 0002.txt"], "Scan 0003": ["Scan 0003.txt"]}

    # The second scan dropped again is still one email.
    path = settings.inbox_incoming / "Scan 0002 again.eml"
    path.write_bytes(second)
    _age(path)
    report = {}
    ingest_folder(store, settings, report=report)
    assert report["already_read"] == 1 and store.counts()["emails"] == 3


# 18. A message that failed on one sync is read on the next -------------------------------------


class _FlakyMailbox(_Mailbox):
    """Lists messages by their received time, as Graph does; ``failures`` says how often each one fails first."""

    def __init__(self, raws, failures):
        super().__init__(raws)
        self.failures = dict(failures)
        self.fetched: list[str] = []

    def list_messages(self, received_after=None):
        return iter([raw for raw in self.raws if received_after is None or raw.received_at >= received_after])

    def get_attachments(self, message_id):
        self.fetched.append(message_id)
        if self.failures.get(message_id, 0):
            self.failures[message_id] -= 1
            from controller_inbox.graph import GraphError

            raise GraphError("Microsoft Graph throttled the request (HTTP 429). Wait and retry.")
        return []


def _graph_raw(message_id: str, received: datetime) -> RawMessage:
    raw = _raw("Invoice attached", received=received)
    raw.id, raw.source, raw.has_attachments = message_id, "graph", True
    return raw


def test_a_message_that_failed_on_a_sync_is_read_on_the_next(store: Store, settings: Settings):
    from datetime import timedelta

    t0 = datetime(2026, 10, 6, 9, 0, tzinfo=timezone.utc)
    mailbox = _FlakyMailbox([_graph_raw("throttled", t0), _graph_raw("fine", t0 + timedelta(minutes=1))], {"throttled": 1})
    report: dict = {}
    ingest_mailbox(mailbox, store, settings, received_after=t0 - timedelta(hours=72), now=t0 + timedelta(minutes=5), report=report)
    assert [row["id"] for row in report["failed"]] == ["throttled"] and store.get_email("fine")
    cursor = datetime.fromisoformat(store.get_state("last_sync_at"))
    assert cursor <= t0, "the cursor waits for the message that failed"

    mailbox.fetched.clear()
    ingest_mailbox(mailbox, store, settings, received_after=cursor, now=t0 + timedelta(hours=1))
    assert store.get_email("throttled") is not None
    assert mailbox.fetched == ["throttled"], "the message read last time is not fetched again"
    assert store.get_state("last_sync_at") == (t0 + timedelta(hours=1)).isoformat()


def test_a_message_that_always_fails_stops_holding_the_cursor(store: Store, settings: Settings, monkeypatch):
    from datetime import timedelta

    from controller_inbox import pipeline

    monkeypatch.setattr(pipeline, "MAX_SYNC_TRIES", 2)
    t0 = datetime(2026, 10, 6, 9, 0, tzinfo=timezone.utc)
    mailbox = _FlakyMailbox([_graph_raw("broken", t0)], {"broken": 99})
    ingest_mailbox(mailbox, store, settings, now=t0 + timedelta(minutes=5))
    assert store.get_state("last_sync_at") == t0.isoformat()
    later = t0 + timedelta(hours=1)
    ingest_mailbox(mailbox, store, settings, received_after=t0, now=later)
    assert store.get_state("last_sync_at") == later.isoformat(), "after the last try the sync moves on"


# 19. watch reads the drop folder and writes the digest while Outlook is down ---------------------


def test_watch_reads_the_drop_folder_while_outlook_is_down(settings: Settings, store: Store, monkeypatch, capsys):
    from controller_inbox import cli
    from controller_inbox.graph import GraphError

    class Down:
        def list_messages(self, received_after=None):
            raise GraphError("Graph 503: Service Unavailable")

    settings.ensure_data_dir()
    settings.azure_client_id = "configured"
    monkeypatch.setattr(cli, "_graph_mailbox", lambda _settings: Down())
    _eml(settings, "Dropped while Outlook is down", message_id="<w1@x>")
    result = cli.watch_tick(settings, store, now=datetime(2026, 10, 6, 9, 0, tzinfo=settings.tz))
    assert [record.subject for record in result["records"]] == ["Dropped while Outlook is down"]
    assert "503" in result["graph_note"] and result["digest"] is not None
    assert not list(settings.inbox_incoming.iterdir())
    assert cli._watch(settings, store, once=True) == 0
    assert "Outlook could not be read (Graph 503" in capsys.readouterr().out


# 20. Reading attachments again never gives a zipped file another file's text --------------------


def test_reading_again_reads_a_zipped_file_from_its_zip(settings: Settings, store: Store):
    from controller_inbox.demo import make_pdf

    settings.ensure_data_dir()
    outer = make_pdf([["Invoice INV-1001 from Alpha Freight", ["Item", "Amount"], ["Freight", "$1,000.00"]]])
    inner = make_pdf([["Invoice INV-2002 from Beta Paper", ["Item", "Amount"], ["Paper", "$2,000.00"]]])
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("invoice.pdf", inner)
    msg = EmailMessage()
    msg["From"] = "AP <ap@vendor.com>"
    msg["Subject"] = "Two invoices"
    msg["Message-ID"] = "<two@x>"
    msg["Date"] = "Mon, 05 Oct 2026 10:00:00 -0400"
    msg.set_content("See attached.")
    # An "invoice.pdf" beside a zip that holds another "invoice.pdf".
    msg.add_attachment(outer, maintype="application", subtype="pdf", filename="invoice.pdf")
    msg.add_attachment(buffer.getvalue(), maintype="application", subtype="zip", filename="older.zip")
    path = settings.inbox_incoming / "two.eml"
    path.write_bytes(bytes(msg))
    _age(path)
    [record] = ingest_folder(store, settings)
    zipped = next(att for att in record.attachments if att.id.endswith("older.zip:invoice.pdf"))
    assert "INV-2002" in zipped.extracted_text

    store.set_attachment_text(zipped.id, "[page 1]\nas an older reader left it")
    store.set_state(folder_mail.READER_KEY, "old")
    assert reread_attachments(store, settings) == 1
    texts = {att.id: att.extracted_text for att in store.get_email(record.id).attachments}
    assert "INV-2002" in texts[zipped.id] and "INV-1001" not in texts[zipped.id]
    assert "INV-1001" in texts[f"{record.id}:invoice.pdf"]


# 21. A long non-Latin attachment name is cut to fit, and files are saved before the email --------


def test_a_long_non_latin_file_name_is_cut_to_fit_and_keeps_its_extension(settings: Settings, store: Store):
    settings.ensure_data_dir()
    name = "請求書" * 30 + ".txt"  # 94 characters but 274 bytes in UTF-8, more than a file name can hold
    _eml(settings, "請求書", message_id="<jp1@x>", attach=[(name, b"first copy"), (name, b"second copy")])
    report: dict = {}
    [record] = ingest_folder(store, settings, report=report)
    assert report["failed"] == [] and not list(settings.inbox_failed.iterdir())
    assert record.source_path.endswith(".eml")
    saved = {path.name: path.read_bytes() for path in (settings.inbox_extracted / record.id).iterdir()}
    assert sorted(saved.values()) == [b"first copy", b"second copy"]
    assert all(len(name.encode("utf-8")) <= 255 and name.endswith(".txt") for name in saved)
    assert any(name.endswith(" (2).txt") for name in saved), "the number of the second copy is not cut off"


def test_an_email_is_not_stored_when_its_files_cannot_be_saved(settings: Settings, store: Store, monkeypatch):
    settings.ensure_data_dir()
    _eml(settings, "Disk full", message_id="<full@x>", attach=[("a.txt", b"data")])

    def full(*_args):
        raise OSError("No space left on device")

    monkeypatch.setattr(folder_mail, "_write_extracted", full)
    report: dict = {}
    assert ingest_folder(store, settings, report=report) == []
    assert len(report["failed"]) == 1 and store.counts()["emails"] == 0
    assert (settings.inbox_failed / "Disk full.eml").exists()


# 22. Outlook sync reads Reply-To, so mail whose replies go elsewhere is flagged ------------------


class _ReplyToClient(_FakeClient):
    def get_json(self, path, params=None):
        self.params.append(params)
        if path.endswith("/attachments"):
            return {"value": []}
        return {"value": [{
            "id": "AAMk1", "subject": "Updated remittance details", "receivedDateTime": "2026-10-05T14:00:00Z",
            "from": {"emailAddress": {"name": "Acme Billing", "address": "billing@acme.com"}},
            "replyTo": [{"emailAddress": {"name": "Acme Billing", "address": "Acme.Billing@protonmail.com"}}],
            "body": {"contentType": "text", "content": "Please use our new bank account for all payments."},
            "hasAttachments": False, "importance": "normal", "isRead": False,
        }]}


def test_graph_sync_reads_reply_to_for_the_fraud_check(store: Store, settings: Settings):
    client = _ReplyToClient()
    [record] = ingest_mailbox(GraphMailbox(client), store, settings, now=NOW)
    assert "replyTo" in client.params[0]["$select"].split(",")
    assert record.reply_to == "acme.billing@protonmail.com"
    assert "reply_to_mismatch" in {signal["key"] for signal in store.fraud_check(record.id)["signals"]}


# 23. Overnight file summaries: flagged mail doesn't use up the night, empty ones aren't retried -----


LONG_FILE = ("Line item freight services rendered in September per contract. " * 60).encode()


def _mail_with_long_file(store: Store, settings: Settings, n: int, *, scam: bool):
    raw = _raw(
        "Our bank details have changed. Please use the new account for all payments from today."
        if scam else "Please find the monthly report attached.",
        received=NOW,
        attachments=[RawAttachment(id="f.txt", filename=f"file{n}.txt", content_type="text/plain", size_bytes=len(LONG_FILE), content=LONG_FILE)],
    )
    raw.id, raw.subject = f"m{n}", "Updated remittance details" if scam else f"Monthly report {n}"
    raw.sender_email = f"billing@acme-pay{n}.net" if scam else "maya@taz.com"
    return process_message(raw, store, settings, now=NOW)


def test_files_on_flagged_mail_do_not_use_up_the_nights_summaries(store: Store, settings: Settings, monkeypatch):
    from controller_inbox import file_summaries

    # Flagged mail scores highest, so it would fill a night's quota of 3 and nothing else would ever be summarized.
    assert all("fraud_risk" in _mail_with_long_file(store, settings, n, scam=True).flags for n in range(3))
    _mail_with_long_file(store, settings, 9, scam=False)
    calls: list[str] = []
    monkeypatch.setattr(file_summaries, "summarize_file", lambda _settings, att: calls.append(att.filename) or "- Freight for September")
    assert file_summaries.summarize_files(store, settings, limit=3, model="m") == 1
    assert calls == ["file9.txt"]


def test_a_file_the_model_cannot_summarize_is_not_tried_every_night(store: Store, settings: Settings, monkeypatch):
    from controller_inbox import file_summaries

    _mail_with_long_file(store, settings, 1, scam=False)
    calls: list[str] = []
    monkeypatch.setattr(file_summaries, "summarize_file", lambda _settings, att: calls.append(att.id) or "")
    for _night in range(3):
        assert file_summaries.summarize_files(store, settings, limit=5, model="small") == 0
    assert calls == ["m1:f.txt"]
    file_summaries.summarize_files(store, settings, limit=5, model="bigger")
    assert len(calls) == 2, "another model gets a try"


# 24. A NUL character in a file's text is not stored -------------------------------------------


NUL_EXPORT = b"2026-09-01  Wire to Alpha Freight   1,250.00\n" * 80 + b"END\x00\x00\n"


def test_a_file_with_a_nul_in_its_text_is_summarized_once(store: Store, settings: Settings, monkeypatch):
    from controller_inbox import file_summaries

    # SQLite's length() stops at a NUL and Python's len() doesn't, so the summary never matched its text.
    export = RawAttachment(id="export.txt", filename="export.txt", content_type="text/plain", size_bytes=len(NUL_EXPORT), content=NUL_EXPORT)
    record = process_message(_raw("Export attached", received=NOW, attachments=[export]), store, settings, now=NOW)
    assert "\x00" not in record.attachments[0].extracted_text
    calls: list[str] = []
    monkeypatch.setattr(file_summaries, "summarize_file", lambda _settings, att: calls.append(att.id) or "- Wires to Alpha Freight")
    for _night in range(3):
        file_summaries.summarize_files(store, settings, limit=20, model="m")
    assert calls == ["raw-1:export.txt"]


def test_text_stored_with_a_nul_before_is_cleaned_once(store: Store, settings: Settings):
    import sqlite3

    export = RawAttachment(id="export.txt", filename="export.txt", content_type="text/plain", size_bytes=3, content=b"abc")
    process_message(_raw("Export attached", received=NOW, attachments=[export]), store, settings, now=NOW)
    with sqlite3.connect(settings.db_path) as conn:  # as an older version left it
        conn.execute("UPDATE attachments SET extracted_text = ?", ("a\x00b",))
        conn.execute("DELETE FROM sync_state WHERE key = 'attachment_text_without_nul'")
    assert Store(settings.db_path).get_email("raw-1").attachments[0].extracted_text == "ab"


# 25. VIP senders can be separated by semicolons, like trusted domains -----------------------------


def test_vip_senders_split_on_semicolons_too():
    settings = Settings(vip_senders="cfo@taz.com; ceo@taz.com,board@taz.com", _env_file=None)
    assert settings.vip_list == ["cfo@taz.com", "ceo@taz.com", "board@taz.com"]


# 26. An invitation's calendar sent inline, with no file name, is kept ----------------------------


INVITE = "BEGIN:VCALENDAR\r\nMETHOD:REQUEST\r\nBEGIN:VEVENT\r\nSUMMARY:Q3 close review\r\nDTSTART:20261008T150000Z\r\nEND:VEVENT\r\nEND:VCALENDAR\r\n"


def _invite(settings: Settings, name: str, *, attached: bool) -> None:
    msg = EmailMessage()
    msg["From"] = "Controller <cfo@taz.com>"
    msg["Subject"] = "Q3 close review"
    msg["Message-ID"] = f"<{name}@x>"
    msg["Date"] = "Mon, 05 Oct 2026 10:00:00 -0400"
    msg.set_content("You have been invited.")
    msg.add_alternative(INVITE, subtype="calendar", params={"method": "REQUEST"})
    if attached:
        msg.add_attachment(INVITE.encode(), maintype="text", subtype="calendar", filename="invite.ics")
    path = settings.inbox_incoming / f"{name}.eml"
    path.write_bytes(bytes(msg))
    _age(path)


def test_an_inline_calendar_is_kept_as_the_invite(settings: Settings, store: Store):
    settings.ensure_data_dir()
    _invite(settings, "inline", attached=False)
    _invite(settings, "both", attached=True)
    records = {record.internet_message_id: record for record in ingest_folder(store, settings)}
    inline = records["<inline@x>"]
    assert [att.filename for att in inline.attachments] == ["invite.ics"]
    assert "Q3 close review" in inline.attachments[0].extracted_text
    assert [att.filename for att in records["<both@x>"].attachments] == ["invite.ics"], "not kept twice"


# 27. The overnight reading is saved for an email whose files didn't arrive in name order --------


def test_a_reading_is_saved_for_an_email_whose_files_are_not_in_name_order(settings: Settings, store: Store):
    from controller_inbox.overnight import read_queue
    from test_bionic import AgreeingReader

    settings.ensure_data_dir()
    # "Statement" comes before "Invoice" in the email; the reading was checked against a copy in name order.
    _eml(settings, "September statement", message_id="<st9@vendor.com>",
         attach=[("Statement.txt", b"Statement of account Sept 2026"), ("Invoice INV-9.txt", b"Invoice INV-9 $310.00")])
    [record] = ingest_folder(store, settings)
    result = read_queue(store, settings, now=NOW, reader=AgreeingReader())
    assert result["read_ids"] == [record.id]
    assert store.get_email(record.id).model_status == "bionic"
    assert store.counts()["waiting_on_bionic"] == 0


# 28. A copy of a read email with new files doesn't undo what the user did while its files were read --


@pytest.mark.parametrize("meanwhile", ["correction", "verdict"])
def test_new_files_on_a_read_email_keep_what_the_user_did_meanwhile(settings: Settings, store: Store, monkeypatch, meanwhile):
    from controller_inbox import pipeline
    from controller_inbox.fraud import record_fraud_verdict
    from controller_inbox.learn import record_correction
    from controller_inbox.reading import apply_bionic_reading

    settings.ensure_data_dir()
    _eml(settings, "Invoice 4410", message_id="<inv4410@vendor.com>")
    _sidecar(settings, "Invoice 4410.txt", b"Invoice INV-4410 Freight $1,250.00")
    [record] = ingest_folder(store, settings)
    reading = {"category": "ap_invoice", "folder": "important", "importance": "high", "summary": "Freight invoice.", "actions": [], "why": "x"}
    apply_bionic_reading(store, record.id, reading)

    real = pipeline.attachment_text

    def slow_read(filename, content_type, data):
        # A scan read with OCR takes a while; the user works on the email in the dashboard meanwhile.
        if meanwhile == "correction":
            record_correction(store, settings, email_id=record.id, corrected_category="newsletter", reason="it is a newsletter")
        else:
            record_fraud_verdict(store, settings, record.id, verdict="fraud", note="phoned the vendor; it is fake")
        return real(filename, content_type, data)

    monkeypatch.setattr(pipeline, "attachment_text", slow_read)
    _eml(settings, "Invoice 4410", message_id="<inv4410@vendor.com>")
    _sidecar(settings, "Invoice 4410.txt", b"Revised invoice INV-4410 Freight $1,300.00")
    ingest_folder(store, settings)

    kept = store.get_email(record.id)
    assert sorted(att.filename for att in kept.attachments) == ["Invoice 4410 (2).txt", "Invoice 4410.txt"]
    if meanwhile == "correction":
        assert kept.model_status == "corrected" and kept.category.value == "newsletter"
    else:
        assert {"fraud_risk", "fraud_confirmed"} <= set(kept.flags)


# 29. A copy of an email with no Subject or no Date, saved again under another name, is still one email --


def _scan(settings: Settings, name: str, *, subject: str | None, date_header: str | None, age: float, data: bytes | None = None) -> bytes:
    """A copier's scan; ``data`` saves the very bytes of an earlier one again under ``name``."""
    if data is None:
        msg = EmailMessage()
        msg["From"] = "Scanner <scan@copier.local>"
        msg["Message-ID"] = "<scan@copier.local>"
        if subject is not None:
            msg["Subject"] = subject
        if date_header is not None:
            msg["Date"] = date_header
        msg.set_content("Scanned document attached.")
        msg.add_attachment(f"Scanned page {name}".encode(), maintype="text", subtype="plain", filename="scan.txt")
        data = bytes(msg)
    path = settings.inbox_incoming / name
    path.write_bytes(data)
    old = time.time() - age
    os.utime(path, (old, old))
    return data


@pytest.mark.parametrize("missing", ["subject", "date"])
def test_a_copy_saved_again_is_one_email_when_the_message_has_no_subject_or_date(settings: Settings, store: Store, missing):
    settings.ensure_data_dir()
    subject = None if missing == "subject" else "Scan"
    date_header = None if missing == "date" else "Mon, 05 Oct 2026 10:00:00 -0400"
    # Its subject would be the file's name, or its date the file's date: both differ for the copy saved later.
    first = _scan(settings, "Scan from copier.eml", subject=subject, date_header=date_header, age=3600)
    ingest_folder(store, settings)
    _scan(settings, "Scan from copier (copy).eml", subject=subject, date_header=date_header, age=60, data=first)
    report: dict = {}
    ingest_folder(store, settings, report=report)
    assert (report["read"], report["already_read"]) == (0, 1) and store.counts()["emails"] == 1


def test_scans_without_a_subject_that_share_a_message_id_are_still_kept_apart(settings: Settings, store: Store):
    settings.ensure_data_dir()
    _scan(settings, "a.eml", subject=None, date_header="Mon, 05 Oct 2026 10:00:00 -0400", age=60)
    _scan(settings, "b.eml", subject=None, date_header="Mon, 05 Oct 2026 11:00:00 -0400", age=60)
    second = (settings.inbox_incoming / "b.eml").read_bytes()
    report: dict = {}
    ingest_folder(store, settings, report=report)
    assert report["read"] == 2 and store.counts()["emails"] == 2
    _scan(settings, "b again.eml", subject=None, date_header=None, age=30, data=second)  # the second, saved again
    report = {}
    ingest_folder(store, settings, report=report)
    assert report["already_read"] == 1 and store.counts()["emails"] == 2


# 30. VIP senders copied from Outlook count by their addresses, not the words of their names ------


def test_vip_senders_copied_from_outlook_count_by_their_addresses(store: Store, settings: Settings):
    pasted = Settings(vip_senders="Chen, Maya <maya@taz.com>; Bob Lee <Bob@taz.com>", _env_file=None)
    assert pasted.vip_list == ["maya@taz.com", "bob@taz.com"]
    assert Settings(vip_senders="cfo@taz.com ceo@taz.com; irs.gov, Maya Chen", _env_file=None).vip_list == [
        "cfo@taz.com", "ceo@taz.com", "irs.gov"
    ]
    assert Settings(vip_senders="treasurer", _env_file=None).vip_list == ["treasurer"]

    settings.vip_senders = pasted.vip_senders
    vip = {}
    for n, sender in enumerate(["maya@taz.com", "colleen@randomvendor.com", "noreply@bobcat-rentals.com"]):
        raw = _raw("Just checking in.", received=NOW)
        raw.id, raw.sender_email = f"v{n}", sender
        vip[sender] = "VIP / elevated sender" in process_message(raw, store, settings, now=NOW).importance_reasons
    assert vip == {"maya@taz.com": True, "colleen@randomvendor.com": False, "noreply@bobcat-rentals.com": False}


# 31. A manual sync of a shorter window leaves the cursor and the message it waits for alone ------


def test_a_shorter_manual_sync_keeps_the_message_the_cursor_waits_for(store: Store, settings: Settings):
    from datetime import timedelta

    t0 = datetime(2026, 10, 6, 9, 0, tzinfo=timezone.utc)
    mailbox = _FlakyMailbox([_graph_raw("throttled", t0)], {"throttled": 1})
    ingest_mailbox(mailbox, store, settings, received_after=t0 - timedelta(hours=72), now=t0 + timedelta(minutes=5))
    held = store.get_state("last_sync_at")
    assert held == t0.isoformat()

    # "closedesk sync --hours 1" three hours later: it doesn't reach back to the message that failed.
    later = t0 + timedelta(hours=3)
    ingest_mailbox(mailbox, store, settings, received_after=later - timedelta(hours=1), now=later)
    assert store.get_state("last_sync_at") == held, "the cursor doesn't jump past mail this sync didn't read"

    ingest_mailbox(mailbox, store, settings, received_after=datetime.fromisoformat(held), now=later + timedelta(hours=1))
    assert store.get_email("throttled") is not None
    assert store.get_state("last_sync_at") == (later + timedelta(hours=1)).isoformat()


# 32. A copy dropped again doesn't rename or redate the email, or bring back snoozed tasks ----------


def _bare_copy(settings: Settings, name: str, body: str, message_id: str) -> Path:
    """The same message saved again by a tool that kept only some of its headers (no Subject, no Date)."""
    msg = EmailMessage()
    msg["From"] = "Vendor AP <ap@vendor.com>"
    msg["Message-ID"] = message_id
    msg.set_content(body)
    path = settings.inbox_incoming / name
    path.write_bytes(bytes(msg))
    _age(path)
    return path


def test_a_copy_dropped_again_keeps_the_name_the_date_and_snoozed_tasks(settings: Settings, store: Store):
    from controller_inbox.reading import apply_bionic_reading

    settings.ensure_data_dir()
    body = "Please approve and pay invoice INV-4410 for $12,400.00. Due date: October 15, 2026."
    subject = "Invoice INV-4410 due October 15, 2026"
    _eml(settings, subject, body, message_id="<inv4410@vendor.com>", date_header="Mon, 21 Sep 2026 10:00:00 +0000")
    [record] = ingest_folder(store, settings)
    assert record.actions
    for task in record.actions:
        store.set_action_status(task.id, "snoozed")
    snoozed = {task.title for task in record.actions}

    def check(email) -> None:
        assert email.subject == subject and email.received_at == record.received_at
        assert {task.title: task.status.value for task in email.actions if task.title in snoozed} == dict.fromkeys(snoozed, "snoozed")
        assert store.counts()["emails"] == 1

    # A copy with nothing new is already read: nothing about the email changes.
    _bare_copy(settings, "scan-export-0001.eml", body, "<inv4410@vendor.com>")
    report: dict = {}
    assert ingest_folder(store, settings, report=report) == [] and report["already_read"] == 1
    check(store.get_email(record.id))

    # A copy that brings a file still adds it, and keeps the name, the date and the snoozed tasks.
    _bare_copy(settings, "scan-export-0002.eml", body, "<inv4410@vendor.com>")
    _sidecar(settings, "scan-export-0002.txt", b"Remittance advice for INV-4410 $12,400.00")
    report = {}
    ingest_folder(store, settings, report=report)
    assert report["already_read"] == 1 and report["failed"] == []
    again = store.get_email(record.id)
    assert [att.filename for att in again.attachments] == ["scan-export-0002.txt"]
    check(again)

    # The model can still read it afterwards.
    reading = {"category": "ap_invoice", "folder": "important", "importance": "high", "summary": "Pay INV-4410.", "actions": [], "why": "x"}
    assert apply_bionic_reading(store, record.id, reading).model_status == "bionic"


def test_a_copy_of_an_email_whose_subject_has_an_account_number_is_still_one_email(settings: Settings, store: Store):
    # The stored subject is masked ("account ****6677"); the copy's own subject isn't yet when they are compared.
    settings.ensure_data_dir()
    _eml(settings, "Remittance - Acct 44556677", "Paid.", message_id="<rem1@vendor.com>")
    [record] = ingest_folder(store, settings)
    assert record.subject == "Remittance - account ****6677"
    copy = _eml(settings, "Remittance - Acct 44556677", "Paid.", message_id="<rem1@vendor.com>")
    copy.rename(settings.inbox_incoming / "Remittance copy.eml")
    report: dict = {}
    ingest_folder(store, settings, report=report)
    assert report["already_read"] == 1 and store.counts()["emails"] == 1


# 33. A Date header past year 9999 in UTC, 8-bit text with no character set ------------------------


def test_an_email_dated_past_what_utc_can_hold_is_read_and_dated_by_its_file(settings: Settings, store: Store):
    settings.ensure_data_dir()
    path = _eml(settings, "Odd date", "Invoice INV-1 for $10.00.", date_header="Fri, 31 Dec 9999 23:30:00 -0100")
    dropped = datetime.fromtimestamp(path.stat().st_mtime, tz=timezone.utc)
    report: dict = {}
    [record] = ingest_folder(store, settings, report=report)
    assert report["failed"] == []
    assert datetime.fromisoformat(record.received_at) == dropped


def test_8bit_text_without_a_character_set_keeps_its_accents(settings: Settings, store: Store):
    settings.ensure_data_dir()
    data = (
        b"From: =?utf-8?q?Jos=C3=A9?= <jose@vendor.com>\r\nTo: ap@co.com\r\n"
        b"Subject: Factura n\xba 4471 \xe9t\xe9\r\nDate: Mon, 21 Sep 2026 10:00:00 +0000\r\nMessage-ID: <a1@vendor.com>\r\n"
        b"MIME-Version: 1.0\r\nContent-Type: text/plain\r\nContent-Transfer-Encoding: 8bit\r\n\r\n"
        b"Montant d\xfb: 1 200,00 EUR. Caf\xe9.\r\n"
    )
    path = settings.inbox_incoming / "factura.eml"
    path.write_bytes(data)
    _age(path)
    [record] = ingest_folder(store, settings)
    assert record.subject == "Factura nº 4471 été"
    assert record.body_text == "Montant dû: 1 200,00 EUR. Café."
    # Declared UTF-8 stays UTF-8.
    assert folder_mail._decode_with("Café".encode(), "utf-8") == "Café"


# 34. An attached email its sender's program encoded in base64 is still read --------------------


def test_an_attached_email_encoded_in_base64_is_read(settings: Settings, store: Store):
    import base64

    settings.ensure_data_dir()
    inner = (
        b"From: Vendor <ar@vendor.com>\r\nSubject: Invoice INV-9001\r\nDate: Mon, 21 Sep 2026 09:00:00 +0000\r\n"
        b"MIME-Version: 1.0\r\nContent-Type: multipart/mixed; boundary=YY\r\n\r\n--YY\r\nContent-Type: text/plain\r\n\r\n"
        b"Invoice INV-9001 for $3,400.00 attached.\r\n--YY\r\nContent-Type: text/csv\r\n"
        b"Content-Disposition: attachment; filename=\"inv.csv\"\r\n\r\nitem,amount\r\nwidgets,3400.00\r\n--YY--\r\n"
    )
    outer = (
        b"From: Ann <ann@co.com>\r\nTo: ap@co.com\r\nSubject: FW invoice\r\nDate: Mon, 21 Sep 2026 10:00:00 +0000\r\n"
        b"Message-ID: <fw-b64@co.com>\r\nMIME-Version: 1.0\r\nContent-Type: multipart/mixed; boundary=XX\r\n\r\n"
        b"--XX\r\nContent-Type: text/plain\r\n\r\nFYI\r\n--XX\r\nContent-Type: message/rfc822; name=\"Invoice.eml\"\r\n"
        b"Content-Disposition: attachment; filename=\"Invoice.eml\"\r\nContent-Transfer-Encoding: base64\r\n\r\n"
        + base64.encodebytes(inner) + b"\r\n--XX--\r\n"
    )
    path = settings.inbox_incoming / "fw.eml"
    path.write_bytes(outer)
    _age(path)
    [record] = ingest_folder(store, settings)
    files = {att.filename: att.extracted_text for att in record.attachments}
    assert set(files) == {"Invoice.txt", "Invoice › inv.csv"}
    assert "Subject: Invoice INV-9001" in files["Invoice.txt"] and "3400.00" in files["Invoice › inv.csv"]
    assert "INV-9001" in record.extracted.invoice_numbers


# 35. A zip with files CloseDesk doesn't unpack keeps a note of them ------------------------------


def test_a_zip_with_more_files_than_are_unpacked_is_kept_with_a_list_of_the_rest(settings: Settings, store: Store):
    settings.ensure_data_dir()
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as packed:
        for number in range(1, 46):
            packed.writestr(f"invoices/INV-{number:04d}.csv", f"invoice,amount\nINV-{number:04d},{number * 100}.00\n")
    path = settings.inbox_incoming / "September AP batch.zip"
    path.write_bytes(archive.getvalue())
    _age(path)
    [record] = ingest_folder(store, settings)
    names = {att.filename: att for att in record.attachments}
    assert len(names) == 41 and "INV-0040.csv" in names and "INV-0041.csv" not in names
    note = names["September AP batch.zip"].extracted_text
    assert "did not unpack" in note and "INV-0041.csv (past the first 40 files)" in note and "INV-0045.csv" in note

    # Dropped again, it adds nothing and stays one email.
    path.write_bytes(archive.getvalue())
    _age(path)
    report: dict = {}
    ingest_folder(store, settings, report=report)
    assert report["already_read"] == 1 and len(store.get_email(record.id).attachments) == 41


def test_a_zip_that_is_unpacked_whole_is_not_kept_beside_its_files(settings: Settings, store: Store):
    settings.ensure_data_dir()
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as packed:
        packed.writestr("INV-1.csv", "invoice,amount\nINV-1,100.00\n")
    path = settings.inbox_incoming / "one.zip"
    path.write_bytes(archive.getvalue())
    _age(path)
    [record] = ingest_folder(store, settings)
    assert [att.filename for att in record.attachments] == ["INV-1.csv"]
