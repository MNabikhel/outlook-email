import io
import zipfile
from email.message import EmailMessage
from pathlib import Path

from controller_inbox.demo import _pdf
from controller_inbox.extract import extract_text_from_bytes, explode_archives
from controller_inbox.folder_mail import ingest_folder
from controller_inbox.learn import record_correction
from controller_inbox.models import DocumentType, RawAttachment


def _eml(path: Path, *, subject: str, sender: str, body: str, attachments: list[tuple[str, bytes, str]] | None = None) -> None:
    message = EmailMessage()
    message["Subject"] = subject
    message["From"] = sender
    message["Date"] = "Tue, 22 Sep 2026 09:00:00 -0400"
    message.set_content(body)
    for name, payload, subtype in attachments or []:
        maintype, subtype = subtype.split("/", 1)
        message.add_attachment(payload, maintype=maintype, subtype=subtype, filename=name)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(message.as_bytes())


def test_drop_folder_reads_eml_pdf_and_sidecar(store, settings):
    settings.ensure_data_dir()
    pdf = _pdf("Invoice INV-555 Amount due $2,200.00 Due date October 1, 2026 Vendor: Harbor Packaging")
    _eml(
        settings.inbox_incoming / "Harbor invoice.eml",
        subject="Invoice INV-555 from Harbor Packaging",
        sender="Harbor AP <ap@harbor.example>",
        body="Please enter invoice INV-555. Amount due $2,200.00. Due October 1, 2026.",
        attachments=[("INV-555.pdf", pdf, "application/pdf")],
    )
    side = settings.inbox_attachments / "Harbor invoice"
    side.mkdir(parents=True)
    side.joinpath("notes.txt").write_text("Vendor: Harbor Packaging invoice support", encoding="utf-8")

    records = ingest_folder(store, settings)
    assert len(records) == 1
    email = records[0]
    assert email.category == DocumentType.AP_INVOICE
    assert email.extracted.primary_invoice == "INV-555"
    names = {att.filename for att in email.attachments}
    assert names == {"INV-555.pdf", "notes.txt"}
    assert any("INV-555" in att.extracted_text for att in email.attachments)
    assert not (settings.inbox_incoming / "Harbor invoice.eml").exists()
    assert list(settings.inbox_processed.rglob("*.eml"))
    extracted = list((settings.inbox_extracted / email.id).iterdir())
    assert {path.name for path in extracted} == names


def test_correction_teaches_the_sender_but_not_fraud(store, settings):
    settings.ensure_data_dir()
    _eml(
        settings.inbox_incoming / "week1.eml",
        subject="Weekly pack",
        sender="Jordan <jordan@vendor.example>",
        body="Please review the attached file.",
        attachments=[("notes.txt", b"hello from jordan", "text/plain")],
    )
    first = ingest_folder(store, settings)[0]
    record_correction(
        store,
        settings,
        email_id=first.id,
        corrected_category="expense_report",
        reason="Jordan's weekly pack is always an expense report.",
    )
    saved = store.get_email(first.id)
    assert saved.category == DocumentType.EXPENSE_REPORT
    assert "user_trained" in saved.flags
    assert settings.training_path.read_text(encoding="utf-8").count("expense_report") == 1

    _eml(
        settings.inbox_incoming / "week2.eml",
        subject="Weekly pack 2",
        sender="Jordan <jordan@vendor.example>",
        body="Another file for you.",
        attachments=[("notes2.txt", b"more notes", "text/plain")],
    )
    _eml(
        settings.inbox_incoming / "wire.eml",
        subject="Updated banking instructions",
        sender="Jordan <jordan@vendor.example>",
        body="Our bank details have changed. Please use the following account and do not use the previous account.",
    )
    second = {item.subject: item for item in ingest_folder(store, settings)}
    assert second["Weekly pack 2"].category == DocumentType.EXPENSE_REPORT
    assert "user_trained" in second["Weekly pack 2"].flags
    assert second["Updated banking instructions"].category == DocumentType.PAYMENT_INSTRUCTION_CHANGE
    assert "fraud_risk" in second["Updated banking instructions"].flags


def test_zip_and_office_files_are_read():
    from openpyxl import Workbook
    from pptx import Presentation

    book = Workbook()
    book.active.append(["Invoice", "INV-900", "Amount due", 4400])
    buf = io.BytesIO()
    book.save(buf)
    text = extract_text_from_bytes("book.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", buf.getvalue())
    assert "INV-900" in text

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Bank reconciliation outstanding checks"
    out = io.BytesIO()
    deck.save(out)
    ppt = extract_text_from_bytes("close.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation", out.getvalue())
    assert "outstanding checks" in ppt

    zipped = io.BytesIO()
    with zipfile.ZipFile(zipped, "w") as archive:
        archive.writestr("inside.txt", "Invoice INV-42 amount due $10.00")
    exploded = explode_archives(
        [RawAttachment(id="z", filename="docs.zip", content_type="application/zip", size_bytes=10, content=zipped.getvalue())]
    )
    assert exploded[0].filename == "inside.txt"
    assert b"INV-42" in exploded[0].content
