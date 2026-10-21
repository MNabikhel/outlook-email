"""The attachment preview (sheet_preview.py, preview.js) in the cases found in review: Excel's number formats,
merges over hidden rows, hidden column ranges, .xls formats, Word merges, nested tables, content controls, text boxes
and moves, CSV quoting, and the thumbnail. Each test states the Excel or Word behaviour it relies on."""

from __future__ import annotations

import io
import json

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from openpyxl import Workbook
from openpyxl.styles import Font

from controller_inbox import sheet_preview
from controller_inbox.sheet_preview import display


def _xlsx(book: Workbook) -> bytes:
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _cells(sheet: dict) -> dict:
    return {(row["n"], cell["c"]): cell for row in sheet["rows"] for cell in row["cells"]}


# display(): Excel number formats --------------------------------------------------------------------


@pytest.mark.parametrize(
    "value, fmt, excel",
    [
        # Built-in 46 / 20-style elapsed time: Excel shows 1.25 days as 30:00, the brackets are not literal.
        (1.25, "[h]:mm", "30:00"),
        (1.25, "[h]:mm:ss", "30:00:00"),
        # Built-in 12 / 13 fractions.
        (0.75, "# ?/?", "3/4"),
        (1.5, "# ?/?", "1 1/2"),
        # Built-in 49 "@" (Text) on a number: Excel shows the number as typed.
        (123, "@", "123"),
        # Conditional sections: 500 fails [>=1000] so Excel uses the second section.
        (500, '[>=1000]#,##0,"K";0', "500"),
        # A quoted "%" is a literal: no x100 scaling.
        (1.234, '0.0"%"', "1.2%"),
        # Empty sections hide the value (";;;" is Excel's standard "hide cell values" format).
        (5, ";;;", ""),
        (-3, "0;", ""),
        (0, "0;-0;;@", ""),
        # Special formats: digit placeholders separated by literals.
        (123456789, "000-00-0000", "123-45-6789"),
        (12, "00-00", "00-12"),
        # General on a small number: Excel writes 1.234E-05, never Python's lowercase "e".
        (1.234e-05, "General", "1.234E-05"),
    ],
)
def test_display_matches_excel(value, fmt, excel):
    assert display(value, fmt)[0] == excel


def test_elapsed_time_cell_in_a_real_workbook_reads_as_hours():
    # A timesheet total: 30 hours in a [h]:mm cell. openpyxl hands it over as a timedelta.
    book = Workbook()
    ws = book.active
    ws["A1"] = 1.25
    ws["A1"].number_format = "[h]:mm"
    cells = _cells(sheet_preview.preview(_xlsx(book), "hours.xlsx")["sheets"][0])
    assert cells[(1, 1)]["v"] == "30:00"


# _sheet(): hidden columns and merges ----------------------------------------------------------------


def test_a_grouped_hidden_column_range_is_hidden_in_full():
    # Excel's "Hide" on B:D (or a collapsed outline group) is stored as one <col min=2 max=4 hidden=1>; Excel hides
    # all three columns. openpyxl keys the dimension under "B" only, so C and D come back shown.
    book = Workbook()
    ws = book.active
    ws.append(["Name", "Cost", "Margin", "Markup", "Price"])
    ws.column_dimensions.group("B", "D", hidden=True)
    cols = sheet_preview.preview(_xlsx(book), "quote.xlsx")["sheets"][0]["cols"]
    assert [c["letter"] for c in cols if c["hidden"]] == ["B", "C", "D"]


def test_a_merge_over_a_hidden_row_spans_only_the_rows_shown():
    # A5:A7 merged, row 6 hidden: the grid has rows 5 and 7 only, so the cell must span 2 table rows. With rs=3 the
    # <td rowspan=3> swallows row 8's first slot and every cell of row 8 moves one column right.
    book = Workbook()
    ws = book.active
    ws["A5"] = "Region"
    ws.merge_cells("A5:A7")
    ws.row_dimensions[6].hidden = True
    ws["B5"], ws["B7"], ws["A8"], ws["B8"] = "x", "y", "after", "z"
    sheet = sheet_preview.preview(_xlsx(book), "regions.xlsx")["sheets"][0]
    shown = [row["n"] for row in sheet["rows"]]
    assert shown[shown.index(5):] == [5, 7, 8]
    assert _cells(sheet)[(5, 1)]["rs"] == 2


# _xls(): Excel 97-2003 ---------------------------------------------------------------------------------


def test_an_xls_time_is_not_shown_as_a_date_in_1899(monkeypatch):
    # An .xls time-only cell (0.5 = 12:00 PM). Excel shows a time; it never shows 12/31/1899. The reader opens the
    # book with formatting_info=False and formats every cell as General, so the cell's format is never used.
    import xlrd
    from xlrd.sheet import Cell

    class Sheet:
        name, nrows, ncols, visibility = "Times", 1, 1, 0

        def cell(self, r, c):
            return Cell(xlrd.XL_CELL_DATE, 0.5)

    class Book:
        nsheets, datemode = 1, 0

        def sheet_by_index(self, i):
            return Sheet()

        def release_resources(self):
            pass

    monkeypatch.setattr(xlrd, "open_workbook", lambda **kw: Book())
    data = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 504
    value = sheet_preview.preview(data, "times.xls")["sheets"][0]["rows"][0]["cells"][0]["v"]
    assert "1899" not in value


# _csv() --------------------------------------------------------------------------------------------------


def test_a_csv_with_apostrophes_is_split_as_excel_splits_it():
    # Excel's CSV import only treats double quotes as quoting; csv.Sniffer guesses "'" and merges the fields.
    data = "Vendor,Note,Amt\nAcme,'Rent, March',5\nBeta,'Power, April',6\n".encode()
    row = sheet_preview.preview(data, "export.csv")["sheets"][0]["rows"][1]
    assert [c["v"] for c in row["cells"]] == ["Acme", "'Rent", " March'", "5"]


# docx_blocks() -----------------------------------------------------------------------------------------


def _docx(document) -> bytes:
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def test_a_horizontally_merged_word_cell_keeps_its_span():
    # "Q3" spans two columns in Word. The block drops the covered cell and says nothing of the span, so the grid
    # draws "Total" over column 2 and the row has 2 cells where the others have 3.
    document = Document()
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text, table.cell(0, 2).text = "Q3", "Total"
    for i, v in enumerate(["Rent", "Power", "1,200"]):
        table.cell(1, i).text = v
    [block] = sheet_preview.docx_blocks(_docx(document))
    widths = {len(row) for row in block["rows"]}
    assert len(widths) == 1 or any(isinstance(cell, dict) for row in block["rows"] for cell in row), block


def test_a_nested_word_table_keeps_its_text():
    document = Document()
    outer = document.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=1, cols=2)
    inner.cell(0, 0).text, inner.cell(0, 1).text = "Retainer", "$5,000"
    text = repr(sheet_preview.docx_blocks(_docx(document)))
    assert "Retainer" in text and "$5,000" in text


def test_text_in_a_content_control_is_shown():
    # Body-level <w:sdt> (cover pages, contract templates, form fields) holds ordinary paragraphs.
    document = Document()
    sdt = parse_xml(f'<w:sdt {nsdecls("w")}><w:sdtContent><w:p><w:r><w:t>Payment due in 30 days</w:t></w:r></w:p></w:sdtContent></w:sdt>')
    document.element.body.insert(0, sdt)
    assert "Payment due in 30 days" in repr(sheet_preview.docx_blocks(_docx(document)))


def test_a_text_box_is_not_shown_twice():
    # Word saves every text box twice: <mc:Choice> (DrawingML) and <mc:Fallback> (VML). Only one is read.
    document = Document()
    p = document.add_paragraph("See ")
    p._p.append(parse_xml(
        f'<w:r {nsdecls("w")} xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" xmlns:v="urn:schemas-microsoft-com:vml">'
        "<mc:AlternateContent><mc:Choice Requires=\"wps\"><w:drawing><wps:txbx><w:txbxContent><w:p><w:r><w:t>BOX</w:t></w:r></w:p>"
        "</w:txbxContent></wps:txbx></w:drawing></mc:Choice><mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent><w:p><w:r>"
        "<w:t>BOX</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r>"
    ))
    [block] = sheet_preview.docx_blocks(_docx(document))
    assert block["text"].count("BOX") == 1, block


def test_moved_text_reads_once_with_changes_accepted():
    # Tracked move: <w:moveFrom> is the old place (gone once accepted), <w:moveTo> the new one.
    document = Document()
    document.element.body.insert(0, parse_xml(
        f'<w:p {nsdecls("w")}><w:moveFrom w:id="1" w:author="a" w:date="2026-01-01T00:00:00Z"><w:r><w:t>Net 30. </w:t></w:r>'
        '</w:moveFrom><w:r><w:t>Terms: </w:t></w:r><w:moveTo w:id="2" w:author="a" w:date="2026-01-01T00:00:00Z"><w:r>'
        "<w:t>Net 30.</w:t></w:r></w:moveTo></w:p>"
    ))
    assert sheet_preview.docx_blocks(_docx(document))[0]["text"] == "Terms: Net 30."


# Front end (Playwright) ----------------------------------------------------------------------------------


def _review_book() -> bytes:
    book = Workbook()
    ws = book.active
    ws.title = "Quote"
    ws.append(["Name", "Cost", "Price"])
    ws.append(["Steel", 900, 1200])
    ws.column_dimensions["B"].hidden = True  # the cost column, hidden by the sender
    ws["A5"] = "Region"
    ws.merge_cells("A5:A7")
    ws.row_dimensions[6].hidden = True
    ws["B5"], ws["B7"], ws["A8"], ws["C8"] = "x", "y", "after", "z"
    ws["C2"].font = Font(color="FF0000FF")
    ws["A3"] = -5
    ws["A3"].number_format = "0;[Red]-0"
    ws["A3"].font = Font(color="FF0000FF")
    return _xlsx(book)


@pytest.fixture
def served(settings, store, browser):
    from controller_inbox import web
    from controller_inbox.folder_mail import ingest_folder
    from liveserver import serving
    from msgfactory import XLSX, write_msg

    settings.trusted_domains = "taz.com"
    settings.ensure_data_dir()
    write_msg(settings.inbox_incoming / "q.msg", "Quote", "Attached.", sender_name="Harbor Steel", sender_email="ar@taz.com",
              attachments=[("quote.xlsx", _review_book(), XLSX)])
    [email] = ingest_folder(store, settings)
    context = browser.new_context()
    context.set_default_timeout(15000)
    with serving(web.create_app(settings, store)) as base:
        yield context.new_page(), base, email
    context.close()


def _open_grid(page, base, email):
    page.set_viewport_size({"width": 1440, "height": 900})
    page.goto(f"{base}/app/mail/{email.id}/file/1?tab=preview")
    page.wait_for_selector(".xl-grid")


def test_grid_cells_after_a_merge_over_a_hidden_row_stay_in_their_columns(served):
    page, base, email = served
    _open_grid(page, base, email)
    # Column A's header and the cell in row 8 that holds "after" (column A in Excel) should line up.
    got = page.evaluate("""() => {
      const head = [...document.querySelectorAll('.xl-grid thead th')].find((th) => th.textContent === 'A').getBoundingClientRect();
      const row = [...document.querySelectorAll('.xl-grid tbody tr')].find((tr) => tr.querySelector('th').textContent === '8');
      const td = [...row.querySelectorAll('td')].find((td) => td.textContent === 'after').getBoundingClientRect();
      return [head.left, td.left];
    }""")
    assert abs(got[0] - got[1]) < 1, got


def test_a_red_negative_format_wins_over_the_font_colour(served):
    # Excel: a [Red] section in the number format overrides the cell's font colour.
    page, base, email = served
    _open_grid(page, base, email)
    colour = page.evaluate("""() => getComputedStyle([...document.querySelectorAll('.xl-grid td')].find((td) => td.textContent === '-5')).color""")
    assert colour == "rgb(192, 0, 0)"


def test_the_attachment_thumbnail_leaves_out_hidden_columns(served):
    page, base, email = served
    page.set_viewport_size({"width": 1440, "height": 900})
    page.goto(f"{base}/app/mail/{email.id}")
    page.wait_for_selector(".file-thumb-grid")
    text = page.text_content(".file-thumb-grid")
    assert "Cost" not in text and "900" not in text, text


# Preview: a big export --------------------------------------------------------------------------------
@pytest.mark.usefixtures("harbor_ocr")
def test_the_thumbnail_of_a_5_mb_ledger_export_doesnt_load_the_whole_workbook_twice(monkeypatch):
    """A 5 MB general-ledger export (200,000 rows) is well inside the preview's limits (25 MB, 150 MB unzipped).
    The email's thumbnail asks for 8 rows: the workbook is streamed (read-only), never loaded in full, which took
    about 25 seconds and 1 GB of memory twice over (values and formulas)."""
    import io

    import openpyxl
    from openpyxl import Workbook

    from controller_inbox import sheet_preview

    book = Workbook(write_only=True)
    ws = book.create_sheet("GL")
    ws.append(["Date", "Account", "Description", "Debit", "Credit"])
    for i in range(200_000):
        ws.append(["2026-09-30", 4000 + i % 50, f"Journal line {i}", i * 1.01, 0])
    out = io.BytesIO()
    book.save(out)
    data = out.getvalue()
    assert len(data) < sheet_preview.MAX_BYTES and sheet_preview._unzipped_size(data) < sheet_preview.MAX_UNZIPPED
    opened = []
    real = openpyxl.load_workbook

    def load_workbook(*args, **kwargs):
        opened.append(bool(kwargs.get("read_only")))
        return real(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", load_workbook)
    got = sheet_preview.preview(data, "GL detail.xlsx", rows=8, cols=12)
    assert len(got["sheets"][0]["rows"]) == 8
    assert opened and all(opened), "the whole workbook was loaded, not streamed"


# Preview: a merge whose top-left cell is hidden ---------------------------------------------------------------------
@pytest.mark.usefixtures("harbor_ocr")
def test_a_merge_whose_first_row_is_hidden_still_shows():
    import io

    from openpyxl import Workbook

    from controller_inbox import sheet_preview

    book = Workbook()
    ws = book.active
    ws["A1"] = "Q3 accruals"
    ws.merge_cells("A1:C2")
    ws.row_dimensions[1].hidden = True
    ws["A3"], ws["B3"], ws["C3"] = "Rent", 4200, 4200
    out = io.BytesIO()
    book.save(out)
    sheet = sheet_preview.preview(out.getvalue(), "accruals.xlsx")["sheets"][0]
    cells = {cell["v"]: (row["n"], cell) for row in sheet["rows"] for cell in row["cells"]}
    assert "Q3 accruals" in cells, sheet["rows"]
    n, cell = cells["Q3 accruals"]
    assert n == 2 and cell["c"] == 1 and cell.get("cs") == 3 and cell.get("rs", 1) == 1


# Spreadsheet previews keep merged cells: an .xls's, and the last row of an .xlsx merge at the bottom of a sheet.
def test_xlsx_preview_keeps_the_last_row_of_a_merge_at_the_bottom(tmp_path):
    """A merge A9:B10 with nothing below it keeps row 10 in the grid, so the merged cell spans two shown rows."""
    from controller_inbox import sheet_preview

    book = Workbook()
    sheet = book.active
    sheet["A1"] = "Item"
    sheet["A9"] = "Approved by Dana"
    sheet.merge_cells("A9:B10")
    out = io.BytesIO()
    book.save(out)
    grid = sheet_preview.preview(out.getvalue(), "book.xlsx")["sheets"][0]
    [merged] = [cell for row in grid["rows"] if row["n"] == 9 for cell in row["cells"]]
    assert merged["rs"] == 2 and merged["cs"] == 2
    assert [row["n"] for row in grid["rows"]][-1] == 10


def test_xls_preview_keeps_merged_cells(monkeypatch):
    """An .xls's merged title (A1:C1) spans three columns and the cells it covers aren't shown on their own."""
    import xlrd
    from xlrd.sheet import Cell

    from controller_inbox import sheet_preview

    values = {(0, 0): "Q3 Budget", (1, 0): "Line", (1, 1): "Q3", (1, 2): "Q4"}

    class Sheet:
        name, nrows, ncols, visibility = "Budget", 2, 3, 0
        merged_cells = [(0, 1, 0, 3)]

        def cell(self, r, c):
            value = values.get((r, c))
            return Cell(xlrd.XL_CELL_TEXT, value) if value else Cell(xlrd.XL_CELL_EMPTY, "")

    class Book:
        nsheets, datemode = 1, 0

        def sheet_by_index(self, i):
            return Sheet()

        def release_resources(self):
            pass

    monkeypatch.setattr(xlrd, "open_workbook", lambda **kw: Book())
    data = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 504
    rows = sheet_preview.preview(data, "book.xls")["sheets"][0]["rows"]
    [title] = rows[0]["cells"]
    assert title["v"] == "Q3 Budget" and title["cs"] == 3 and title["rs"] == 1
    assert [cell["v"] for cell in rows[1]["cells"]] == ["Line", "Q3", "Q4"]


# A web page attachment's tables are shown as tables (text only); a calendar file as its events.
def test_an_html_attachment_preview_keeps_its_tables_as_text_only():
    """statement.html previews as a heading and a table of text cells; its script, hidden text and markup are gone."""
    from controller_inbox import sheet_preview

    page = (
        b"<html><head><title>Statement</title><script>window.parent.__xss='html'</script></head><body><h1>Statement</h1>"
        b"<p>Balance <b>due</b><span style='display:none'>zz</span></p><table><tr><th>Item</th><th>Amount</th></tr>"
        b"<tr><td>Balance</td><td>1,234.00</td></tr><tr><td colspan=2><img src=x onerror=alert(5)>Thanks</td></tr></table></body></html>"
    )
    got = sheet_preview.preview(page, "statement.html")
    assert got["kind"] == "doc"
    assert got["blocks"] == [
        {"t": "h2", "text": "Statement"},
        {"t": "p", "text": "Balance due"},
        {"t": "table", "rows": [["Item", "Amount"], ["Balance", "1,234.00"], [{"v": "Thanks", "cs": 2}]]},
    ]
    assert "xss" not in json.dumps(got) and "onerror" not in json.dumps(got)


def test_an_ics_attachment_previews_as_its_event():
    """meeting.ics shows its summary, start, end, place and organizer instead of the raw iCalendar lines."""
    from controller_inbox import sheet_preview

    ics = (
        b"BEGIN:VCALENDAR\r\nVERSION:2.0\r\nBEGIN:VEVENT\r\nUID:1@x\r\nDTSTART:20261024T140000Z\r\nDTEND:20261024T150000Z\r\n"
        b"SUMMARY:Q4 planning\r\nLOCATION:Room 4\\, 2nd floor\r\nORGANIZER;CN=Dana Cole:mailto:dana@taz.com\r\nEND:VEVENT\r\nEND:VCALENDAR\r\n"
    )
    got = sheet_preview.preview(ics, "meeting.ics")
    assert got["kind"] == "doc"
    assert got["blocks"][:2] == [
        {"t": "h2", "text": "Q4 planning"},
        {"t": "table", "rows": [["Starts", "Sat, Oct 24, 2026, 2:00 PM UTC"], ["Ends", "Sat, Oct 24, 2026, 3:00 PM UTC"],
                                ["Where", "Room 4, 2nd floor"], ["Organizer", "Dana Cole dana@taz.com"]]},
    ]
    assert "BEGIN:VEVENT" not in json.dumps(got["blocks"]) and "SUMMARY:Q4 planning" in got["source"]
