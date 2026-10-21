"""The attachment preview's review fixes beyond test_preview_edges.py: the number formats around the reviewed ones,
.xls formats, a workbook too big to load, a web page saved as .xls, Word lists and tab stops, and the browser freeing
the pictures it draws."""

from __future__ import annotations

import io
from datetime import datetime

import pytest
from docx import Document
from openpyxl import Workbook

from controller_inbox import sheet_preview
from controller_inbox.sheet_preview import display


def _xlsx(book: Workbook) -> bytes:
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


@pytest.mark.parametrize(
    "value, fmt, excel",
    [
        (1500, '[>=1000]#,##0,"K";0', "2K"),
        (-1500, '[>=1000]#,##0,"K";0', "-1500"),
        (5551234567, "[<=9999999]###-####;(###) ###-####", "(555) 123-4567"),
        (1234567, "[<=9999999]###-####;(###) ###-####", "123-4567"),
        (2.25, "# ??/??", "2  1/4"),
        (0.3333, "?/?", "1/3"),
        (0.5, "# ?/8", "4/8"),
        (3, "# ?/?", "3"),
        (-1.25, "# ?/?", "-1 1/4"),
        (12345, "##0.0E+0", "12.3E+3"),
        (0.000123, "0.00E+00", "1.23E-04"),
        (5, "0,000", "0,005"),
        (0.5, "0.0\\%", "0.5%"),
        (123456789012.5, "General", "1.23457E+11"),
        (1 / 3, "General", "0.333333333"),
        (0.0001234, "General", "0.0001234"),
        (-5, "General;[Red]-General", "-5"),
        (5, '"Yes";"No"', "Yes"),
        (0.06277777777777778, "[mm]:ss", "90:24"),
        (0.0001, "mm:ss.0", "00:08.6"),
        (45000.99999999, "m/d/yyyy h:mm:ss", "3/16/2023 0:00:00"),
        ("abc", '0;0;0;"["@"]"', "[abc]"),
        ("abc", '"Ref: "@', "Ref: abc"),
    ],
)
def test_more_formats_match_excel(value, fmt, excel):
    assert display(value, fmt)[0] == excel


def test_a_date_sits_at_the_right_of_its_cell_like_a_number():
    book = Workbook()
    book.active["A1"] = datetime(2026, 9, 30)
    cell = sheet_preview.preview(_xlsx(book), "dates.xlsx")["sheets"][0]["rows"][0]["cells"][0]
    assert cell["num"] and "2026" in cell["v"]


def test_a_workbook_that_unzips_too_big_is_refused_with_a_sentence(monkeypatch):
    # A zip of a few MB can hold gigabytes of sheet XML; openpyxl would load all of it.
    monkeypatch.setattr(sheet_preview, "MAX_UNZIPPED", 1000)
    book = Workbook()
    book.active["A1"] = "x"
    got = sheet_preview.preview(_xlsx(book), "huge.xlsx")
    assert got["sheets"] == [] and "Download" in got["note"]


def test_a_web_page_saved_as_xls_says_so_instead_of_showing_markup():
    data = b"\xef\xbb\xbf  <html><body><table><tr><td>Amount</td></tr></table></body></html>"
    got = sheet_preview.preview(data, "report.xls")
    assert got["sheets"] == [] and "web page" in got["note"]
    # A real CSV named .csv is still read as one.
    assert sheet_preview.preview(b"a,b\n1,2\n", "export.csv")["sheets"][0]["rows"][1]["cells"][0]["v"] == "1"


def test_an_xls_cell_uses_its_own_number_format(monkeypatch):
    import xlrd
    from xlrd.sheet import Cell

    class Format:
        def __init__(self, text):
            self.format_str = text

    class XF:
        def __init__(self, key):
            self.format_key = key

    cells = [Cell(xlrd.XL_CELL_NUMBER, 1234.5, 0), Cell(xlrd.XL_CELL_NUMBER, -5.0, 1), Cell(xlrd.XL_CELL_DATE, 46295.0, 2), Cell(xlrd.XL_CELL_NUMBER, 7.0, 99)]

    class Sheet:
        name, nrows, ncols, visibility = "S", 1, len(cells), 0

        def cell(self, r, c):
            return cells[c]

    class Book:
        nsheets, datemode = 1, 0
        xf_list = [XF(10), XF(11), XF(12)]
        format_map = {10: Format("#,##0.00"), 11: Format("0;[Red]-0"), 12: Format("mm/dd/yyyy")}

        def sheet_by_index(self, i):
            return Sheet()

        def release_resources(self):
            pass

    asked = []
    monkeypatch.setattr(xlrd, "open_workbook", lambda **kw: asked.append(kw["formatting_info"]) or Book())
    row = sheet_preview.preview(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 504, "book.xls")["sheets"][0]["rows"][0]["cells"]
    assert asked == [True]
    assert [c["v"] for c in row] == ["1,234.50", "-5", "09/30/2026", "7"], "an unknown XF falls back to General"
    assert row[1]["red"] and all(c["num"] for c in row)


def _docx(document) -> bytes:
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def test_a_numbered_word_list_shows_its_numbers():
    document = Document()
    for text in ("Sign the lease", "Pay the deposit", "Collect the keys"):
        document.add_paragraph(text, style="List Number")
    document.add_paragraph("A bullet", style="List Bullet")
    blocks = sheet_preview.docx_blocks(_docx(document))
    assert [b.get("num") for b in blocks] == ["1.", "2.", "3.", None]
    assert all(b["t"] == "li" for b in blocks)


def test_tab_stops_are_not_read_as_tabs():
    document = Document()
    paragraph = document.add_paragraph("Total")
    paragraph.paragraph_format.tab_stops.add_tab_stop(914400)
    assert sheet_preview.docx_blocks(_docx(document))[0]["text"] == "Total"


# Front end (Playwright) ----------------------------------------------------------------------------------

@pytest.fixture(scope="module")
def browser():
    pytest.importorskip("playwright")
    from liveserver import chromium_path
    from playwright.sync_api import sync_playwright

    path = chromium_path()
    if path is None:
        pytest.skip("no Chromium for Playwright on this machine")
    with sync_playwright() as p:
        b = p.chromium.launch(executable_path=path)
        yield b
        b.close()


@pytest.fixture
def page(browser):
    context = browser.new_context()
    context.set_default_timeout(15000)
    errors = []
    context.on("weberror", lambda error: errors.append(str(error.error)))
    yield context.new_page()
    context.close()
    assert errors == []


COUNT_URLS = """
  window.__urls = { made: 0, freed: 0 };
  const make = URL.createObjectURL.bind(URL), free = URL.revokeObjectURL.bind(URL);
  URL.createObjectURL = (b) => { window.__urls.made += 1; return make(b); };
  URL.revokeObjectURL = (u) => { window.__urls.freed += 1; return free(u); };
"""


@pytest.fixture
def pdf_email(settings, store):
    from msgfactory import DOCX, PDF, XLSX, write_msg
    from test_page_view import _invoice

    from controller_inbox.folder_mail import ingest_folder

    document = Document()
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text, table.cell(0, 2).text = "Q3", "Total"
    for i, v in enumerate(["Rent", "Power", "1,200"]):
        table.cell(1, i).text = v
    book = Workbook()
    ws = book.active
    ws["A5"] = "Region"
    ws.merge_cells("A5:A7")
    ws.row_dimensions[6].hidden = True
    ws["B5"], ws["B7"], ws["B8"] = "x", "y", "z"
    settings.trusted_domains = "taz.com"
    settings.ensure_data_dir()
    write_msg(settings.inbox_incoming / "invoice.msg", "Invoice INV-1042", "Attached.", sender_name="Harbor Steel", sender_email="ar@taz.com",
              attachments=[("invoice.pdf", _invoice(pages=1), PDF), ("terms.docx", _docx(document), DOCX), ("regions.xlsx", _xlsx(book), XLSX)])
    [email] = ingest_folder(store, settings)
    return email


def _n(email, name: str) -> int:
    return 1 + [a.filename for a in email.attachments].index(name)


def test_a_thumbnail_frees_its_picture_once_drawn(page, settings, store, pdf_email):
    from liveserver import serving

    from controller_inbox import web

    page.add_init_script(COUNT_URLS)
    with serving(web.create_app(settings, store)) as base:
        page.set_viewport_size({"width": 1440, "height": 900})
        page.goto(f"{base}/app/mail/{pdf_email.id}")
        page.wait_for_function("() => { const img = document.querySelector('.file-thumb img'); return img && img.naturalWidth > 0; }")
        page.wait_for_function("() => window.__urls.made > 0 && window.__urls.freed === window.__urls.made")


def test_a_thumbnail_taken_off_the_page_before_its_picture_arrives_makes_none(page, settings, store, pdf_email):
    from liveserver import serving

    from controller_inbox import web

    page.add_init_script(COUNT_URLS)
    held = []
    with serving(web.create_app(settings, store)) as base:
        page.route("**/pages/1.png", lambda route: held.append(route))
        page.set_viewport_size({"width": 1440, "height": 900})
        page.goto(f"{base}/app/mail/{pdf_email.id}")
        page.wait_for_function("() => document.querySelector('.file-thumb.pdf')")
        page.wait_for_timeout(300)
        assert held, "the thumbnail asked for its picture"
        page.evaluate("() => document.querySelector('.file-thumb.pdf').remove()")
        for route in held:
            route.continue_()
        page.wait_for_timeout(500)
        assert page.evaluate("() => window.__urls.made") == 0


def test_a_merged_word_cell_spans_its_columns_in_the_preview(page, settings, store, pdf_email):
    from liveserver import serving

    from controller_inbox import web

    with serving(web.create_app(settings, store)) as base:
        page.set_viewport_size({"width": 1440, "height": 900})
        page.goto(f"{base}/app/mail/{pdf_email.id}/file/{_n(pdf_email, 'terms.docx')}?tab=preview")
        page.wait_for_selector(".doc-table")
        spans = page.evaluate("() => [...document.querySelectorAll('.doc-table tr')].map((tr) => [...tr.cells].map((td) => td.colSpan))")
        assert spans == [[2, 1], [1, 1, 1]]


def test_the_row_under_a_merge_over_a_hidden_row_keeps_its_columns(page, settings, store, pdf_email):
    # A5:A7 merged with row 6 hidden: row 7 is the merge's second shown row, so its "y" stays in column B.
    from liveserver import serving

    from controller_inbox import web

    with serving(web.create_app(settings, store)) as base:
        page.set_viewport_size({"width": 1440, "height": 900})
        page.goto(f"{base}/app/mail/{pdf_email.id}/file/{_n(pdf_email, 'regions.xlsx')}?tab=preview")
        page.wait_for_selector(".xl-grid")
        got = page.evaluate("""() => {
          const head = [...document.querySelectorAll('.xl-grid thead th')].find((th) => th.textContent === 'B').getBoundingClientRect();
          const td = [...document.querySelectorAll('.xl-grid td')].find((td) => td.textContent === 'y').getBoundingClientRect();
          return [head.left, td.left];
        }""")
        assert abs(got[0] - got[1]) < 1, got

