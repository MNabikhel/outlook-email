"""The Preview tab for the rest of the files CloseDesk reads: PowerPoint slide by slide, text, RTF, web pages and
calendar files as their text (markup never rendered), and zips as the list of what's in them."""

from __future__ import annotations

import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

from controller_inbox import sheet_preview, web
from controller_inbox.folder_mail import ingest_folder
from msgfactory import write_msg

PAGE = {"X-CloseDesk": "1"}
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
PNG_1PX = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d49444154789c6360000002000154a24f5d0000000049454e44ae426082"
)


def _deck() -> bytes:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])  # title and content
    slide.shapes.title.text = "Q3 close status"
    body = slide.placeholders[1].text_frame
    body.text = "AP is closed"
    point = body.add_paragraph()
    point.text = "Two accruals still open"
    point.level = 1
    slide.notes_slide.notes_text_frame.text = "Mention the rent accrual."
    second = deck.slides.add_slide(deck.slide_layouts[5])  # title only
    second.shapes.title.text = "Variance"
    table = second.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Account", "Variance"
    table.cell(1, 0).text, table.cell(1, 1).text = "Rent", "$1,200"
    second.shapes.add_picture(io.BytesIO(PNG_1PX), Inches(5), Inches(2))
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def test_a_deck_is_shown_slide_by_slide():
    slides = sheet_preview.preview(_deck(), "status.pptx")["slides"]
    assert [s["title"] for s in slides] == ["Q3 close status", "Variance"]
    assert slides[0]["blocks"] == [{"t": "li", "text": "AP is closed", "level": 0}, {"t": "li", "text": "Two accruals still open", "level": 1}]
    assert slides[0]["notes"] == "Mention the rent accrual."
    assert slides[1]["blocks"] == [{"t": "table", "rows": [["Account", "Variance"], ["Rent", "$1,200"]]}]
    assert slides[1]["pictures"] == 1


@pytest.mark.parametrize(
    "name, data, shown, hidden",
    [
        ("notes.txt", b"Remit to 123 Main St.\r\nThanks", "Remit to 123 Main St.\nThanks", ""),
        ("memo.rtf", rb"{\rtf1\ansi Payment {\b due} Friday.}", "Payment due Friday.", "rtf1"),
        ("invite.ics", b"BEGIN:VCALENDAR\nSUMMARY:Close review\nEND:VCALENDAR", "SUMMARY:Close review", ""),
    ],
)
def test_text_like_files_are_shown_as_their_text(name, data, shown, hidden):
    got = sheet_preview.preview(data, name)
    assert got["kind"] == "text" and shown in got["text"]
    if hidden:
        assert hidden not in got["text"], "markup is never shown or run"


def test_a_zip_is_listed():
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("invoices/INV-1.pdf", b"%PDF-1.4 x" * 50)
        archive.writestr("invoices/", b"")
        archive.writestr("readme.txt", b"hi")
    got = sheet_preview.preview(out.getvalue(), "pack.zip")
    assert got["kind"] == "zip" and [f["name"] for f in got["files"]] == ["invoices/INV-1.pdf", "readme.txt"]
    assert got["files"][0]["size"] == 500
    assert sheet_preview.preview(b"not a zip", "bad.zip")["note"]


def test_the_preview_api_serves_the_new_kinds(store, settings):
    settings.ensure_data_dir()
    write_msg(
        settings.inbox_incoming / "pack.msg", "Close pack", "See attached.", sender_email="maya@taz.com",
        attachments=[("status.pptx", _deck(), PPTX), ("notes.txt", b"Remit to 123 Main St.", "text/plain")],
    )
    [email] = ingest_folder(store, settings)
    client = TestClient(web.create_app(settings, store))
    files = {f["name"]: f for f in client.get(f"/api/mail/{email.id}", headers=PAGE).json()["attachments"]}
    assert files["status.pptx"]["show"] == "slides" and files["notes.txt"]["show"] == "text"
    deck = client.get(f"/api/mail/{email.id}/files/{files['status.pptx']['n']}/sheet", headers=PAGE).json()
    assert deck["slides"][0]["title"] == "Q3 close status"
    text = client.get(f"/api/mail/{email.id}/files/{files['notes.txt']['n']}/sheet", headers=PAGE).json()
    assert text["text"] == "Remit to 123 Main St."
