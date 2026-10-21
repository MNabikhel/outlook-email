"""The caches that keep a chat question, a page view and Setup quick never serve an old result: each is kept by
what it was worked out from (the text, the file's SHA-256), so a changed file or text is worked out again."""

from __future__ import annotations

import copy
import dataclasses
import io
import os

import httpx
from openpyxl import Workbook, load_workbook

from controller_inbox import agent, answer_check, assistant, documents, local_llm, memo, page_view, semantic, table_lookup, vision
from controller_inbox.demo import make_pdf

TABLE = "\n".join([
    "[page 1]", "[table]",
    "Item: Ads | Amount: 1,000.00",
    "Item: Travel | Amount: 250.00",
    "Item: Total | Amount: 1,250.00",
])


def test_lru_drops_the_least_recently_used():
    cache = memo.LRU(2)
    cache.put("a", 1)
    cache.put("b", 2)
    assert cache.get("a") == 1  # "b" is now the oldest
    cache.put("c", 3)
    assert cache.get("b") is None and cache.get("a") == 1 and cache.get("c") == 3
    assert cache.get_or("d", lambda: 4) == 4 and len(cache) == 2


def test_tables_of_a_changed_text_are_read_again():
    first = table_lookup.tables_in(TABLE)
    changed = TABLE.replace("250.00", "300.00").replace("1,250.00", "1,300.00")
    again = table_lookup.tables_in(changed)
    assert [row.value("Amount") for row in again[0].rows] == ["1,000.00", "300.00", "1,300.00"]
    assert [row.value("Amount") for row in first[0].rows] == ["1,000.00", "250.00", "1,250.00"]
    # The list handed out is the caller's own: changing it doesn't change what the next caller gets.
    first.clear()
    assert len(table_lookup.tables_in(TABLE)) == 1


def test_prompt_size_counts_ascii_and_other_text_alike():
    def by_character(text: str) -> int:
        words = agent._LETTERS.findall(text)
        digits = sum(ch.isdigit() for ch in text)
        marks = sum(1 for ch in text if not ch.isalnum() and not ch.isspace())
        long = sum(max(0, len(word) - 15) for word in words) / 3
        tokens = 0.8 * len(words) + 0.07 * sum(map(len, words)) + long + 1.3 * digits + 1.2 * marks
        return int(tokens * agent.CHARS_PER_TOKEN) + 1

    samples = ["", "Item: Ads | Amount: $1,000.00\n\tx\x1c\x1f_", "".join(map(chr, range(128))), "Café 1² ½ — résumé", TABLE * 300]
    for text in samples:
        assert agent.prompt_size(text) == by_character(text), text[:40]
    # Kept by the text: a changed long text is counted again.
    assert agent.prompt_size(TABLE * 300 + "1") == by_character(TABLE * 300 + "1")


def test_figures_of_changed_material_are_found_again():
    long = "Notes. " * 200 + "Total 1,250.00"
    assert answer_check.Grounding([long]).has(answer_check.numbers_in("1,250.00")[0])
    changed = long.replace("1,250.00", "1,300.00")
    grounding = answer_check.Grounding([changed])
    assert grounding.has(answer_check.numbers_in("1,300.00")[0]) and not grounding.has(answer_check.numbers_in("1,250.00")[0])


def test_the_files_are_worked_out_once_a_question(store, settings, mail, monkeypatch):
    budget = mail["Q4 budget draft"]
    budget.attachments[0].extracted_text = TABLE
    calls = []
    real = table_lookup.answer
    monkeypatch.setattr(table_lookup, "answer", lambda text, question: (text == TABLE and calls.append(question)) or real(text, question))
    ws = agent.Workspace(store, settings, [budget], question="q", current_id=budget.id)
    question = "What is the amount for Travel?"
    agent.question_points(ws, question)
    first = agent.file_context(ws, question, 6000)
    second = agent.file_context(ws, question, 6000)
    assert first == second and "Travel" in first[budget.id]
    assert len(calls) == 1, calls
    # Another question, or the next question on its own Workspace, works it out again.
    agent.file_context(ws, "What is the amount for Ads?", 6000)
    assert len(calls) == 2


def test_found_in_follows_the_email_as_it_is_now(mail):
    email = copy.deepcopy(mail["Q4 budget draft"])
    email.body_text = "Northwind"
    assert assistant._found_in(email, ["northwind", "harbor"]) == ["northwind"]
    email.body_text = "Harbor Steel"
    assert assistant._found_in(email, ["northwind", "harbor"]) == ["harbor"]


def test_scanned_pages_are_kept_by_the_file_hash_only_when_given():
    calls = []
    real = vision._scanned_pages
    vision._scanned_pages = lambda data: calls.append(len(data)) or real(data)
    try:
        pdf = make_pdf([["page one has enough words to count as text"], [" "]])
        sha = page_view.file_hash(pdf)
        assert vision.scanned_pages(pdf, sha256=sha) == vision.scanned_pages(pdf, sha256=sha) == [2]
        assert len(calls) == 1
        other = make_pdf([[" "], ["page two has enough words to count as text"]])
        assert vision.scanned_pages(other, sha256=page_view.file_hash(other)) == [1]
        assert vision.scanned_pages(other) == [1] and len(calls) == 3  # without a hash, read each time
    finally:
        vision._scanned_pages = real


def test_doubtful_pages_follow_the_text():
    good = TABLE
    bad = TABLE.replace("Total | Amount: 1,250.00", "Total | Amount: 9,999.00")
    assert vision.doubtful_pages(good) == []
    assert vision.doubtful_pages(bad) == [1]
    assert vision.doubtful_pages(good) == []


def test_an_original_changed_on_disk_is_not_taken_for_the_attachment(store, settings, mail):
    quote = mail["Updated remittance details"]
    att = quote.attachments[0]
    data = vision.original_bytes(settings, quote, att)
    assert data is not None and page_view.file_hash(data) == att.sha256
    path = agent.original_file(settings, quote, att)
    stat = path.stat()
    path.write_bytes(data[:-1] + b"X")
    os.utime(path, ns=(stat.st_atime_ns, stat.st_mtime_ns + 1_000_000))
    assert vision.original_bytes(settings, quote, att) is None
    path.write_bytes(data)
    assert vision.original_bytes(settings, quote, att) == data


def test_a_page_drawn_by_its_known_hash_is_the_one_drawn_from_its_bytes(settings):
    pdf = make_pdf([["Invoice 1001 total 1,250.00 due now"], ["Second page of the invoice here"]])
    sha = page_view.file_hash(pdf)
    by_hash = page_view.page_png(settings, pdf, "invoice.pdf", 2, sha256=sha)
    assert (page_view.folder(settings) / f"{sha}-2.png").is_file()
    assert page_view.page_png(settings, pdf, "invoice.pdf", 2) == by_hash
    assert page_view.regions(settings, pdf, "invoice.pdf", 1, sha256=sha) == page_view.regions(settings, pdf, "invoice.pdf", 1)


def _book(q4: int) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Budget"
    sheet.append(["Line", "Q3", "Q4", "Change"])
    sheet.append(["Ads", 1000, q4, "=C2-B2"])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_workbooks_loaded_for_a_question_read_as_fresh_ones(store, settings, mail):
    budget = mail["Q4 budget draft"]
    ws = agent.Workspace(store, settings, [budget], question="q", current_id=budget.id)
    args = {"email": "1", "file": "Q4 budget.xlsx", "sheet": "Budget"}
    first = agent.run_tool(ws, "read_cells", {**args, "cells": "A1:D4"}, limit=4000)
    agent.run_tool(ws, "trace_cell", {**args, "cell": "D4"}, limit=4000)
    agent.run_tool(ws, "compare_columns", {**args, "from": "Q3", "to": "Q4"}, limit=4000)
    # Tracing and comparing make cells in the workbook as they look; what is read next is as from the file.
    assert agent.run_tool(ws, "read_cells", {**args, "cells": "A1:D4"}, limit=4000) == first
    book = ws.file(budget, "Q4 budget.xlsx")
    assert first.startswith('Sheet "Budget"') and first == documents.read_cells(ws.original(budget, book), "Budget", "A1:D4")


def test_books_are_kept_by_the_workbook_hash(mail):
    one, two = _book(1500), _book(2000)
    att = next(att for att in mail["Q4 budget draft"].attachments if att.filename == "Q4 budget.xlsx")
    holder = agent.Workspace(None, None, [])
    first = agent._books(holder, dataclasses.replace(att, sha256=page_view.file_hash(one)), one)
    second = agent._books(holder, dataclasses.replace(att, sha256=page_view.file_hash(two)), two)
    assert agent._books(holder, dataclasses.replace(att, sha256=page_view.file_hash(one)), one) is first
    assert first is not second
    assert "C2: 1,500" in documents.read_cells(one, "Budget", "A1:D2", books=first)
    assert "C2: 2,000" in documents.read_cells(two, "Budget", "A1:D2", books=second)
    assert load_workbook(io.BytesIO(two))["Budget"]["C2"].value == 2000


def test_setup_count_follows_a_file_read_again(store, settings, mail):
    settings.embedding_model = "fake-embed"
    settings.llm = True
    budget = mail["Q4 budget draft"]
    rows = [(key, email.id, semantic._key(text), b"\0" * 16) for email in mail.values() for key, text in semantic._items(email)]
    store.save_embeddings("fake-embed", rows)
    done = semantic.coverage(store, settings)
    assert done["sections_done"] == done["sections"] > 0
    att = budget.attachments[0]
    store.set_attachment_text(att.id, (att.extracted_text or "") + "\nA9: read again")
    after = semantic.coverage(store, settings)
    assert after["sections_done"] < after["sections"] and after["waiting"] == 1
    assert semantic.file_states(store, settings, store.get_email(budget.id))[att.id] == "changed"


def test_model_calls_share_one_client_unless_httpx_is_replaced(monkeypatch):
    assert local_llm._http("post") == local_llm.shared_client().post
    assert local_llm.shared_client() is local_llm.shared_client()

    def fake(*_a, **_k):
        return httpx.Response(200, json={})

    monkeypatch.setattr(local_llm.httpx, "post", fake)
    assert local_llm._http("post") is fake


def test_the_shared_client_keeps_no_cookies():
    seen = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append(request.headers.get("cookie"))
        return httpx.Response(200, headers={"set-cookie": "session=abc; Path=/"}, json={})

    client = httpx.Client(transport=httpx.MockTransport(handler), cookies=local_llm._no_cookies())
    client.get("http://127.0.0.1:1234/v1/models")
    client.get("http://127.0.0.1:1234/v1/models")
    assert seen == [None, None]
