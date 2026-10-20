"""Seeing a file as it came (the Preview tab, and the small preview on each attachment), and teaching CloseDesk a field
by showing it on the page: the email uses it at once, and the sender's next invoice is read the same way."""

from __future__ import annotations

import io
from datetime import date, datetime

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from controller_inbox import sheet_preview, teach, web
from controller_inbox.folder_mail import ingest_folder
from controller_inbox.models import ExtractedFields
from controller_inbox.sheet_preview import display
from msgfactory import PDF, XLSX, DOCX, write_msg
from pdffactory import Text, build_pdf

PAGE = {"X-CloseDesk": "1"}
AS_OF = date(2026, 10, 20)


# How Excel shows a value -------------------------------------------------------------------------------


@pytest.mark.parametrize(
    "value, fmt, shown",
    [
        (1234.5, "General", "1234.5"),
        (1234, "General", "1234"),
        (0.1 + 0.2, "General", "0.3"),
        (1234.5, "#,##0.00", "1,234.50"),
        (1234567.8, "#,##0", "1,234,568"),
        (-1234.5, "$#,##0.00_);($#,##0.00)", "($1,234.50)"),
        (1234.5, "$#,##0.00_);($#,##0.00)", "$1,234.50"),
        (0.125, "0%", "13%"),
        (0.125, "0.0%", "12.5%"),
        (1234, "000000", "001234"),
        (1234567890, "#,##0,,", "1,235"),
        (1234.5, "0.00E+00", "1.23E+03"),
        (1234.5, '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)', "$1,234.50"),
        (0, '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)', "$-"),
        (-5, "#,##0.00", "-5.00"),
        (2.5, '#,##0.00 "hrs"', "2.50 hrs"),
        (datetime(2026, 9, 30), "mm/dd/yyyy", "09/30/2026"),
        (datetime(2026, 3, 1), "mmm-yy", "Mar-26"),
        (datetime(2026, 9, 30), "d-mmm-yyyy", "30-Sep-2026"),
        (datetime(2026, 9, 30), "dddd, mmmm d, yyyy", "Wednesday, September 30, 2026"),
        (datetime(2026, 9, 30, 14, 5), "m/d/yyyy h:mm AM/PM", "9/30/2026 2:05 PM"),
        (datetime(2026, 9, 30), "General", "9/30/2026"),
        (46295, "m/d/yyyy", "9/30/2026"),
        (True, "General", "TRUE"),
        ("Accrued rent", "General", "Accrued rent"),
        (None, "#,##0.00", ""),
    ],
)
def test_values_show_as_excel_shows_them(value, fmt, shown):
    assert display(value, fmt)[0] == shown


def test_a_red_negative_section_says_so():
    assert display(-50, "#,##0.00;[Red](#,##0.00)") == ("(50.00)", True)
    assert display(50, "#,##0.00;[Red](#,##0.00)") == ("50.00", False)


# Previews -------------------------------------------------------------------------------------------


def _workbook() -> bytes:
    book = Workbook()
    ws = book.active
    ws.title = "Summary"
    ws["A1"] = "Q3 accruals"
    ws["A1"].font = Font(bold=True)
    ws.merge_cells("A1:C1")
    ws.append(["Account", "Amount", "Share"])
    ws.append(["Rent", 12500, 0.625])
    ws.append(["Utilities", 7500, 0.375])
    ws["A5"] = "Total"
    ws["B5"] = "=SUM(B3:B4)"
    for row in (3, 4, 5):
        ws[f"B{row}"].number_format = "#,##0.00"
        ws[f"C{row}"].number_format = "0.0%"
    ws["A2"].fill = PatternFill("solid", fgColor="FFDDEBF7")
    ws.column_dimensions["A"].width = 24
    hidden = book.create_sheet("Workings")
    hidden["A1"] = "scratch"
    hidden.sheet_state = "hidden"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_a_workbook_is_shown_sheet_by_sheet_as_excel_shows_it():
    got = sheet_preview.preview(_workbook(), "accruals.xlsx")
    assert got["kind"] == "sheet" and [s["name"] for s in got["sheets"]] == ["Summary", "Workings"]
    summary, workings = got["sheets"]
    assert workings["hidden"] and not summary["hidden"]
    cells = {(row["n"], cell["c"]): cell for row in summary["rows"] for cell in row["cells"]}
    assert cells[(1, 1)]["v"] == "Q3 accruals" and cells[(1, 1)]["b"] and cells[(1, 1)]["cs"] == 3
    assert (1, 2) not in cells, "covered by the merge"
    assert cells[(3, 2)] == {"c": 2, "v": "12,500.00", "num": True}
    assert cells[(3, 3)]["v"] == "62.5%"
    assert cells[(5, 2)]["f"] == "=SUM(B3:B4)"
    assert cells[(5, 2)]["v"] == "=SUM(B3:B4)" and cells[(5, 2)]["nocalc"], "saved without its result: the formula shows"
    assert cells[(2, 1)]["bg"] == "#ddebf7"
    assert summary["cols"][0]["w"] > summary["cols"][1]["w"]
    small = sheet_preview.preview(_workbook(), "accruals.xlsx", rows=2, cols=2)["sheets"][0]
    assert len(small["rows"]) == 2 and small["more_rows"] == 3 and small["more_cols"] == 1


def test_a_csv_is_shown_with_its_own_delimiter():
    got = sheet_preview.preview("Vendor;Amount\nNorthwind;1.250,00\n".encode(), "export.csv")
    rows = got["sheets"][0]["rows"]
    assert [c["v"] for c in rows[0]["cells"]] == ["Vendor", "Amount"] and rows[1]["cells"][1]["v"] == "1.250,00"


def _docx() -> bytes:
    document = Document()
    document.add_heading("Services agreement", level=1)
    document.add_paragraph("The term is twelve months.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Fee", "$4,000"
    table.cell(1, 0).text, table.cell(1, 1).text = "Term", "12 months"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def test_a_word_document_is_shown_as_headings_paragraphs_and_tables():
    blocks = sheet_preview.preview(_docx(), "agreement.docx")["blocks"]
    assert blocks[0] == {"t": "h2", "text": "Services agreement"}
    assert blocks[1] == {"t": "p", "text": "The term is twelve months."}
    assert blocks[2] == {"t": "table", "rows": [["Fee", "$4,000"], ["Term", "12 months"]]}


# The vendor's invoices ------------------------------------------------------------------------------

SENDER = "billing@keystonepaper.com"


def _invoice(ref: str, balance: str, pay_by: str, credit: str = "$10,000.00") -> bytes:
    """Keystone Paper's invoice: its own labels ("Ref #", "Balance payable", "Pay by"), which the scripts don't know,
    and a credit limit larger than the total."""
    items = [Text(54, 740, "Keystone Paper Co", size=17, bold=True), Text(54, 722, "12 Mill Road, Albany NY", size=9.5)]
    for i, (label, value) in enumerate([("Ref #", ref), ("Pay by", pay_by), ("Credit limit", credit)]):
        items += [Text(380, 700 - 16 * i, label, size=10, bold=True), Text(558, 700 - 16 * i, value, size=10, right=True)]
    items += [Text(54, 600, "Copy paper, 40 cases", size=10), Text(558, 600, balance, size=10, right=True)]
    items += [Text(380, 560, "Balance payable", size=10, bold=True), Text(558, 560, balance, size=10, bold=True, right=True)]
    return build_pdf([items])


def _drop(settings, name: str, subject: str, pdf: bytes, filename: str, **extra) -> None:
    write_msg(
        settings.inbox_incoming / f"{name}.msg", subject, "Please see the attached.", sender_name="Keystone Paper",
        sender_email=SENDER, attachments=[(filename, pdf, PDF), *extra.get("more", [])], message_id=f"<{name}@keystonepaper.com>",
    )


@pytest.fixture
def mail(store, settings):
    settings.trusted_domains = "keystonepaper.com"
    settings.ensure_data_dir()
    _drop(settings, "a", "Your Keystone invoice", _invoice("KP-55120", "$2,118.40", "11/15/2026"), "KP-55120.pdf",
          more=[("accruals.xlsx", _workbook(), XLSX), ("agreement.docx", _docx(), DOCX)])
    _drop(settings, "c", "Another Keystone invoice", _invoice("KP-55190", "$640.00", "11/20/2026"), "KP-55190.pdf")
    ingest_folder(store, settings)
    return {e.subject: e for e in store.list_emails(limit=50)}


@pytest.fixture
def client(store, settings, mail):
    return TestClient(web.create_app(settings, store))


def _region(view: dict, text: str) -> dict:
    """The box with this text lowest on the page: the balance, not the line item that shows the same figure."""
    return max((r for r in view["regions"] if r["text"] == text), key=lambda r: r["y"])


def _box(region: dict, pad: float = 0.004) -> dict:
    return {"x": region["x"] - pad, "y": region["y"] - pad, "w": region["w"] + 2 * pad, "h": region["h"] + 2 * pad}


def _teach(client, email, field: str, text: str, status: int = 200) -> dict:
    view = client.get(f"/api/mail/{email.id}/files/1/pages/1/regions", headers=PAGE).json()
    reply = client.post(f"/api/mail/{email.id}/files/1/pages/1/fixes", headers=PAGE, json={"kind": "field", "field": field, "box": _box(_region(view, text))})
    assert reply.status_code == status, reply.text
    return reply.json()


def test_preview_hints_and_endpoints(client, mail):
    email = mail["Your Keystone invoice"]
    detail = client.get(f"/api/mail/{email.id}", headers=PAGE).json()
    assert [f["show"] for f in detail["attachments"]] == ["pdf", "sheet", "doc"]
    assert client.get(f"/api/mail/{email.id}/files/1", headers=PAGE).json()["file"]["show"] == "pdf"
    sheet = client.get(f"/api/mail/{email.id}/files/2/sheet", headers=PAGE)
    assert sheet.status_code == 200 and sheet.json()["sheets"][0]["name"] == "Summary"
    assert client.get(f"/api/mail/{email.id}/files/3/sheet", headers=PAGE).json()["blocks"][0]["text"] == "Services agreement"
    assert client.get(f"/api/mail/{email.id}/files/1/sheet", headers=PAGE).status_code == 404, "a PDF opens in its viewer"
    assert client.get(f"/api/mail/{email.id}/files/2/sheet").status_code == 403, "only CloseDesk's own page"
    view = client.get(f"/inbox/{email.id}/files/1/view")
    assert view.status_code == 200 and view.headers["content-type"] == "application/pdf"
    assert view.headers["content-security-policy"] == "frame-ancestors 'self'" and view.headers["x-frame-options"] == "SAMEORIGIN"


def test_a_locked_files_preview_isnt_shown(client, store, mail, monkeypatch):
    from controller_inbox import fraud

    email = mail["Your Keystone invoice"]
    monkeypatch.setattr(fraud, "attachments_locked", lambda _email: True)
    assert client.get(f"/api/mail/{email.id}/files/2/sheet", headers=PAGE).status_code == 403


# Teaching ---------------------------------------------------------------------------------------------


def test_teaching_fields_fixes_this_email_and_reads_the_vendors_next_invoice(client, store, settings, mail):
    email = mail["Your Keystone invoice"]
    assert store.get_email(email.id).extracted.primary_amount == 10000.0, "the scripts take the largest amount"

    reply = _teach(client, email, "invoice_number", "KP-55120")
    assert reply["message"].startswith("Invoice number set to KP-55120. Learned for Keystone Paper: their next files are read after “Ref #”.")
    assert "1 other email" in reply["message"], "the vendor's other stored invoice is read again"
    _teach(client, email, "total", "$2,118.40")
    _teach(client, email, "due_date", "11/15/2026")

    stored = store.get_email(email.id)
    assert stored.extracted.primary_invoice == "KP-55120"
    assert stored.extracted.primary_amount == 2118.40
    assert stored.extracted.primary_due == "2026-11-15"
    assert stored.extracted.taught["total"] == {"value": "2118.4", "by": "you", "label": "Balance payable"}
    # The email page marks them.
    fields = client.get(f"/api/mail/{email.id}", headers=PAGE).json()["fields"]
    assert fields["taught"]["total"]["shown"] == "$2,118.40" and fields["taught"]["total"]["by"] == "you"
    # The page draws them where they were shown.
    view = client.get(f"/api/mail/{email.id}/files/1/pages/1/regions", headers=PAGE).json()
    assert {t["field"]: t["value"] for t in view["taught"]} == {"invoice_number": "KP-55120", "total": "$2,118.40", "due_date": "2026-11-15"}

    # The other invoice already here was read again with what was taught.
    other = store.get_email(mail["Another Keystone invoice"].id)
    assert other.extracted.primary_invoice == "KP-55190" and other.extracted.primary_amount == 640.0
    assert other.extracted.taught["total"]["by"] == "learnt"

    # A new invoice from them is read the same way as it arrives.
    _drop(settings, "b", "Keystone invoice KP-55377", _invoice("KP-55377", "$980.10", "12/01/2026", credit="$25,000.00"), "KP-55377.pdf")
    ingest_folder(store, settings)
    new = next(e for e in store.list_emails(limit=50) if e.subject == "Keystone invoice KP-55377")
    assert new.extracted.primary_invoice == "KP-55377"
    assert new.extracted.primary_amount == 980.10
    assert new.extracted.primary_due == "2026-12-01"
    assert {key: entry["by"] for key, entry in new.extracted.taught.items()} == {"invoice_number": "learnt", "total": "learnt", "due_date": "learnt"}


def test_taking_a_taught_field_back(client, store, mail):
    email = mail["Your Keystone invoice"]
    reply = _teach(client, email, "total", "$2,118.40")
    assert store.get_email(email.id).extracted.primary_amount == 2118.40
    undone = client.post(f"/api/mail/{email.id}/files/1/pages/1/fixes/{reply['id']}/undo", headers=PAGE)
    assert undone.status_code == 200 and "isn't learnt" in undone.json()["message"]
    stored = store.get_email(email.id)
    assert stored.extracted.primary_amount == 10000.0 and "total" not in stored.extracted.taught
    assert teach.templates(store, SENDER) == []


def test_words_that_dont_read_as_the_field_are_refused(client, mail):
    email = mail["Your Keystone invoice"]
    reply = _teach(client, email, "due_date", "Keystone Paper Co", status=400)
    assert "don't read as a due date" in reply["detail"]
    reply = client.post(f"/api/mail/{email.id}/files/1/pages/1/fixes", headers=PAGE, json={"kind": "field", "field": "colour", "box": {"x": 0, "y": 0, "w": 1, "h": 1}})
    assert reply.status_code == 400


def test_value_of_each_field():
    assert teach.value_of("total", "Balance payable $2,118.40", as_of=AS_OF) == "2118.4"
    assert teach.value_of("total", "(1,250.00)", as_of=AS_OF) == "1250.0"
    assert teach.value_of("invoice_number", "Ref # KP-55120", as_of=AS_OF) == "KP-55120"
    assert teach.value_of("invoice_number", "Invoice date 10/01/2026", as_of=AS_OF) is None
    assert teach.value_of("due_date", "Pay by Nov 15, 2026", as_of=AS_OF) == "2026-11-15"
    assert teach.value_of("invoice_date", "2026-09-30", as_of=AS_OF) == "2026-09-30"
    assert teach.value_of("vendor", "Keystone Paper Co", as_of=AS_OF) == "Keystone Paper Co"
    assert teach.value_of("due_date", "Keystone", as_of=AS_OF) is None


def test_a_label_above_the_value_is_learnt_and_read_from_the_line_below():
    found = [
        {"x": 0.60, "y": 0.10, "w": 0.15, "h": 0.015, "text": "Amount Due"},
        {"x": 0.60, "y": 0.13, "w": 0.12, "h": 0.015, "text": "$1,980.00"},
        {"x": 0.10, "y": 0.13, "w": 0.20, "h": 0.015, "text": "Thank you"},
    ]
    fix = teach.field_fix(found, {"x": 0.595, "y": 0.125, "w": 0.13, "h": 0.025}, "total", as_of=AS_OF)
    assert fix["anchor"] == {"label": "Amount Due", "position": "below", "value": "1980.0"}
    text = "Keystone Paper\n             Invoice           Amount Due\n  Thank you      KP-1         $2,450.75\n"
    got = teach.read_with(text, [{"field": "total", "label": "Amount Due", "position": "below"}], as_of=AS_OF)
    assert got == {"total": {"value": "2450.75", "label": "Amount Due"}}


def test_a_label_and_value_read_as_one_box():
    found = [{"x": 0.1, "y": 0.2, "w": 0.4, "h": 0.02, "text": "Total due $4,210.00"}]
    fix = teach.field_fix(found, {"x": 0.3, "y": 0.195, "w": 0.21, "h": 0.03}, "total", as_of=AS_OF)
    assert fix["anchor"]["label"] == "Total due" and fix["was"] == "$4,210.00"


def test_a_label_found_elsewhere_with_the_wrong_kind_of_value_fills_nothing():
    text = "Pay by cheque or wire\nPay by 12/01/2026"
    got = teach.read_with(text, [{"field": "due_date", "label": "Pay by", "position": "right"}], as_of=AS_OF)
    assert got == {"due_date": {"value": "2026-12-01", "label": "Pay by"}}


def test_taught_values_win_over_guesses_and_yours_over_learnt():
    fields = ExtractedFields(amounts=[10000.0, 2118.4], invoice_numbers=["X1"], due_dates=["2026-11-01"])
    learnt = teach.with_taught(fields, "total", "2118.4", by="learnt", label="Balance payable")
    assert learnt.primary_amount == 2118.4
    mine = ExtractedFields(taught={"total": {"value": "999.0", "by": "you", "label": ""}})
    merged = learnt.merged_with(mine)
    assert merged.primary_amount == 999.0
    assert ExtractedFields.from_dict(merged.to_dict()).primary_amount == 999.0
    assert ExtractedFields.from_dict({"amounts": [5.0]}).taught == {}, "fields stored before this keep working"


def test_a_short_label_isnt_found_inside_a_longer_word():
    text = "Reference 99812 for your records\nRef # KP-1001"
    got = teach.read_with(text, [{"field": "invoice_number", "label": "Ref #", "position": "right"}], as_of=AS_OF)
    assert got == {"invoice_number": {"value": "KP-1001", "label": "Ref #"}}


def test_a_sheet_formatted_to_its_last_row_is_shown_quickly():
    import time

    book = Workbook()
    ws = book.active
    ws["A1"], ws["B1"] = "Vendor", "Amount"
    ws["A2"], ws["B2"] = "Northwind", 1250
    ws["A900000"] = "far down"
    out = io.BytesIO()
    book.save(out)
    started = time.monotonic()
    sheet = sheet_preview.preview(out.getvalue(), "big.xlsx", rows=50)["sheets"][0]
    assert time.monotonic() - started < 10
    assert [row["n"] for row in sheet["rows"]] == [1, 2], "blank rows after the last shown one aren't sent"
    assert sheet["more_rows"] == 1


def test_teaching_keeps_the_local_models_reading(client, store, mail):
    email = store.get_email(mail["Your Keystone invoice"].id)
    email.folder, email.summary, email.model_status = "important", "Keystone invoice for copy paper, due mid-November.", "bionic"
    store.upsert_email(email)
    _teach(client, email, "total", "$2,118.40")
    stored = store.get_email(email.id)
    assert (stored.folder, stored.summary, stored.model_status) == ("important", "Keystone invoice for copy paper, due mid-November.", "bionic")
    assert stored.extracted.primary_amount == 2118.40
