"""Browser and security sweep: no framing by other sites, safe links, a coherent task list after "Not fraud", amounts
without a currency mark, blocked mail kept out of AP coding and the classic file page, previews of merges, HTML and .ics."""

from __future__ import annotations

import io
import json
import re
import shutil
import subprocess
from datetime import datetime, timedelta, timezone
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from controller_inbox import cost_codes, web
from controller_inbox.folder_mail import ingest_folder
from msgfactory import PDF, XLSX, write_msg

UI = Path(__file__).resolve().parent.parent / "src" / "controller_inbox" / "static" / "ui"
PAGE = {"X-CloseDesk": "1"}
SENT = datetime(2026, 10, 19, 13, 0, tzinfo=timezone.utc)


def _bolt_invoice(number: str, amount: str) -> bytes:
    """A letterhead, a block of details on the right, a parts table and "Amount Due" with no currency mark."""
    out = io.BytesIO()
    c = canvas.Canvas(out, pagesize=letter)
    c.setFont("Helvetica-Bold", 20)
    c.drawString(72, 720, "Bolt Fasteners Inc")
    c.setFont("Helvetica", 10)
    c.drawString(72, 704, "77 Rivet Lane, Akron OH 44301")
    c.setFont("Helvetica-Bold", 16)
    c.drawString(420, 720, "INVOICE")
    c.setFont("Helvetica", 11)
    for i, (key, value) in enumerate([("Invoice #", number), ("Date", "10/14/2026"), ("Due", "11/13/2026"), ("PO", "PO-9001"), ("Job Site", "Riverside Plant")]):
        c.drawString(380, 690 - 16 * i, key)
        c.drawString(460, 690 - 16 * i, value)
    y = 580
    c.setFont("Helvetica-Bold", 11)
    for x, text in [(72, "Part"), (200, "Description"), (400, "Qty"), (450, "Price"), (510, "Ext")]:
        c.drawString(x, y, text)
    c.setFont("Helvetica", 11)
    for row in [("B-1", "Lag bolts", "100", "2.00", "200.00"), ("B-2", "Washers", "500", "0.10", "50.00")]:
        y -= 18
        for x, text in zip([72, 200, 400, 450, 510], row):
            c.drawString(x, y, text)
    y -= 30
    c.setFont("Helvetica-Bold", 12)
    c.drawString(400, y, "Amount Due")
    c.drawString(510, y, amount)
    c.showPage()
    c.save()
    return out.getvalue()


def _subledger() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "AP subledger"
    sheet.append(["Vendor", "Invoice", "Amount"])
    for row in [("Bolt Fasteners", "BF-5001", 250), ("Acme", "INV-88123", 4319.5), ("Northwind", "INV-10482", 12850)]:
        sheet.append(list(row))
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _simple_pdf(lines: list[str]) -> bytes:
    out = io.BytesIO()
    c = canvas.Canvas(out, pagesize=letter)
    c.setFont("Helvetica", 12)
    for i, line in enumerate(lines):
        c.drawString(72, 720 - 20 * i, line)
    c.showPage()
    c.save()
    return out.getvalue()


def _eml(folder: Path, name: str, *, subject: str, body: str, sender: str, address: str, files, when=SENT) -> None:
    message = EmailMessage()
    message["From"] = f"{sender} <{address}>"
    message["To"] = "ap@horizongoods.example"
    message["Subject"] = subject
    message["Date"] = format_datetime(when)
    message["Message-ID"] = f"<{name}@{address.split('@')[1]}>"
    message.set_content(body)
    for filename, data, mime in files:
        main, sub = mime.split("/")
        message.add_attachment(data, maintype=main, subtype=sub, filename=filename)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / f"{name}.eml").write_bytes(bytes(message))


def _ingest(store, settings) -> dict:
    return {email.subject: store.get_email(email.id) for email in ingest_folder(store, settings)}


def _fraud_with_files(settings) -> None:
    _eml(
        settings.inbox_incoming,
        "fraud-att",
        subject="URGENT change of bank details - Bolt Fasteners",
        body="Our bank details have changed. Please update your records and send all future payments to the new account: "
        "routing 021000021 account 99887766. Please pay invoice BF-5001 $250.00 today to the new account. Keep this confidential.",
        sender="Bolt Fasteners Accounts",
        address="accounts@b0ltfasteners.example",
        files=[("new-bank-letter.pdf", _bolt_invoice("BF-5001", "250.00"), PDF), ("bank.xlsx", _subledger(), XLSX)],
    )


@pytest.fixture
def client(store, settings, mail):
    return TestClient(web.create_app(settings, store))


# 1. Every page refuses to be framed by another site; the PDF preview may still be framed by CloseDesk itself.


def test_every_page_refuses_to_be_framed_by_another_site(client, mail):
    """The workspace, the classic pages and the API answer with X-Frame-Options DENY and frame-ancestors 'none'."""
    budget = mail["Q4 budget draft"]
    for path in ("/app", "/", "/inbox", f"/inbox/{budget.id}", "/settings", "/static/ui/main.js"):
        response = client.get(path)
        assert response.headers.get("x-frame-options") == "DENY", path
        assert "frame-ancestors 'none'" in response.headers.get("content-security-policy", ""), path
    api = client.get("/api/mail", headers=PAGE)
    assert api.headers.get("x-frame-options") == "DENY"


def test_the_original_pdf_is_still_framed_by_closedesk_itself(store, settings):
    """/files/{n}/view keeps its own same-origin framing rule, which the Preview tab's iframe needs."""
    settings.ensure_data_dir()
    _eml(settings.inbox_incoming, "bolt-1", subject="Invoice BF-5001 from Bolt Fasteners", body="Invoice attached.",
         sender="Bolt Fasteners AR", address="ar@boltfasteners.example", files=[("BF-5001.pdf", _bolt_invoice("BF-5001", "250.00"), PDF)])
    [email] = _ingest(store, settings).values()
    response = TestClient(web.create_app(settings, store)).get(f"/inbox/{email.id}/files/1/view")
    assert response.status_code == 200
    assert response.headers["x-frame-options"] == "SAMEORIGIN"
    assert response.headers["content-security-policy"] == "frame-ancestors 'self'"


# 2. A link starting "/\" is a link to another site (browsers read it as "//"); the page drops it.


def test_safe_url_refuses_backslash_protocol_relative_links():
    """dom.js's SAFE_URL keeps same-site paths, anchors and mailto: and refuses "//x", "/\\x" and "\\\\x"."""
    node = shutil.which("node")
    if not node:
        pytest.skip("Node.js isn't installed")
    source = (UI / "dom.js").read_text(encoding="utf-8")
    pattern = re.search(r"const SAFE_URL = (/.*/[a-z]*);", source).group(1)
    tries = ["/app", "/inbox/x?y=1", "#top", "mailto:a@b.c", "//evil.example", "/\\evil.example", "/\t/evil.example", "\\\\evil", "javascript:alert(1)", "https://evil.example"]
    script = f"const re = {pattern};\nconsole.log(JSON.stringify({json.dumps(tries)}.map(u => re.test(u))));\n"
    result = json.loads(subprocess.run([node, "-e", script], capture_output=True, text=True, check=True).stdout)
    assert result == [True, True, True, True, False, False, False, False, False, False]


# 5. An email held as possible payment fraud is not offered for AP coding.


def test_a_blocked_email_is_not_offered_for_ap_coding(store, settings):
    """A bank-change email carrying an invoice is held: it gets no cost-code suggestion and isn't on the coding list."""
    settings.ensure_data_dir()
    _fraud_with_files(settings)
    scam = _ingest(store, settings)["URGENT change of bank details - Bolt Fasteners"]
    assert "fraud_risk" in scam.flags
    client = TestClient(web.create_app(settings, store))
    assert store.cost_coding(scam.id) is None
    listed = client.get("/api/coding", params={"status": "all"}, headers=PAGE).json()
    assert scam.id not in [item["id"] for item in listed["items"]]
    assert listed["counts"]["all"] == 0


# 6. The classic file page doesn't show a blocked email's file text, as the workspace doesn't.


def test_classic_file_page_refuses_a_blocked_emails_file(client, mail):
    """GET /inbox/<blocked>/files/1 is refused with 403 and shows none of the file's text until the email is cleared."""
    scam = mail["Updated remittance details"]
    refused = client.get(f"/inbox/{scam.id}/files/1")
    assert refused.status_code == 403
    assert "8899" not in refused.text and "Remit to account" not in refused.text
    client.post(f"/inbox/{scam.id}/fraud", data={"choice": "safe:email"})
    opened = client.get(f"/inbox/{scam.id}/files/1")
    assert opened.status_code == 200 and "Remit to account" in opened.text


# 7. Invoice numbers listed in a workbook on an email are not that email's own invoices.


def test_an_invoice_listed_in_a_subledger_is_no_duplicate(store, settings):
    """Email A's bank.xlsx lists INV-10482; email B's invoice INV-10482 is not flagged as a possible duplicate."""
    settings.ensure_data_dir()
    _fraud_with_files(settings)
    first = _ingest(store, settings)["URGENT change of bank details - Bolt Fasteners"]
    assert "INV-10482" not in first.extracted.invoice_numbers
    _eml(
        settings.inbox_incoming,
        "nw-10482",
        subject="Invoice INV-10482 from Northwind Supplies",
        body="Please find invoice INV-10482 attached. Total $12,850.00.",
        sender="Northwind Billing",
        address="billing@northwind.example",
        files=[("INV-10482.pdf", _simple_pdf(["Northwind Supplies", "INVOICE", "Invoice No: INV-10482", "Total Due: $12,850.00"]), PDF)],
        when=SENT + timedelta(hours=2),
    )
    second = _ingest(store, settings)["Invoice INV-10482 from Northwind Supplies"]
    assert "duplicate_invoice" not in second.flags
    assert not [a for a in second.actions if "duplicate" in a.title.lower()]


# 8. A loose dropped file asking for a bank change is scored as the same words in an email are.


def test_a_loose_bank_change_note_is_held_as_an_email_would_be(store, settings):
    """A dropped .txt saying the bank details changed, wire today, keep it confidential is held as fraud."""
    settings.ensure_data_dir()
    (settings.inbox_incoming / "loose-bank-change.txt").write_bytes(
        b"From the CFO: our bank details have changed. Send the $48,000 wire today to routing 021000021 "
        b"account 11223344. Keep this confidential, I'm in meetings."
    )
    [note] = _ingest(store, settings).values()
    check = store.fraud_check(note.id)
    assert check["level"] == "high", check
    assert "fraud_risk" in note.flags


# 4. "Amount Due 250.00" with no currency mark is the email's amount.


def test_amount_due_without_a_currency_mark_is_the_emails_amount(store, settings):
    """The invoice's printed "Amount Due 250.00" becomes the email's amount, as the Page tab reads it."""
    settings.ensure_data_dir()
    _eml(
        settings.inbox_incoming,
        "bolt-1",
        subject="Invoice BF-5001 from Bolt Fasteners",
        body="Invoice attached.\nBolt AR",
        sender="Bolt Fasteners AR",
        address="ar@boltfasteners.example",
        files=[("BF-5001.pdf", _bolt_invoice("BF-5001", "250.00"), PDF)],
    )
    email = _ingest(store, settings)["Invoice BF-5001 from Bolt Fasteners"]
    assert email.extracted.primary_amount == 250.0
    client = TestClient(web.create_app(settings, store))
    [card] = [item for item in client.get("/api/mail", headers=PAGE).json()["items"] if item["id"] == email.id]
    assert "250.00" in json.dumps(card)


def test_an_invoice_table_keeps_qty_and_price_apart_and_checks_its_amount_due(store, settings):
    """"Amount Due" written across the Qty and Price columns doesn't merge them, and is checked as the table's total."""
    settings.ensure_data_dir()
    _eml(settings.inbox_incoming, "bolt-1", subject="Invoice BF-5001 from Bolt Fasteners", body="Invoice attached.",
         sender="Bolt Fasteners AR", address="ar@boltfasteners.example", files=[("BF-5001.pdf", _bolt_invoice("BF-5001", "250.00"), PDF)])
    [email] = _ingest(store, settings).values()
    data = TestClient(web.create_app(settings, store)).get(f"/api/mail/{email.id}/files/1", headers=PAGE).json()
    [parts] = [table for table in data["tables"] if "Part" in table["labels"]]
    assert parts["labels"] == ["Part", "Description", "Qty", "Price", "Ext"]
    assert parts["check"]["matched"] == 1 and not parts["check"]["mismatched"]


# 3. "Not fraud" on the sample's bank-change wire leaves one coherent list: the phone call first, then the payment.


def test_not_fraud_on_the_sample_wire_puts_the_phone_call_first(store, settings):
    """After "Not fraud" on demo-bec-wire, the call is the most urgent task, no "process the wire today" task
    contradicts it, and the new tasks are dated on the sample board's day (Sep 22), not today."""
    from controller_inbox.cli import load_sample

    settings.ensure_data_dir()
    load_sample(store, settings)
    client = TestClient(web.create_app(settings, store))
    assert client.post("/api/mail/demo-bec-wire/fraud", json={"choice": "safe:email"}, headers=PAGE).json()["ok"]
    email = store.get_email("demo-bec-wire")
    open_tasks = [item for item in email.actions if item.status.value == "open"]
    titles = [item.title for item in open_tasks]
    assert titles[0] == "Confirm the bank-detail change by phone before paying"
    assert not [t for t in titles if "process" in t.lower()], titles
    rank = ["low", "medium", "high", "critical"]
    call = open_tasks[0]
    assert all(rank.index(call.priority.value) >= rank.index(item.priority.value) for item in open_tasks)
    assert {item.due_date for item in open_tasks if item.due_date} == {"2026-09-22"}


# 9. Spreadsheet previews keep merged cells: an .xls's, and the last row of an .xlsx merge at the bottom of a sheet.


def test_xlsx_preview_keeps_the_last_row_of_a_merge_at_the_bottom(tmp_path):
    """A merge A9:B10 with nothing below it keeps row 10 in the grid, so the merged cell spans two shown rows."""
    from controller_inbox import sheet_preview

    book = Workbook()
    sheet = book.active
    sheet["A1"] = "Item"
    sheet["A9"] = "Approved by Dana"
    sheet.merge_cells("A9:B10")
    out = io.BytesIO()
    book.save(out)
    grid = sheet_preview.preview(out.getvalue(), "book.xlsx")["sheets"][0]
    [merged] = [cell for row in grid["rows"] if row["n"] == 9 for cell in row["cells"]]
    assert merged["rs"] == 2 and merged["cs"] == 2
    assert [row["n"] for row in grid["rows"]][-1] == 10


def test_xls_preview_keeps_merged_cells(monkeypatch):
    """An .xls's merged title (A1:C1) spans three columns and the cells it covers aren't shown on their own."""
    import xlrd
    from xlrd.sheet import Cell

    from controller_inbox import sheet_preview

    values = {(0, 0): "Q3 Budget", (1, 0): "Line", (1, 1): "Q3", (1, 2): "Q4"}

    class Sheet:
        name, nrows, ncols, visibility = "Budget", 2, 3, 0
        merged_cells = [(0, 1, 0, 3)]

        def cell(self, r, c):
            value = values.get((r, c))
            return Cell(xlrd.XL_CELL_TEXT, value) if value else Cell(xlrd.XL_CELL_EMPTY, "")

    class Book:
        nsheets, datemode = 1, 0

        def sheet_by_index(self, i):
            return Sheet()

        def release_resources(self):
            pass

    monkeypatch.setattr(xlrd, "open_workbook", lambda **kw: Book())
    data = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 504
    rows = sheet_preview.preview(data, "book.xls")["sheets"][0]["rows"]
    [title] = rows[0]["cells"]
    assert title["v"] == "Q3 Budget" and title["cs"] == 3 and title["rs"] == 1
    assert [cell["v"] for cell in rows[1]["cells"]] == ["Line", "Q3", "Q4"]


# 10. A web page attachment's tables are shown as tables (text only); a calendar file as its events.


def test_an_html_attachment_preview_keeps_its_tables_as_text_only():
    """statement.html previews as a heading and a table of text cells; its script, hidden text and markup are gone."""
    from controller_inbox import sheet_preview

    page = (
        b"<html><head><title>Statement</title><script>window.parent.__xss='html'</script></head><body><h1>Statement</h1>"
        b"<p>Balance <b>due</b><span style='display:none'>zz</span></p><table><tr><th>Item</th><th>Amount</th></tr>"
        b"<tr><td>Balance</td><td>1,234.00</td></tr><tr><td colspan=2><img src=x onerror=alert(5)>Thanks</td></tr></table></body></html>"
    )
    got = sheet_preview.preview(page, "statement.html")
    assert got["kind"] == "doc"
    assert got["blocks"] == [
        {"t": "h2", "text": "Statement"},
        {"t": "p", "text": "Balance due"},
        {"t": "table", "rows": [["Item", "Amount"], ["Balance", "1,234.00"], [{"v": "Thanks", "cs": 2}]]},
    ]
    assert "xss" not in json.dumps(got) and "onerror" not in json.dumps(got)


def test_an_ics_attachment_previews_as_its_event():
    """meeting.ics shows its summary, start, end, place and organizer instead of the raw iCalendar lines."""
    from controller_inbox import sheet_preview

    ics = (
        b"BEGIN:VCALENDAR\r\nVERSION:2.0\r\nBEGIN:VEVENT\r\nUID:1@x\r\nDTSTART:20261024T140000Z\r\nDTEND:20261024T150000Z\r\n"
        b"SUMMARY:Q4 planning\r\nLOCATION:Room 4\\, 2nd floor\r\nORGANIZER;CN=Dana Cole:mailto:dana@taz.com\r\nEND:VEVENT\r\nEND:VCALENDAR\r\n"
    )
    got = sheet_preview.preview(ics, "meeting.ics")
    assert got["kind"] == "doc"
    assert got["blocks"][:2] == [
        {"t": "h2", "text": "Q4 planning"},
        {"t": "table", "rows": [["Starts", "Sat, Oct 24, 2026, 2:00 PM UTC"], ["Ends", "Sat, Oct 24, 2026, 3:00 PM UTC"],
                                ["Where", "Room 4, 2nd floor"], ["Organizer", "Dana Cole dana@taz.com"]]},
    ]
    assert "BEGIN:VEVENT" not in json.dumps(got["blocks"]) and "SUMMARY:Q4 planning" in got["source"]


# 11. Ask CloseDesk without a model: "due this week" keeps the matches due then; a letterhead isn't read as a table.


def test_a_lookup_for_invoices_due_this_week_leaves_out_later_and_undated_ones(store, settings):
    """On the sample board (Sep 22) no invoice is due this week; next week's is Harbor Packaging's (Oct 1) only."""
    from controller_inbox.assistant import answer_stream
    from controller_inbox.cli import load_sample

    settings.ensure_data_dir()
    load_sample(store, settings)

    def ask(question: str) -> str:
        return "".join(e.get("text", "") for e in answer_stream(store, settings, question, today="2026-09-22") if e["type"] == "delta")

    this_week = ask("Which invoices are due this week?")
    assert "None of the emails that match is due this week" in this_week and "INV-10482" not in this_week
    next_week = ask("Which invoices are due next week?")
    assert "Harbor Packaging" in next_week and "INV-10482" not in next_week


def test_a_letterhead_beside_the_invoice_details_is_not_read_as_a_table():
    """The vendor's name and address beside "INVOICE" and its details read as lines, not "Bolt Fasteners Inc: not listed"."""
    from controller_inbox import documents

    text = documents.pdf_text(_bolt_invoice("BF-5001", "250.00"))
    assert "not listed | INVOICE" not in text and "Bolt Fasteners Inc: not listed" not in text
    assert "Invoice # BF-5001" in text and "Due 11/13/2026" in text
